"""The campaign archive workbook, format version 1.0 (#1018).

One XLSX workbook per campaign holds its Family-level results, every source
submission, the day-by-day series and provenance. This module is the one
owner of the format: the table of sheets, columns and metadata keys below is
what the README sheet's dictionary is generated from, what the reader checks
headers and cells against, and what a test compares with the tables in
docs/specs/stewardship/campaign-history/spec.md ("Archive workbook format").

``read_archive`` is the strict reader: it refuses an unsafe package
(``package``), then checks the schema, every cell and the cross-sheet
consistency rules, and returns the validated rows or every located problem
(up to 50). A problem names a sheet, a cell and the rule, never the value it
found, the header text it found or an unknown key's name, because a value
may be a Family identifier or an amount. ``build_workbook`` and
``write_archive`` are the minimal writer the tests use; the end-of-campaign
export (#1018 part (a)) will fill them with real data later.

The pledge band set and the quartile calculation of the comparison report
live here too, beside the format, so a band-set change is one edit.
"""

import bisect
import json
import math
import re
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from enum import StrEnum
from io import BytesIO
from zoneinfo import ZoneInfo

from .package import PackageRefused, check_package

SCHEMA = "stewardship-campaign-archive"
# The version this reader implements; a file of the same major version and
# any minor version is read, ignoring what a later minor version added.
READER_VERSION = (1, 0)
SCHEMA_VERSION = "1.0"
MAX_PROBLEMS = 50
MAX_ROWS = 500_000
# Bounds no real value comes near, so arithmetic and storage stay bounded:
# an id or count below 10**15, an amount below a trillion dollars.
MAX_INTEGER = 10**15
MAX_MONEY = Decimal(10**12)
CENT = Decimal("0.01")
# The stored columns' limits (migration 0059), so a file the reader accepts
# never fails on save: an integer column, a numeric(18, 2) amount, and the
# archive guard's 256 KB bound on the note.* and kept count.* values (as
# PostgreSQL prints the JSON object). Text keys carry their own max_length.
STORED_INTEGER_MAX = 2**31 - 1
STORED_MONEY_LIMIT = Decimal(10**16)
STORED_NOTES_BYTES = 256 * 1024
PERCENT_TOLERANCE = 0.0005
# Forms that asked the financial question; 2020's ministry form did not.
FINANCIAL_FORMS = ("combined", "pledge")


class Kind(StrEnum):
    """A cell type of the format; the value is its name in the specification."""

    TEXT = "Text"
    INTEGER = "Integer"
    MONEY = "Money"
    NUMBER = "Number"
    BOOLEAN = "Boolean"
    INSTANT = "Instant"
    DATE = "Local date"
    ENUM = "Enumeration"


@dataclass(frozen=True)
class Column:
    """One column of a data sheet: its name, type, range and description."""

    name: str
    kind: Kind
    description: str
    optional: bool = False
    minimum: int | None = None
    choices: tuple[str, ...] = ()

    @property
    def type_text(self):
        """The type as the specification's column tables word it."""
        text = str(self.kind)
        if self.minimum is not None:
            text += f" ≥ {self.minimum}"
        return text + (" or blank" if self.optional else "")


@dataclass(frozen=True)
class MetadataKey:
    """One defined metadata key: its value's type, range and description."""

    name: str
    kind: Kind
    description: str
    optional: bool = False
    minimum: int | None = None
    maximum: int | None = None
    max_length: int | None = None


FREQUENCIES = ("weekly", "monthly", "quarterly", "annual", "none", "other")
PLEDGE_STATUSES = ("pledged", "zero", "cannot_give", "no_answer", "no_pledge_form")
# A family_key's prefix, and the family_id_system it belongs to.
FAMILY_ID_SYSTEMS = {
    "duid": "parishsoft_duid",
    "pds_fid": "pds_fid",
    "pds_env": "pds_envelope",
}
# Envelope numbers are text and may keep leading zeros; ids never do.
FAMILY_KEY = re.compile(r"(duid|pds_fid):([1-9][0-9]{0,14})|pds_env:[0-9]{1,15}")

_T, _I, _M = Kind.TEXT, Kind.INTEGER, Kind.MONEY
_B, _A, _D, _E = Kind.BOOLEAN, Kind.INSTANT, Kind.DATE, Kind.ENUM

# The format table. Order matters: it is the header row of each sheet, and
# the writer's sheet order (after README). Descriptions come from the
# one-time converter's table and become the README dictionary.
SHEETS: dict[str, tuple[Column, ...]] = {
    "families": (
        Column(
            "family_key",
            _T,
            "Key of the Family in this file: 'duid:<n>' (ParishSoft Family "
            "DUID), 'pds_fid:<n>' (old PDS family id) or 'pds_env:<n>' (old "
            "PDS envelope number, only when no family id exists).",
        ),
        Column(
            "family_id",
            _I,
            "The number in family_key; blank for an envelope-only key.",
            optional=True,
        ),
        Column(
            "family_id_system",
            _E,
            "parishsoft_duid, pds_fid or pds_envelope. PDS ids (2020-2024) and "
            "ParishSoft DUIDs (2025 on) are different id spaces.",
            choices=tuple(FAMILY_ID_SYSTEMS.values()),
        ),
        Column(
            "envelope_id",
            _T,
            "Old PDS envelope number (2020-2024 only).",
            optional=True,
        ),
        Column(
            "submission_count",
            _I,
            "Attributed submissions the Family made, all forms.",
            minimum=1,
        ),
        Column("first_submitted_at", _A, "First submission, ISO 8601 UTC."),
        Column("first_local_date", _D, "Campaign-local date of the first submission."),
        Column(
            "first_campaign_day",
            _I,
            "Campaign day of the first submission; day 1 is campaign_start_date.",
            minimum=1,
        ),
        Column(
            "last_submitted_at",
            _A,
            "Last submission, ISO 8601 UTC. Only this submission counts.",
        ),
        Column("last_local_date", _D, "Campaign-local date of the last submission."),
        Column(
            "last_campaign_day", _I, "Campaign day of the last submission.", minimum=1
        ),
        Column(
            "pledge_status",
            _E,
            "pledged (annual_pledge greater than zero), zero (0 entered), "
            "cannot_give (the Family said it cannot contribute), no_answer "
            "(financial question left blank) or no_pledge_form (2020 only: the "
            "Family sent only the separate ministry form).",
            choices=PLEDGE_STATUSES,
        ),
        Column(
            "annual_pledge",
            _M,
            "Annual pledge in US dollars from the Family's last financial answer.",
            optional=True,
        ),
        Column(
            "frequency",
            _E,
            "weekly, monthly, quarterly, annual, none or other.",
            optional=True,
            choices=FREQUENCIES,
        ),
        Column(
            "frequency_detail",
            _T,
            "The choice as the form worded it; free-text answers only as "
            "'other (free text omitted)'.",
            optional=True,
        ),
        Column(
            "installment_amount",
            _M,
            "Per-installment amount: the form's own computed column, or the "
            "annual pledge divided by 52, 12, 4 or 1.",
            optional=True,
            minimum=0,
        ),
        Column(
            "installment_source",
            _E,
            "form or computed; blank exactly when installment_amount is blank.",
            optional=True,
            choices=("form", "computed"),
        ),
        Column(
            "cannot_give",
            _B,
            "TRUE when the Family checked the 'cannot contribute financially' "
            "box; blank where the form had none.",
            optional=True,
        ),
        Column(
            "prior_year_pledge",
            _M,
            "The previous year's pledge as pre-filled on the form (old system only).",
            optional=True,
            minimum=0,
        ),
        Column(
            "prior_year_contributions",
            _M,
            "The previous year's contributions as pre-filled on the form (old "
            "system only); may be negative after refunds or adjustments.",
            optional=True,
        ),
        Column(
            "pledge_parse_note",
            _T,
            "Set when the entered pledge text needed interpretation.",
            optional=True,
        ),
    ),
    "submissions": (
        Column("row_id", _I, "1, 2, 3 ... in row order, oldest submission first."),
        Column(
            "family_key",
            _T,
            "As in families; blank exactly when attributed is FALSE.",
            optional=True,
        ),
        Column(
            "form",
            _E,
            "combined, pledge or ministry (2020 used two forms).",
            choices=("combined", "pledge", "ministry"),
        ),
        Column("submitted_at", _A, "ISO 8601 UTC."),
        Column("local_date", _D, "Campaign-local date."),
        Column(
            "campaign_day", _I, "Campaign day (day 1 = campaign_start_date).", minimum=1
        ),
        Column(
            "annual_pledge",
            _M,
            "The annual pledge in this submission.",
            optional=True,
            minimum=0,
        ),
        Column(
            "frequency",
            _E,
            "As in families.",
            optional=True,
            choices=FREQUENCIES,
        ),
        Column("cannot_give", _B, "As in families.", optional=True),
        Column(
            "is_last_for_family",
            _B,
            "TRUE for the Family's last submission of this form.",
        ),
        Column(
            "attributed",
            _B,
            "FALSE for a row with no Family identity, which no count or total "
            "includes.",
        ),
    ),
    "daily": (
        Column("campaign_day", _I, "1, 2, 3 ... with no gap."),
        Column("local_date", _D, "campaign_start_date + campaign_day - 1."),
        Column(
            "first_submissions",
            _I,
            "Families whose first submission fell on this day.",
            minimum=0,
        ),
        Column(
            "cumulative_families",
            _I,
            "Distinct Families that had submitted by the end of this day.",
            minimum=0,
        ),
        Column(
            "cumulative_pct_of_active",
            Kind.NUMBER,
            "cumulative_families / active_families as a fraction from 0 to 1.",
        ),
        Column(
            "cumulative_pledging_families",
            _I,
            "Families whose latest answer by the end of this day pledged more "
            "than zero.",
            minimum=0,
        ),
        Column(
            "cumulative_annual_pledged",
            _M,
            "Sum of each Family's latest annual pledge at the end of this day; "
            "a revision replaces, never adds.",
            minimum=0,
        ),
        Column(
            "source",
            _T,
            "derived (recomputed from submissions), or a note naming the system "
            "table it equals.",
        ),
    ),
    "ministry_responses": (
        Column("family_key", _T, "A families key."),
        Column(
            "member_slot",
            _I,
            "Member position on the Family's form (1 = first listed).",
            minimum=1,
        ),
        Column(
            "member_id",
            _I,
            "Old PDS member id or ParishSoft Member DUID.",
            optional=True,
        ),
        Column(
            "ministry",
            _T,
            "Ministry as labelled on the form (old system) or ParishSoft "
            "Ministry id (this system).",
        ),
        Column(
            "ministry_code",
            _T,
            "Leading three-digit Ministry number when the label has one.",
            optional=True,
        ),
        Column(
            "response",
            _T,
            "participate, interested or not_involved (old system); join or leave "
            "(this system, which records changes only); other text verbatim.",
        ),
    ),
    "metadata": (
        Column("key", _T, "Metadata key; see the list below."),
        Column(
            "value",
            _T,
            "Value as text: numbers as plain digits, dates and instants in ISO 8601.",
            optional=True,
        ),
        Column("note", _T, "Provenance or explanation.", optional=True),
    ),
}
SHEET_ORDER = ("README", *SHEETS)

METADATA_KEYS: tuple[MetadataKey, ...] = (
    MetadataKey("schema", _T, "Always 'stewardship-campaign-archive'."),
    MetadataKey(
        "schema_version",
        _T,
        "Format version; a reader refuses a major version it does not know.",
    ),
    MetadataKey(
        "campaign_name",
        _T,
        "Such as 'Stewardship 2027'; the name is the pledge year.",
        max_length=200,
    ),
    MetadataKey(
        "campaign_year",
        _I,
        "The pledge year (the year the pledges are for).",
        minimum=2000,
        maximum=2100,
    ),
    MetadataKey(
        "is_final",
        _B,
        "TRUE when the campaign had closed when exported; FALSE for a partial export.",
    ),
    MetadataKey(
        "as_of",
        _A,
        "Instant the data describes (UTC); for a final file, the last submission.",
    ),
    MetadataKey(
        "timezone",
        _T,
        "IANA zone used for local dates and day numbers.",
        max_length=64,
    ),
    MetadataKey("campaign_start_date", _D, "Local date of day 1."),
    MetadataKey(
        "campaign_last_submission_date",
        _D,
        "Local date of the last attributed submission.",
    ),
    MetadataKey(
        "active_families",
        _I,
        "Active parishioner Families, the percentage denominator.",
        minimum=1,
        maximum=STORED_INTEGER_MAX,
    ),
    MetadataKey(
        "active_families_provenance",
        _T,
        "Where active_families came from.",
        max_length=10_000,
    ),
    MetadataKey(
        "source_system",
        _T,
        "The system that collected the responses.",
        max_length=10_000,
    ),
    MetadataKey("source_files", _T, "Names of the original files converted."),
    MetadataKey("converted_at", _A, "When this file was produced (UTC)."),
    MetadataKey("converter", _T, "The program that produced it."),
    MetadataKey(
        "dedup_rule",
        _T,
        "How repeat submissions were reduced to one per Family.",
        max_length=10_000,
    ),
)
METADATA = {key.name: key for key in METADATA_KEYS}
# Optional, repeatable keys: the prefix, the accepted suffix and the README
# description. Their values are free text, except the counts below.
METADATA_PATTERNS = (
    (
        "column_mapping.",
        re.compile(r"[0-9]{1,4}"),
        "One per source column used, with the archive column it fed.",
    ),
    ("note.", re.compile(r"[0-9]{1,4}"), "Conversion notes and assumptions."),
    (
        "count.",
        re.compile(r"[a-z][a-z0-9_]{0,59}"),
        "Aggregates for quick checks; the reader recomputes the ones it knows.",
    ),
)
# The counts the reader recomputes; any other count.* is kept as a note.
KNOWN_COUNTS = (
    "count.families",
    "count.rows_read",
    "count.unattributed_rows",
    "count.ministry_response_rows",
)
INSTANT = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}Z")
# A C0 control character other than tab and newline (NUL included): text
# holding one is refused, since PostgreSQL cannot store NUL in text or jsonb
# and the others are never meaningful in a cell or a note.
CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f]")
LOCAL_DATE = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}")
VERSION = re.compile(r"([1-9][0-9]{0,2})\.(0|[1-9][0-9]{0,2})")
TIME_ZONE = re.compile(r"[A-Za-z][A-Za-z0-9_+-]*(?:/[A-Za-z0-9_+-]+){0,2}")

README_LINES = (
    f"{SCHEMA} version {SCHEMA_VERSION}",
    "One Stewardship campaign's results: one row per Family in 'families' (its "
    "last submission only), every source submission in 'submissions', the "
    "day-by-day series in 'daily', the Ministry answers of each Family's last "
    "submission in 'ministry_responses', and provenance in 'metadata'.",
    "PRIVATE: this file holds Family-level giving. Keep it only in the parish's "
    "restricted archive. It holds no names, email or postal addresses, phone "
    "numbers, IP addresses, comments, talents, free-text answers, Family codes, "
    "credentials or edit links.",
    "Conventions: instants are ISO 8601 UTC text such as 2026-10-03T14:05:09Z; "
    "local dates are ISO 8601 dates in the campaign's time zone; money is US "
    "dollars with at most two decimals; booleans are TRUE/FALSE; blank means "
    "unknown or not applicable, never zero; cells hold values, never formulas.",
    "A reader checks the schema and version in 'metadata', every header and "
    "cell, and the sums across sheets; a corrected file must keep them in step.",
)


# --- Problems ---------------------------------------------------------------


@dataclass(frozen=True)
class Problem:
    """One located problem: a closed ``code`` and fixed ``rule`` text.

    ``sheet`` and ``cell`` locate it when it has a place. Nothing here is
    ever taken from the file's own text.
    """

    code: str
    rule: str
    sheet: str | None = None
    cell: str | None = None

    @property
    def text(self):
        """The problem as one line, such as ``families!J17: <rule>``."""
        if self.sheet is None:
            return self.rule
        where = self.sheet if self.cell is None else f"{self.sheet}!{self.cell}"
        return f"{where}: {self.rule}"

    def as_dict(self):
        """The problem as JSON-ready fields."""
        return {
            "code": self.code,
            "sheet": self.sheet,
            "cell": self.cell,
            "rule": self.rule,
        }


class ArchiveRefused(Exception):
    """The workbook is refused: up to 50 problems, and how many more there were."""

    def __init__(self, problems, more=0):
        super().__init__("The archive workbook is refused.")
        self.problems = tuple(problems)
        self.more = more


class _Stop(Exception):
    """Reading cannot usefully continue; the problem is already recorded."""


class _Bad(Exception):
    """A cell value breaks its column's rule; ``rule`` follows the column name."""

    def __init__(self, code, rule):
        super().__init__(rule)
        self.code = code
        self.rule = rule


class _Problems:
    """The first MAX_PROBLEMS problems, and a count of the ones cut off."""

    def __init__(self):
        self.items = []
        self.more = 0

    def add(self, code, rule, sheet=None, cell=None):
        """Record one problem, or count it once the list is full."""
        if len(self.items) < MAX_PROBLEMS:
            self.items.append(Problem(code, rule, sheet, cell))
        else:
            self.more += 1

    def __bool__(self):
        return bool(self.items)


def _letter(index):
    """The column letter of a zero-based column index (A, B, ... Z, AA ...)."""
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _cell(sheet, column, row):
    """The cell reference of a named column of a sheet in one row."""
    names = [c.name for c in SHEETS[sheet]]
    return f"{_letter(names.index(column))}{row}"


# --- Cell values --------------------------------------------------------------

# Cells that hold no literal value, by openpyxl's data type: a formula,
# whose cached value may be stale or absent, and an error such as #N/A.
NOT_LITERAL = {
    "f": ("formula", "the cell holds a formula; only values are allowed"),
    "e": ("error_value", "the cell holds an error value; only values are allowed"),
}


def _blank(value):
    """Whether a cell is blank: empty, or an empty string."""
    return value is None or value == ""


def _whole(value):
    """The int of an integer cell or a whole-valued float; None otherwise.

    A boolean is not a number here, though Python treats it as one.
    """
    if type(value) is int:
        return value
    if type(value) is float and math.isfinite(value) and value.is_integer():
        return int(value)
    return None


def _money(value):
    """The exact Decimal of a money cell, through its shortest decimal text.

    A float's ``repr`` is the shortest text that reads back as the same
    double, so ``1234.56`` stored inexactly in binary is still exactly
    1234.56 here, and ``0.30000000000000004`` has more than two decimals.
    """
    if type(value) is int:
        amount = Decimal(value)
    elif type(value) is float and math.isfinite(value):
        amount = Decimal(repr(value))
    else:
        raise _Bad("type", "must be an amount in dollars")
    if amount.as_tuple().exponent < -2:
        raise _Bad("type", "must have at most two decimals")
    if abs(amount) >= MAX_MONEY:
        raise _Bad("range", "is out of range")
    return amount.quantize(CENT)


def _instant(value):
    """A UTC datetime from ISO 8601 text with a Z suffix and whole seconds."""
    if type(value) is not str or not INSTANT.fullmatch(value):
        raise _Bad("type", "must be a UTC instant such as 2026-10-03T14:05:09Z")
    try:
        return datetime.strptime(value, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)
    except ValueError:
        raise _Bad("type", "must be a real date and time") from None


def _local_date(value):
    """A date from ISO 8601 date text."""
    if type(value) is not str or not LOCAL_DATE.fullmatch(value):
        raise _Bad("type", "must be a date such as 2026-10-03")
    try:
        return date.fromisoformat(value)
    except ValueError:
        raise _Bad("type", "must be a real date") from None


def convert(column, value):
    """The typed value of one non-blank data cell, or raise _Bad.

    Text columns take text without C0 control characters (tab and newline
    excepted), or a whole number read as its digits; integers
    are whole numbers (a stored 12.0 reads as 12), never text; booleans are
    boolean cells. Excel date cells are refused everywhere, since they carry
    no time zone.
    """
    kind = column.kind
    if kind in (Kind.TEXT, Kind.ENUM):
        if type(value) is str:
            text = value
        elif kind is Kind.TEXT and _whole(value) is not None:
            text = str(_whole(value))
        else:
            raise _Bad("type", "must be text")
        if CONTROL.search(text):
            raise _Bad("control", "must not hold control characters")
        if kind is Kind.ENUM and text not in column.choices:
            raise _Bad("choice", "must be one of " + ", ".join(column.choices))
        return text
    if kind is Kind.INTEGER:
        number = _whole(value)
        if number is None:
            raise _Bad("type", "must be a whole number")
        if abs(number) >= MAX_INTEGER:
            raise _Bad("range", "is out of range")
        result = number
    elif kind is Kind.MONEY:
        result = _money(value)
    elif kind is Kind.NUMBER:
        if type(value) not in (int, float) or not math.isfinite(value):
            raise _Bad("type", "must be a number")
        result = float(value)
    elif kind is Kind.BOOLEAN:
        if type(value) is not bool:
            raise _Bad("type", "must be TRUE or FALSE")
        return value
    elif kind is Kind.INSTANT:
        return _instant(value)
    else:
        return _local_date(value)
    if column.minimum is not None and result < column.minimum:
        raise _Bad("range", f"must be at least {column.minimum}")
    return result


def _metadata_value(key, value):
    """The typed value of one non-blank defined metadata value, or raise _Bad.

    Values are text; a numeric key also takes a number cell and a boolean key
    a boolean cell.
    """
    if key.kind is Kind.INTEGER:
        if type(value) is str and re.fullmatch(r"[0-9]{1,15}", value):
            number = int(value)
        else:
            number = _whole(value)
            if number is None:
                raise _Bad("type", "must be a whole number")
        if (key.minimum is not None and number < key.minimum) or (
            key.maximum is not None and number > key.maximum
        ):
            raise _Bad("range", f"must be from {key.minimum} to {key.maximum}")
        if number >= MAX_INTEGER:
            raise _Bad("range", "is out of range")
        return number
    if key.kind is Kind.BOOLEAN:
        if type(value) is bool:
            return value
        if value in ("TRUE", "FALSE"):
            return value == "TRUE"
        raise _Bad("type", "must be TRUE or FALSE")
    column = Column(key.name, key.kind, key.description)
    typed = convert(column, value)
    if key.max_length is not None and len(typed) > key.max_length:
        raise _Bad("length", f"must be at most {key.max_length:,} characters")
    return typed


def stored_notes_bytes(notes):
    """The size of the stored notes object as PostgreSQL prints its jsonb.

    jsonb prints ``{"key": "value", ...}`` with these separators, raw UTF-8
    and the same escapes as ``json.dumps``; key order does not change the
    size.
    """
    return len(json.dumps(notes, ensure_ascii=False).encode())


def _time_zone(text):
    """The ZoneInfo of an IANA zone name, or raise _Bad."""
    if not TIME_ZONE.fullmatch(text):
        raise _Bad("type", "must be an IANA time zone such as America/New_York")
    try:
        return ZoneInfo(text)
    except (ValueError, OSError, KeyError):
        raise _Bad(
            "type", "must be an IANA time zone such as America/New_York"
        ) from None


# --- The reader -----------------------------------------------------------------


@dataclass
class Archive:
    """A validated archive: its typed metadata and the rows a caller may keep.

    ``families`` and ``daily`` are row dicts keyed by column name (with the
    sheet row in ``_row``); submissions and Ministry answers are only
    counted, since nothing past the checks uses them.
    """

    metadata: dict
    minor_version: int
    notes: dict[str, str]
    other_counts: dict[str, str]
    families: list[dict]
    daily: list[dict]
    submission_rows: int
    unattributed_rows: int
    ministry_rows: int
    family_id_systems: tuple[str, ...] = field(default=())


def read_archive(data):
    """Validate one workbook's bytes completely; return the Archive.

    Raises ArchiveRefused with the located problems. A MemoryError is let
    through, so the helper process can report "too large to check".
    """
    try:
        check_package(data)
    except PackageRefused as refused:
        raise ArchiveRefused([Problem(refused.code, refused.rule)]) from None
    problems = _Problems()
    archive = None
    with warnings.catch_warnings():
        # openpyxl warns about styles and extensions it does not read; the
        # warning text could echo file content, and none of it matters here.
        warnings.simplefilter("ignore")
        try:
            import openpyxl

            book = openpyxl.load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
        except MemoryError:
            raise
        except Exception:
            raise ArchiveRefused(
                [Problem("unreadable", "the file is not a readable .xlsx")]
            ) from None
        try:
            archive = _Reader(book, problems).read()
        except _Stop:
            pass
        except MemoryError:
            raise
        except Exception:
            # A worksheet that cannot be parsed part-way through. Nothing of
            # the exception is kept: its text could quote the file.
            problems.add("unreadable", "the file is not a readable .xlsx")
        finally:
            book.close()
    if problems or archive is None:
        raise ArchiveRefused(problems.items, problems.more)
    return archive


class _Reader:
    """One pass over an opened workbook, collecting problems as it goes."""

    def __init__(self, book, problems):
        self.book = book
        self.problems = problems
        self.rows = 0
        self.minor = 0
        self.metadata_extras = []
        # Sheets some of whose rows were not read (a wrong header, or a row
        # with a problem): a key missing from them proves nothing.
        self.incomplete = set()

    def add(self, *args):
        """Record a problem (see _Problems.add)."""
        self.problems.add(*args)

    def read(self):
        """Read and check every sheet, then the cross-sheet rules."""
        self._sheets()
        metadata, rows = self._metadata()
        self._unknown_sheets()
        self._readme()
        families = self._families()
        submissions = self._submissions()
        daily = self._daily()
        ministry_rows = self._ministry(families)
        if self.problems:
            raise _Stop
        _Consistency(metadata, rows, families, submissions, daily, self.add).check(
            ministry_rows
        )
        if self.problems:
            raise _Stop
        notes, other_counts = {}, {}
        for key, (_, value) in rows.items():
            prefix = _pattern_prefix(key)
            if prefix not in ("note.", "count.") or key in KNOWN_COUNTS:
                continue
            text = None if _blank(value) else convert(SHEETS["metadata"][1], value)
            (notes if prefix == "note." else other_counts)[key] = text
        if stored_notes_bytes({**notes, **other_counts}) > STORED_NOTES_BYTES:
            self.add(
                "notes_size",
                "the note.* and count.* values must total at most 256 KB",
                "metadata",
            )
            raise _Stop
        return Archive(
            metadata=metadata,
            minor_version=self.minor,
            notes=notes,
            other_counts=other_counts,
            families=list(families.values()),
            daily=daily,
            submission_rows=len(submissions),
            unattributed_rows=sum(not s["attributed"] for s in submissions),
            ministry_rows=ministry_rows,
            family_id_systems=tuple(
                sorted({f["family_id_system"] for f in families.values()})
            ),
        )

    # Sheets and rows

    def _sheets(self):
        """README first and every defined sheet present; missing ones are fatal."""
        names = self.book.sheetnames
        if not names or names[0] != "README":
            self.add("sheet_order", "the first sheet must be README", "workbook")
        missing = [name for name in SHEET_ORDER if name not in names]
        for name in missing:
            self.add("sheet_missing", f"the sheet {name} is missing", "workbook")
        if missing:
            raise _Stop

    def _unknown_sheets(self):
        """Refuse a sheet version 1.0 does not define; a later minor's are ignored."""
        if self.minor > READER_VERSION[1]:
            return
        for position, name in enumerate(self.book.sheetnames, start=1):
            if name not in SHEET_ORDER:
                self.add(
                    "sheet_not_defined",
                    f"sheet {position} is not defined in version {SCHEMA_VERSION}",
                    "workbook",
                )

    def _count(self):
        """Count one row read; past MAX_ROWS the file is too large."""
        self.rows += 1
        if self.rows > MAX_ROWS:
            self.add("too_many_rows", "the file has more than 500,000 rows")
            raise _Stop

    def _readme(self):
        """README needs no header; only its row count and formulas matter."""
        sheet = self.book["README"]
        sheet.reset_dimensions()
        for number, row in enumerate(sheet.iter_rows(), start=1):
            self._count()
            for index, cell in enumerate(row):
                if cell.data_type in NOT_LITERAL:
                    self.add(
                        *NOT_LITERAL[cell.data_type],
                        "README",
                        f"{_letter(index)}{number}",
                    )

    def _header(self, name, header):
        """Check a header row; return the sheet's width, or None when wrong.

        The known columns must come first, in order. Columns after them are
        allowed only in a file declaring a later minor version; the metadata
        sheet is read before the version is known, so its extra columns are
        checked once it is.
        """
        columns = SHEETS[name]
        names = [
            None if cell.data_type in NOT_LITERAL else cell.value for cell in header
        ]
        while names and _blank(names[-1]):
            names.pop()
        for index, column in enumerate(columns):
            if index >= len(names) or names[index] != column.name:
                self.add(
                    "header",
                    f"expected column {column.name}",
                    name,
                    f"{_letter(index)}1",
                )
                return None
        extras = [
            f"{_letter(index)}1"
            for index in range(len(columns), len(names))
            if not _blank(names[index])
        ]
        if name == "metadata":
            self.metadata_extras = extras
        elif extras and self.minor <= READER_VERSION[1]:
            for cell in extras:
                self._column_not_defined(name, cell)
            return None
        return len(names)

    def _column_not_defined(self, name, cell):
        """Record a header cell this reader's version does not define."""
        self.add(
            "column_not_defined",
            f"column not defined in version {SCHEMA_VERSION}",
            name,
            cell,
        )

    def _rows(self, name):
        """Yield (row number, cell values) of a data sheet's data rows.

        Data ends at the first fully blank row; a non-blank row after it is a
        problem, as is a formula or a value outside the header's columns.
        Only rows without such a problem are yielded, cut to the known
        columns (a later minor version's extra columns are ignored).
        """
        sheet = self.book[name]
        # Never trust the sheet's own <dimension>: read what is there.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        width = self._header(name, next(rows, ()))
        if width is None:
            self.incomplete.add(name)
            return
        known = len(SHEETS[name])
        ended = False
        for number, row in enumerate(rows, start=2):
            self._count()
            cells = [(cell.value, cell.data_type) for cell in row]
            if all(_blank(value) for value, _ in cells):
                ended = True
                continue
            if ended:
                self.add(
                    "row_after_blank",
                    "a row follows the first blank row, which ends the data",
                    name,
                    f"A{number}",
                )
                return
            usable = True
            for index, (value, data_type) in enumerate(cells):
                if data_type in NOT_LITERAL:
                    usable = False
                    self.add(
                        *NOT_LITERAL[data_type],
                        name,
                        f"{_letter(index)}{number}",
                    )
                elif index >= width and not _blank(value):
                    usable = False
                    self.add(
                        "cell_outside_columns",
                        "the cell is outside the sheet's columns",
                        name,
                        f"{_letter(index)}{number}",
                    )
            if usable:
                values = [value for value, _ in cells[:known]]
                yield number, values + [None] * (known - len(values))
            else:
                self.incomplete.add(name)

    def _typed(self, name, number, values):
        """The row's typed values by column name (``_row`` is its number).

        Returns None when a cell breaks its column's rule; each such cell is
        a problem.
        """
        row, usable = {"_row": number}, True
        for index, (column, value) in enumerate(zip(SHEETS[name], values, strict=True)):
            if _blank(value):
                if not column.optional:
                    usable = False
                    self.add(
                        "blank",
                        f"{column.name} must not be blank",
                        name,
                        f"{_letter(index)}{number}",
                    )
                row[column.name] = None
                continue
            try:
                row[column.name] = convert(column, value)
            except _Bad as bad:
                usable = False
                self.add(
                    bad.code,
                    f"{column.name} {bad.rule}",
                    name,
                    f"{_letter(index)}{number}",
                )
        return row if usable else None

    # Metadata

    def _metadata(self):
        """Read the metadata sheet: schema and version first, then every key.

        Returns the typed values of the defined keys and every key's (row,
        raw value). A wrong schema or version is fatal: nothing else can be
        read against a format the file does not declare.
        """
        rows = {}
        for number, (key, value, note) in self._rows("metadata"):
            if _blank(key) or type(key) is not str:
                self.add("blank", "key must be text", "metadata", f"A{number}")
                continue
            if key in rows:
                self.add(
                    "metadata_key_repeated",
                    "key repeats an earlier key",
                    "metadata",
                    f"A{number}",
                )
                continue
            rows[key] = (number, value)
            if not _blank(note):
                try:
                    convert(SHEETS["metadata"][2], note)
                except _Bad as bad:
                    self.add(bad.code, f"note {bad.rule}", "metadata", f"C{number}")
        if "metadata" in self.incomplete and not rows:
            # A wrong header: the keys cannot be read, and saying each
            # required key is missing would only repeat that.
            raise _Stop
        self._schema(rows)
        for cell in self.metadata_extras if self.minor <= READER_VERSION[1] else ():
            self._column_not_defined("metadata", cell)
        metadata = {}
        for key in METADATA_KEYS:
            if key.name not in rows:
                self.add(
                    "metadata_key_missing",
                    f"required key {key.name} is missing",
                    "metadata",
                )
                continue
            number, value = rows[key.name]
            metadata[key.name] = self._defined_value(key, number, value)
        for name, (number, value) in rows.items():
            if name not in METADATA:
                self._pattern_value(name, number, value)
        return metadata, rows

    def _schema(self, rows):
        """The schema name and a major version this reader reads, or stop."""
        for name in ("schema", "schema_version"):
            if name not in rows:
                self.add(
                    "metadata_key_missing",
                    f"required key {name} is missing",
                    "metadata",
                )
                raise _Stop
        number, value = rows["schema"]
        if value != SCHEMA:
            self.add("schema", f"schema must be {SCHEMA}", "metadata", f"B{number}")
            raise _Stop
        number, value = rows["schema_version"]
        version = VERSION.fullmatch(value) if type(value) is str else None
        if version is None or int(version[1]) != READER_VERSION[0]:
            self.add(
                "version",
                f"schema_version must be a version {READER_VERSION[0]}.x",
                "metadata",
                f"B{number}",
            )
            raise _Stop
        self.minor = int(version[2])

    def _defined_value(self, key, number, value):
        """The typed value of a defined key, or None (blank or a problem)."""
        if _blank(value):
            if not key.optional:
                self.add(
                    "blank", f"{key.name} must have a value", "metadata", f"B{number}"
                )
            return None
        try:
            typed = _metadata_value(key, value)
            if key.name == "timezone":
                _time_zone(typed)
            return typed
        except _Bad as bad:
            self.add(bad.code, f"{key.name} {bad.rule}", "metadata", f"B{number}")
            return None

    def _pattern_value(self, name, number, value):
        """Check an optional repeatable key; an unknown key is a problem in 1.0."""
        if _pattern_prefix(name) is not None:
            if name in KNOWN_COUNTS:
                key = MetadataKey(name, _I, "", minimum=0)
                self._defined_value(key, number, value)
            elif not _blank(value):
                try:
                    convert(SHEETS["metadata"][1], value)
                except _Bad as bad:
                    self.add(bad.code, f"value {bad.rule}", "metadata", f"B{number}")
            return
        if self.minor <= READER_VERSION[1]:
            self.add(
                "metadata_key_not_defined",
                f"key not defined in version {SCHEMA_VERSION}",
                "metadata",
                f"A{number}",
            )

    # Data sheets

    def _families(self):
        """The families rows by key, with each row's own rules checked."""
        families = {}
        for number, values in self._rows("families"):
            row = self._typed("families", number, values)
            if row is None or not self._family_rules(row):
                self.incomplete.add("families")
                continue
            if row["family_key"] in families:
                self.add(
                    "family_key",
                    "family_key repeats an earlier row's key",
                    "families",
                    _cell("families", "family_key", number),
                )
                continue
            families[row["family_key"]] = row
        return families

    def _family_rules(self, row):
        """The families sheet's conditional rules for one typed row."""
        number, ok = row["_row"], True

        def fail(column, rule, code="rule"):
            nonlocal ok
            ok = False
            self.add(code, rule, "families", _cell("families", column, number))

        key = FAMILY_KEY.fullmatch(row["family_key"])
        if key is None:
            fail(
                "family_key",
                "family_key must be duid:<n>, pds_fid:<n> or pds_env:<n>",
                "family_key",
            )
        else:
            prefix = key[1] or "pds_env"
            if FAMILY_ID_SYSTEMS[prefix] != row["family_id_system"]:
                fail("family_id_system", "family_id_system must match family_key")
            expected = None if key[1] is None else int(key[2])
            if row["family_id"] != expected:
                fail(
                    "family_id",
                    "family_id must be the number in family_key, blank for pds_env",
                )
        if row["last_submitted_at"] < row["first_submitted_at"]:
            fail("last_submitted_at", "last_submitted_at must not be before the first")
        status, amount = row["pledge_status"], row["annual_pledge"]
        allowed = {
            "pledged": amount is not None and amount > 0,
            "zero": amount == 0,
            "cannot_give": amount is None or amount == 0,
        }.get(status, amount is None)
        if not allowed:
            fail("annual_pledge", ANNUAL_PLEDGE_RULES[status])
        if (row["installment_amount"] is None) != (row["installment_source"] is None):
            fail(
                "installment_source",
                "installment_source must be blank exactly when installment_amount is",
            )
        return ok

    def _submissions(self):
        """The submissions rows, with each row's own rules checked."""
        submissions = []
        for number, values in self._rows("submissions"):
            row = self._typed("submissions", number, values)
            if row is None:
                continue
            if (row["family_key"] is None) == row["attributed"]:
                self.add(
                    "rule",
                    "family_key must be blank exactly when attributed is FALSE",
                    "submissions",
                    _cell("submissions", "family_key", number),
                )
                continue
            submissions.append(row)
        return submissions

    def _daily(self):
        """The daily rows, each typed, with the percentage within 0 to 1."""
        daily = []
        for number, values in self._rows("daily"):
            row = self._typed("daily", number, values)
            if row is None:
                continue
            share = row["cumulative_pct_of_active"]
            if not 0 <= share <= 1:
                self.add(
                    "range",
                    "cumulative_pct_of_active must be a fraction from 0 to 1",
                    "daily",
                    _cell("daily", "cumulative_pct_of_active", number),
                )
                continue
            daily.append(row)
        return daily

    def _ministry(self, families):
        """Count the Ministry answers, checking each cell and its Family key.

        The rows are not kept: nothing past the checks uses them, and they
        are most of the file.
        """
        count = 0
        for number, values in self._rows("ministry_responses"):
            count += 1
            row = self._typed("ministry_responses", number, values)
            if (
                row is not None
                and "families" not in self.incomplete
                and row["family_key"] not in families
            ):
                self.add(
                    "family_key",
                    "family_key must be a key of the families sheet",
                    "ministry_responses",
                    f"A{number}",
                )
        return count


def _pattern_prefix(name):
    """The prefix of a repeatable metadata key (such as ``note.``), or None."""
    for prefix, suffix, _ in METADATA_PATTERNS:
        if name.startswith(prefix) and suffix.fullmatch(name[len(prefix) :]):
            return prefix
    return None


ANNUAL_PLEDGE_RULES = {
    "pledged": "annual_pledge must be greater than zero when pledge_status is pledged",
    "zero": "annual_pledge must be 0 when pledge_status is zero",
    "cannot_give": "annual_pledge must be 0 or blank when pledge_status is cannot_give",
    "no_answer": "annual_pledge must be blank when pledge_status is no_answer",
    "no_pledge_form": (
        "annual_pledge must be blank when pledge_status is no_pledge_form"
    ),
}


def financial_answer(submission):
    """(pledge_status, annual_pledge, frequency, cannot_give) of a last answer.

    ``submission`` is a Family's last financial submission, or None when it
    sent no financial form (2020's ministry-only Families).
    """
    if submission is None:
        return "no_pledge_form", None, None, None
    amount = submission["annual_pledge"]
    if amount is not None and amount > 0:
        status = "pledged"
    elif submission["cannot_give"]:
        status = "cannot_give"
    elif amount is not None:
        status = "zero"
    else:
        status = "no_answer"
    return status, amount, submission["frequency"], submission["cannot_give"]


def submission_order(submission):
    """The last-submission ordering: submitted_at, then row_id for a tie."""
    return submission["submitted_at"], submission["row_id"]


class _Consistency:
    """The cross-sheet rules (check 4), over rows whose cells are all valid."""

    def __init__(self, metadata, rows, families, submissions, daily, add):
        self.metadata = metadata
        self.rows = rows
        self.families = families
        self.submissions = submissions
        self.daily = daily
        self.add = add
        self.attributed = sorted(
            (s for s in submissions if s["attributed"]), key=submission_order
        )

    def check(self, ministry_rows):
        """Every rule, each recording its own problems."""
        self._times()
        self._submission_order()
        self._families()
        self._daily()
        self._metadata(ministry_rows)

    def _times(self):
        """Each instant gives its row's local date, which gives its day."""
        zone = ZoneInfo(self.metadata["timezone"])
        start = self.metadata["campaign_start_date"]
        checks = (
            (
                "submissions",
                self.submissions,
                "submitted_at",
                "local_date",
                "campaign_day",
            ),
            (
                "families",
                self.families.values(),
                "first_submitted_at",
                "first_local_date",
                "first_campaign_day",
            ),
            (
                "families",
                self.families.values(),
                "last_submitted_at",
                "last_local_date",
                "last_campaign_day",
            ),
        )
        for sheet, rows, instant, local, day in checks:
            for row in rows:
                number = row["_row"]
                if row[instant].astimezone(zone).date() != row[local]:
                    self.add(
                        "date",
                        f"{local} must be the local date of {instant}",
                        sheet,
                        _cell(sheet, local, number),
                    )
                elif (row[local] - start).days + 1 != row[day]:
                    self.add(
                        "day",
                        f"{day} must be the campaign day of {local}",
                        sheet,
                        _cell(sheet, day, number),
                    )

    def _submission_order(self):
        """row_id runs 1, 2, 3 ... and the rows are oldest first."""
        previous = None
        for index, row in enumerate(self.submissions, start=1):
            number = row["_row"]
            if row["row_id"] != index:
                self.add(
                    "order",
                    "row_id must run 1, 2, 3 ... in row order",
                    "submissions",
                    _cell("submissions", "row_id", number),
                )
            if previous is not None and row["submitted_at"] < previous:
                self.add(
                    "order",
                    "submitted_at must not be before the row above (oldest first)",
                    "submissions",
                    _cell("submissions", "submitted_at", number),
                )
            previous = row["submitted_at"]

    def _families(self):
        """Each Family matches its submissions by the last-submission rule."""
        by_family = defaultdict(list)
        for row in self.attributed:
            if row["family_key"] in self.families:
                by_family[row["family_key"]].append(row)
            else:
                self.add(
                    "family_key",
                    "family_key must be a key of the families sheet",
                    "submissions",
                    _cell("submissions", "family_key", row["_row"]),
                )
        for row in self.submissions:
            if not row["attributed"] and row["is_last_for_family"]:
                self._not_last(row)
        for key, family in self.families.items():
            self._family(family, by_family.get(key, []))

    def _not_last(self, row):
        """Record an is_last_for_family that breaks its rule."""
        self.add(
            "family",
            "is_last_for_family must be TRUE exactly on a Family's last "
            "submission of each form",
            "submissions",
            _cell("submissions", "is_last_for_family", row["_row"]),
        )

    def _family(self, family, submissions):
        """One Family against its attributed submissions, in order."""
        number = family["_row"]

        def differs(column, rule):
            self.add("family", rule, "families", _cell("families", column, number))

        if family["submission_count"] != len(submissions):
            differs(
                "submission_count",
                "submission_count must equal the Family's attributed submissions",
            )
        if not submissions:
            return
        if family["first_submitted_at"] != submissions[0]["submitted_at"]:
            differs(
                "first_submitted_at",
                "first_submitted_at must be the Family's first submission",
            )
        if family["last_submitted_at"] != submissions[-1]["submitted_at"]:
            differs(
                "last_submitted_at",
                "last_submitted_at must be the Family's latest attributed submission",
            )
        financial = [s for s in submissions if s["form"] in FINANCIAL_FORMS]
        expected = financial_answer(financial[-1] if financial else None)
        for column, value in zip(
            ("pledge_status", "annual_pledge", "frequency", "cannot_give"),
            expected,
            strict=True,
        ):
            if family[column] != value:
                differs(
                    column, f"{column} must equal the Family's last financial answer"
                )
        last = {s["form"]: s for s in submissions}
        for submission in submissions:
            if submission["is_last_for_family"] != (
                last[submission["form"]] is submission
            ):
                self._not_last(submission)

    def _daily(self):
        """The daily series against families and submissions, day by day."""
        start = self.metadata["campaign_start_date"]
        last_date = self.metadata["campaign_last_submission_date"]
        active = self.metadata["active_families"]
        days = (last_date - start).days + 1
        if len(self.daily) != days:
            self.add(
                "daily",
                "daily must run from day 1 to the day of campaign_last_submission_date",
                "daily",
            )
        firsts = Counter(f["first_campaign_day"] for f in self.families.values())
        latest, position, running = {}, 0, 0
        for index, row in enumerate(self.daily, start=1):
            number = row["_row"]

            def wrong(column, rule, number=number):
                self.add("daily", rule, "daily", _cell("daily", column, number))

            if row["campaign_day"] != index:
                wrong("campaign_day", "campaign_day must run 1, 2, 3 ... with no gap")
                return
            if row["local_date"] != start + timedelta(days=index - 1):
                wrong("local_date", "local_date must be the date of its campaign_day")
            if row["first_submissions"] != firsts[index]:
                wrong(
                    "first_submissions",
                    "first_submissions must equal the Families whose first "
                    "submission fell on this day",
                )
            running += row["first_submissions"]
            if row["cumulative_families"] != running:
                wrong(
                    "cumulative_families",
                    "cumulative_families must be the running sum of first_submissions",
                )
            # Each Family's latest financial answer by the end of this day.
            while (
                position < len(self.attributed)
                and self.attributed[position]["campaign_day"] <= index
            ):
                submission = self.attributed[position]
                if submission["form"] in FINANCIAL_FORMS:
                    latest[submission["family_key"]] = submission[
                        "annual_pledge"
                    ] or Decimal(0)
                position += 1
            if row["cumulative_pledging_families"] != sum(
                v > 0 for v in latest.values()
            ):
                wrong(
                    "cumulative_pledging_families",
                    "cumulative_pledging_families must equal the value recomputed "
                    "from submissions",
                )
            if row["cumulative_annual_pledged"] != sum(latest.values(), Decimal(0)):
                wrong(
                    "cumulative_annual_pledged",
                    "cumulative_annual_pledged must equal the value recomputed "
                    "from submissions",
                )
            share = row["cumulative_pct_of_active"]
            if abs(share - row["cumulative_families"] / active) > PERCENT_TOLERANCE:
                wrong(
                    "cumulative_pct_of_active",
                    "cumulative_pct_of_active must be cumulative_families divided "
                    "by active_families",
                )
        if self.daily:
            self._totals(self.daily[-1])

    def _totals(self, last):
        """The last day equals the totals of the families sheet."""
        pledged = [
            f["annual_pledge"]
            for f in self.families.values()
            if f["pledge_status"] == "pledged"
        ]
        expected = {
            "cumulative_families": len(self.families),
            "cumulative_pledging_families": len(pledged),
            "cumulative_annual_pledged": sum(pledged, Decimal(0)),
        }
        # The last day's own cell is below MAX_MONEY, so a consistent file
        # never reaches this; it keeps the stored total's bound explicit.
        if expected["cumulative_annual_pledged"] >= STORED_MONEY_LIMIT:
            self.add("range", "the pledged total is out of range", "families")
        for column, value in expected.items():
            if last[column] != value:
                self.add(
                    "daily",
                    f"the last day's {column} must equal the families sheet's total",
                    "daily",
                    _cell("daily", column, last["_row"]),
                )

    def _metadata(self, ministry_rows):
        """Metadata against the rows: last date, as_of, active count, counts."""
        metadata = self.metadata

        def wrong(key, rule):
            number = self.rows[key][0]
            self.add("metadata", rule, "metadata", f"B{number}")

        if metadata["campaign_last_submission_date"] < metadata["campaign_start_date"]:
            wrong(
                "campaign_last_submission_date",
                "campaign_last_submission_date must not be before campaign_start_date",
            )
        if self.attributed:
            zone = ZoneInfo(metadata["timezone"])
            last = max(s["submitted_at"] for s in self.attributed)
            if (
                last.astimezone(zone).date()
                != metadata["campaign_last_submission_date"]
            ):
                wrong(
                    "campaign_last_submission_date",
                    "campaign_last_submission_date must be the local date of the "
                    "last attributed submission",
                )
            if metadata["as_of"] < last:
                wrong(
                    "as_of", "as_of must not be before the last attributed submission"
                )
        counts = {
            "count.families": len(self.families),
            "count.rows_read": len(self.submissions),
            "count.unattributed_rows": len(self.submissions) - len(self.attributed),
            "count.ministry_response_rows": ministry_rows,
        }
        for key, value in counts.items():
            if key not in self.rows:
                continue
            number, raw = self.rows[key]
            stated = _metadata_value(MetadataKey(key, _I, ""), raw)
            if stated != value:
                self.add(
                    "count",
                    f"{key} must equal the recomputed count",
                    "metadata",
                    f"B{number}",
                )


# --- What a validated archive yields -----------------------------------------

# The pledge distribution's band set (campaign-history spec, "Pledge
# distribution"). A band includes its lower edge and excludes its upper one;
# changing an edge is a new BAND_SET_VERSION.
BAND_SET_VERSION = 1
BAND_EDGES = tuple(Decimal(edge) for edge in (250, 500, 1000, 1500, 2500, 5000, 10000))
BAND_COUNT = len(BAND_EDGES) + 1
# Notes longer than this, or holding a dollar amount or five digits in a row
# (a pledge or an identifier), are flagged for the Administrator.
NOTE_LIMIT = 500
NOTE_FLAG = re.compile(r"\$\s*[0-9]|[0-9]{5,}")


def pledge_band(amount):
    """The band (1 to 8) of one annual pledge greater than zero."""
    return bisect.bisect_right(BAND_EDGES, amount) + 1


def pledge_distribution(pledges):
    """Per band, in order: the band, its number of Families and their total."""
    families, totals = [0] * BAND_COUNT, [Decimal(0)] * BAND_COUNT
    for amount in pledges:
        band = pledge_band(amount) - 1
        families[band] += 1
        totals[band] += amount
    return [
        {"band": index + 1, "families": families[index], "total": totals[index]}
        for index in range(BAND_COUNT)
    ]


def quartiles(pledges):
    """(first quartile, median, third quartile), Excel's QUARTILE.INC method.

    For the n pledges sorted as x1 ... xn and p = 1/4, 1/2, 3/4: h = (n - 1)p
    + 1 and the quartile is x[floor h] + (h - floor h)(x[floor h + 1] -
    x[floor h]), in exact decimals, rounded to the cent half up. None for no
    pledges.
    """
    values = sorted(pledges)
    if not values:
        return None
    result = []
    for p in (Decimal("0.25"), Decimal("0.5"), Decimal("0.75")):
        h = (len(values) - 1) * p + 1
        whole = int(h)
        low = values[whole - 1]
        high = values[whole] if whole < len(values) else low
        result.append((low + (h - whole) * (high - low)).quantize(CENT, ROUND_HALF_UP))
    return tuple(result)


def pledged_amounts(archive):
    """The annual pledges of the Families with pledge_status pledged."""
    return [
        f["annual_pledge"] for f in archive.families if f["pledge_status"] == "pledged"
    ]


def note_flagged(text):
    """Whether a note needs the Administrator's look before Apply."""
    return text is not None and (len(text) > NOTE_LIMIT or bool(NOTE_FLAG.search(text)))


def flagged_notes(archive):
    """The keys of the free-text values the Review card flags."""
    texts = dict(archive.notes)
    for key in ("active_families_provenance", "dedup_rule"):
        texts[key] = archive.metadata.get(key)
    return sorted(key for key, text in texts.items() if note_flagged(text))


# --- The writer ------------------------------------------------------------------


def utc_text(instant):
    """An aware datetime as ISO 8601 UTC text with a Z suffix."""
    return instant.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def _cell_value(value):
    """A Python value as a data cell holds it (money as a float, text times)."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return utc_text(value)
    if isinstance(value, date):
        return value.isoformat()
    return value


def _metadata_text(value):
    """A metadata value as the text the format stores."""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(_cell_value(value))


def readme_rows():
    """The README sheet's rows: the description, then the generated dictionary."""
    rows = [[line] for line in README_LINES]
    for name, columns in SHEETS.items():
        rows += [[], [f"Sheet '{name}'"]]
        rows += [[c.name, c.type_text, c.description] for c in columns]
    rows += [[], ["Metadata keys"]]
    rows += [[k.name, str(k.kind), k.description] for k in METADATA_KEYS]
    rows += [[prefix + "*", "", text] for prefix, _, text in METADATA_PATTERNS]
    return rows


def build_workbook(metadata, sheets):
    """An openpyxl Workbook in this format, which a test may adjust and save.

    ``metadata`` is a list of (key, value, note) rows; ``sheets`` maps a data
    sheet name to a list of row dicts keyed by column name (a missing column
    is blank). Money is written as a number, instants and dates as text.
    """
    import openpyxl

    book = openpyxl.Workbook()
    readme = book.active
    readme.title = "README"
    for row in readme_rows():
        readme.append(row)
    rows = dict(sheets)
    rows["metadata"] = [
        {"key": key, "value": _metadata_text(value), "note": note}
        for key, value, note in metadata
    ]
    for name, columns in SHEETS.items():
        sheet = book.create_sheet(name)
        names = [column.name for column in columns]
        sheet.append(names)
        for row in rows.get(name, ()):
            unknown = set(row) - set(names)
            if unknown:
                raise ValueError(f"{name}: unknown columns {sorted(unknown)}")
            sheet.append([_cell_value(row.get(column)) for column in names])
    return book


def save_workbook(book):
    """The bytes of a saved workbook."""
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def write_archive(metadata, sheets):
    """The bytes of one archive workbook (see ``build_workbook``)."""
    return save_workbook(build_workbook(metadata, sheets))
