"""The shared look of every Admin portal XLSX workbook.

Writers put their values in the cells exactly as before (``xlsx_cell``
keeps text literal, dates native and money numeric); these helpers only
style what is there:

- ``style_table`` turns a header row and the rows below it into a readable
  table: a bold white-on-teal header, Arial body text, a frozen header (and
  optionally first column), autofilter, light zebra banding, widths fitted to
  the content within limits, and a landscape print setup that repeats the
  header and numbers the pages.
- ``style_details`` gives label/value blocks (the "Report information" sheet,
  a Ministry's details above its table) one label style.

Colors are the Admin portal's (``ui-v1.css``); banding is a conditional
format, so no cell gains a formula or a changed value.
"""

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from parishkit.stewardship.web.design_tokens import COLORS, hex6


def _color(token):
    """A design-token color as the six hex digits a workbook stores."""
    return hex6(COLORS[token]).lstrip("#").upper()


FONT_NAME = "Arial"
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=_color("accent"))
HEADER_BORDER = Border(bottom=Side("medium", color=_color("accent-strong")))
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_FONT = Font(name=FONT_NAME, size=10, color=_color("ink"))
BODY_ALIGNMENT = Alignment(vertical="top")
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
BAND = _color("paper-soft")
BAND_FILL = PatternFill("solid", fgColor=BAND, bgColor=BAND)
LABEL_FONT = Font(name=FONT_NAME, size=10, bold=True, color=_color("accent-strong"))
LABEL_FILL = PatternFill("solid", fgColor=_color("accent-soft"))
RULE_BORDER = Border(bottom=Side("thin", color=_color("border")))

# Fitted column widths (in characters) stay within these limits; longer
# text wraps at the maximum. Only the first rows are measured, which is
# plenty to size a column and keeps large workbooks fast.
MIN_WIDTH = 8
MAX_WIDTH = 60
SAMPLE_ROWS = 500
# Points per wrapped line of 10-point Arial, for rows whose height is set.
LINE_HEIGHT = 13.0
# Excel's tallest row, in points; a taller height is invalid in the file.
MAX_ROW_HEIGHT = 409.0


def _length(value):
    """The widest line of a cell's displayed text, in characters."""
    if value is None:
        return 0
    if hasattr(value, "year"):
        return 19  # a date-time in Excel's built-in format
    return max(len(line) for line in str(value).split("\n"))


def style_table(
    sheet,
    *,
    header_row=1,
    widths=None,
    freeze_first_column=False,
    title=None,
):
    """Style one table whose headings are in ``header_row``; data below it.

    ``widths`` maps a heading to a fixed width for known long-text columns;
    every other column fits its heading and its first ``SAMPLE_ROWS`` values
    within ``MIN_WIDTH``..``MAX_WIDTH``. A column that hits the maximum wraps.
    ``title`` (the report name) goes in the printed page header.
    """
    widths = widths or {}
    last_row = max(sheet.max_row, header_row)
    last_column = sheet.max_column
    fitted = {}
    for column in range(1, last_column + 1):
        heading = sheet.cell(header_row, column)
        heading.font = HEADER_FONT
        heading.fill = HEADER_FILL
        heading.border = HEADER_BORDER
        heading.alignment = HEADER_ALIGNMENT
        width = widths.get(heading.value)
        if width is None:
            sample = (
                sheet.cell(row, column).value
                for row in range(
                    header_row + 1, min(last_row, header_row + SAMPLE_ROWS) + 1
                )
            )
            # Bold headings run wider; their words may still wrap.
            longest = max(
                [len(word) * 1.15 for word in str(heading.value or "").split()]
                + [_length(value) for value in sample]
            )
            width = min(max(longest + 2, MIN_WIDTH), MAX_WIDTH)
        fitted[column] = width
        sheet.column_dimensions[get_column_letter(column)].width = width
    # Tall enough for the heading that wraps to the most lines.
    lines = max(
        _line_count(sheet.cell(header_row, column).value, fitted[column], 1.15)
        for column in range(1, last_column + 1)
    )
    sheet.row_dimensions[header_row].height = LINE_HEIGHT * max(lines, 2) + 4
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=last_row):
        for cell in row:
            cell.font = BODY_FONT
            # Keep wrapping a writer asked for; wrap any column at the limit.
            wrap = cell.alignment.wrap_text or fitted[cell.column] >= MAX_WIDTH
            cell.alignment = WRAP_ALIGNMENT if wrap else BODY_ALIGNMENT
    end = get_column_letter(last_column)
    if last_row > header_row:
        # Zebra banding as a display rule: values and types stay untouched.
        sheet.conditional_formatting.add(
            f"A{header_row + 1}:{end}{last_row}",
            FormulaRule(formula=["MOD(ROW(),2)=0"], fill=BAND_FILL),
        )
    sheet.freeze_panes = sheet.cell(header_row + 1, 2 if freeze_first_column else 1)
    sheet.auto_filter.ref = f"A{header_row}:{end}{last_row}"
    _print_setup(sheet, title)
    sheet.print_title_rows = f"{header_row}:{header_row}"


def style_details(sheet, rows, *, widths=True, span=None):
    """Style label/value pairs in columns A and B of ``rows`` (1-based).

    ``widths`` sets the usual label and value column widths; a sheet whose
    columns a table below already sized passes False, and ``span`` (that
    table's last column) to merge each value across columns B..``span`` so
    it wraps at the table's width, not column B's. Excel does not fit a
    merged cell's height, so each such row's height is set from its text.
    """
    for row in rows:
        label, value = sheet.cell(row, 1), sheet.cell(row, 2)
        label.font = LABEL_FONT
        label.fill = LABEL_FILL
        label.alignment = WRAP_ALIGNMENT
        label.border = RULE_BORDER
        value.font = BODY_FONT
        value.alignment = WRAP_ALIGNMENT
        value.border = RULE_BORDER
    if widths:
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 90
    if span and span > 2:
        width = sum(
            sheet.column_dimensions[get_column_letter(column)].width or MIN_WIDTH
            for column in range(2, span + 1)
        )
        for row in rows:
            sheet.merge_cells(
                start_row=row, start_column=2, end_row=row, end_column=span
            )
            lines = _line_count(sheet.cell(row, 2).value, width)
            if lines > 1:
                # Very long text scrolls inside a row at Excel's maximum.
                sheet.row_dimensions[row].height = min(
                    LINE_HEIGHT * lines + 2, MAX_ROW_HEIGHT
                )


def _line_count(value, width, factor=1.0):
    """About how many lines a cell's text wraps to in a ``width`` column.

    Words wrap greedily by character count (``factor`` widens bold text);
    a word longer than the column breaks across lines.
    """
    if value is None:
        return 1
    usable = max(width - 1, 1)
    count = 0
    for paragraph in str(value).split("\n"):
        lines, used = 1, 0.0
        for word in paragraph.split():
            length = len(word) * factor
            if used and used + 1 + length > usable:
                lines, used = lines + 1, 0.0
            used += (1 if used else 0) + length
            while used > usable:
                lines, used = lines + 1, used - usable
        count += lines
    return max(count, 1)


def style_information(sheet, *, title=None):
    """The shared "Report information" sheet: labeled rows, printed portrait."""
    style_details(sheet, range(1, sheet.max_row + 1))
    _print_setup(sheet, title, landscape=False)


def _print_setup(sheet, title, *, landscape=True):
    """Fit one page wide, title the page header and number the pages."""
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.gridLines = False
    if title:
        # "&&" is a literal ampersand in a header; "&B" toggles bold.
        sheet.oddHeader.left.text = "&B" + str(title).replace("&", "&&")
    sheet.oddHeader.right.text = "&A"
    sheet.oddFooter.center.text = "Page &P of &N"
