"""Safe complete-text exports shared by information and Family-directory reports."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from itertools import chain
from unicodedata import category

from parishkit.stewardship.web import dates
from parishkit.stewardship.web.exports import csv_cell

from .money import MoneyAmount

# Excel's number format for exact USD: "-$50.00" for a negative amount, the
# same text the reports display, while the cell stays a summable number.
MONEY_FORMAT = '"$"#,##0.00'
# Excel keeps only 15 significant digits; a longer amount would round silently.
EXCEL_DIGITS = 15
# Internal references no download carries in any format (#966): row
# versions, reference UUIDs and generations mean nothing to a reader, and
# troubleshooting uses the audit log and pk-admin. The
# export guard test checks every export's columns and report information
# against this list; System logs is the recorded exception.
INTERNAL_FIELDS = frozenset(
    {
        "Row type",
        "Version",
        "Campaign reference",
        "Source reference",
        "Source generation",
        "Response reference",
        "Proposed Member reference",
        "Email revision",
        "Family version",
        "Campaign UUID",
        "Fact generation UUID",
        "Submission cutoff",
    }
)
FORMAT_NOTE = (
    "Unsupported characters use Unicode escapes (\\uXXXX or \\UXXXXXXXX); "
    "literal backslashes are doubled. CSV retains the original Unicode text."
)


def visible_text(value, *, supported=None):
    """Reversibly represent format/font exclusions; never drop a private character.

    XLSX permits XML 1.0 characters. PDF uses the exact bundled font's character
    map and escapes controls as well, avoiding missing glyphs and invisible
    terminal controls. Escaping backslashes makes the notation unambiguous.
    """
    result = []
    for character in value:
        code = ord(character)
        allowed = (
            code in {9, 10, 13}
            or 0x20 <= code <= 0xD7FF
            or 0xE000 <= code <= 0xFFFD
            or 0x10000 <= code <= 0x10FFFF
        )
        if supported is not None:
            allowed = code == 10 or (
                category(character) not in {"Cc", "Cf"} and code in supported
            )
        if character == "\\":
            result.append("\\\\")
        elif allowed:
            result.append(character)
        else:
            result.append(f"\\u{code:04x}" if code <= 0xFFFF else f"\\U{code:08x}")
    return "".join(result)


def plain(value):
    """Money as its report text for PDF; every other value unchanged.

    The XLSX writer keeps money numeric (see ``xlsx_cell``) and CSV writes
    its canonical amount (``csv_cell_value``); PDF says exactly what the page says.
    """
    return value.display if isinstance(value, MoneyAmount) else value


def csv_cell_value(value):
    """One CSV cell: known money as its canonical amount, else the guarded text.

    A known amount is the plain signed decimal (``1200.00``, ``-50.00``,
    ``0.00``): no dollar sign, thousands separator or formula-guard
    apostrophe, so a spreadsheet reads it as a number and a script parses
    it without stripping text (#388 L5). Such a value is only an optional
    minus sign, digits, a point and two digits, so it cannot be a formula
    and needs no neutralizing. Unavailable money stays the word, and every
    other value goes through ``csv_cell`` as before.
    """
    if isinstance(value, MoneyAmount) and value.available:
        return value.canonical
    return csv_cell(plain(value))


def excel_amount(value):
    """An exact Decimal Excel can hold, or None when it would lose a digit."""
    amount = Decimal(value)
    return amount if len(amount.as_tuple().digits) <= EXCEL_DIGITS else None


def xlsx_cell(sheet, row, column, value):
    """Write one cell: a native date/timestamp or amount, or literal text.

    Timestamps arrive already converted to the export's stated display time
    zone; Excel datetimes are naive, so the zone is dropped only here. Excel's
    built-in formats 14 and 22 follow each viewer's own regional settings.

    Known money is a number cell with a dollar format, so staff can sum it.
    The value is built exactly from whole cents as a Decimal, but a
    spreadsheet number is an IEEE double: openpyxl saves it with "%.16g", so
    $97.57 may be stored as 97.56999999999999. Within 15 significant digits
    that is still accurate to the cent when opened, displayed or summed.
    Unavailable money stays the word, never zero, and an amount beyond
    those 15 digits stays its exact text rather than rounding. A whole
    number (an ``int``, never a ``bool``) is a number cell, such as a count.
    None is a truly empty cell. Text is never a formula.
    """
    from openpyxl.styles.numbers import BUILTIN_FORMATS

    if value is None or type(value) is int:
        return sheet.cell(row, column, value)

    if isinstance(value, MoneyAmount):
        amount = excel_amount(value.canonical) if value.available else None
        if amount is not None:
            cell = sheet.cell(row, column, amount)
            cell.number_format = MONEY_FORMAT
            return cell
        value = value.display
    if isinstance(value, date):
        timestamp = isinstance(value, datetime)
        # Set the format before the value: openpyxl otherwise registers its own
        # custom ISO pattern for the value, which would stay in the file.
        cell = sheet.cell(row, column)
        cell.number_format = BUILTIN_FORMATS[
            dates.XLSX_DATETIME if timestamp else dates.XLSX_DATE
        ]
        cell.value = value.replace(tzinfo=None) if timestamp else value
        return cell
    cell = sheet.cell(row, column, visible_text(dates.display_text(value)))
    cell.data_type = "s"
    return cell


def information_csv(document, output):
    """One header row and one row per record: exactly the table.

    Every format carries the same columns (#966): the report information
    lives in the XLSX "Report information" sheet and the PDF's report card,
    never as extra CSV columns or rows. The one exception is a document's
    ``header_rows`` (a report with the Families counted choice, R59 on
    #2324): label/value rows, then a blank row, above the table.
    """
    wrapper = io.TextIOWrapper(output, encoding="utf-8", newline="", write_through=True)
    try:
        writer = csv.writer(wrapper, lineterminator="\r\n")
        header = tuple(getattr(document, "header_rows", ()))
        for label, value in header:
            writer.writerow((label, csv_cell_value(value)))
        if header:
            writer.writerow(())
        writer.writerow(document.headings)
        for row in document.rows:
            writer.writerow(tuple(csv_cell_value(value) for value in row))
        wrapper.flush()
    finally:
        wrapper.detach()


def information_xlsx(document, output):
    """Literal cells retain complete values; metadata survives empty reports too.

    A document about one subject may have ``title_rows`` (label, value pairs,
    the Family timeline's Family, code, campaign and time): they open the
    table's sheet, then a blank row, then the table, as a packet's Ministry
    details do (reports spec, "File design"; #1841). A document's
    ``header_rows`` (its CSV's header block, R59 on #2324) open it the same
    way.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    from .xlsx_design import describe, style_details, style_information, style_table

    titles = tuple(
        getattr(document, "title_rows", ()) or getattr(document, "header_rows", ())
    )
    # The table's heading row: right after the title rows and a blank row.
    head = len(titles) + 2 if titles else 1
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = document.sheet_name
        for row_index, (label, value) in enumerate(titles, 1):
            xlsx_cell(sheet, row_index, 1, label)
            xlsx_cell(sheet, row_index, 2, value)
        for row_index, values in enumerate(
            chain((document.headings,), document.rows), head
        ):
            for column, value in enumerate(values, 1):
                cell = xlsx_cell(sheet, row_index, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        # Long free text gets a wide column; a directory's Family column
        # stays in view while scrolling across its contact columns.
        style_table(
            sheet,
            header_row=head,
            widths={"Submitted text": 70, "Staff notes": 70},
            freeze_first_column=document.headings[0] == "Family",
            title=document.title,
        )
        if titles:
            # Styled after the table, which sized the columns the values span.
            style_details(
                sheet,
                range(1, len(titles) + 1),
                widths=False,
                span=len(document.headings),
            )
        metadata = book.create_sheet("Report information")
        for index, (key, value) in enumerate(
            chain(document.metadata, (("Text representation", FORMAT_NOTE),)), 1
        ):
            xlsx_cell(metadata, index, 1, key)
            xlsx_cell(metadata, index, 2, value)
        style_information(metadata, title=document.title)
        describe(book, document.title, created=document.requested_at)
        book.save(output)
    finally:
        book.close()


def information_records(document):
    """The report-information card, then one card per record (lazily).

    A field/value card avoids compressing many columns into unreadable
    widths. A card shows the same columns as the CSV and XLSX, leaving out
    only blank fields (``pdf_design.row_records``).
    """
    from .pdf_design import report_record, row_records

    return chain(
        (report_record(document.metadata, values=chain.from_iterable(document.rows)),),
        row_records(document.headings, document.rows),
    )


def information_pdf(document, output):
    """Paginate before drawing so every page states its position.

    Count without retaining wrapped lines, then stream one bounded page at a
    time; both passes lay out the same detached document identically. A
    document's ``frame_details`` are the lines under every page's title
    (Financial stewardship's Families counted, R59 on #2324).
    """
    from .pdf_design import PdfFrame, write_records

    return write_records(
        output,
        PdfFrame.for_document(document, details=getattr(document, "frame_details", ())),
        lambda: information_records(document),
        requested_at=document.requested_at,
    )


def render_information(document, output, *, format):
    """Only three explicitly compiled report formats are executable."""
    renderers = {
        "csv": information_csv,
        "xlsx": information_xlsx,
        "pdf": information_pdf,
    }
    try:
        renderer = renderers[format]
    except (KeyError, TypeError) as error:
        raise ValueError("Unsupported information export format.") from error
    return renderer(document, output)
