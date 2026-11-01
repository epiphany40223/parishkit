"""Typed, formula-free participation workbooks over the shared frozen document.

The workbook holds the Response dashboard's sections, a sheet each (#1557;
UX conventions, "Downloaded files"): Summary, when the export captured one,
then Daily participation, then the "Report information" sheet.
"""

from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from .information_rendering import MONEY_FORMAT, excel_amount, xlsx_cell
from .participation import participation_columns, participation_table
from .participation_summary import COLUMNS as SUMMARY_COLUMNS
from .participation_summary import DAILY, SUMMARY
from .xlsx_design import describe, style_details, style_information, style_table


def _cell(sheet, row, column, value):
    """Keep text literal even when a parish-controlled label starts with '='."""
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def header_block(sheet, document):
    """Open a table sheet with the file's header block; return the heading row.

    The report, Families counted and Source as of rows (R59 on #2324: a
    printed or forwarded sheet still says which Families it counts), then a
    blank row; the table's headings follow. ``finish_header`` styles the
    block once the table has sized the columns.
    """
    for index, (label, value) in enumerate(document.header_rows, 1):
        _cell(sheet, index, 1, label)
        xlsx_cell(sheet, index, 2, value)
    return len(document.header_rows) + 2


def finish_header(sheet, document, columns):
    """Style the header block as label/value rows spanning the table's width."""
    style_details(
        sheet, range(1, len(document.header_rows) + 1), widths=False, span=columns
    )


def summary_sheet(sheet, document):
    """The Summary sheet: the page's Summary tiles under the Summary columns.

    Counts are number cells, an amount a dollar-formatted number (or its
    exact text past Excel's precision), Data as of a native date-time in
    the requester's display time zone, and a share a percentage cell of the
    exact ratio, which Excel shows rounded as the page does. An unknown
    value reads "Unavailable", never zero.
    """
    sheet.title = SUMMARY
    head = header_block(sheet, document)
    for column, heading in enumerate(SUMMARY_COLUMNS, 1):
        _cell(sheet, head, column, heading)
    zone = ZoneInfo(document.browser_timezone)
    for index, row in enumerate(document.summary.rows(), head + 1):
        group, figure, value, ratio, note = row
        _cell(sheet, index, 1, group)
        _cell(sheet, index, 2, figure)
        if value is None:
            _cell(sheet, index, 3, "Unavailable")
        else:
            if isinstance(value, datetime):
                value = value.astimezone(zone)
            cell = xlsx_cell(sheet, index, 3, value)
            if type(value) is int:
                cell.number_format = "#,##0"
        if ratio is not None:
            sheet.cell(index, 4, ratio[0] / ratio[1]).number_format = "0%"
        _cell(sheet, index, 5, note or None)
    style_table(sheet, header_row=head, title=SUMMARY)
    finish_header(sheet, document, len(SUMMARY_COLUMNS))


def participation_xlsx(document, output):
    """Export every daily row, preserving unavailable versus zero and provenance.

    Excel only preserves 15 significant decimal digits. Very large aggregate
    pledges therefore use their exact decimal text instead of silently rounding.
    Ordinary amounts are numeric cells with a dollar number format.
    """
    rows = participation_table(document)
    book = Workbook()
    sheet = book.active
    if document.summary is not None:
        summary_sheet(sheet, document)
        sheet = book.create_sheet(DAILY)
    else:
        sheet.title = DAILY
    columns = participation_columns(document)
    keys = [key for key, _ in columns]
    head = header_block(sheet, document)
    for column, (_, heading) in enumerate(columns, 1):
        _cell(sheet, head, column, heading)
    for index, (day, row) in enumerate(zip(document.days, rows, strict=True), head + 1):
        for column, key in enumerate(keys, 1):
            value = row[key]
            if key == "date":
                value = day.local_date
            elif key == "source_as_of" and value is not None:
                # In the stated display zone; xlsx_cell writes a native cell.
                value = datetime.fromisoformat(value).astimezone(
                    ZoneInfo(document.browser_timezone)
                )
            elif key == "pledge_usd" and value is not None:
                # Test None explicitly: a zero Decimal is falsy but still a number.
                amount = excel_amount(value)
                value = value if amount is None else amount
            if isinstance(value, date):
                # Excel's built-in date (14) and date-time (22) formats.
                xlsx_cell(sheet, index, column, value)
                continue
            cell = _cell(
                sheet, index, column, "Unavailable" if value is None else value
            )
            if key == "pledge_usd" and isinstance(value, Decimal):
                cell.number_format = MONEY_FORMAT
            elif type(value) is int:
                cell.number_format = "#,##0"
    style_table(sheet, header_row=head, title=DAILY)
    finish_header(sheet, document, len(columns))
    metadata = book.create_sheet("Report information")
    for index, (label, value) in enumerate(
        (
            ("Parish", document.parish_name),
            ("Campaign", document.campaign_name),
            ("Population", document.scope_label),
            ("Campaign timezone", document.campaign_timezone),
            ("Display timezone", document.browser_timezone),
            ("Source and original request", document.as_of_label),
            ("Daily rows", len(rows)),
            ("Missing observations", "Unavailable is not zero."),
            (
                "Large amounts",
                "Amounts exceeding Excel's 15-digit precision are exact text.",
            ),
        ),
        1,
    ):
        _cell(metadata, index, 1, label)
        _cell(metadata, index, 2, value)
    style_information(metadata, title="Participation")
    # Titled as the chart PDF's page title is, naming the population (the
    # chart files' own Title property stays "Family participation").
    describe(
        book,
        f"Family participation — {document.scope_label}",
        created=document.requested_at,
    )
    book.save(output)
    book.close()
