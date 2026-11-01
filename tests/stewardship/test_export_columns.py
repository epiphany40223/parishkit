"""One column set in every format of every Admin download (#966).

The Administrator's decision B1: a download's CSV, XLSX and PDF carry the
same columns, and none of them carries internal references (row versions,
reference UUIDs, generations). This guard renders every registered export
in each of its formats from one document and compares the columns:

- CSV: its header row (a Ministry packet repeats it per Ministry);
- XLSX: the table sheet's header row;
- PDF: the table's headings, or, for a card report, the fields its cards
  show plus the fields a card puts in its title (name, DUID, status).

It also checks that no heading or report-information label is in the
shared internal-field list, and that every export kind the export worker
accepts has a case here, so a new export cannot skip the guard. Page
downloads (CSV, XLSX or PDF written straight from a page) come from the
``web.exports.page_download`` registry, cross-checked against the URL
patterns and against the modules that name a download file, so a new one
fails here until it has a case. A case without a PDF table must be on the
recorded list below, so a new export cannot quietly skip its PDF. System
logs is the recorded exception in the UX spec's Downloaded files section:
CSV and XLSX share plain headings, it has no PDF, and it keeps its
cross-link identifiers.
"""

import ast
import csv
import io
import re
from dataclasses import replace
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from django.urls import get_resolver
from openpyxl import load_workbook

import parishkit.stewardship
from parishkit.stewardship.audit import log_reads
from parishkit.stewardship.reports import (
    family_test_names,
    limitations,
    ministry_roster_views,
    pdf_design,
    pledge_review_views,
)
from parishkit.stewardship.reports.census_changes import (
    CensusQuery,
    census_csv,
    census_pdf,
    census_xlsx,
    select,
)
from parishkit.stewardship.reports.charts import render_participation
from parishkit.stewardship.reports.comparison_documents import render_comparison
from parishkit.stewardship.reports.directory_rendering import (
    address_blocks,
    render_directory,
)
from parishkit.stewardship.reports.export_models import ExportRequest
from parishkit.stewardship.reports.information_rendering import (
    INTERNAL_FIELDS,
    render_information,
)
from parishkit.stewardship.reports.ministry_packets import render_packet
from parishkit.stewardship.reports.participation import participation_csv
from parishkit.stewardship.reports.spreadsheets import participation_xlsx
from parishkit.stewardship.web.exports import PAGE_DOWNLOADS

from .test_census_changes import SUBMITTED
from .test_census_changes import rows as census_rows
from .test_comparison_exports import document as comparison
from .test_directory_rendering import document as directory
from .test_directory_rendering import item as family
from .test_family_test_names import request as names_request
from .test_financial_documents import document as financial
from .test_financial_documents import row as financial_row
from .test_information_rendering import document as information
from .test_limitations_report import result as limitations_result
from .test_ministry_exports import document as ministry
from .test_ministry_followup_exports import document as followup_requests
from .test_ministry_packets import item as member
from .test_ministry_packets import packet
from .test_participation_rendering import document as participation
from .test_participation_rendering import summary as participation_summary
from .test_pledge_review import pledge_row
from .test_roster import roster_row
from .test_timeline_exports import document as timeline

# The export worker's report kinds, each with the cases below that render
# it. "ministry" covers both lists and the multi-Ministry packet.
WORKER_CASES = {
    "participation": ("participation",),
    "additional_information": ("additional_information",),
    "family_directory": ("family_directory",),
    "postal_outreach": ("postal_outreach",),
    "ministry": ("ministry_join", "ministry_leave", "packet"),
    "financial": ("financial",),
    "family_test_names": ("family_test_names",),
    "family_timeline": ("family_timeline",),
    "campaign_comparison": ("campaign_comparison",),
}


def routed_page_downloads(patterns):
    """The ``page_download`` name of every routed view, through includes."""
    names = set()
    for pattern in patterns:
        if hasattr(pattern, "url_patterns"):
            names |= routed_page_downloads(pattern.url_patterns)
        elif name := getattr(pattern.callback, "page_download", None):
            names.add(name)
    return names


# Loading the URL patterns imports every routed view, which registers its
# page download before PAGE_CASES is read from the registry.
ROUTED_PAGE_DOWNLOADS = routed_page_downloads(get_resolver().url_patterns)
# The synchronous page downloads (``web.exports.page_download``) this guard
# does not render, each with its documented reason. None today: System logs
# has a case (its CSV and XLSX; JSON Lines has no columns).
PAGE_EXCEPTIONS = frozenset()
# The cases whose PDF is not a table of their columns, each a recorded
# reason in the UX spec's Downloaded files section: the postal mail merge's
# PDF is address labels (checked by value in its case) and System logs has
# no PDF. Every other case must have one; the participation PDF's tables
# follow its chart since #1557.
NO_PDF_TABLE = frozenset({"postal_outreach", "system_logs"})
# Every other registered page download is a case below of the same name, so
# a new page download fails the guard until it is checked or excepted.
PAGE_CASES = tuple(sorted(set(PAGE_DOWNLOADS) - PAGE_EXCEPTIONS))
# The postal mail merge's PDF is a sheet of address labels, not a table:
# each card holds the mailing columns, so it is checked by value instead.
ADDRESS_CARD_COLUMNS = (
    "ParishSoft DUID",
    "Family",
    "Addressee",
    "Address line 1",
    "Address line 2",
    "Address line 3",
    "City",
    "State",
    "ZIP",
    "Family code",
)


def worker_report_kinds():
    """Every ``report`` value the export request constraint accepts."""
    [constraint] = [
        item
        for item in ExportRequest._meta.constraints
        if item.name == "export_report_known"
    ]
    kinds = set()

    def walk(node):
        """Collect ``report=`` and ``report__in=`` values from the Q tree."""
        for child in node.children:
            if isinstance(child, tuple):
                key, value = child
                if key == "report":
                    kinds.add(value)
                elif key == "report__in":
                    kinds.update(value)
            else:
                walk(child)

    walk(constraint.condition)
    return kinds


def filled(document):
    """The document with every blank cell given a value, except change cells.

    A PDF card leaves blank fields out, so a guard row has none; the CSV
    and XLSX are rendered from the same filled document. Changed at and
    Changed by stay as they are: they are blank on an item row and set on
    its earlier-workflow rows, and that is how the PDF nests a revision
    under its item, so filling them would hide the real item/revision shape.
    """
    keep = {
        index
        for index, heading in enumerate(document.headings)
        if heading in pdf_design.CHANGE_FIELDS
    }
    return replace(
        document,
        rows=tuple(
            tuple(
                "Filled" if value in (None, "") and index not in keep else value
                for index, value in enumerate(row)
            )
            for row in document.rows
        ),
    )


def render(renderer, document, format):
    """The bytes ``renderer(document, output, format=...)`` writes."""
    output = io.BytesIO()
    renderer(document, output, format=format)
    return output.getvalue()


def csv_headers(body, first=None):
    """Every header row of a CSV: its first row, plus any later row equal
    to it. A packet's report details come first and it repeats the headings
    for each Ministry, so its header rows are those starting with ``first``.
    """
    rows = list(csv.reader(io.StringIO(body.decode())))
    # A report with the Families counted choice opens with its header block
    # (R59 on #2324): label/value rows, then a blank row, then the table.
    if rows and rows[0][:1] == ["Report"] and [] in rows:
        rows = rows[rows.index([]) + 1 :]
    if first is not None:
        return [row for row in rows if row[:1] == [first]]
    return [row for row in rows if row == rows[0]] if rows else []


def xlsx_headers(body):
    """Each table sheet's header row: its first row with every cell filled.

    A packet's Ministry details sit above its table in two columns, so they
    are skipped; the "Report information" sheet is not a table.
    """
    book = load_workbook(io.BytesIO(body))
    try:
        headers = []
        for sheet in book.worksheets:
            if sheet.title == "Report information":
                continue
            headers.append(
                next(
                    list(row)
                    for row in sheet.iter_rows(values_only=True)
                    if all(value not in (None, "") for value in row)
                )
            )
        labels = (
            [row[0] for row in book["Report information"].iter_rows(values_only=True)]
            if "Report information" in book.sheetnames
            else []
        )
        return headers, labels
    finally:
        book.close()


def pdf_capture(monkeypatch, draw):
    """Run ``draw()`` and return what its PDF laid out: table headings, or
    the record cards (``write_records``)."""
    seen = {"tables": []}
    pages, records = pdf_design.table_pages, pdf_design.write_records

    def table_pages(table, rows, frame):
        """Keep the table's headings (each table's, in order), then paginate."""
        seen["table"] = list(table.headings)
        seen["tables"].append(list(table.headings))
        return pages(table, rows, frame)

    def write_records(output, frame, factory, **options):
        """Keep the cards, then write them as usual."""
        seen["cards"] = list(factory())
        return records(output, frame, factory, **options)

    monkeypatch.setattr(pdf_design, "table_pages", table_pages)
    monkeypatch.setattr(pdf_design, "write_records", write_records)
    body = draw()
    assert body.startswith(b"%PDF")
    return seen


def card_columns(cards, headings):
    """The columns a card PDF shows, in heading order, and any other label.

    A card's title holds its name, DUID and status fields; a nested
    workflow revision's heading holds its Changed at and Changed by. The
    report-information and Ministry-details cards are not records.
    """
    shown = {field for field in pdf_design._identity(dict.fromkeys(headings)) if field}
    for card in cards:
        if card.style == "info":
            continue
        shown.update(label for label, _ in card.fields)
        for _, fields in card.children:
            shown.update(label for label, _ in fields)
            shown.update(pdf_design.CHANGE_FIELDS)
    return [heading for heading in headings if heading in shown], shown - set(headings)


def shared_case(monkeypatch, document, renderer=render_information):
    """A report rendered through one shared three-format renderer."""
    document = filled(document)
    headers = csv_headers(render(renderer, document, "csv"))
    sheets, labels = xlsx_headers(render(renderer, document, "xlsx"))
    seen = pdf_capture(monkeypatch, lambda: render(renderer, document, "pdf"))
    headings = list(document.headings)
    if "table" in seen:
        pdf = seen["table"]
    else:
        pdf, extra = card_columns(seen["cards"], headings)
        assert not extra, extra
    labels = [*labels, *(label for label, _ in document.metadata)]
    return headings, headers, sheets, pdf, labels


def case_participation(monkeypatch):
    """Participation: the Response dashboard's two sections in every format.

    The XLSX has a sheet per section and the PDF a table per section (with
    the chart between them); the CSV is one table, a Section column and
    then each section's columns in order (#1557; UX conventions,
    "Downloaded files"), so each section is checked against its own slice
    of the CSV header.
    """
    from parishkit.stewardship.reports.participation_summary import (
        COLUMNS as SUMMARY_COLUMNS,
    )

    value = replace(participation(), summary=participation_summary())
    output = io.BytesIO()
    participation_csv(value, output)
    [header] = csv_headers(output.getvalue())
    output = io.BytesIO()
    participation_xlsx(value, output)
    sheets, labels = xlsx_headers(output.getvalue())
    seen = pdf_capture(monkeypatch, lambda: render(render_participation, value, "pdf"))
    summary, daily = sheets
    assert header == ["Section", *summary, *daily]
    width = 1 + len(SUMMARY_COLUMNS)
    return [
        (summary, [header[1:width]], [summary], seen["tables"][0], labels),
        (daily, [header[width:]], [daily], seen["tables"][1], labels),
    ]


def case_additional_information(monkeypatch):
    """Additional information, with workflow history."""
    return shared_case(monkeypatch, information(count=2))


def case_family_directory(monkeypatch):
    """The Family-code directory (a table PDF)."""
    return shared_case(monkeypatch, directory([family()]), render_directory)


def case_postal_outreach(monkeypatch):
    """The postal mail merge: its PDF is address cards, checked by value."""
    document = directory([family()], postal=True)
    headers = csv_headers(render(render_directory, document, "csv"))
    sheets, labels = xlsx_headers(render(render_directory, document, "xlsx"))
    assert render(render_directory, document, "pdf").startswith(b"%PDF")
    values = dict(zip(document.headings, document.rows[0], strict=True))
    [block] = address_blocks(document)
    text = "\n".join(line for line, _ in block)
    for heading in ADDRESS_CARD_COLUMNS:
        assert str(values[heading]) in text, heading
    labels = [*labels, *(label for label, _ in document.metadata)]
    return list(document.headings), headers, sheets, None, labels


def case_ministry_join(monkeypatch):
    """Prospective Ministry joiners."""
    return shared_case(monkeypatch, ministry(action="join", count=1))


def case_ministry_leave(monkeypatch):
    """Requested Ministry leavers."""
    return shared_case(monkeypatch, ministry(action="leave", count=1))


NEW_MEMBER = dict(
    member_duid=None,
    proposed_id="00000000-0000-4000-8000-000000000099",
    entity_kind="proposed_member",
)


def xlsx_values(body):
    """Every cell value of a workbook, as one string, for absence checks."""
    book = load_workbook(io.BytesIO(body))
    try:
        return str([list(sheet.iter_rows(values_only=True)) for sheet in book])
    finally:
        book.close()


def case_packet(monkeypatch):
    """The multi-Ministry follow-up packet: one table per Ministry."""
    sections = [
        dict(duid=9, name="Choir", chairs=["Pat Lee"], rows=[member()]),
        dict(duid=10, name="Ushers", chairs=[], rows=[member(member_duid=7)]),
        # A Member added on the form: "New Member", never its proposed id.
        dict(duid=11, name="Greeters", chairs=[], rows=[member(**NEW_MEMBER)]),
    ]
    document = packet(sections)
    body = render(render_packet, document, "csv")
    headers = csv_headers(body, first=document.headings[0])
    assert len(headers) == len(sections)
    xlsx = render(render_packet, document, "xlsx")
    sheets, labels = xlsx_headers(xlsx)
    seen = pdf_capture(monkeypatch, lambda: render(render_packet, document, "pdf"))
    proposed = NEW_MEMBER["proposed_id"]
    assert ",New Member," in body.decode() and proposed not in body.decode()
    assert proposed not in xlsx_values(xlsx)
    assert "=Example Member (New Member)" in [card.title for card in seen["cards"]]
    assert proposed not in str(seen["cards"])
    # A packet prints blank fields as write-in lines, so every column shows.
    pdf, extra = card_columns(seen["cards"], list(document.headings))
    assert not extra, extra
    labels = [*labels, *(label for label, _ in document.metadata)]
    return list(document.headings), headers, sheets, pdf, labels


def case_financial(monkeypatch):
    """Financial stewardship detail."""
    return shared_case(monkeypatch, financial([financial_row()]))


def case_family_test_names(monkeypatch):
    """The chosen-Family test names (``pk-admin test families-preview``)."""
    monkeypatch.setattr(family_test_names, "campaign_name", lambda request: "C")
    document = family_test_names.names_document(names_request([(12, "A")]))
    return shared_case(monkeypatch, document)


def case_family_timeline(monkeypatch):
    """One Family's timeline (``export family-timeline``, ADM-11 PR 8g).

    Its PDF is an infographic (#926), not a table or cards: its columns
    are the column headings it draws over the timeline, in drawing order.
    """
    from parishkit.stewardship.reports.timeline_pdf import render_timeline

    document = filled(timeline())
    headers = csv_headers(render(render_timeline, document, "csv"))
    sheets, labels = xlsx_headers(render(render_timeline, document, "xlsx"))
    drawn, text = [], pdf_design.Canvas.text

    def record(canvas, x, y, value, *args, **kwargs):
        """Keep each drawn string, then draw it."""
        drawn.append(value)
        return text(canvas, x, y, value, *args, **kwargs)

    monkeypatch.setattr(pdf_design.Canvas, "text", record)
    assert render(render_timeline, document, "pdf").startswith(b"%PDF")
    headings = list(document.headings)
    pdf = [value for value in dict.fromkeys(drawn) if value in headings]
    labels = [*labels, *(label for label, _ in document.metadata)]
    return headings, headers, sheets, pdf, labels


def case_campaign_comparison(monkeypatch):
    """The Campaign comparison's Every number table (#1018, slice 6).

    Its PDF is the one-page summary and then this table, so the guard
    reads the table's headings.
    """
    return shared_case(monkeypatch, comparison(), render_comparison)


def case_census_changes(monkeypatch):
    """Pending census changes."""
    zone = ZoneInfo("America/Chicago")
    result = select(census_rows(), CensusQuery(), administrator=False)
    headers = csv_headers(census_csv(result, zone))
    sheets, labels = xlsx_headers(
        census_xlsx(result, zone, parish="P", campaign="C", as_of=SUBMITTED)
    )
    seen = pdf_capture(
        monkeypatch,
        lambda: census_pdf(result, zone, parish="P", campaign="C", as_of=SUBMITTED),
    )
    return headers[0], headers, sheets, seen["table"], labels


def case_ministry_roster(monkeypatch):
    """The roster list's file (#528)."""
    zone = ZoneInfo("America/Chicago")
    rows = [roster_row()]
    headers = csv_headers(ministry_roster_views.roster_csv(rows, zone))
    sheets, labels = xlsx_headers(ministry_roster_views.roster_xlsx(rows, zone))
    seen = pdf_capture(
        monkeypatch,
        lambda: ministry_roster_views.roster_pdf(
            rows, zone, parish="P", campaign="C", as_of=SUBMITTED
        ),
    )
    return headers[0], headers, sheets, seen["table"], labels


def case_pledges_to_enter(monkeypatch):
    """Pledges (#1686)."""
    zone = ZoneInfo("America/Chicago")
    rows = [pledge_row()]
    headers = csv_headers(pledge_review_views.pledges_csv(rows, zone))
    sheets, labels = xlsx_headers(pledge_review_views.pledges_xlsx(rows, zone))
    seen = pdf_capture(
        monkeypatch,
        lambda: pledge_review_views.pledges_pdf(
            rows, zone, parish="P", campaign="C", as_of=SUBMITTED
        ),
    )
    return headers[0], headers, sheets, seen["table"], labels


def case_ministry_followup(monkeypatch):
    """Export these requests, the follow-up queue's file (#1157)."""
    return shared_case(monkeypatch, followup_requests())


def case_limitations(monkeypatch):
    """The Limitations report: a Members table, then a Families table.

    Its CSV holds both tables, one after the other, its XLSX one sheet
    each and its PDF one table each; each table's columns are checked on
    their own.
    """
    zone = ZoneInfo("America/Chicago")
    result = limitations.shape_result(limitations_result())
    body = limitations.limitations_csv(result, zone).decode()
    rows = list(csv.reader(io.StringIO(body)))
    blank = rows.index([])
    sheets, labels = xlsx_headers(
        limitations.limitations_xlsx(
            result, zone, parish="P", campaign="C", as_of=SUBMITTED
        )
    )
    members, families = pdf_capture(
        monkeypatch,
        lambda: limitations.limitations_pdf(
            result, zone, parish="P", campaign="C", as_of=SUBMITTED
        ),
    )["tables"]
    return [
        (limitations.MEMBER_HEADINGS, [rows[0]], sheets[:1], members, labels),
        (limitations.FAMILY_HEADINGS, [rows[blank + 1]], sheets[1:], families, labels),
    ]


def case_system_logs(monkeypatch):
    """System logs: CSV and XLSX share plain headings; it has no PDF.

    Its cross-link identifier columns are the recorded exception to leaving
    internal identifiers out (UX spec, Downloaded files), and are named in
    plain words, so the internal-field check below still holds for it.
    """
    rows = [
        {
            "created_at": SUBMITTED,
            "source": "audit",
            "level": None,
            "event": "system_logs_exported",
            "actor": "admin@example.org",
            "actor_id": "a",
            "correlation_id": "c",
            "campaign_id": None,
            "subject_id": None,
            "details": {"count": "1"},
        }
    ]
    zone = ZoneInfo("America/Chicago")
    headers = csv_headers(log_reads.export_body(rows, "csv", zone, as_of=SUBMITTED))
    sheets, labels = xlsx_headers(
        log_reads.export_body(rows, "xlsx", zone, as_of=SUBMITTED)
    )
    return list(log_reads.EXPORT_HEADINGS.values()), headers, sheets, None, labels


# A page download without a case maps to None, which the page-download
# test reports by name instead of failing collection here.
CASES = {
    name: globals().get(f"case_{name}")
    for name in (
        *(name for names in WORKER_CASES.values() for name in names),
        *PAGE_CASES,
    )
}


# Export kinds that are not report downloads. The campaign archive workbook
# is the archive format (campaign-history spec), whose sheets and columns
# the format table fixes and test_archive_export covers; the reports spec's
# File design excludes it.
NOT_REPORT_DOWNLOADS = frozenset({"campaign_archive"})


def test_every_export_kind_has_a_guard_case():
    """A report kind the export worker accepts cannot skip this guard."""
    assert worker_report_kinds() == set(WORKER_CASES) | NOT_REPORT_DOWNLOADS


def test_every_page_download_has_a_guard_case():
    """A routed page download is registered and either checked or excepted.

    The exceptions must still be registered and routed, so a stale one is
    noticed when its download goes away.
    """
    assert set(PAGE_DOWNLOADS) == ROUTED_PAGE_DOWNLOADS
    assert set(PAGE_DOWNLOADS) >= PAGE_EXCEPTIONS
    assert not [name for name in PAGE_CASES if CASES[name] is None]


def calls(tree, name):
    """The calls in ``tree`` to a function named ``name``."""
    return [
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and getattr(node.func, "id", getattr(node.func, "attr", None)) == name
    ]


def test_every_download_writer_is_classified():
    """Code that names a download file is a registered page download or known.

    ``download_headers`` names a download file; ``campaign_response`` does so
    when given a ``filename``. A module that does either must register its
    page download, or be one of the known non-page writers below: the shared
    helpers themselves, the export worker's artifact download (checked by
    kind above) and the daily digest's PNG image (not a table).
    """
    known = {
        "web/exports.py",
        "web/responses.py",
        "reports/export_views.py",
        "reports/digest_views.py",
    }
    root = Path(parishkit.stewardship.__file__).parent
    unclassified = []
    for path in sorted(root.rglob("*.py")):
        name = path.relative_to(root).as_posix()
        tree = ast.parse(path.read_text(encoding="utf-8"))
        writes = calls(tree, "download_headers") or [
            call
            for call in calls(tree, "campaign_response")
            if any(keyword.arg == "filename" for keyword in call.keywords)
        ]
        if writes and name not in known and not calls(tree, "page_download"):
            unclassified.append(name)
    assert not unclassified


@pytest.mark.parametrize("name", sorted(name for name in CASES if CASES[name]))
def test_every_format_carries_the_same_columns(monkeypatch, name):
    """CSV header == XLSX header == PDF columns, with no internal field.

    ``None`` for the PDF means it is not a table of these columns: the
    postal labels (checked in their case) and System logs, which has no
    PDF. Only a recorded case may have none.
    """
    results = CASES[name](monkeypatch)
    for headings, headers, sheets, pdf, labels in (
        results if isinstance(results, list) else [results]
    ):
        headings = list(headings)
        assert headers and all(header == headings for header in headers)
        assert sheets and all(sheet == headings for sheet in sheets)
        if pdf is not None:
            assert pdf == headings
        else:
            assert name in NO_PDF_TABLE, f"{name} has no PDF table"
        assert not INTERNAL_FIELDS & {*headings, *labels}
        # Plain names, as on the page: no snake_case program keys.
        assert not [
            heading for heading in headings if re.fullmatch(r"[a-z_]+", heading)
        ]
