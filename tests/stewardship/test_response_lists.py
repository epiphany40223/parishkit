"""The response lists' choices, rows, sorting, files and page, without a database.

The funnel rows are the response-metrics test rows (#477); the views'
admission, read guard, purge gate and audit are covered against PostgreSQL
in ``database/test_response_lists_postgresql.py``.
"""

import csv
import io
import re
from dataclasses import replace
from datetime import timedelta
from types import SimpleNamespace
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

import pytest
from django.http import QueryDict
from django.template.loader import render_to_string
from django.urls import reverse

from parishkit.stewardship.reports import pdf_design, response_lists
from parishkit.stewardship.reports.response_list_views import page_context
from parishkit.stewardship.reports.response_lists import (
    EVERYONE,
    LISTS,
    ListQuery,
    candidates,
    list_counts,
    list_csv,
    list_details,
    list_file,
    listed,
)
from parishkit.stewardship.reports.response_metrics import (
    ResponseScope,
    response_families,
    stage_counts,
)
from parishkit.stewardship.source.snapshot_names import FamilyFacts
from parishkit.stewardship.web.dates import using
from parishkit.stewardship.web.tables import paginate

from .test_response_metrics import ROWS, START

CAMPAIGN = SimpleNamespace(
    pk=UUID(int=477), active_configuration=SimpleNamespace(name="Sample campaign")
)
BASE = reverse("admin:response_dashboard")
NEW_YORK = ZoneInfo("America/New_York")
# Facts for every test row: Family 3's mailing name is blank and Family 4's
# envelope number is 0; Family 6 is missing from the snapshot.
FACTS = {
    1: FamilyFacts("Adams, Ann", 101, "Ann Adams"),
    2: FamilyFacts("=Baker, Bob", 102, "Bob Baker"),
    3: FamilyFacts("Cole, Cy", 103, ""),
    4: FamilyFacts("Diaz, Dee", 0, "Dee Diaz"),
    5: FamilyFacts("Evans, Eve", None, "Eve Evans"),
}


def rows_of(key, show=EVERYONE, active=None, search=""):
    """A list's rows over the test rows, with every Family active by default."""
    spec = LISTS[key]
    active = {row.family_id for row in ROWS} if active is None else active
    chosen = candidates(spec, ROWS, frozenset(active))
    return listed(spec, chosen, FACTS, ListQuery(show=show, search=search))


def duids(rows):
    """The DUIDs of listed rows, in order."""
    return [row.duid for row in rows]


def test_funnel_lists_match_the_dashboard_figures():
    """Each funnel list holds exactly the Families its dashboard count counts."""
    stages = {stage.key: stage.count for stage in stage_counts(ROWS)}
    counts = list_counts(ROWS)
    assert counts == {
        "submitted": stages["submitted"],
        "started": stages["form_opened"] - stages["submitted"],
        "not-opened": 1,
        "more-than-once": sum(1 for row in ROWS if row.submissions > 1),
    }
    assert duids(rows_of("submitted")) == [1, 2, 5]
    assert duids(rows_of("started")) == [4]
    # Family 6 was never invited, so it is not "invited but never opened".
    assert duids(rows_of("not-opened")) == [3]
    assert duids(rows_of("more-than-once")) == [1]
    assert "data-quality" not in counts


def test_filters_split_each_list():
    """Every filter choice keeps a subset; together the choices cover the list."""
    assert duids(rows_of("started", "progressed")) == []
    assert duids(rows_of("started", "opened")) == [4]
    assert duids(rows_of("not-opened", "followed")) == [3]
    assert duids(rows_of("not-opened", "unfollowed")) == []
    with pytest.raises(ValueError):
        LISTS["submitted"].choice("progressed")


def test_submitted_has_no_filter_and_refuses_old_show_values():
    """Families that submitted lists everyone (#860); its old Show values are gone."""
    spec = LISTS["submitted"]
    assert [choice.value for choice in spec.choices] == [EVERYONE]
    query, _ = ListQuery.parse(spec, QueryDict(f"show={EVERYONE}"))
    assert query == ListQuery()
    assert duids(rows_of("submitted", EVERYONE)) == [1, 2, 5]
    # No backward compatibility for old addresses: a retired value is just
    # an unknown one.
    for old in ("invited", "uninvited", "followed"):
        with pytest.raises(ValueError):
            ListQuery.parse(spec, QueryDict(f"show={old}"))


def test_search_matches_part_of_a_name_or_an_exact_duid_or_envelope():
    """The directory's name rule, any case; all digits: a DUID or an envelope."""
    assert duids(rows_of("submitted", search="adams")) == [1]
    assert duids(rows_of("submitted", search="BOB")) == [2]
    # Part of the name, anywhere in it (surname or head).
    assert duids(rows_of("submitted", search="ve")) == [5]
    assert duids(rows_of("submitted", search="a")) == [1, 2, 5]
    # The envelope number must match exactly; a part of one does not.
    assert duids(rows_of("submitted", search="102")) == [2]
    assert duids(rows_of("submitted", search="10")) == []
    assert duids(rows_of("data-quality", search="0")) == [4]
    # So must the Family DUID, on every list (#884); these names hold no
    # digits, so only the number can match.
    assert duids(rows_of("submitted", search="2")) == [2]
    assert duids(rows_of("submitted", search="5")) == [5]
    assert duids(rows_of("data-quality", search="4")) == [4]
    assert duids(rows_of("started", "opened", search="4")) == [4]
    assert duids(rows_of("submitted", search="99")) == []
    # Nothing found, and a search combines with the list's filter.
    assert duids(rows_of("submitted", search="zzz")) == []
    assert duids(rows_of("started", "opened", search="diaz")) == [4]
    assert duids(rows_of("started", "progressed", search="diaz")) == []
    # A Family missing from the snapshot has no name to match.
    missing = ListQuery(search="family")
    assert not missing.matches(
        listed(LISTS["not-opened"], ROWS[2:3], {}, ListQuery())[0]
    )


def test_data_quality_lists_active_families_with_a_problem():
    """Blank mailing name or envelope 0, among the campaign's active Families."""
    assert duids(rows_of("data-quality")) == [3, 4]
    assert duids(rows_of("data-quality", "mailing-name")) == [3]
    assert duids(rows_of("data-quality", "envelope")) == [4]
    assert LISTS["data-quality"].choices[0].label == "Everything to check"
    # An inactive Family is not listed; a Family missing from the snapshot
    # (6) has nothing to check.
    active = {row.family_id for row in ROWS if row.family_duid != 4}
    assert duids(rows_of("data-quality", active=active)) == [3]
    problems = {row.duid: row.problems for row in rows_of("data-quality")}
    assert problems == {3: ("Blank mailing name",), 4: ("Envelope number 0",)}


def test_every_column_sorts_with_missing_values_last():
    """Times newest first on the first click; text ignores case; blanks last."""
    spec = LISTS["submitted"]
    rows = rows_of("submitted")
    assert spec.sorting.parse({}) == "submitted"
    chronological = paginate(rows, {}, sorting=spec.sorting)
    assert duids(chronological.rows) == [1, 2, 5]
    newest = paginate(rows, {"sort": "-submitted"}, sorting=spec.sorting)
    assert duids(newest.rows) == [5, 2, 1]
    # Family 5 has no envelope number: last in both directions.
    for token in ("envelope", "-envelope"):
        assert (
            duids(paginate(rows, {"sort": token}, sorting=spec.sorting).rows)[-1] == 5
        )
    by_name = paginate(rows, {"sort": "family"}, sorting=spec.sorting)
    assert duids(by_name.rows) == [2, 1, 5]  # "=Baker" sorts before letters
    # A count sorts largest first on its first click.
    assert spec.sorting.tokens["-submissions"] == ("submissions", True)
    with pytest.raises(ValueError):
        paginate(rows, {"sort": "name"}, sorting=spec.sorting)
    # A Family missing from the snapshot has no name and sorts last by name.
    missing = listed(LISTS["not-opened"], ROWS[2:3], {}, ListQuery())
    assert missing[0].name is None and missing[0].envelope is None


def test_query_accepts_only_closed_choices():
    """Mode, show and the table's own parameters; nothing identifying."""
    spec = LISTS["not-opened"]
    query, values = ListQuery.parse(spec, QueryDict("show=unfollowed&sort=-family"))
    assert query == ListQuery("production", "unfollowed")
    assert values == {"sort": "-family"}
    for invalid in (
        "mode=live",
        "show=envelope",
        "search=smith",
        "show=followed&show=unfollowed",
    ):
        with pytest.raises(ValueError):
            ListQuery.parse(spec, QueryDict(invalid))
    assert ListQuery().url("submitted") == BASE + "submitted/"
    assert ListQuery("testing", "followed").url("not-opened") == (
        BASE + "not-opened/?mode=testing&show=followed"
    )
    assert ListQuery().url("started", sort="", size="25") == (BASE + "started/?size=25")


def test_search_comes_only_from_a_post_body():
    """A search is private: parsed from a POST body, trimmed, never in a URL."""
    spec = LISTS["submitted"]
    body = QueryDict("search=+Adams+&mode=testing&sort=family&page=2")
    query, values = ListQuery.parse(spec, body, private=True)
    assert query == ListQuery("testing", EVERYONE, "Adams")
    assert values == {"sort": "family", "page": "2"}
    assert query.posted() == [("mode", "testing"), ("search", "Adams")]
    # Links and URLs carry only the closed choices.
    assert query.carried() == [("mode", "testing")]
    assert "Adams" not in query.url("submitted") and "Adams" not in repr(query)
    assert ListQuery.parse(spec, QueryDict("search=+"), private=True)[0] == (
        ListQuery()
    )
    for invalid in (
        "search=a&search=b",
        "search=" + "x" * 201,
        "search=" + "1" * 5000,
        "search=a%00",
    ):
        with pytest.raises(ValueError):
            ListQuery.parse(spec, QueryDict(invalid), private=True)


def test_csv_is_complete_neutralized_and_in_the_chosen_zone():
    """Header row, CRLF, formula-neutralized text, ISO times, blanks for missing."""
    spec = LISTS["submitted"]
    body = list_csv(spec, rows_of("submitted"), NEW_YORK)
    assert body.count(b"\r\n") == 4
    table = list(csv.reader(io.StringIO(body.decode("utf-8"))))
    # Family leads (the most relevant column, #932), then its dates.
    assert table[0] == [
        "Family",
        "Family DUID",
        "Envelope number",
        "Submissions",
        "First submitted",
    ]
    first = ROWS[0].submitted_at.astimezone(NEW_YORK).isoformat(" ", "seconds")
    assert table[1] == ["Adams, Ann", "1", "101", "2", first]
    assert first.endswith("-04:00")
    # A name starting with "=" can never run as a formula.
    assert table[2][0] == "'=Baker, Bob"
    # No envelope number is a blank cell, never a word.
    assert table[3][2] == ""


def test_data_quality_csv_says_what_to_check():
    """The file names the problem; a blank mailing name is an empty cell."""
    body = list_csv(LISTS["data-quality"], rows_of("data-quality"), NEW_YORK)
    table = list(csv.reader(io.StringIO(body.decode("utf-8"))))
    # Family, Family DUID and envelope lead; the date comes last (#932).
    assert table[0] == [
        "Family",
        "Family DUID",
        "Envelope number",
        "Mailing name",
        "What to check",
        "First submitted",
    ]
    assert table[1][0:5] == ["Cole, Cy", "3", "103", "", "Blank mailing name"]
    assert table[2][2] == "0"


def test_response_families_checks_inputs_before_any_read():
    """A naive cutoff or a bare UUID never reaches PostgreSQL."""
    scope = ResponseScope(uuid4())
    with pytest.raises(ValueError, match="timezone-aware"):
        response_families(scope, START.replace(tzinfo=None))
    with pytest.raises(TypeError):
        response_families(scope.campaign_id, START)


def render(key="submitted", query=None, rows=None, values=None, **options):
    """A list page as the view renders it, from the test rows."""
    spec = LISTS[key]
    query = query or ListQuery()
    table = paginate(
        []
        if rows is False
        else rows_of(key, query.show, search=query.search)
        if rows is None
        else rows,
        values or {},
        carry=query.posted(),
        sorting=spec.sorting,
    )
    options.setdefault("no_rehearsal", rows is False)
    options.setdefault("can_test", True)
    options.setdefault("can_export", True)
    with using("us_long"):
        context = page_context(
            CAMPAIGN, spec, query, table, START + timedelta(days=1), **options
        )
        return render_to_string("stewardship/response-list.html", context)


def test_page_shows_the_table_filter_and_download():
    """Heading and About first, then the filter, the table and the download."""
    page = render()
    assert "<h1>Families that submitted</h1>" in page
    # The breadcrumb names the list, and the way back to the dashboard is
    # replaced on every in-place refresh, so it follows the mode.
    assert '<p id="list-dashboard-link" data-table-sync>' in page
    assert page.index("<h1>") < page.index('data-about-page="response-list"')
    assert "Sample campaign — Production" in page
    assert '<div id="table" data-table-region>' in page
    # The Family cell is the row's heading and opens the Family's timeline
    # by its opaque campaign record id; identifiers are never grouped.
    adams = next(row for row in ROWS if row.family_duid == 1)
    timeline = reverse("admin:family_timeline", args=[adams.family_id])
    assert f'<th scope="row"><a href="{timeline}">Adams, Ann</a></th>' in page
    assert '<td class="numeric">101</td>' in page
    assert re.search(r'<th scope="row"><a href="[^"?]+/">=Baker, Bob</a></th>', page)
    # Sort headings are links; the search posts privately (CSRF form), and
    # Families that submitted offers no Show selector (#860).
    assert 'aria-sort="ascending" data-sort-column="submitted"' in page
    assert 'href="?size=50&amp;sort=family#table"' in page
    assert (
        f'<form id="table-filters" data-table-sync method="post" '
        f'action="{BASE}submitted/" class="filter-bar">'
    ) in page
    assert '<input type="search" id="list-search" name="search"' in page
    assert 'id="list-show"' not in page and "delivered invitation" not in page
    assert ">Search</button>" in page
    # The address an in-place answer leaves: closed choices only.
    assert f'<a href="{BASE}submitted/" data-page-address hidden></a>' in page
    # The download posts the page's filter and order, with a CSRF form.
    assert (
        f'<form id="table-export" data-table-sync method="post" '
        f'action="{BASE}submitted/csv/">' in page
    )
    assert '<input type="hidden" name="sort" value="submitted">' in page
    assert "The file holds the 3 Families on this list.</p>" in page
    assert "Sensitive parish report" in page
    assert "Downloads are paused" not in page
    # No inline script or style under the CSP.
    assert not re.search(r"<script(?![^>]*\bsrc=)", page)
    assert " style=" not in page and "<style" not in page


def test_page_keeps_filter_in_links_and_download():
    """A chosen filter travels in every table link and in the download."""
    query = ListQuery("testing", "followed")
    page = render("not-opened", query=query)
    assert "Sample campaign — Testing rehearsal" in page
    assert "never counted in Production" in page
    assert '<input type="hidden" name="show" value="followed">' in page
    assert '<input type="hidden" name="mode" value="testing">' in page
    assert "mode=testing&amp;show=followed" in page
    assert '<option value="followed" selected>Link followed</option>' in page
    assert ">Apply</button>" in page
    # A Testing list's Family opens the Family's Testing timeline.
    assert re.search(
        r'<th scope="row"><a href="[^"]+/families/[^"]+/\?mode=testing">', page
    )
    # Switching mode keeps the filter and a chosen order; it is a link, so
    # the filter form then shows what the fresh page applied.
    resorted = render("not-opened", query=query, values={"sort": "-family"})
    assert (
        f'href="{BASE}not-opened/?show=followed&amp;sort=-family#table" '
        'data-in-place="mode-production" data-in-place-filters'
    ) in resorted
    assert "The file holds the 1 Family on this list, with the filter chosen." in page


def test_search_makes_a_private_post_table():
    """With a search, headings, navigators and the download post it privately."""
    query = ListQuery("testing", EVERYONE, "Adams")
    page = render(query=query, values={"sort": "family", "size": "25"})
    assert '<input type="search" id="list-search" name="search"' in page
    assert 'value="Adams"' in page
    # Headings are POST forms carrying the search; no link carries it.
    assert 'class="inline-form sort-form"' in page
    assert '<input type="hidden" name="search" value="Adams">' in page
    assert "Adams&" not in page and "=Adams" not in page
    assert 'href="?' not in page
    # The address keeps only the closed choices; the mode links drop it.
    assert (
        f'<a href="{BASE}submitted/?mode=testing&amp;sort=family&amp;size=25" '
        "data-page-address hidden></a>"
    ) in page
    assert f'href="{BASE}submitted/?size=25&amp;sort=family#table"' in page
    assert (
        "The file holds the 1 Family on this list that matches the search.</p>" in page
    )
    export = page[page.index('id="table-export"') :]
    assert '<input type="hidden" name="search" value="Adams">' in export
    # Nothing found says so, and nothing can be downloaded.
    none = render(query=ListQuery(search="zzz"))
    assert "No Families on this list match the search." in none
    assert "Nothing to download: no Families on this list match the search." in none
    # With a filter as well, the download sentence names both.
    both = render("not-opened", query=ListQuery(show="followed", search="cole"))
    assert "that matches the search, with the filter chosen." in both


def test_missing_values_read_as_words_on_the_page():
    """No envelope, no progress, no name: each says so in plain words."""
    page = render("started", rows=listed(LISTS["started"], ROWS[3:4], {}, ListQuery()))
    assert "Not in the latest ParishSoft data" in page
    assert "Not yet" in page


def test_list_without_a_filter_has_only_the_search():
    """Submitted more than once has one choice, so only the search is offered."""
    page = render("more-than-once")
    assert 'id="table-filters"' in page and 'id="list-show"' not in page
    assert 'aria-sort="descending" data-sort-column="submissions"' in page


def test_staff_and_paused_and_empty_testing_pages():
    """No Testing switch for Staff; a closed purge gate disables the download."""
    staff = render(can_test=False)
    assert "mode=testing" not in staff
    paused = render(paused=True)
    assert "Downloads are paused while this campaign is prepared for purge." in paused
    assert '<button type="submit" disabled>Download</button>' in paused
    empty = render(query=ListQuery("testing"), rows=False, values={"sort": "family"})
    assert "no Testing responses to show" in empty
    assert "<table" not in empty
    assert "Nothing to download" in empty
    assert '<button type="submit" disabled>Download</button>' in empty
    # Its switch back to Production keeps the chosen order and lands on the
    # table; each switch is an in-place link (#519).
    production = f'<a href="{BASE}submitted/?sort=family#table" data-in-place='
    assert production + '"mode-production" data-in-place-filters>Production</a>' in (
        empty
    )
    no_export = render(can_export=False)
    assert 'id="table-export"' not in no_export


def test_empty_list_says_so():
    """A list with no Families says so in its one row."""
    page = render(rows=[])
    assert "No Families on this list." in page
    assert "Nothing to download: no Families are on this list." in page
    assert '<button type="submit" disabled>Download</button>' in page
    # A list without a filter does not mention one.
    assert "The file holds the 1 Family on this list.</p>" in render("more-than-once")


def test_dashboard_links_each_list_with_its_count(monkeypatch):
    """The dashboard's panel names every list, with its length when known."""
    from parishkit.stewardship.reports.response_dashboard import (
        DashboardQuery,
        family_lists,
    )

    from .test_chart_specs import metrics

    shown = family_lists(DashboardQuery("testing"), metrics())
    assert [(item["url"], item["count"]) for item in shown] == [
        (BASE + "submitted/?mode=testing", 3),
        (BASE + "started/?mode=testing", 1),
        (BASE + "not-opened/?mode=testing", 1),
        (BASE + "more-than-once/?mode=testing", 1),
        (BASE + "data-quality/?mode=testing", None),
    ]
    # Only the list without a length says what it holds.
    assert [bool(item["description"]) for item in shown] == [False] * 4 + [True]
    assert response_lists.LISTS is LISTS


def test_rows_keep_the_funnel_order_for_ties():
    """Equal values keep DUID order, so pages never swap rows between reads."""
    same = tuple(replace(row, submitted_at=START) for row in ROWS if row.submitted_at)
    spec = LISTS["submitted"]
    rows = listed(spec, same, FACTS, ListQuery())
    assert duids(paginate(rows, {}, sorting=spec.sorting).rows) == [1, 2, 5]


def test_breadcrumb_names_the_list():
    """The last breadcrumb is the list's own title, not a generic label."""
    spec = LISTS["not-opened"]
    table = paginate([], {}, sorting=spec.sorting)
    context = page_context(CAMPAIGN, spec, ListQuery(), table, START)
    assert context["breadcrumb_label"] == (
        "Invited Families that never opened the form"
    )


def details_of(key, query=None, count=3):
    """The report details a download of ``key`` carries (``list_details``)."""
    return list_details(
        LISTS[key],
        query or ListQuery(),
        parish="Sample parish",
        campaign="Sample campaign",
        as_of=START + timedelta(days=1),
        zone=NEW_YORK,
        count=count,
    )


def test_xlsx_has_the_csv_rows_as_native_cells_and_the_details():
    """Same headings and rows as the CSV; times, counts native; text literal."""
    from openpyxl import load_workbook

    spec = LISTS["submitted"]
    rows = rows_of("submitted")
    query = ListQuery(search="Adams")
    body = list_file(
        spec,
        rows,
        NEW_YORK,
        "xlsx",
        details=details_of("submitted", query),
        as_of=START + timedelta(days=1),
    )
    book = load_workbook(io.BytesIO(body))
    try:
        sheet = book["Families"]
        table = [[cell.value for cell in row] for row in sheet.iter_rows()]
        csv_rows = list(
            csv.reader(io.StringIO(list_csv(spec, rows, NEW_YORK).decode("utf-8")))
        )
        assert table[0] == csv_rows[0]
        assert len(table) == len(csv_rows)
        first = ROWS[0].submitted_at.astimezone(NEW_YORK).replace(tzinfo=None)
        # A native date and time in the chosen zone; identifiers as text; the
        # count a number; a missing envelope number a truly empty cell.
        assert table[1] == ["Adams, Ann", "1", "101", 2, first]
        assert sheet.cell(2, 5).is_date
        # "=Baker, Bob" is literal text, never a formula.
        assert sheet.cell(3, 1).data_type == "s"
        assert table[2][0] == "=Baker, Bob"
        assert table[3][2] is None
        information = {
            row[0].value: row[1].value for row in book["Report information"].iter_rows()
        }
    finally:
        book.close()
    assert information["Report"] == "Families that submitted"
    assert information["Campaign"] == "Sample campaign"
    assert information["Responses"] == "Production"
    assert information["Display time zone"] == "America/New_York"
    assert information["Counted at"] == (START + timedelta(days=1)).astimezone(
        NEW_YORK
    ).replace(tzinfo=None)
    assert information["Families in this file"] == "3"
    # Whether a search was applied, never its text (#849).
    assert information["Search applied"] == "Yes"
    assert "Adams" not in repr(sorted(map(str, information.values())))
    assert information["Privacy"] == response_lists.PRIVACY


def test_file_text_is_neutralized_in_every_format():
    """A control character pasted into ParishSoft cannot break XLSX or PDF."""
    from openpyxl import load_workbook

    spec = LISTS["submitted"]
    row = rows_of("submitted")[0]
    odd = replace(row, facts=FamilyFacts("Tab\x0bName", 7, "x"))
    book = load_workbook(
        io.BytesIO(
            list_file(
                spec, [odd], NEW_YORK, "xlsx", details=details_of("submitted", count=1)
            )
        )
    )
    try:
        assert book["Families"].cell(2, 1).value == "Tab\\u000bName"
    finally:
        book.close()
    drawn = []
    original = pdf_design.Canvas.text

    def capture(canvas, x, y, text, *args, **kwargs):
        """Keep each drawn string, then draw it for real."""
        drawn.append(text)
        return original(canvas, x, y, text, *args, **kwargs)

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(pdf_design.Canvas, "text", capture)
        body = list_file(
            spec,
            [odd],
            NEW_YORK,
            "pdf",
            details=details_of("submitted", count=1),
            as_of=START,
        )
    assert body.startswith(b"%PDF")
    assert any("Tab\\u000bName" in text for text in drawn)
    assert csv_cell_text(list_csv(spec, [odd], NEW_YORK)) == "Tab\x0bName"


def csv_cell_text(body):
    """The Family cell of a one-row CSV."""
    return list(csv.reader(io.StringIO(body.decode("utf-8"))))[1][0]


def test_pdf_reads_as_the_page_with_its_details_on_every_page():
    """Missing values in the page's words; counts grouped; details in the header."""
    drawn = {}

    def capture(document, output):
        """Record the document and its page frame instead of drawing them."""
        drawn.update(document=document, frame=response_lists.list_frame(document))
        return 1

    spec = LISTS["started"]
    query = ListQuery(show="opened")
    with pytest.MonkeyPatch.context() as patch, using("us_long"):
        patch.setattr(response_lists, "list_pdf", capture)
        list_file(
            spec,
            rows_of("started", "opened"),
            NEW_YORK,
            "pdf",
            details=details_of("started", query, count=1),
            as_of=START,
        )
    document = drawn["document"]
    assert document.headings == tuple(str(c.heading) for c in spec.columns)
    family, duid, envelope, opened, progressed = document.rows[0]
    # Got past the first step is not reached yet: the page's "Not yet".
    assert progressed == "Not yet"
    assert family == "Diaz, Dee" and duid == "4" and envelope == "0"
    assert opened.startswith("Oct ") and "UTC" not in opened
    frame = drawn["frame"]
    assert frame.title == str(spec.title)
    assert frame.eyebrow == "Sample parish · Sample campaign"
    assert frame.stamp.startswith("Counted at October 4, 2026")
    assert frame.stamp.endswith("(America/New_York)")
    assert frame.details == (
        "1 Family in this file. Production responses. Show: Opened the form only.",
    )
    assert frame.notice == response_lists.PRIVACY
    assert response_lists.list_table(document).headings[0] == "Family"


def test_pdf_columns_fit_the_page_and_their_headings():
    """Every list's PDF table spans the page; times and heading words fit.

    A compact time never wraps, and no heading word is broken across lines
    (``pdf_design`` wraps a heading only between words).
    """
    time = "Sep 30, 2026 12:04 PM"
    for spec in LISTS.values():
        headings = tuple(str(column.heading) for column in spec.columns)
        document = SimpleNamespace(
            headings=headings,
            weights={
                heading: response_lists.PDF_WEIGHTS[column.key]
                for heading, column in zip(headings, spec.columns, strict=True)
            },
        )
        table = response_lists.list_table(document)
        assert sum(table.widths) == pytest.approx(pdf_design.BODY_WIDTH)
        times = table.cells([time] * len(headings))
        for column, heading, cell in zip(
            spec.columns, table.heading_cells, times, strict=True
        ):
            assert " ".join(heading).split() == str(column.heading).split()
            if column.kind == "instant":
                assert cell == [time], column.key


def test_counts_group_in_the_pdf_and_stay_numbers_in_xlsx():
    """A count over a thousand reads "1,234" in the PDF; XLSX keeps the int."""
    spec = LISTS["more-than-once"]
    count, envelope = (
        next(column for column in spec.columns if column.key == key)
        for key in ("submissions", "envelope")
    )
    assert response_lists.pdf_text(count, 1234, NEW_YORK) == "1,234"
    assert response_lists.xlsx_value(count, 1234, NEW_YORK) == 1234
    assert response_lists.pdf_text(envelope, 4711, NEW_YORK) == "4711"
    assert response_lists.xlsx_value(envelope, 4711, NEW_YORK) == "4711"
    assert response_lists.pdf_text(envelope, None, NEW_YORK) == ""
    assert response_lists.xlsx_value(envelope, None, NEW_YORK) is None


def test_an_unknown_format_is_refused():
    """Only CSV, XLSX and PDF render."""
    with pytest.raises(ValueError):
        list_file(LISTS["submitted"], [], NEW_YORK, "docx")


def test_page_offers_each_format_with_one_download_button():
    """The format is a choice beside the time zone; the button never changes."""
    page = render()
    export = page[page.index('id="table-export"') :]
    assert (
        '<select id="list-format" name="format"><option value="csv">CSV</option>'
        '<option value="xlsx">XLSX</option><option value="pdf">PDF</option></select>'
    ) in export
    assert export.index('id="list-format"') < export.index('id="list-timezone"')
    assert '<button type="submit">Download</button>' in export
