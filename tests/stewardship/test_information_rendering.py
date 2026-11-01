"""Complete, safe information exports without database or provider startup."""

import csv
import io
import warnings
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from itertools import accumulate

import pytest
from openpyxl import load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS

from parishkit.stewardship.reports.information_documents import (
    HEADINGS,
    information_document,
)
from parishkit.stewardship.reports.information_rendering import (
    information_csv,
    information_pdf,
    information_records,
    information_xlsx,
    plain,
    visible_text,
)
from parishkit.stewardship.reports.pdf_design import (
    BODY_SIZE,
    HIDDEN_FIELDS,
    VALUE_WIDTH,
    pdf_glyphs,
    record_body,
    safe,
    text_width,
)
from parishkit.stewardship.source.snapshot_names import FAMILY_NAMES_FULL


def card_text(records):
    """Each card as drawn: its title and tag, then (label, text) per line.

    Count tiles give their label and figure; a line holding two short
    fields gives both.
    """
    for record in records:
        yield "Title", record.title
        if record.tag:
            yield "Tag", record.tag
        for line in record_body(record):
            if line.kind == "stats":
                yield from ((label, safe(plain(value))) for label, value in line.text)
                continue
            yield line.label, line.text
            if line.right:
                yield line.right[1], line.right[2]


def document_metadata():
    """The captured campaign and source details every sample report shares."""
    instant = "2026-10-02T12:00:00+00:00"
    return dict(
        id="campaign",
        name="Sample campaign",
        source_id="source",
        source_generation=1234,
        source_as_of=instant,
        family_names=FAMILY_NAMES_FULL,
        observed_at=instant,
        timezone="America/Detroit",
    )


def document(*, count=1, history=True, **changes):
    """Use hostile literal values and enough text to cross a PDF page boundary.

    ``changes`` replace the item's values.
    """
    instant = "2026-10-02T12:00:00+00:00"
    item = (
        dict(
            id="item",
            version=3,
            family_name="=Sample <Family>",
            family_duid=12345,
            submitted_at=instant,
            disposition="current_actionable",
            replacement_id=None,
            text="First line\n\n" + "x" * 4200 + "\nFINAL TEXT MARKER",
            follow_up_needed=True,
            followed_up_at=instant,
            completed_by="staff@example.org",
            notes="+Literal notes",
            previously_reported=True,
            correction_resolved=False,
            history=[
                dict(
                    version=2,
                    follow_up_needed=True,
                    followed_up_at=None,
                    notes="@Older notes",
                    created_at=instant,
                    actor_label="staff@example.org",
                )
            ],
        )
        | changes
    )
    return information_document(
        dict(
            metadata=document_metadata(),
            total=count,
            rows=[item] * count,
        ),
        dict(filters={"search": "private"}, history=history),
        parish_name="Sample Parish",
        requested_at=datetime(2026, 10, 2, 13, tzinfo=UTC),
        timezone="America/Detroit",
    )


def _information_details(filters):
    """The report information of a sample capture with these filters."""
    report = information_document(
        dict(metadata=document_metadata(), total=0, rows=[]),
        dict(filters=filters, history=False),
        parish_name="Sample Parish",
        requested_at=datetime(2026, 10, 2, 13, tzinfo=UTC),
        timezone="America/Detroit",
    )
    return dict(report.metadata)


def test_an_undated_capture_names_no_date_range():
    """The page has no submitted-date filter (#1888): its capture's empty
    days and zone, and the date-filter zone, are left out of the file."""
    details = _information_details(
        {"search": "", "start": "", "end": "", "zone": "UTC", "sort": "name"}
    )
    assert "Campaign date-filter time zone" not in details
    assert details["Filters and sort"] == '{"search": "", "sort": "name"}'


def test_a_dated_capture_keeps_its_date_range():
    """A ``pk-admin export information --start`` capture keeps both."""
    details = _information_details(
        {"search": "", "start": "2026-11-01", "end": "", "zone": "UTC"}
    )
    assert details["Campaign date-filter time zone"] == "America/Detroit"
    assert details["Filters and sort"] == (
        '{"end": "", "search": "", "start": "2026-11-01", "zone": "UTC"}'
    )


@pytest.mark.parametrize("history", [True, False])
def test_csv_full_results_and_formula_safety(history):
    """More than a UI page, complete text/history, immutable local-time metadata."""
    report = document(count=51, history=history)
    stream = io.BytesIO()
    information_csv(report, stream)
    assert not stream.closed and b"\r\n" in stream.getvalue()
    reader = csv.DictReader(io.StringIO(stream.getvalue().decode()))
    rows = list(reader)
    # Exactly the table (#966): the headings, then one row per record, with
    # no report-information columns or row.
    assert tuple(reader.fieldnames) == HEADINGS
    assert len(rows) == 51 * (2 if history else 1)
    items = [row for row in rows if not row["Changed at"]]
    assert len(items) == 51
    assert all(row["Family"] == "'=Sample <Family>" for row in items)
    assert all(row["Staff notes"] == "'+Literal notes" for row in items)
    assert items[-1]["Submitted text"].endswith("FINAL TEXT MARKER")
    # CSV timestamps are ISO 8601 in the stated display time zone (#221).
    assert items[-1]["Submitted"] == "2026-10-02 08:00:00-04:00"
    # No internal references (PR #929, #966): an earlier workflow row
    # follows its item and repeats the item's identity, so a reader can
    # match them; it alone has a Changed at time.
    assert not {"Row type", "Version", "Item reference"} & set(reader.fieldnames)
    assert all("item" not in row.values() for row in rows)
    if history:
        identity = ("Family", "Family DUID", "Submitted", "Submitted text")
        for item, revision in zip(rows[0::2], rows[1::2], strict=True):
            assert not item["Changed at"] and revision["Changed at"]
            assert all(revision[key] == item[key] for key in identity)
            assert revision["Staff notes"] == "'@Older notes"
    assert len(report.rows[0]) == len(HEADINGS)


def test_replacement_is_named_by_its_submission_time():
    """A superseded item names its replacement's row, or says it is not exported."""
    base = document(count=0)
    payload_item = dict(
        version=1,
        family_name="Smith",
        family_duid=7,
        disposition="current_actionable",
        replacement_id=None,
        text="Request",
        follow_up_needed=False,
        followed_up_at=None,
        completed_by="",
        notes="",
        previously_reported=False,
        correction_resolved=False,
        history=[],
    )
    rows = [
        payload_item
        | dict(
            id="old",
            submitted_at="2026-10-01T12:00:00+00:00",
            disposition="superseded",
            replacement_id="new",
        ),
        payload_item | dict(id="new", submitted_at="2026-10-03T12:00:00+00:00"),
        payload_item
        | dict(
            id="older",
            submitted_at="2026-09-30T12:00:00+00:00",
            disposition="superseded",
            replacement_id="filtered-out",
        ),
    ]
    report = information_document(
        dict(metadata=document_metadata(), total=3, rows=rows),
        dict(filters={}, history=False),
        parish_name="Sample Parish",
        requested_at=base.requested_at,
        timezone="America/Detroit",
    )
    stream = io.BytesIO()
    information_csv(report, stream)
    old, new, older = list(csv.DictReader(io.StringIO(stream.getvalue().decode())))
    assert old["Disposition"] == "Superseded by a later response"
    assert old["Replaced by request submitted"] == new["Submitted"]
    assert new["Replaced by request submitted"] == ""
    assert older["Replaced by request submitted"] == "Not in this export"
    # The PDF card says the same in the parish date format.
    _, card, *_ = information_records(report)
    assert dict(card.fields)["Replaced by request submitted"] == datetime(
        2026, 10, 3, 12, tzinfo=UTC
    )


def test_xlsx_values_are_literal_complete_and_structured():
    """Test ParishKit's cell construction, not the spreadsheet library itself."""
    stream = io.BytesIO()
    information_xlsx(document(), stream)
    book = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        sheet = book["Information"]
        # The same headings as the CSV, with no internal references.
        assert [cell.value for cell in sheet[1]] == list(HEADINGS)
        # The weekly digest's columns left with it (#2102); the Timeline
        # and the request's history keep the record.
        assert HEADINGS[-2:] == ("Changed at", "Changed by")
        for retired in ("Previously reported", "Correction resolved"):
            assert retired not in HEADINGS
        details = [row[0].value for row in book["Report information"].iter_rows()]
        assert "Digest resolution" not in details
        assert "item" not in [cell.value for cell in sheet[2]]
        # The earlier workflow row (the one with a Changed at time) repeats
        # its item's identity.
        # ParishSoft status (#1216) is column B, so the rest move one right.
        assert sheet["B2"].value is None
        assert sheet["L2"].value is None and sheet["L3"].value is not None
        assert all(sheet[f"{c}3"].value == sheet[f"{c}2"].value for c in "ABCF")
        assert sheet["A2"].value == "=Sample <Family>" and sheet["A2"].data_type == "s"
        assert sheet["G2"].value.endswith("FINAL TEXT MARKER")
        assert sheet["K3"].value == "@Older notes" and sheet["K3"].data_type == "s"
        # A Family-keyed sheet keeps its Family column in view too.
        assert sheet.freeze_panes == "B2" and sheet.print_title_rows == "$1:$1"
        assert sheet.auto_filter.ref == "A1:M3"
        # Native date-time cells in Excel's locale-aware built-in format 22.
        requested = book["Report information"]["B7"]
        assert requested.value == datetime(2026, 10, 2, 9, 0)
        assert requested.is_date and requested.number_format == BUILTIN_FORMATS[22]
    finally:
        book.close()


def test_pdf_pagination_preserves_all_long_text_and_history():
    """The exact line plan drawn by PDF preserves paragraphs, tails and long words."""
    report = document()
    lines = tuple(text for _, text in card_text(information_records(report)))
    assert any("FINAL TEXT MARKER" in line for line in lines)
    assert (
        sum(len(line.strip()) for line in lines if set(line.strip()) == {"x"}) == 4200
    )
    assert any("@Older notes" in line for line in lines)
    # PDFs use the parish date format (the default US long style here).
    assert any("October 2, 2026 at 9:00 AM EDT" in line for line in lines)
    # Every line fits the value column by the font's own glyph widths.
    assert all(text_width(line, BODY_SIZE) <= VALUE_WIDTH for line in lines)
    stream = io.BytesIO()
    pages = information_pdf(report, stream)
    assert pages > 1 and stream.getvalue().startswith(b"%PDF")
    assert not stream.closed


def test_empty_reports_still_include_source_and_request_provenance():
    """No-result exports remain self-identifying without inventing an item row.

    The CSV is just the headings; the XLSX's report information still
    names the source and request times, with no internal references.
    """
    report = document(count=0)
    stream = io.BytesIO()
    information_csv(report, stream)
    assert stream.getvalue().decode() == ",".join(HEADINGS) + "\r\n"
    stream = io.BytesIO()
    information_xlsx(report, stream)
    book = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        details = dict(book["Report information"].iter_rows(values_only=True))
    finally:
        book.close()
    assert details["Matching items"] == "0"
    assert details["Source as of"] == details["Requested at"] - timedelta(hours=1)
    assert not {"Source reference", "Campaign reference"} & set(details)
    assert report.item_count == 0 and not report.rows


@pytest.mark.parametrize(
    ("changes", "cells"),
    [
        ({}, ("=Sample <Family>", "")),
        ({"active": True}, ("=Sample <Family>", "")),
        ({"active": False}, ("=Sample <Family>", "No longer in this campaign")),
        ({"family_name": None, "active": None}, ("", "Not in ParishSoft")),
    ],
)
def test_parishsoft_status_follows_the_family(changes, cells):
    """The page's badge is a column of its own, on every row of an item (#1216).

    A capture made before the selection returned ``active`` reads as active.
    """
    report = document(**changes)
    assert HEADINGS[:3] == ("Family", "ParishSoft status", "Family DUID")
    assert {row[:2] for row in report.rows} == {cells}


def test_format_exclusions_have_lossless_visible_notation():
    """XML controls and unavailable glyphs cannot poison an immutable export."""
    value = "Éλληνικά 中文 🙂\x0b\x1b\ufffe literal \\u000b"
    report = document()
    row = list(report.rows[0])
    row[HEADINGS.index("Submitted text")] = value
    report = replace(report, rows=(tuple(row),))
    expected = "Éλληνικά 中文 🙂\\u000b\\u001b\\ufffe literal \\\\u000b"
    assert visible_text(value) == expected
    stream = io.BytesIO()
    information_xlsx(report, stream)
    book = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        assert book["Information"]["G2"].value == expected
        notes = [str(cell.value) for cell in book["Report information"]["B"]]
        assert any("Unsupported characters" in note for note in notes)
    finally:
        book.close()
    lines = "\n".join(text for _, text in card_text(information_records(report)))
    assert "Éλληνικά" in lines
    # Only a report that needs escapes explains them.
    assert "Unsupported characters" in lines
    assert "\\u4e2d\\u6587" in lines and "\\U0001f642" in lines
    assert "\\u000b\\u001b\\ufffe" in lines and "\\\\u000b" in lines
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        information_pdf(report, io.BytesIO())
    assert not any("Glyph" in str(w.message) for w in caught)


def test_pdf_draws_pages_without_accumulating_wrapped_report(monkeypatch):
    """After the counting pass, drawing reads at most one card past each page."""
    from matplotlib.backends.backend_pdf import PdfPages

    from parishkit.stewardship.reports import information_rendering, pdf_design

    passes, consumed, saved = 0, 0, 0
    savefig = PdfPages.savefig
    frame = pdf_design.PdfFrame.for_document(document())
    total = 60

    def cards():
        """Short cards, so several share each page."""
        return (
            pdf_design.Record(title=f"Card {index}", fields=(("Line", str(index)),))
            for index in range(total)
        )

    ends = list(
        accumulate(len(page) for page in pdf_design.record_pages(cards(), frame))
    )

    def records(document):
        """Expose card consumption without relying on platform memory accounting."""
        nonlocal passes, consumed
        passes += 1
        for card in cards():
            if passes == 2:
                consumed += 1
            yield card

    def save_page(pdf, figure):
        """A whole-report list would consume every card before the first save."""
        nonlocal saved
        assert consumed <= min(ends[saved] + 1, total)
        saved += 1
        return savefig(pdf, figure)

    monkeypatch.setattr(information_rendering, "information_records", records)
    monkeypatch.setattr(PdfPages, "savefig", save_page)
    assert information_pdf(document(), io.BytesIO()) == saved == len(ends) > 2
    assert passes == 2


def test_pdf_cards_are_titled_nested_and_leave_out_internal_fields():
    """The PDF reads as cards holding the CSV's and XLSX's columns."""
    report = document()
    report_card, item = information_records(report)
    assert report_card.title == "About this report"
    assert dict(report_card.stats) == {"Matching items": "1"}
    details = dict(report_card.fields)
    assert details["Filters and sort"] == "Search: private"
    # The page frame states the parish, campaign, capture time and privacy.
    assert not {"Parish", "Campaign", "Captured at", "Privacy"} & set(details)
    assert not HIDDEN_FIELDS & set(details)
    assert "Text representation" not in details
    assert item.title == "=Sample <Family> (DUID 12345)"
    assert item.tag == "Current actionable request"
    labels = {label for label, _ in item.fields}
    assert not HIDDEN_FIELDS & labels
    # Blank fields are not drawn; the workflow revision nests under its item.
    drawn = {label for label, _ in card_text([item])}
    assert "Changed at" not in drawn and "Completed by" in drawn
    [(heading, history)] = item.children
    assert heading.startswith("Earlier workflow · changed October 2, 2026")
    assert heading.endswith("by staff@example.org")
    assert {label: value for label, value in history if value} == {
        "Follow-up needed": "Yes",
        "Staff notes": "@Older notes",
    }
    # Without history the item has no nested blocks.
    _, item = information_records(document(history=False))
    assert not item.children


def test_pdf_escapes_mapped_invisible_format_characters():
    """A cmap entry is not proof that a Unicode format/control glyph is visible."""
    supported = pdf_glyphs()
    assert 0xFEFF in supported
    assert (
        visible_text("before\ufeffafter", supported=supported) == "before\\ufeffafter"
    )
    assert (
        visible_text("a\u200db\x7fc\nd", supported=supported) == "a\\u200db\\u007fc\nd"
    )
