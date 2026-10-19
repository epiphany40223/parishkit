"""The roster changes to enter: eligibility, Ministry requests' roster rows
and the file.

Roster changes to enter (#528 step 4, #1687) was folded into Ministry
requests (R39 finding 2): Staff mark roster changes at its Stage To enter by
hand, where the rows a review stage leaves out (a Member no longer in
ParishSoft, a request a later response replaced) offer Entered by hand
alone, with no note. ``report roster`` and ``export roster`` still read the
list and write its file.
"""

import csv
import io
from datetime import UTC, datetime, timedelta
from uuid import UUID
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy import Principal
from parishkit.stewardship.pipeline import Decision, Placement, Stage
from parishkit.stewardship.reports import pdf_design
from parishkit.stewardship.reports.ministry_roster_views import (
    HEADINGS,
    PDF_WEIGHTS,
    roster_csv,
    roster_pdf,
    roster_xlsx,
)
from parishkit.stewardship.workflows.roster import eligible

WHEN = datetime(2054, 10, 5, 14, 30, tzinfo=UTC)


def roster_row(**values):
    """One listed roster change, as ``roster_rows`` and ``with_decisions``
    shape it, with overrides; ``roster_entered`` sets its stage."""
    row = {
        "id": "1",
        "member_name": "Sample, Pat",
        "member_label": "Sample, Pat",
        "member_duid": 3,
        "member_status": "current",
        "source_mark": None,
        "family_name": "Sample, Pat and Lee",
        "family_duid": 1,
        "ministry_name": "Food pantry",
        "family_ministry_name": "Food pantry",
        "ministry_duid": 9,
        "action": "join",
        "state": "resolved",
        "outcome": "joined",
        "source_resolved": False,
        "resolved_at": WHEN,
        "submitted_at": WHEN - timedelta(days=2),
        "decided_by": "admin@example.org",
        "decided_at": WHEN,
        "note": "Called the Family",
        "latest": True,
        "roster_entered": True,
        "roster_sequence": 1,
        "roster_by": "staff@example.org",
        "roster_at": WHEN,
    } | values
    stage = Stage.ENTERED_BY_HAND if row["roster_entered"] else Stage.TO_ENTER_BY_HAND
    decision = (
        Decision.ADDED
        if row.get("parish_added")
        else Decision.ACCEPTED_EDITED
        if row.get("reviewed_ministry_duid")
        else Decision.ACCEPTED
    )
    row.setdefault("placement", Placement(stage, decision))
    row.setdefault("stage", row["placement"].stage.value)
    return row


def test_only_person_resolved_roster_changes_take_the_tick():
    """Resolved joins and leaves with no resolution source; nothing else."""
    base = {"state": "resolved", "outcome": "joined", "source_resolved": False}
    assert eligible(base)
    assert eligible(base | {"outcome": "leave_confirmed"})
    assert not eligible(base | {"source_resolved": True})
    assert not eligible(base | {"outcome": "declined"})
    assert not eligible(base | {"state": "superseded"})


def test_download_times_are_in_the_chosen_zone_and_formulas_are_neutral():
    """Instants carry the zone's offset; a formula-like name stays text."""
    row = roster_row(member_name="=HYPERLINK(1)")
    text = roster_csv([row], ZoneInfo("America/Chicago")).decode()
    table = list(csv.reader(io.StringIO(text)))
    assert tuple(table[0]) == HEADINGS
    # The Ministry is named; its DUID is not a column (#1123).
    assert "Ministry" in HEADINGS and "Ministry DUID" not in HEADINGS
    line = dict(zip(HEADINGS, table[1], strict=True))
    assert not line["Member"].startswith("=")
    assert line["Last decision at"] == "2054-10-05T09:30:00-05:00"
    assert line["Last decision"] == "admin@example.org"
    assert line["Note"] == "Called the Family"
    assert line["Entered at"] == "2054-10-05T09:30:00-05:00"
    assert line["Submitted"] == "2054-10-03T09:30:00-05:00"
    assert line["Family"] == "Sample, Pat and Lee" and line["Family DUID"] == "1"
    assert line["Stage"] == "Entered by hand" and line["Decision"] == "Accepted"
    assert line["Family's answer"] == "Join" and line["Reviewer's value"] == ""


def test_download_words_a_new_member_and_an_unentered_change_as_the_page():
    """A Member added on the form has "New Member" for a DUID; a change not
    yet entered is To enter by hand, with no one and no time; an edited one
    names the Family's Ministry and the one to enter."""
    row = roster_row(
        member_duid=None, roster_entered=False, roster_by="", roster_at=None
    )
    text = roster_csv([row], ZoneInfo("UTC")).decode()
    line = dict(zip(HEADINGS, list(csv.reader(io.StringIO(text)))[1], strict=True))
    assert line["Member DUID"] == "New Member"
    assert line["Stage"] == "To enter by hand"
    assert line["Replaced by a later response"] == ""
    assert line["Entered by"] == "" and line["Entered at"] == ""
    # A Family that answered again after the change was decided.
    replaced = roster_csv([roster_row(latest=False)], ZoneInfo("UTC")).decode()
    line = dict(zip(HEADINGS, list(csv.reader(io.StringIO(replaced)))[1], strict=True))
    assert line["Replaced by a later response"] == "Yes"
    edited = roster_row(
        reviewed_ministry_duid=4,
        reviewed_ministry_name="Adult Choir",
        ministry_name="Adult Choir",
        family_ministry_name="Altar Servers",
    )
    text = roster_csv([edited], ZoneInfo("UTC")).decode()
    line = dict(zip(HEADINGS, list(csv.reader(io.StringIO(text)))[1], strict=True))
    assert line["Ministry"] == "Adult Choir"
    assert line["Family's answer"] == "Join Altar Servers"
    assert line["Reviewer's value"] == "Join Adult Choir"
    assert line["Decision"] == "Accepted (edited)"


def test_csv_xlsx_and_pdf_carry_the_same_columns_and_rows():
    """One column set in every format (#966): the XLSX cells are the CSV's,
    and the PDF table has the same headings and values, only its times worded
    in the parish's compact date format."""
    zone = ZoneInfo("America/Chicago")
    rows = [roster_row(), roster_row(id="2", action="leave", member_duid=None)]
    table = list(csv.reader(io.StringIO(roster_csv(rows, zone).decode())))
    book = load_workbook(io.BytesIO(roster_xlsx(rows, zone)))
    sheet = [
        ["" if value is None else value for value in values]
        for values in book["Roster changes"].iter_rows(values_only=True)
    ]
    assert sheet == table
    drawn = {}
    original = pdf_design.table_pages

    def capture(pdf_table, pdf_rows, frame):
        """Keep the table and rows the PDF lays out, then lay them out."""
        pdf_rows = list(pdf_rows)
        drawn.update(table=pdf_table, rows=pdf_rows, frame=frame)
        return original(pdf_table, pdf_rows, frame)

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(pdf_design, "table_pages", capture)
        body = roster_pdf(rows, zone, parish="P", campaign="C", as_of=WHEN)
    assert body.startswith(b"%PDF-")
    assert drawn["table"].headings == HEADINGS == tuple(table[0])
    times = {
        HEADINGS.index(name) for name in ("Entered at", "Submitted", "Last decision at")
    }

    def untimed(values):
        """The row's cells other than its times."""
        return [cell for index, cell in enumerate(values) if index not in times]

    assert [untimed(values) for values in drawn["rows"]] == [
        untimed(values) for values in table[1:]
    ]
    assert drawn["rows"][0][HEADINGS.index("Last decision at")] == (
        "Oct 5, 2054 9:30 AM"
    )
    assert drawn["frame"].details == ("2 changes in this file.",)
    assert set(PDF_WEIGHTS) == set(HEADINGS) and sum(PDF_WEIGHTS.values()) == 100
    # A whole DUID stays on one line; a compact time keeps its date whole.
    for heading, text, lines in (
        ("Member DUID", "1234567", ["1234567"]),
        ("Family DUID", "1234567", ["1234567"]),
        ("Entered at", "Sep 30, 2026 12:04 PM", ["Sep 30, 2026", "12:04 PM"]),
    ):
        values = [text if name == heading else "" for name in HEADINGS]
        cell = drawn["table"].cells(values)[HEADINGS.index(heading)]
        assert [line.strip() for line in cell] == lines, heading


def tick_html(row, **context):
    """The roster tick template rendered for one row and context."""
    from django.template.loader import render_to_string

    return render_to_string(
        "stewardship/roster-tick.html", {"row": row, "roster_key": "k"} | context
    ).strip()


def test_queue_status_cell_says_one_short_line():
    """The follow-up queue's Status cell says only Entered or Not yet
    entered, and never draws the form; who and when are on the request
    page."""

    def cell(**values):
        """The tick as the queue's Status cell draws it for Staff."""
        row = roster_row(roster_eligible=True, source_resolved=False) | values
        return tick_html(row, can_tick=True, compact=True, mutable=True)

    assert cell() == "<br>Entered"
    assert cell(roster_entered=False) == "<br>Not yet entered"
    assert "staff@example.org" not in cell()


def test_a_leader_sees_no_tick_at_all():
    """A Ministry leader (no can_tick) gets nothing: no state, who, when or
    form, on the request page or in the queue (#1160)."""
    for entered in (True, False):
        row = roster_row(
            roster_eligible=True, source_resolved=False, roster_entered=entered
        )
        for compact in (True, False):
            assert tick_html(row, can_tick=False, compact=compact, mutable=True) == ""


def test_download_names_a_member_no_longer_in_parishsoft_in_its_own_column():
    """The Member's ParishSoft status is blank while ParishSoft lists them,
    and otherwise #1216's phrase, beside the Member, not in the name (#1217)."""
    zone = ZoneInfo("UTC")
    rows = [
        roster_row(),
        roster_row(id="2", source_mark=("Deceased in ParishSoft", "reason")),
    ]
    table = list(csv.reader(io.StringIO(roster_csv(rows, zone).decode())))
    lines = [dict(zip(HEADINGS, line, strict=True)) for line in table[1:]]
    assert [line["ParishSoft status"] for line in lines] == [
        "",
        "Deceased in ParishSoft",
    ]
    assert lines[1]["Member"] == "Sample, Pat"


STAFF = Principal(UUID(int=97), frozenset({"staff"}))
LEADER = Principal(UUID(int=98), frozenset({"ministry_leader"}), frozenset({9}))
GONE = ("Not in ParishSoft", "ParishSoft no longer has this Member")


def request_row(**values):
    """One follow-up selection row of an accepted join, with overrides."""
    return {
        "id": str(UUID(int=1)),
        "version": 2,
        "member_name": "Sample, Pat",
        "member_label": "Sample, Pat",
        "member_duid": 3,
        "member_status": "current",
        "source_mark": None,
        "ministry_name": "Food pantry",
        "ministry_duid": 9,
        "action": "join",
        "state": "resolved",
        "outcome": "joined",
        "source_resolved": False,
        "latest": True,
        "in_campaign": True,
        "cannot_serve": False,
        "reviewed_ministry_duid": None,
        "reviewed_ministry_name": None,
        "parish_added": False,
        "added_in_place_of_id": None,
        "last_contact_at": None,
        "submitted_at": WHEN,
    } | values


def requests_worklist(rows, *, principal=STAFF, stage="to_enter_by_hand", marks=None):
    """Ministry requests' worklist for detached selection ``rows``, at ``stage``."""
    from django.urls import reverse

    from parishkit.stewardship.reports.ministry_worklist import (
        WorklistQuery,
        place,
        worklist_context,
    )

    result = {
        "rows": rows,
        "ministries": [dict(duid=9, name="Food pantry", in_campaign=True)],
        "metadata": {},
    }
    place(result, principal, marks=marks or {})
    return worklist_context(
        UUID(int=92),
        WorklistQuery(stage=stage),
        principal,
        result,
        token="",
        mutable=True,
        action=reverse("admin:ministry_report"),
        decisions={},
        asked={},
        families={},
    )["worklist"]


def test_staff_views_keep_the_roster_rows_the_review_stages_leave_out():
    """The every-revision read keeps a Staff view's current requests for
    current Members, and every accepted join or leave still to enter, for a
    Member no longer in ParishSoft or one a later response replaced; the
    "No longer in ParishSoft" choice keeps only such Members' rows."""
    from parishkit.stewardship.reports.ministry_worklist import listed, roster_only

    current = request_row(state="in_progress", outcome=None)
    replaced_open = request_row(state="in_progress", outcome=None, latest=False)
    withdrawn = request_row(state="cancelled", outcome=None)
    gone_open = request_row(state="new", outcome=None, source_mark=GONE)
    gone_roster = request_row(source_mark=GONE)
    replaced_roster = request_row(latest=False)
    declined_replaced = request_row(outcome="declined", latest=False)
    keep = [
        listed(row, "current")
        for row in (
            current,
            replaced_open,
            withdrawn,
            gone_open,
            gone_roster,
            replaced_roster,
            declined_replaced,
        )
    ]
    assert keep == [True, False, False, False, True, True, False]
    gone = [listed(row, "gone") for row in (current, gone_open, gone_roster)]
    assert gone == [False, True, True]
    assert [roster_only(row) for row in (current, gone_roster, replaced_roster)] == [
        False,
        True,
        True,
    ]


def test_a_roster_row_offers_entered_by_hand_alone_with_no_note():
    """At To enter by hand an ordinary accepted request offers Back to review
    and Entered by hand; a Member no longer in ParishSoft's, and one a later
    response replaced, offer Entered by hand alone. Its dialog hides the
    note (the ``no_note`` choice), and the bulk bar's Entered by hand does
    too. A replaced row says so, in its name and under its Stage."""
    rows = [
        request_row(id=str(UUID(int=1))),
        request_row(id=str(UUID(int=2)), source_mark=GONE, member_name=""),
        request_row(id=str(UUID(int=3)), latest=False),
    ]
    worklist = requests_worklist(rows)
    ordinary, gone, replaced = worklist.rows
    offered = [
        [icon["action"] for icon in row["icons"] if not icon["selected"]]
        for row in worklist.rows
    ]
    assert offered == [
        ["back_to_review", "entered_by_hand"],
        ["entered_by_hand"],
        ["entered_by_hand"],
    ]
    assert gone["kinds"] == replaced["kinds"] == "entered_by_hand"
    assert {icon["choice"] for icon in replaced["icons"] if not icon["selected"]} == {
        "no_note"
    }
    bulk = {button["value"]: button["choice"] for button in worklist.bulk}
    assert bulk["entered_by_hand"] == "no_note"
    assert replaced["name"] == (
        "Sample, Pat — Food pantry, replaced by a later response"
    )
    assert "Replaced by a later response" in str(replaced["decision_label"])
    assert "replaced" not in ordinary["name"]


def test_an_entered_row_says_who_marked_it_and_a_leader_sees_accepted():
    """At Entered by hand the line under the Stage says who marked it and
    when (a removed user as Former portal user); a Ministry leader's row
    reads Accepted, with no mark and no Entered by hand."""
    marks = {
        str(UUID(int=1)): {
            "entered": True,
            "sequence": 1,
            "by": "staff@example.org",
            "at": WHEN,
        },
        str(UUID(int=2)): {"entered": True, "sequence": 1, "by": "", "at": None},
    }
    rows = [request_row(id=str(UUID(int=1))), request_row(id=str(UUID(int=2)))]
    worklist = requests_worklist(rows, stage="entered_by_hand", marks=marks)
    named, former = (str(row["decision_label"]) for row in worklist.rows)
    assert named.startswith("Accepted<br>Entered by staff@example.org, ")
    assert 'datetime="2054-10-05T14:30:00+00:00"' in named
    assert former == "Accepted<br>Entered by Former portal user"
    leader = requests_worklist(
        [request_row(id=str(UUID(int=1)))],
        principal=LEADER,
        stage="accepted",
        marks=marks,
    )
    (row,) = leader.rows
    assert "Entered by" not in str(row["decision_label"])
    assert all(icon["action"] != "entered_by_hand" for icon in row["icons"])
