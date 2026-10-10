"""The shared look of Admin portal PDF and XLSX files (#926).

Styling only: these tests pin the design (header styles, frozen panes,
widths, print setup, the PDF page frame) so a writer cannot drift from it,
while each format's own tests keep pinning the values.
"""

import io
import random
import re
from datetime import UTC, datetime
from functools import partial
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook

from parishkit.stewardship.reports import pdf_design, xlsx_design
from parishkit.stewardship.reports.census_changes import HEADINGS, census_xlsx
from parishkit.stewardship.reports.directory_documents import TESTING_NOTE
from parishkit.stewardship.reports.directory_rendering import (
    address_blocks,
    directory_frame,
    directory_table,
    render_directory,
    table_rows,
)
from parishkit.stewardship.reports.information_rendering import (
    information_pdf,
    information_records,
    information_xlsx,
)
from parishkit.stewardship.reports.ministry_packets import packet_records, render_packet
from parishkit.stewardship.reports.response_lists import (
    LISTS,
    PRIVACY,
    ListQuery,
    list_file,
)
from parishkit.stewardship.reports.spreadsheets import participation_xlsx
from parishkit.stewardship.reports.talents import shape_result, talents_xlsx

from .test_directory_rendering import document as directory
from .test_directory_rendering import item as family
from .test_information_rendering import document as information
from .test_ministry_packets import item as member
from .test_ministry_packets import packet
from .test_participation_rendering import document as participation
from .test_response_lists import NEW_YORK, details_of, rows_of
from .test_talents_report import result as talents_result

# --- XLSX -----------------------------------------------------------------


def assert_styled_table(sheet, header_row=1, *, first_column=False):
    """One table sheet carries the whole shared table style."""
    heading = sheet.cell(header_row, 1)
    assert heading.font.b and heading.font.name == "Arial"
    assert heading.font.color.rgb.endswith("FFFFFF")
    assert heading.fill.fgColor.rgb.endswith("115E56")
    assert heading.border.bottom.style == "medium"
    assert heading.alignment.wrap_text
    if sheet.max_row > header_row:
        body = sheet.cell(header_row + 1, 1)
        assert body.font.name == "Arial" and not body.font.b
        assert body.alignment.vertical == "top"
        # Zebra banding is a display rule, never a cell formula or value.
        [rule] = [
            rule for ranges in sheet.conditional_formatting for rule in ranges.rules
        ]
        assert rule.formula == ["MOD(ROW(),2)=0"]
    assert sheet.freeze_panes == f"{'B' if first_column else 'A'}{header_row + 1}"
    assert sheet.auto_filter.ref.startswith(f"A{header_row}:")
    assert sheet.print_title_rows == f"${header_row}:${header_row}"
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.oddFooter.center.text == "Page &P of &N"
    for column in range(1, sheet.max_column + 1):
        letter = sheet.cell(header_row, column).column_letter
        width = sheet.column_dimensions[letter].width
        assert xlsx_design.MIN_WIDTH <= width <= 70


def assert_information_sheet(sheet):
    """The shared "Report information" style: teal-tinted labels, wrapped values."""
    label, value = sheet["A1"], sheet["B1"]
    assert label.font.b and label.fill.fgColor.rgb.endswith("E4F0EE")
    assert value.alignment.wrap_text and not value.font.b
    assert sheet.column_dimensions["A"].width == 32
    assert sheet.page_setup.orientation == "portrait"


def test_style_table_fits_widths_wraps_long_text_and_keeps_values():
    """Widths fit the content within limits; values and types are untouched."""
    book = Workbook()
    sheet = book.active
    when = datetime(2026, 10, 9, 9, 30)
    rows = [
        ("ID", "Long notes", "When"),
        ("7", "word " * 40, when),
        ("8", "short", when),
    ]
    for values in rows:
        sheet.append(values)
    xlsx_design.style_table(sheet, freeze_first_column=True, title="R&D list")
    assert_styled_table(sheet, first_column=True)
    assert sheet.column_dimensions["A"].width == xlsx_design.MIN_WIDTH
    assert sheet.column_dimensions["B"].width == xlsx_design.MAX_WIDTH
    assert sheet["B2"].alignment.wrap_text and not sheet["A2"].alignment.wrap_text
    assert 19 <= sheet.column_dimensions["C"].width < xlsx_design.MAX_WIDTH
    assert sheet.auto_filter.ref == "A1:C3"
    # A literal ampersand in a title cannot become a header control code.
    assert sheet.oddHeader.left.text == "&BR&&D list"
    assert [tuple(row) for row in sheet.iter_rows(values_only=True)] == rows
    # Two heading lines fit the default height.
    assert sheet.row_dimensions[1].height == 2 * xlsx_design.LINE_HEIGHT + 4


def test_heading_row_grows_with_its_wrapped_lines():
    """A heading that wraps to more lines than two gets a taller row."""
    book = Workbook()
    sheet = book.active
    sheet.append(("ID", "Approximate installment amount per scheduled gift"))
    sheet.append(("7", "x"))
    xlsx_design.style_table(sheet)
    width = sheet.column_dimensions["B"].width
    lines = xlsx_design._line_count(sheet["B1"].value, width, 1.15)
    assert lines > 2
    assert sheet.row_dimensions[1].height == lines * xlsx_design.LINE_HEIGHT + 4


def saved(writer, *arguments):
    """Render a workbook through a writer that takes an output stream."""
    output = io.BytesIO()
    writer(*arguments, output)
    return load_workbook(io.BytesIO(output.getvalue()))


def test_information_and_directory_workbooks_share_the_style():
    """Information keeps column A scrolling; a directory freezes its Family."""
    book = saved(information_xlsx, information(count=3))
    assert_styled_table(book["Information"])
    # Submitted text (column H) gets the wide free-text column.
    assert book["Information"].column_dimensions["H"].width == 70
    assert_information_sheet(book["Report information"])
    output = io.BytesIO()
    render_directory(directory([family()]), output, format="xlsx")
    book = load_workbook(output)
    assert_styled_table(book["Families"], first_column=True)
    assert_information_sheet(book["Report information"])


def test_response_list_workbook_shares_the_style():
    """A response list's XLSX (#850) is the shared styled workbook.

    Its Family column leads (#932), so it stays in view while the sheet
    scrolls sideways, as a directory's does.
    """
    query = ListQuery(search="adams")
    body = list_file(
        LISTS["submitted"],
        rows_of("submitted"),
        NEW_YORK,
        "xlsx",
        details=details_of("submitted", query),
    )
    book = load_workbook(io.BytesIO(body))
    assert book["Families"]["A1"].value == "Family"
    assert_styled_table(book["Families"], first_column=True)
    assert_information_sheet(book["Report information"])


def test_packet_workbook_labels_ministry_details_above_its_table():
    """The details block uses the label style; the table header is frozen."""
    sections = [dict(duid=9, name="Choir", chairs=["Pat Lee"], rows=[member()])]
    output = io.BytesIO()
    render_packet(packet(sections), output, format="xlsx")
    book = load_workbook(output)
    sheet = book["Choir"]
    head = next(
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, 1).value == "Member"
    )
    assert_styled_table(sheet, head)
    assert sheet["A1"].fill.fgColor.rgb.endswith("E4F0EE") and sheet["A1"].font.b
    # Each detail value spans the table's width instead of wrapping in
    # column B, which the table sized for Member DUIDs.
    last = sheet.cell(head, sheet.max_column).column_letter
    merged = {str(cells) for cells in sheet.merged_cells.ranges}
    assert merged == {f"B{row}:{last}{row}" for row in range(1, head - 1)}
    assert sheet["B1"].value == "Choir"
    assert_information_sheet(book["Report information"])


def test_packet_detail_rows_stop_at_excels_tallest_row():
    """A very long Ministry detail is clamped to Excel's 409-point row limit."""
    sections = [dict(duid=9, name="Choir", chairs=["Q" * 30000], rows=[member()])]
    output = io.BytesIO()
    render_packet(packet(sections), output, format="xlsx")
    sheet = load_workbook(output)["Choir"]
    heights = [sheet.row_dimensions[row].height or 0 for row in range(1, 6)]
    assert max(heights) == xlsx_design.MAX_ROW_HEIGHT


def test_participation_census_and_talents_workbooks_share_the_style():
    """The previously unstyled census and talents workbooks match the rest."""
    book = saved(participation_xlsx, participation())
    assert_styled_table(book["Participation"])
    assert_information_sheet(book["Report information"])
    row = dict(
        family_name="Example",
        family_duid=10,
        who="Alex Example",
        label="Email",
        parishsoft_now="old@example.org",
        family_answer="new@example.org",
        edited="",
        route_label="Staff update",
        status_label="Pending",
        submitted_at=datetime(2026, 10, 5, 14, tzinfo=UTC),
    )
    book = load_workbook(io.BytesIO(census_xlsx({"rows": [row]}, ZoneInfo("UTC"))))
    sheet = book["Census changes"]
    assert tuple(cell.value for cell in sheet[1]) == HEADINGS
    assert_styled_table(sheet, first_column=True)
    shaped = shape_result(talents_result({}, {}), configuration={"modules": []})
    book = load_workbook(io.BytesIO(talents_xlsx(shaped, ZoneInfo("UTC"))))
    for title in ("Members", "Families"):
        assert_styled_table(book[title], first_column=True)


# --- PDF ------------------------------------------------------------------


def test_wrap_text_is_lossless_and_fits_the_width():
    """Every paragraph rejoins exactly; every line fits unless one glyph cannot."""
    generator = random.Random(926)
    alphabet = "abc defghij  klmnopqrstuvwxyzÉλ-—.,\t" + "W" * 3
    for _ in range(200):
        text = "".join(
            generator.choice(alphabet) for _ in range(generator.randint(0, 300))
        )
        text += "\n\n" + "x" * generator.randint(0, 400)
        width = generator.uniform(4, 300)
        lines = pdf_design.wrap_text(text, width, 9)
        assert "".join(lines) == "".join(text.expandtabs().split("\n"))
        assert len(lines) >= text.count("\n") + 1
        for line in lines:
            visible = line.rstrip()
            assert (
                len(visible) <= 1 or pdf_design.text_width(visible, 9) <= width + 1e-9
            )


def frames(monkeypatch, render):
    """Render one PDF and return the text drawn on each page."""
    from matplotlib.backends.backend_pdf import PdfPages

    pages, current = [], []
    text = pdf_design.Canvas.text
    savefig = PdfPages.savefig

    def record(canvas, x, y, value, *args, **kwargs):
        """Collect each drawn string for the page being drawn."""
        current.append(value)
        return text(canvas, x, y, value, *args, **kwargs)

    def save(pdf, figure, **kwargs):
        """Close the page's text list when the page is saved."""
        pages.append(list(current))
        current.clear()
        return savefig(pdf, figure, **kwargs)

    monkeypatch.setattr(pdf_design.Canvas, "text", record)
    monkeypatch.setattr(PdfPages, "savefig", save)
    output = io.BytesIO()
    count = render(output)
    assert output.getvalue().startswith(b"%PDF")
    assert count == len(pages)
    return pages


def assert_framed(pages, title, notice=None):
    """Every page repeats the title, parish line, notice and "Page N of M"."""
    total = len(pages)
    for number, drawn in enumerate(pages, 1):
        assert title in drawn
        assert any(value.startswith("SAMPLE PARISH") for value in drawn)
        assert f"Page {number:,} of {total:,}" in drawn
        if notice:
            assert notice in drawn


@pytest.mark.parametrize("postal", [False, True])
def test_directory_pdfs_frame_every_page(monkeypatch, postal):
    """Table and address-card pages share the frame; the table heading repeats."""
    report = directory(
        [family(family_duid=index) for index in range(60)], postal=postal
    )
    pages = frames(
        monkeypatch, lambda output: render_directory(report, output, format="pdf")
    )
    assert len(pages) > 1
    assert_framed(
        pages, report.title, "Sensitive: Family codes. Authorized recipients only."
    )
    if not postal:
        assert all("Family code" in drawn for drawn in pages)
    drawn = [value for page in pages for value in page]
    assert sum("ABCDEFGH" in value for value in drawn) == 60
    # The Testing-mode note is on every page, below the privacy line.
    report = directory(
        [family(family_duid=index) for index in range(60)],
        postal=postal,
        testing=True,
    )
    pages = frames(
        monkeypatch, lambda output: render_directory(report, output, format="pdf")
    )
    note = pdf_design.wrap_text(
        TESTING_NOTE, pdf_design.BODY_WIDTH - 90, pdf_design.FOOTER_SIZE
    )[0]
    assert len(pages) > 1 and all(note in drawn for drawn in pages)


def response_list_pdf(spec, rows, details, output):
    """Write one response list's PDF and return its page count."""
    body = list_file(spec, rows, NEW_YORK, "pdf", details=details)
    output.write(body)
    return len(re.findall(rb"/Type /Page\b", body))


def test_response_list_pdfs_frame_every_page(monkeypatch):
    """Each response list's PDF is the shared framed table (#850).

    The frame states the counts and filters and that a search was applied,
    never its text; the table heading repeats on every page.
    """
    for key, spec in LISTS.items():
        listed = rows_of(key)
        rows = listed * (60 // len(listed) + 1)  # enough rows for two pages
        details = details_of(key, ListQuery(search="adams"), count=len(rows))
        pages = frames(monkeypatch, partial(response_list_pdf, spec, rows, details))
        assert len(pages) > 1, key
        assert_framed(pages, str(spec.title), PRIVACY)
        heading = pdf_design.wrap_text(
            str(spec.columns[0].heading), 40, pdf_design.LABEL_SIZE, "bold"
        )[0]
        for drawn in pages:
            assert any(value.endswith("Search applied.") for value in drawn)
            assert any(value.startswith(heading) for value in drawn), key
            assert not any("adams" in value for value in drawn)


def test_record_pdfs_frame_every_page(monkeypatch):
    """Information and packet PDFs draw cards inside the same frame."""
    report = information(count=4)
    pages = frames(monkeypatch, lambda output: information_pdf(report, output))
    assert len(pages) > 1
    assert_framed(
        pages,
        report.title,
        "Sensitive parish information. Share only with authorized recipients.",
    )
    sections = [
        dict(duid=4, name="Altar Servers", chairs=[], rows=[member()] * 30),
        dict(duid=9, name="Choir", chairs=[], rows=[member()]),
    ]
    document = packet(sections)
    pages = frames(
        monkeypatch, lambda output: render_packet(document, output, format="pdf")
    )
    assert len(pages) >= 3
    assert_framed(pages, document.title)


def test_frame_escapes_glyphs_the_fonts_cannot_draw():
    """A title or detail outside the fonts' coverage becomes a visible escape."""
    frame = pdf_design.PdfFrame(title="Family 中", details=("Note \x1b",))
    assert frame.header_lines == ["Note \\u001b"]
    assert pdf_design.safe(frame.title) == "Family \\u4e2d"
    # Every escaped character is drawable in both body faces.
    assert all(ord(c) in pdf_design.pdf_glyphs() for c in "Family \\u4e2d")


def draw_record_page(canvas, page, *, frame):
    """``pdf_design.draw_record_page`` with the frame bound by keyword."""
    pdf_design.draw_record_page(canvas, page, frame)


class Recorder:
    """A stand-in canvas that records where a page body draws."""

    def __init__(self):
        """Start with nothing drawn."""
        self.extents = []

    def text(self, x, y, text, size, *args, **kwargs):
        """A text line spans about its size above and a quarter below its baseline."""
        if text:
            self.extents.append((y - size, y + size / 4))

    def rect(self, x, y, width, height, *args, **kwargs):
        """A rectangle spans its own height."""
        self.extents.append((y, y + height))

    def hline(self, x0, x1, y, color, width):
        """A rule spans its thickness."""
        self.extents.append((y - width / 2, y + width / 2))


def body_extents(frame, pages, draw):
    """Draw each page body on a recorder and return what it covered."""
    recorder = Recorder()
    for page in pages:
        draw(recorder, page)
    return recorder.extents


def assert_inside(frame, extents):
    """Nothing a body draws reaches the header or the footer."""
    assert extents
    assert min(top for top, _ in extents) >= frame.body_top - 1e-6
    assert max(bottom for _, bottom in extents) <= frame.body_bottom + 1e-6


def test_every_body_stays_between_the_header_and_the_footer():
    """Tables, record cards and address cards never draw past body_bottom."""
    phones = [
        dict(owner=f"Head {number}", kind="cell", value="202-555-0123")
        for number in range(80)
    ]
    rows = [family(family_duid=index) for index in range(30)]
    rows.append(family(family_duid=99, phones=phones))
    report = directory(rows, reach="neither", testing=True)
    frame = directory_frame(report)
    table = directory_table(report)
    pages = list(pdf_design.table_pages(table, table_rows(report), frame))
    assert_inside(
        frame, body_extents(frame, pages, pdf_design.draw_table(table, frame))
    )

    postal = directory(rows, postal=True)
    frame = directory_frame(postal)
    cards = [pdf_design.card_lines(block) for block in address_blocks(postal)]
    pages = list(pdf_design.card_pages(cards, frame))
    assert_inside(
        frame,
        body_extents(frame, pages, lambda c, p: pdf_design.draw_cards(c, p, frame)),
    )

    for document, records in (
        (information(count=5), information_records),
        (
            packet(
                [
                    dict(
                        duid=4,
                        name="Altar Servers",
                        chairs=["Q" * 9000],
                        rows=[member()] * 9,
                    )
                ]
            ),
            packet_records,
        ),
    ):
        frame = pdf_design.PdfFrame.for_document(document)
        pages = list(pdf_design.record_pages(records(document), frame))
        assert_inside(
            frame,
            body_extents(frame, pages, partial(draw_record_page, frame=frame)),
        )


def test_record_cards_move_whole_and_only_a_page_tall_card_splits():
    """A card never splits when it fits a page; a taller one repeats its title."""
    frame = pdf_design.PdfFrame(title="Cards")
    cards = [
        pdf_design.Record(
            title=f"Card {index}",
            fields=tuple((f"Field {line}", "value") for line in range(index % 9 + 1)),
        )
        for index in range(60)
    ]
    tall = pdf_design.Record(title="Tall", fields=(("Text", "line\n" * 120),))
    pages = list(pdf_design.record_pages([*cards, tall], frame))
    pieces = [piece for page in pages for piece in page]
    for card in cards:
        [piece] = [piece for piece in pieces if piece.record is card]
        assert len(piece.lines) == len(pdf_design.record_body(card))
    tall_titles = [piece.titles for piece in pieces if piece.record is tall]
    assert len(tall_titles) > 1
    assert tall_titles[0] == ("Tall",)
    assert all(titles == ("Tall (continued)",) for titles in tall_titles[1:])
    for page in pages:
        used = sum(pdf_design.piece_height(piece) for piece in page)
        used += pdf_design.CARD_SPACING * (len(page) - 1)
        assert used <= frame.body_height + 1e-6


def test_a_page_tall_title_is_cut_and_stays_above_the_footer():
    """A title taller than a page ends in an ellipsis; nothing passes body_bottom.

    Each piece keeps room for a body line, and a "(continued)" piece carries
    only the title's first line.
    """
    frame = pdf_design.PdfFrame(title="Cards", notice="Private.", note="Testing.")
    name = "Squyres " * 600
    tall = pdf_design.Record(title=name, fields=(("Text", "line\n" * 200),))
    pages = list(pdf_design.record_pages([tall], frame))
    pieces = [piece for page in pages for piece in page]
    assert len(pieces) > 1 and all(piece.lines for piece in pieces)
    assert pieces[0].titles[-1].endswith("…")
    assert all(
        len(piece.titles) == 1 and piece.titles[0].endswith("… (continued)")
        for piece in pieces[1:]
    )
    assert_inside(
        frame, body_extents(frame, pages, partial(draw_record_page, frame=frame))
    )


def test_eyebrow_escapes_after_uppercasing_and_wraps():
    """Uppercasing an escape would change it; a long eyebrow pushes the body down."""
    frame = pdf_design.PdfFrame(title="T", eyebrow="Parish \x1b")
    assert frame.eyebrow_lines == ["PARISH \\u001b"]
    short = pdf_design.PdfFrame(title="T", eyebrow="Parish")
    long = pdf_design.PdfFrame(title="T", eyebrow="Parish " * 60)
    assert len(long.eyebrow_lines) > 1
    assert long.body_top == short.body_top + pdf_design.EYEBROW_PITCH * (
        len(long.eyebrow_lines) - 1
    )


def test_filters_read_as_plain_words():
    """Search and sort are always stated; other filters only when they narrow."""
    assert (
        pdf_design.humanize_filters(
            '{"completed": "any", "disposition": "current_actionable", '
            '"search": "", "sort": "newest", "start": ""}'
        )
        == "Disposition: current actionable · Search: none · Sort: newest"
    )
    assert (
        pdf_design.humanize_filters('{"action": "join", "filters": {"search": "Lee"}}')
        == "Action: join · Search: Lee"
    )
    assert pdf_design.humanize_filters("not json") == "not json"
