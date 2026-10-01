"""Talents report shaping and exports (#247 follow-up)."""

import csv
import io
from datetime import UTC
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from parishkit.stewardship.reports.talents import (
    UNAVAILABLE,
    TalentQuery,
    shape_result,
    talents_csv,
    talents_xlsx,
)
from parishkit.stewardship.responses.service import DEFAULT_TALENTS

PAINTER, OTHER = DEFAULT_TALENTS[0].id, DEFAULT_TALENTS[-1].id
RETIRED = "11111111-1111-4111-8111-111111111111"
WHEN = "2026-10-05T14:00:00+00:00"


def result(talents, counts):
    """A raw SQL projection with one Member and one Family."""
    return {
        "members": [
            {
                "family_name": "Example",
                "family_duid": 10,
                "member_name": "Alex Example",
                "proposed": False,
                "cannot_serve": False,
                "talents": talents,
                "submitted_at": WHEN,
            }
        ],
        "families": [
            {"family_name": "Example", "family_duid": 10, "submitted_at": WHEN}
        ],
        "summary": {
            "members": 1,
            "cannot_serve": 0,
            "cannot_attend": 1,
            "talents": counts,
        },
    }


def test_removed_talents_are_worded_and_counted_as_unavailable():
    """A talent the parish removed still appears in rows and the summary."""
    shaped = shape_result(
        result({PAINTER: "", RETIRED: ""}, {PAINTER: 1, RETIRED: 1}),
        configuration={"modules": ["ministry"]},
    )
    assert shaped["members"][0]["talents"] == ["Painter", UNAVAILABLE]
    summary = dict(shaped["summary"]["talents"])
    assert summary["Painter"] == 1 and summary[UNAVAILABLE] == 1
    plain = shape_result(result({}, {}), configuration={"modules": ["ministry"]})
    assert UNAVAILABLE not in dict(plain["summary"]["talents"])


def test_exports_keep_formula_text_literal():
    """An Other text such as =HYPERLINK(...) is never a spreadsheet formula."""
    formula = '=HYPERLINK("https://example.org","x")'
    shaped = shape_result(
        result({OTHER: formula}, {OTHER: 1}), configuration={"modules": ["ministry"]}
    )
    rows = list(csv.reader(io.StringIO(talents_csv(shaped, UTC).decode())))
    cell = rows[1][3]
    # The cell starts with the talent label, so it cannot begin a formula.
    assert cell == f"Other: {formula}" and not cell.startswith(("=", "+", "-", "@"))
    book = load_workbook(io.BytesIO(talents_xlsx(shaped, ZoneInfo("UTC"))))
    value = book["Members"]["D2"]
    assert value.data_type == "s" and value.value == f"Other: {formula}"
    # Parish-typed names are neutralized too.
    shaped["members"][0]["member_name"] = "=cmd"
    rows = list(csv.reader(io.StringIO(talents_csv(shaped, UTC).decode())))
    assert rows[1][2] == "'=cmd"


def test_xlsx_escapes_control_characters_in_source_names():
    """A vertical tab pasted into a ParishSoft name cannot break the workbook.

    openpyxl rejects characters XML 1.0 forbids, so a raw write raised and the
    download failed. The shared writer escapes them reversibly instead.
    """
    raw = result({PAINTER: ""}, {PAINTER: 1})
    raw["members"][0]["member_name"] = "Alex\x0bExample"
    raw["families"][0]["family_name"] = "Example\x0b"
    shaped = shape_result(raw, configuration={"modules": ["ministry"]})
    book = load_workbook(io.BytesIO(talents_xlsx(shaped, ZoneInfo("UTC"))))
    member = book["Members"]["C2"]
    assert member.data_type == "s" and member.value == "Alex\\u000bExample"
    assert book["Families"]["A2"].value == "Example\\u000b"


def test_campaign_without_talents_reports_only_limitations():
    """An emptied talent list drops every talent column, count and filter.

    A Member listed only for a talent an earlier response chose is left out;
    a Member who cannot participate stays, with no talents shown.
    """
    emptied = {"modules": ["ministry"], "talent_options": []}
    raw = result({PAINTER: ""}, {PAINTER: 2})
    raw["members"].append(
        raw["members"][0] | {"member_name": "Sam Example", "cannot_serve": True}
    )
    raw["summary"] |= {"members": 2, "cannot_serve": 1}
    shaped = shape_result(raw, configuration=emptied)
    assert shaped["collects_talents"] is False
    assert [row["member_name"] for row in shaped["members"]] == ["Sam Example"]
    assert shaped["members"][0]["talents"] == []
    assert shaped["summary"]["talents"] == [] and shaped["summary"]["members"] == 1
    assert shaped["talent_choices"] == []
    rows = list(csv.reader(io.StringIO(talents_csv(shaped, UTC).decode())))
    assert rows[0] == [
        "Family",
        "Family DUID",
        "Member",
        "Cannot participate in ministries",
        "Latest response",
    ]
    assert rows[1][3] == "Yes"
    book = load_workbook(io.BytesIO(talents_xlsx(shaped, ZoneInfo("UTC"))))
    headings = [cell.value for cell in book["Members"][1]]
    assert "Talents" not in headings and len(headings) == 5
    # A campaign that does collect talents keeps its Talents column.
    kept = shape_result(
        result({PAINTER: ""}, {PAINTER: 1}), configuration={"modules": ["ministry"]}
    )
    assert kept["collects_talents"] is True
    rows = list(csv.reader(io.StringIO(talents_csv(kept, UTC).decode())))
    assert rows[0][3] == "Talents"


def test_report_page_says_plainly_that_no_talents_are_collected():
    """The page explains the empty list instead of showing talent columns."""
    from django.template.loader import render_to_string

    from parishkit.stewardship.reports.talent_views import tables

    def page(configuration, raw):
        """Render the report page from a shaped result, as the view does."""
        shaped = shape_result(raw, configuration=configuration)
        query = TalentQuery()
        members, families = tables(shaped, query, {}, "/report")
        return render_to_string(
            "stewardship/talents-report.html",
            shaped
            | {
                "metadata": {"name": "Renewal"},
                "members_table": members,
                "families_table": families,
                "campaign_id": "00000000-0000-4000-8000-000000000000",
                "query": query,
                "query_fields": query.form_values(),
                "export_timezones": ["UTC"],
            },
        )

    html = page({"modules": ["ministry"], "talent_options": []}, result({}, {}))
    assert "This campaign does not collect talents" in html
    assert "Talents</" not in html and "Members with talents" not in html
    assert "Painter" not in html and 'colspan="4"' in html
    html = page({"modules": ["ministry"]}, result({PAINTER: ""}, {PAINTER: 1}))
    assert "does not collect talents" not in html
    assert "Members with talents or limitations" in html and "Painter" in html
