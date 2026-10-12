"""The campaign archive workbook format and its strict reader (#1018).

Every workbook here is synthetic (``archive_factory``). Each refusal case
changes one thing in a consistent file and asserts the problem's code, and
that no problem text repeats a value, a found header or an unknown key.
"""

import re
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

import pytest

from parishkit.stewardship.archive import workbook
from parishkit.stewardship.archive.workbook import (
    BAND_EDGES,
    METADATA_KEYS,
    SHEETS,
    ArchiveRefused,
    Column,
    Kind,
    convert,
    flagged_notes,
    note_flagged,
    pledge_band,
    pledge_distribution,
    pledged_amounts,
    quartiles,
    read_archive,
    readme_rows,
)

from .archive_factory import (
    at,
    campaign,
    set_metadata,
    submission,
    workbook_bytes,
)

SPEC = (
    Path(__file__).resolve().parents[2]
    / "docs/specs/stewardship/campaign-history/spec.md"
)
SECRET = "Zq-secret-7731"
SECRET_AMOUNT = 98765.43


def refused(data):
    """The problems of a workbook the reader must refuse."""
    with pytest.raises(ArchiveRefused) as caught:
        read_archive(data)
    return caught.value


def codes(error):
    """The set of problem codes of a refusal."""
    return {problem.code for problem in error.problems}


def assert_nothing_echoed(error, *texts):
    """No problem repeats a Family key, an amount or any given text."""
    _, sheets = campaign()
    found = " ".join(problem.text for problem in error.problems)
    for family in sheets["families"]:
        assert family["family_key"] not in found
    for text in (SECRET, "98765", *texts):
        assert text not in found


# --- The format table against the specification --------------------------


def spec_table(heading):
    """The rows (lists of cells) of the first table under a spec heading."""
    lines = SPEC.read_text().splitlines()
    start = lines.index(heading)
    rows = []
    for line in lines[start + 1 :]:
        if line.startswith("#"):
            break
        if line.startswith("|"):
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
        elif rows:
            break
    return rows[2:]


@pytest.mark.parametrize(
    "sheet", ["families", "submissions", "daily", "ministry_responses"]
)
def test_the_column_tables_match_the_specification(sheet):
    """Names, order and types agree; every enumeration word is in its rule."""
    rows = spec_table(f"#### `{sheet}` sheet")
    assert [row[0].strip("`") for row in rows] == [c.name for c in SHEETS[sheet]]
    assert [row[1] for row in rows] == [c.type_text for c in SHEETS[sheet]]
    for row, column in zip(rows, SHEETS[sheet], strict=True):
        # "As in `families`" refers to that sheet's list instead.
        for choice in () if row[2].startswith("As in") else column.choices:
            assert f"`{choice}`" in row[2]


def test_the_metadata_keys_match_the_specification():
    """The required keys, in order, are the specification's table."""
    rows = spec_table("#### `metadata` sheet")
    assert [row[0].strip("`") for row in rows] == [k.name for k in METADATA_KEYS]
    columns = "Columns `key`, `value`, `note`"
    assert columns in SPEC.read_text()
    assert [c.name for c in SHEETS["metadata"]] == ["key", "value", "note"]


def test_the_band_edges_match_the_specification():
    """The band set's edges are the dollar amounts of the band table."""
    text = SPEC.read_text()
    table = text[text.index("| Band | Annual pledge |") :].split("\n\n")[0]
    amounts = re.findall(r"\$([0-9,]+)", table)
    edges = sorted({Decimal(a.replace(",", "")) for a in amounts})
    assert tuple(edges) == BAND_EDGES


def test_the_readme_dictionary_is_generated_from_the_table():
    """A written README holds the description, then every column and key."""
    from io import BytesIO

    import openpyxl

    book = openpyxl.load_workbook(BytesIO(workbook_bytes()), read_only=True)
    rows = [
        [value for value in row if value is not None]
        for row in book["README"].iter_rows(values_only=True)
    ]
    assert rows == [[c for c in row if c] for row in readme_rows()]
    flat = [row[0] for row in rows if row]
    for name, columns in SHEETS.items():
        assert f"Sheet '{name}'" in flat
        assert all(column.name in flat for column in columns)
    assert all(key.name in flat for key in METADATA_KEYS)
    assert "PRIVATE" in " ".join(rows[2])


# --- Reading a valid workbook ----------------------------------------------


def test_a_written_workbook_reads_back_to_the_same_values():
    """Round trip: instants, money to the cent, booleans, blank versus zero."""
    metadata, sheets = campaign()
    archive = read_archive(workbook_bytes(metadata, sheets))
    assert archive.metadata["timezone"] == "America/New_York"
    assert archive.metadata["is_final"] is True
    assert archive.metadata["as_of"] == at(6)
    assert archive.metadata["active_families"] == 40
    families = {f["family_key"]: f for f in archive.families}
    assert families["duid:1001"]["annual_pledge"] == Decimal("2400.50")
    assert families["duid:1001"]["first_submitted_at"] == at(1)
    assert families["duid:1002"]["annual_pledge"] == Decimal(0)
    assert families["duid:1002"]["pledge_status"] == "zero"
    assert families["duid:1003"]["annual_pledge"] is None
    assert families["duid:1003"]["cannot_give"] is True
    assert families["duid:1006"]["annual_pledge"] == Decimal(10000)
    assert families["duid:1007"]["pledge_status"] == "no_pledge_form"
    assert families["duid:1007"]["cannot_give"] is None
    assert archive.daily[-1]["cumulative_annual_pledged"] == Decimal("12660.75")
    assert archive.submission_rows == 10 and archive.unattributed_rows == 1
    assert archive.ministry_rows == 7
    assert archive.notes == {"note.01": "Synthetic data for tests"}
    assert archive.other_counts == {"count.active_now": "41"}
    assert archive.family_id_systems == ("parishsoft_duid",)


def test_whole_floats_and_numbers_in_text_columns_are_accepted():
    """A stored 12.0 is 12; a number in a text column reads as its digits."""
    metadata, sheets = campaign()
    sheets["families"][0]["submission_count"] = 2.0
    sheets["families"][0]["envelope_id"] = 4521
    sheets["ministry_responses"][0]["ministry_code"] = 101
    archive = read_archive(workbook_bytes(metadata, sheets))
    family = next(f for f in archive.families if f["family_key"] == "duid:1001")
    assert family["submission_count"] == 2 and family["envelope_id"] == "4521"


def test_binary_inexact_two_decimal_money_reads_exactly():
    """1234.56 has no exact double, yet reads as exactly 1234.56."""
    column = Column("annual_pledge", Kind.MONEY, "")
    assert convert(column, 1234.56) == Decimal("1234.56")
    assert convert(column, 0.1) == Decimal("0.10")
    assert convert(column, 17) == Decimal("17.00")


@pytest.mark.parametrize(
    ("kind", "value"),
    [
        (Kind.MONEY, 12.345),
        (Kind.MONEY, 0.1 + 0.2),
        (Kind.MONEY, "12.00"),
        (Kind.MONEY, True),
        (Kind.INTEGER, "12"),
        (Kind.INTEGER, 12.5),
        (Kind.INTEGER, False),
        (Kind.BOOLEAN, "TRUE"),
        (Kind.BOOLEAN, 1),
        (Kind.TEXT, 1.5),
        (Kind.INSTANT, "2054-10-04 14:00:00"),
        (Kind.INSTANT, "2054-10-04T14:00:00.5Z"),
        (Kind.INSTANT, "2054-02-30T14:00:00Z"),
        (Kind.DATE, "10/04/2054"),
    ],
)
def test_cells_of_the_wrong_type_are_refused(kind, value):
    """Digit text is not an integer, TRUE text not a boolean, and so on."""
    with pytest.raises(workbook._Bad):
        convert(Column("c", kind, ""), value)


def test_metadata_takes_boolean_text_or_cells_and_numeric_cells():
    """is_final as a boolean cell and campaign_year as a number both read."""
    metadata, sheets = campaign()
    metadata = set_metadata(metadata, "is_final", "FALSE")
    metadata = set_metadata(metadata, "campaign_year", 2055)

    def as_cells(book):
        sheet = book["metadata"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == "is_final":
                row[1].value = False
            if row[0].value == "campaign_year":
                row[1].value = 2055.0

    archive = read_archive(workbook_bytes(metadata, sheets, as_cells))
    assert archive.metadata["is_final"] is False
    assert archive.metadata["campaign_year"] == 2055


def test_a_later_minor_version_may_add_columns_keys_and_sheets():
    """1.1 extras are read past and never returned."""
    metadata, sheets = campaign()
    metadata = set_metadata(metadata, "schema_version", "1.1")
    metadata = set_metadata(metadata, "new_key_in_1_1", SECRET)

    def extras(book):
        families = book["families"]
        families.cell(1, len(SHEETS["families"]) + 1, "added_in_1_1")
        families.cell(2, len(SHEETS["families"]) + 1, SECRET)
        book.create_sheet("added_sheet").append([SECRET])

    archive = read_archive(workbook_bytes(metadata, sheets, extras))
    assert archive.minor_version == 1
    assert "new_key_in_1_1" not in archive.metadata
    assert SECRET not in repr(archive)


# --- Refusals ----------------------------------------------------------------


def _formula(book):
    book["daily"]["C2"] = "=1+1"


def _reorder(book):
    sheet = book["families"]
    sheet["B1"], sheet["C1"] = sheet["C1"].value, sheet["B1"].value


def _drop_column(book):
    book["submissions"].delete_cols(11)


def _unknown_column(book):
    book["families"].cell(1, len(SHEETS["families"]) + 1, SECRET)


def _error_cell(book):
    cell = book["daily"]["C3"]
    cell.value = "#N/A"
    cell.data_type = "e"


def _readme_second(book):
    book.move_sheet("README", offset=1)


def _row_after_blank(book):
    sheet = book["daily"]
    sheet.cell(sheet.max_row + 2, 1, 99)


def _unknown_sheet(book):
    book.create_sheet(SECRET)


def _outside_cell(book):
    book["daily"].cell(2, 20, SECRET)


def _daily_gap(book):
    book["daily"].delete_rows(3)


def _cell(sheet, column, row, value):
    """A workbook adjustment that sets one data cell by column name."""

    def adjust(book):
        index = [c.name for c in SHEETS[sheet]].index(column) + 1
        book[sheet].cell(row, index).value = value

    return adjust


def _family(key, column, value):
    """Change one Family's column in the derived sheets."""

    def change(metadata, sheets):
        next(f for f in sheets["families"] if f["family_key"] == key)[column] = value
        return metadata, sheets

    return change


def _metadata(key, value):
    """Change one metadata key's value."""

    def change(metadata, sheets):
        return set_metadata(metadata, key, value), sheets

    return change


def _daily_value(day, column, value):
    """Change one daily row's column."""

    def change(metadata, sheets):
        sheets["daily"][day - 1][column] = value
        return metadata, sheets

    return change


def _duplicate_family(metadata, sheets):
    sheets["families"].append(dict(sheets["families"][0]))
    return metadata, sheets


def _drop_key(metadata, sheets):
    return [row for row in metadata if row[0] != "active_families"], sheets


def _long_notes(metadata, sheets):
    """Nine 30,000-character notes: each fits a cell, together past 256 KB."""
    for number in range(2, 11):
        metadata = set_metadata(metadata, f"note.{number:02}", "x" * 30_000)
    return metadata, sheets


def _many_problems(metadata, sheets):
    for row in sheets["ministry_responses"] * 10:
        sheets["ministry_responses"].append({**row, "member_slot": 0})
    return metadata, sheets


CHANGES = {
    "formula": (None, _formula, "formula"),
    "reordered column": (None, _reorder, "header"),
    "missing column": (None, _drop_column, "header"),
    "unknown column": (None, _unknown_column, "column_not_defined"),
    "README not first": (None, _readme_second, "sheet_order"),
    "row after blank": (None, _row_after_blank, "row_after_blank"),
    "unknown sheet": (None, _unknown_sheet, "sheet_not_defined"),
    "cell outside columns": (None, _outside_cell, "cell_outside_columns"),
    "daily gap": (None, _daily_gap, "daily"),
    "digit text integer": (
        None,
        _cell("families", "submission_count", 2, "1"),
        "type",
    ),
    "money with three decimals": (
        None,
        _cell("families", "prior_year_pledge", 2, 10.125),
        "type",
    ),
    "TRUE text in a data sheet": (
        None,
        _cell("submissions", "attributed", 2, "TRUE"),
        "type",
    ),
    "Excel date cell": (
        None,
        _cell("daily", "local_date", 2, at(1).replace(tzinfo=None)),
        "type",
    ),
    "unknown enumeration": (
        None,
        _cell("families", "frequency", 2, SECRET),
        "choice",
    ),
    "required blank": (None, _cell("submissions", "form", 3, None), "blank"),
    "error cell": (None, _error_cell, "error_value"),
    "negative installment": (
        _family("duid:1005", "installment_amount", Decimal("-1")),
        None,
        "range",
    ),
    "negative prior-year pledge": (
        _family("duid:1005", "prior_year_pledge", Decimal("-5")),
        None,
        "range",
    ),
    "negative submission pledge": (
        None,
        _cell("submissions", "annual_pledge", 2, -12.5),
        "range",
    ),
    "unknown metadata key": (_metadata("favourite_" + SECRET, "1"), None, None),
    "estimate key": (_metadata("active_families_estimate", 400), None, None),
    "blank active_families": (_metadata("active_families", None), None, "blank"),
    "missing active_families": (_drop_key, None, "metadata_key_missing"),
    "blank percentage": (
        _daily_value(2, "cumulative_pct_of_active", None),
        None,
        "blank",
    ),
    "wrong major version": (_metadata("schema_version", "2.0"), None, "version"),
    "wrong schema": (_metadata("schema", SECRET), None, "schema"),
    "bad time zone": (_metadata("timezone", "Mars/" + SECRET), None, "type"),
    "year out of range": (_metadata("campaign_year", 1999), None, "range"),
    "wrong count": (_metadata("count.families", 99), None, "count"),
    "as_of too early": (_metadata("as_of", at(1)), None, "metadata"),
    "pledged without amount": (
        _family("duid:1005", "annual_pledge", None),
        None,
        "rule",
    ),
    "installment without source": (
        _family("duid:1005", "installment_source", None),
        None,
        "rule",
    ),
    "key and id disagree": (_family("duid:1005", "family_id", 4), None, "rule"),
    "financial columns disagree": (
        _family("duid:1001", "frequency", "weekly"),
        None,
        "family",
    ),
    "first_submissions disagree": (
        _family("duid:1005", "first_campaign_day", 3),
        None,
        "day",
    ),
    "middle day pledges disagree": (
        _daily_value(3, "cumulative_annual_pledged", Decimal("1300")),
        None,
        "daily",
    ),
    "duplicate Family": (_duplicate_family, None, "family_key"),
    "problem cap": (_many_problems, None, "range"),
    # The stored columns' limits (migration 0059): an accepted file never
    # fails on save.
    "long campaign name": (_metadata("campaign_name", "n" * 201), None, "length"),
    "active_families past an integer": (
        _metadata("active_families", 2**31),
        None,
        "range",
    ),
    "long provenance": (
        _metadata("active_families_provenance", "p" * 10_001),
        None,
        "length",
    ),
    "long source system": (_metadata("source_system", "s" * 10_001), None, "length"),
    "long dedup rule": (_metadata("dedup_rule", "d" * 10_001), None, "length"),
    "notes past 256 KB": (_long_notes, None, "notes_size"),
}


@pytest.mark.parametrize("case", CHANGES)
def test_each_rule_refuses_without_echoing_values(case):
    """One change, one expected code; the text names cells, never values."""
    change, adjust, code = CHANGES[case]
    metadata, sheets = campaign()
    if change is not None:
        metadata, sheets = change(metadata, sheets)
    error = refused(workbook_bytes(metadata, sheets, adjust))
    expected = code or "metadata_key_not_defined"
    assert expected in codes(error), [p.text for p in error.problems]
    assert_nothing_echoed(error, "favourite_", "Mars/", "estimate")
    if code is None:
        assert any(
            re.fullmatch(r"metadata!A\d+: key not defined in version 1\.0", p.text)
            for p in error.problems
        )


def test_values_at_the_stored_limits_are_accepted():
    """The limits are inclusive: each value at its column's limit reads."""
    metadata, sheets = campaign()
    for key, value in (
        ("campaign_name", "n" * 200),
        ("active_families_provenance", "p" * 10_000),
        ("source_system", "s" * 10_000),
        ("dedup_rule", "d" * 10_000),
    ):
        metadata = set_metadata(metadata, key, value)
    for number in range(2, 10):
        metadata = set_metadata(metadata, f"note.{number:02}", "x" * 30_000)
    archive = read_archive(workbook_bytes(metadata, sheets))
    assert len(archive.metadata["campaign_name"]) == 200
    stored = {**archive.notes, **archive.other_counts}
    assert workbook.stored_notes_bytes(stored) <= workbook.STORED_NOTES_BYTES


def test_a_pledged_total_past_the_stored_amount_is_refused(monkeypatch):
    """The total is checked against numeric(18, 2), here lowered to reach it."""
    monkeypatch.setattr(workbook, "STORED_MONEY_LIMIT", Decimal(1000))
    error = refused(workbook_bytes())
    assert "families: the pledged total is out of range" in [
        p.text for p in error.problems
    ]


def test_problems_stop_at_fifty_and_count_the_rest():
    """Seventy bad cells give fifty problems and a count of twenty more."""
    error = refused(workbook_bytes(*_many_problems(*campaign())))
    assert len(error.problems) == 50 and error.more == 20


def test_a_wrong_header_names_the_expected_column_only():
    """The problem says which column belongs there, not what was found."""
    error = refused(workbook_bytes(adjust=_reorder))
    assert [p.text for p in error.problems] == [
        "families!B1: expected column family_id"
    ]


def test_a_two_form_family_takes_its_financial_answer_from_the_pledge_form():
    """2020 style: a later ministry form does not replace the pledge."""
    rows = [
        submission("pds_fid:5", "pledge", at(1), "300", "annual"),
        submission("pds_fid:5", "ministry", at(2)),
    ]
    metadata, sheets = campaign(rows)
    for family in sheets["families"]:
        family["family_id_system"] = "pds_fid"
    archive = read_archive(workbook_bytes(metadata, sheets))
    [family] = archive.families
    assert family["annual_pledge"] == Decimal(300)
    assert family["last_submitted_at"] == at(2)
    sheets["families"][0]["annual_pledge"] = None
    sheets["families"][0]["pledge_status"] = "no_pledge_form"
    sheets["families"][0]["frequency"] = None
    sheets["families"][0]["cannot_give"] = None
    assert "family" in codes(refused(workbook_bytes(metadata, sheets)))


def test_a_tie_in_submitted_at_is_settled_by_row_id():
    """Claiming the earlier row of a same-second pair as last is refused."""
    metadata, sheets = campaign()
    family = next(f for f in sheets["families"] if f["family_key"] == "duid:1006")
    family["annual_pledge"] = Decimal(500)
    error = refused(workbook_bytes(metadata, sheets))
    assert "family" in codes(error)


def test_a_local_date_that_disagrees_with_its_instant_is_refused():
    """The submission's local date must be its instant's date in the zone."""
    metadata, sheets = campaign()
    sheets["submissions"][0]["local_date"] += timedelta(days=1)
    assert "date" in codes(refused(workbook_bytes(metadata, sheets)))


def test_totals_that_disagree_with_families_are_refused():
    """The last day must equal the families sheet's pledged total."""
    metadata, sheets = campaign()
    family = next(f for f in sheets["families"] if f["family_key"] == "duid:1005")
    family["annual_pledge"] = Decimal("261.25")
    error = refused(workbook_bytes(metadata, sheets))
    assert any("families sheet's total" in p.rule for p in error.problems)


def test_the_row_limit_refuses_a_large_file(monkeypatch):
    """Past MAX_ROWS rows in all, the file is too large."""
    monkeypatch.setattr(workbook, "MAX_ROWS", 30)
    assert codes(refused(workbook_bytes())) == {"too_many_rows"}


def test_an_unreadable_file_is_refused():
    """Bytes that are not a workbook give one fixed problem."""
    error = refused(b"PK\x03\x04" + b"\x00" * 100)
    assert codes(error) == {"unreadable"}


# --- Pledge distribution, quartiles and notes ------------------------------


def test_a_pledge_on_an_edge_falls_in_the_higher_band():
    """Bands include their lower edge and exclude their upper one."""
    assert pledge_band(Decimal("249.99")) == 1
    assert pledge_band(Decimal(250)) == 2
    assert pledge_band(Decimal("9999.99")) == 7
    assert pledge_band(Decimal(10000)) == 8


def test_the_distribution_counts_and_sums_pledging_families_only():
    """Zero pledges are left out; each band has its count and total."""
    archive = read_archive(workbook_bytes())
    pledges = pledged_amounts(archive)
    assert sorted(pledges) == [Decimal("260.25"), Decimal("2400.50"), Decimal(10000)]
    bands = {b["band"]: b for b in pledge_distribution(pledges)}
    assert bands[2] == {"band": 2, "families": 1, "total": Decimal("260.25")}
    assert bands[5]["families"] == 1 and bands[8]["total"] == Decimal(10000)
    assert sum(b["families"] for b in bands.values()) == 3


def test_quartiles_follow_quartile_inc_rounded_half_up():
    """Hand-computed QUARTILE.INC values in exact decimals."""
    values = [Decimal(v) for v in ("100", "200", "300", "400")]
    # h = 1.75, 2.5, 3.25: 175, 250, 325.
    assert quartiles(values) == (
        Decimal("175.00"),
        Decimal("250.00"),
        Decimal("325.00"),
    )
    odd = [Decimal("0.01"), Decimal("0.02")]
    # h = 1.25 gives 0.0125, which rounds half up to 0.01; h = 1.5 gives 0.015.
    assert quartiles(odd) == (Decimal("0.01"), Decimal("0.02"), Decimal("0.02"))
    assert quartiles([Decimal(5)]) == (Decimal("5.00"),) * 3
    assert quartiles([]) is None


def test_notes_with_amounts_digits_or_length_are_flagged():
    """A dollar amount, five digits in a row, or over 500 characters."""
    assert note_flagged("Pledges of $1,200 were rounded")
    assert note_flagged("Family 12345 merged")
    assert note_flagged("x" * 501)
    assert not note_flagged("Converted from the 2026 export")
    metadata, sheets = campaign()
    metadata = set_metadata(metadata, "note.02", "Envelope 55555 skipped")
    archive = read_archive(workbook_bytes(metadata, sheets))
    assert flagged_notes(archive) == ["note.02"]


def test_negative_prior_year_contributions_are_accepted():
    """Old-system contributions can be net negative after a refund."""
    metadata, sheets = campaign()
    sheets["families"][0]["prior_year_contributions"] = Decimal("-20.00")
    archive = read_archive(workbook_bytes(metadata, sheets))
    assert archive.families[0]["prior_year_contributions"] == Decimal("-20.00")
