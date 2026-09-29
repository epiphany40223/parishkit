"""The parish date format: every style, exports, templates and the JS twin (#221)."""

import io
import re
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from django.template import Context, Template
from openpyxl import Workbook, load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS

from parishkit.config import ConfigError
from parishkit.stewardship.accounts.parish_views import ParishForm
from parishkit.stewardship.accounts.templatetags.stewardship import parish_time
from parishkit.stewardship.jobs.campaign_mail_values import campaign_values
from parishkit.stewardship.reports.information_rendering import xlsx_cell
from parishkit.stewardship.responses.page_content import public_values
from parishkit.stewardship.web import dates
from parishkit.stewardship.web.exports import csv_cell

from .campaign_factory import campaign
from .configuration_factory import configuration_document, configuration_version

PACKAGE = Path(dates.__file__).parents[1]
DAY = date(2027, 1, 5)
# 20:04 UTC is 3:04 PM in New York (EST) and still January 5.
MOMENT = datetime(2027, 1, 5, 20, 4, tzinfo=UTC)
ZONE = "America/New_York"
# style: (full date, compact date, full timestamp, compact timestamp)
EXPECTED = {
    "us_long": (
        "January 5, 2027",
        "Jan 5, 2027",
        "January 5, 2027 at 3:04 PM EST",
        "Jan 5, 2027 3:04 PM",
    ),
    "us_medium": (
        "Jan 5, 2027",
        "Jan 5, 2027",
        "Jan 5, 2027 at 3:04 PM EST",
        "Jan 5, 2027 3:04 PM",
    ),
    "us_numeric": (
        "01/05/2027",
        "01/05/27",
        "01/05/2027 3:04 PM EST",
        "01/05/27 3:04 PM",
    ),
    "eu_long": (
        "5 January 2027",
        "5 Jan 2027",
        "5 January 2027, 15:04 EST",
        "5 Jan 2027 15:04",
    ),
    "eu_medium": (
        "5 Jan 2027",
        "5 Jan 2027",
        "5 Jan 2027, 15:04 EST",
        "5 Jan 2027 15:04",
    ),
    "eu_numeric": (
        "05/01/2027",
        "05/01/27",
        "05/01/2027 15:04 EST",
        "05/01/27 15:04",
    ),
    "eu_dot": (
        "05.01.2027",
        "05.01.27",
        "05.01.2027 15:04 EST",
        "05.01.27 15:04",
    ),
    "iso": (
        "2027-01-05",
        "2027-01-05",
        "2027-01-05 15:04 EST",
        "2027-01-05 15:04",
    ),
}


def test_every_style_is_tested_and_offered():
    """The Admin choice list, the table and this test cover the same styles."""
    assert set(EXPECTED) == set(dates.FORMATS) == {code for code, _ in dates.CHOICES}
    assert dates.DEFAULT == "us_long"


@pytest.mark.parametrize("style", sorted(EXPECTED))
def test_each_style_formats_dates_and_instants(style):
    """Full and compact dates and timestamps, with the paired clock."""
    full, compact, stamp, compact_stamp = EXPECTED[style]
    assert dates.format_date(DAY, style) == full
    assert dates.format_date(DAY, style, compact=True) == compact
    assert dates.format_instant(MOMENT, ZONE, style) == stamp
    assert dates.format_instant(MOMENT, ZONE, style, compact=True) == compact_stamp
    local = MOMENT.astimezone(ZoneInfo(ZONE))
    assert dates.format_local(local, style) == stamp
    with dates.using(style):
        assert dates.display_text(local, compact=True) == compact_stamp
        assert dates.display_text(DAY) == full


def test_twelve_hour_edges_and_leading_zeros():
    """Midnight and noon read as 12 AM/PM; 24-hour clocks keep two digits."""
    midnight = datetime(2027, 3, 9, 5, 7, tzinfo=UTC)
    assert dates.format_instant(midnight, ZONE, "us_long").endswith("12:07 AM EST")
    assert dates.format_instant(midnight, ZONE, "iso") == "2027-03-09 00:07 EST"
    noon = datetime(2027, 3, 9, 17, 0, tzinfo=UTC)
    assert dates.format_instant(noon, ZONE, "us_medium").endswith("12:00 PM EST")


def test_default_request_style_and_explicit_pinning():
    """No resolver means the default; a lent resolver or block overrides it."""
    assert dates.current() == "us_long"
    token = dates.use(lambda: "iso")
    try:
        assert dates.current() == "iso"
        assert dates.format_date(DAY) == "2027-01-05"
        with dates.using("eu_dot"):
            assert dates.format_date(DAY) == "05.01.2027"
        assert dates.current() == "iso"
    finally:
        dates.reset(token)
    # Unknown or unset codes (older snapshots) fall back to the default.
    with dates.using(None):
        assert dates.current() == "us_long"
    assert dates.normalized("fr_FR") == "us_long"


def test_rejects_untyped_values():
    """Datetimes are not calendar dates, and naive instants are refused."""
    with pytest.raises(ValueError):
        dates.format_date(MOMENT, "iso")
    with pytest.raises(ValueError):
        dates.format_instant(datetime(2027, 1, 5, 12), ZONE, "iso")


def test_csv_is_iso_whatever_the_parish_style():
    """CSV dates are YYYY-MM-DD; timestamps are ISO 8601 with their zone offset."""
    local = MOMENT.astimezone(ZoneInfo(ZONE))
    with dates.using("us_long"):
        assert csv_cell(DAY) == "2027-01-05"
        assert csv_cell(local) == "2027-01-05 15:04:00-05:00"
        assert csv_cell(dates.Span(DAY, date(2027, 6, 30))) == (
            "2027-01-05 through 2027-06-30"
        )
        assert csv_cell("=formula") == "'=formula"


def test_xlsx_dates_are_native_cells_with_builtin_formats():
    """Excel's built-in short date (14) and date-time (22), shown per viewer locale."""
    book = Workbook()
    sheet = book.active
    xlsx_cell(sheet, 1, 1, DAY)
    xlsx_cell(sheet, 1, 2, MOMENT.astimezone(ZoneInfo(ZONE)))
    xlsx_cell(sheet, 1, 3, "2027-01-05")
    stream = io.BytesIO()
    book.save(stream)
    loaded = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        day, stamp, text = loaded.active[1]
        assert day.is_date and day.value == datetime(2027, 1, 5)
        assert day.number_format == BUILTIN_FORMATS[14]
        # Converted to the stated display zone before the zone is dropped.
        assert stamp.is_date and stamp.value == datetime(2027, 1, 5, 15, 4)
        assert stamp.number_format == BUILTIN_FORMATS[22]
        assert text.data_type == "s" and text.value == "2027-01-05"
    finally:
        loaded.close()
    styles = zipfile.ZipFile(stream).read("xl/styles.xml").decode()
    # Built-in formats are referenced by id, not written as custom patterns.
    assert 'numFmtId="14"' in styles and 'numFmtId="22"' in styles
    assert "<numFmt " not in styles


def test_pdf_text_follows_the_chosen_style():
    """PDFs are read by people, so they use the parish style."""
    local = MOMENT.astimezone(ZoneInfo(ZONE))
    with dates.using("eu_long"):
        assert dates.display_text(local) == "5 January 2027, 15:04 EST"
        assert dates.display_text(dates.Span(DAY, date(2027, 6, 30))) == (
            "5 January 2027 through 30 June 2027"
        )
        assert dates.display_text("text") == "text"


def test_template_filter_uses_the_request_style_and_iso_text():
    """parish_date accepts dates and ISO text; unparseable text is shown as-is."""
    template = Template(
        "{% load stewardship %}{{ day|parish_date }}|{{ text|parish_date }}|"
        "{{ bad|parish_date }}|{{ missing|parish_date }}"
    )
    context = Context({"day": DAY, "text": "2027-06-30", "bad": "soon"})
    with dates.using("eu_dot"):
        assert template.render(context) == "05.01.2027|30.06.2027|soon|"


def test_email_and_page_placeholders_use_the_parish_style():
    """Worker email values and page values read the parish's own setting."""
    parish = {
        "name": "P",
        "website": "https://parish.example.org/",
        "phone": "+12125551234",
        "date_format": "iso",
    }
    values = campaign()["values"]
    for rendered in (
        campaign_values(parish=parish, campaign=values),
        public_values(parish, values),
    ):
        assert rendered["campaign_start"] == values["start_date"]
        assert rendered["campaign_end"] == values["end_date"]
    expected = dates.format_date(date.fromisoformat(values["start_date"]), "us_long")
    unset = parish | {"date_format": None}
    assert campaign_values(parish=unset, campaign=values)["campaign_start"] == expected


@pytest.mark.parametrize("value", ["eu_dot", "iso"])
def test_schema_accepts_known_styles(value):
    """The key is optional; when present it names a known style."""
    document = configuration_document()
    document["sections"]["parish"][0]["values"]["date_format"] = value
    configuration_version(document)


@pytest.mark.parametrize("value", [None, "", "dd/mm/yyyy", 7])
def test_schema_rejects_unknown_styles(value):
    """An unset style is spelled by omitting the key."""
    document = configuration_document()
    document["sections"]["parish"][0]["values"]["date_format"] = value
    with pytest.raises(ConfigError):
        configuration_version(document)


def test_parish_form_offers_examples_and_defaults_a_blank_post():
    """Admins pick by example; an older post without the field keeps the default."""
    base = {
        "name": "Sample Parish",
        "website": "https://parish.example.org/",
        "timezone": "America/New_York",
        "phone": "+12125551234",
        "base_digest": "a" * 64,
    }
    form = ParishForm(base)
    assert form.is_valid(), form.errors
    assert form.cleaned_data["date_format"] == "us_long"
    assert ParishForm(base | {"date_format": "eu_long"}).is_valid()
    assert not ParishForm(base | {"date_format": "dd/mm/yyyy"}).is_valid()
    assert "1 January 2027" in str(form["date_format"])


def test_templates_never_format_dates_themselves():
    """Guard: dates in templates go through parish_date or data-local-instant.

    ``|date:'c'`` (a machine-readable ISO value for a <time> element) is the
    only allowed use of Django's own date formatting.
    """
    offenders = []
    for path in sorted(PACKAGE.rglob("templates/**/*.html")):
        text = path.read_text()
        for match in re.finditer(r"\|(date|time)\b(?::(['\"])(.*?)\2)?", text):
            if match.group(1) != "date" or match.group(3) != "c":
                offenders.append(f"{path.name}: {match.group(0)}")
        offenders += [f"{path.name}: now tag" for _ in re.finditer(r"{%\s*now\b", text)]
    assert offenders == []


def test_scripts_format_dates_only_through_the_shared_formatter():
    """Guard: only date-format-v1.js builds date text from a Date."""
    pattern = re.compile(
        r"dateStyle|timeStyle|toLocaleDateString|toLocaleTimeString|"
        r"toLocaleString\(\s*\)|month:\s*\"(?:short|long|numeric)\""
    )
    static = PACKAGE / "accounts/static/stewardship"
    offenders = [
        path.name
        for path in sorted(static.glob("*.js"))
        if path.name != "date-format-v1.js" and pattern.search(path.read_text())
    ]
    assert offenders == []


def test_script_style_table_matches_python():
    """date-format-v1.js mirrors dates.py's patterns and code list exactly."""
    script = (PACKAGE / "accounts/static/stewardship/date-format-v1.js").read_text()
    table = dict(
        (code, (full, compact))
        for code, full, compact in re.findall(
            r'^\s+(\w+): \["([^"]+)", "([^"]+)"\]', script, re.MULTILINE
        )
    )
    assert table == {code: style[1:] for code, style in dates._STYLES.items()}


def test_the_repeated_fall_back_hour_stays_distinct():
    """01:30 EDT and 01:30 EST on 2026-11-01 differ in the zone name and in CSV."""
    first = datetime(2026, 11, 1, 5, 30, tzinfo=UTC).astimezone(ZoneInfo(ZONE))
    second = datetime(2026, 11, 1, 6, 30, tzinfo=UTC).astimezone(ZoneInfo(ZONE))
    assert dates.format_instant(first, ZONE, "us_long") == (
        "November 1, 2026 at 1:30 AM EDT"
    )
    assert dates.format_instant(second, ZONE, "us_long") == (
        "November 1, 2026 at 1:30 AM EST"
    )
    assert csv_cell(first) == "2026-11-01 01:30:00-04:00"
    assert csv_cell(second) == "2026-11-01 01:30:00-05:00"


def test_an_instant_after_utc_midnight_can_be_the_previous_local_day():
    """03:00 UTC on January 6 is still the evening of January 5 in New York."""
    moment = datetime(2027, 1, 6, 3, 0, tzinfo=UTC)
    assert dates.format_instant(moment, ZONE, "iso") == "2027-01-05 22:00 EST"
    assert dates.format_instant(moment, ZONE, "us_numeric", compact=True) == (
        "01/05/27 10:00 PM"
    )


@pytest.mark.parametrize(
    ("value", "style", "expected"),
    [
        ("09:00:00", "us_long", "9:00 AM"),
        ("21:05:00", "us_long", "9:05 PM"),
        ("21:05:00", "eu_long", "21:05"),
        ("09:00:30", "us_long", "9:00:30 AM"),
        ("", "us_long", ""),
        ("not a time", "us_long", "not a time"),
    ],
)
def test_parish_time_uses_the_style_clock(value, style, expected):
    """Schedule times follow the parish style's 12- or 24-hour clock."""
    with dates.using(style):
        assert parish_time(value) == expected
