"""``export census``: the Census changes download (#463, ADM-11 PR 8h).

Process admission is replaced by the test's own Django assembly; everything
after it is the real command line on the restricted web login, reading the
list through the page's download admission, filters and ``census_changes``
under the campaign read guard and writing it through the page's own writer.
Its CSV is the page's download byte for byte, and it records the page's
``census_changes_exported`` with the same outcomes and context members.
"""

import csv
import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.accounts import sessions
from parishkit.stewardship.audit.models import AuditEvent
from parishkit.stewardship.reports import (
    census_change_views,
    census_changes,
    export_services,
)
from parishkit.stewardship.responses.models import ProposedChange

from .auth_builders import signed_in
from .automation_builders import as_staff
from .test_admin_census_cli_postgresql import submitted, with_roles
from .test_admin_digest_list_cli_postgresql import revoke
from .test_admin_export_cli_postgresql import read_only_session
from .test_admin_logs_cli_postgresql import counted, stream
from .test_admin_report_cli_postgresql import cli, parity  # noqa: F401
from .test_admin_status_cli_postgresql import admin, session  # noqa: F401
from .test_information_followup_postgresql import search

pytestmark = pytest.mark.django_db(transaction=True)

ROUTE = "/admin/reports/census/exports/"
EVENT = "census_changes_exported"


def export(run, *argv, secret=None):
    """``export census`` with ``argv``: (exit code, file bytes, document)."""
    return stream(None, "export", "census", *argv, secret=secret or run.secret)


def workbook(body):
    """An XLSX file's sheets (rows of values, by title) and its properties."""
    book = load_workbook(io.BytesIO(body))
    sheets = {
        sheet.title: [list(values) for values in sheet.iter_rows(values_only=True)]
        for sheet in book.worksheets
    }
    return sheets, book.properties


def created_at_download(properties, information, zone):
    """Whether Created is the sheet's Downloaded at (stored without a zone)."""
    downloaded = dict(information)["Downloaded at"]
    assert isinstance(downloaded, datetime)
    instant = downloaded.replace(tzinfo=ZoneInfo(zone)).astimezone(UTC)
    # Excel stores the cell to the millisecond; the property to the second.
    return abs(instant.replace(tzinfo=None) - properties.created) < timedelta(seconds=1)


def lines(body):
    """A JSON Lines file, one record per line."""
    return [json.loads(line) for line in body.decode().splitlines()]


def test_the_page_s_files_and_json_lines(response_service, google, cli):  # noqa: F811
    """CSV is the page's download; XLSX holds the same table and is the
    page's, cell for cell but for its download time; PDF is a PDF;
    JSON Lines adds each change's id, version and actions. Each records the
    page's event with the rows in the file."""
    harness = submitted(response_service)
    browser, _ = signed_in()
    run = cli()
    form = {"format": "csv", "timezone": "America/New_York", "status": "all"}
    argv = ("--status", "all", "--timezone", "America/New_York")
    found = {}

    def page():
        """The page's download, as the browser posts it (in memory, so not
        on the download pool)."""
        response, found["page"] = search(browser, ROUTE, form)
        assert response.status_code == 200

    def command():
        """The command's CSV of the same list."""
        found["command"] = export(run, *argv, "--format", "csv")

    parity(EVENT, page, command)
    code, body, document = found["command"]
    assert code == 0, document
    assert body == found["page"]
    table = list(csv.reader(io.StringIO(body.decode())))
    assert len(table) == 3 and b"Requested" in body
    result = document["result"]
    assert result["campaign_id"] == str(harness.campaign.pk)
    assert result["format"] == "csv" and result["content_type"] == "text/csv"
    assert result["timezone"] == "America/New_York" and result["count"] == 2
    assert result["size"] == len(body)
    assert result["sha256"] == hashlib.sha256(body).hexdigest()
    assert result["file_name"].startswith("stewardship-census-changes-")
    assert result["file_name"].endswith(".csv")
    # The page counted its two rows, and so did the command.
    assert counted(EVENT)[-2:] == [2, 2]

    code, body, document = export(run, *argv, "--format", "xlsx")
    assert code == 0, document
    command_book, command_properties = workbook(body)
    sheet = [
        ["" if value is None else value for value in values]
        for values in command_book["Census changes"]
    ]
    assert sheet == table
    # The page's XLSX is the command's, sheet for sheet, cell for cell and
    # property for property, but for the moment each was downloaded (#926).
    response, page_body = search(browser, ROUTE, form | {"format": "xlsx"})
    assert response.status_code == 200
    page_book, page_properties = workbook(page_body)
    sheets = ["Census changes", "Report information"]
    assert list(page_book) == list(command_book) == sheets
    assert page_book["Census changes"] == command_book["Census changes"]
    assert page_properties.title == command_properties.title == "Census changes"
    assert page_properties.creator == command_properties.creator
    page_information = page_book["Report information"]
    command_information = command_book["Report information"]
    for properties, information in (
        (page_properties, page_information),
        (command_properties, command_information),
    ):
        assert created_at_download(properties, information, "America/New_York")
    # Row by row in order; only the Downloaded at row (the fourth) differs,
    # and the command downloaded first.
    assert [row[0] for row in page_information] == [
        row[0] for row in command_information
    ]
    downloaded = 3
    assert command_information[downloaded][0] == "Downloaded at"
    assert command_information[downloaded][1] <= page_information[downloaded][1]
    del page_information[downloaded], command_information[downloaded]
    assert (
        page_information
        == command_information
        == [
            ["Report", "Census changes"],
            ["Parish", "Example Parish"],
            ["Campaign", "Annual campaign"],
            ["Display time zone", "America/New_York"],
            ["Changes in this file", "2"],
            ["Privacy", census_changes.PRIVACY],
        ]
    )
    code, body, document = export(run, *argv, "--format", "pdf")
    assert code == 0 and body.startswith(b"%PDF-"), document
    assert document["result"]["content_type"] == "application/pdf"

    code, body, document = export(run, *argv, "--format", "jsonl")
    assert code == 0, document
    assert document["result"]["count"] == 2
    records = {record["what"]: record for record in lines(body)}
    proposals = {
        str(proposal.pk): proposal.version for proposal in ProposedChange.objects.all()
    }
    for record in records.values():
        assert proposals[record["proposal_id"]] == record["version"]
    # The By-hand address may be entered or ignored; the automatic rename
    # is publication's, so it offers nothing.
    assert records["Home address"]["actions"] == ["entered", "ignored"]
    assert records["First name"]["actions"] == []
    assert records["First name"]["family_answer"] == "Requested"
    # The cells are the CSV's.
    assert sorted(
        [record[key] for key in ("family", "family_duid", "who", "what")]
        for record in records.values()
    ) == sorted([cells[0], cells[1], cells[2], cells[4]] for cells in table[1:])

    # The filters narrow the file as they do the page's. The command
    # records the page's event and nothing else: no admin_cmd_* event.
    before = set(AuditEvent.objects.values_list("pk", flat=True))
    code, body, document = export(run, "--route", "by_hand", "--format", "jsonl")
    assert code == 0 and document["result"]["count"] == 1, document
    assert [record["what"] for record in lines(body)] == ["Home address"]
    assert counted(EVENT)[-1] == 1
    added = AuditEvent.objects.exclude(pk__in=before)
    assert list(added.values_list("event_type", flat=True)) == [EVENT]

    # While other campaign work holds the campaign, no outcome is offered.
    real = export_services.admit_campaign

    def held(campaign_id, *, mutating=False):
        """The page's check, refusing a change as held work does; reads pass."""
        if mutating:
            raise PermissionError("held")
        return real(campaign_id, mutating=False)

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(export_services, "admit_campaign", held)
        code, body, document = export(run, "--status", "all", "--format", "jsonl")
    assert code == 0, document
    assert [record["actions"] for record in lines(body)] == [[], []]


def test_refusals_print_nothing_and_record_nothing(response_service, google, cli):  # noqa: F811
    """A choice the page refuses is invalid; a read-only session is refused;
    neither writes a byte or records the page's event."""
    submitted(response_service)
    signed_in()
    run = cli()
    before = AuditEvent.objects.filter(event_type=EVENT).count()
    for argv in (
        ("--timezone", "Mars/Olympus"),
        ("--status", "bad"),
        ("--start", "2000-01-01"),
    ):
        code, body, document = export(run, *argv)
        assert code == 1 and document["error"]["code"] == "invalid", argv
        assert body == b"", argv
    reader = read_only_session(run.service)
    code, body, document = export(run, secret=reader)
    assert code == 1 and document["error"]["code"] == "denied", document
    assert body == b""
    assert AuditEvent.objects.filter(event_type=EVENT).count() == before


def test_staff_export_and_a_leader_is_denied(
    response_service,
    google,
    cli,  # noqa: F811
    monkeypatch,
):
    """Staff export as on the page, offered no Reopen; a Ministry leader is
    refused with nothing written or recorded."""
    submitted(response_service)
    signed_in()
    run = cli()
    with monkeypatch.context() as patch:
        as_staff(patch)
        # The page's module holds its own reference to the admission.
        patch.setattr(
            census_change_views, "authenticated_admin", sessions.authenticated_admin
        )
        code, body, document = export(run, "--status", "all", "--format", "jsonl")
        assert code == 0, document
        actions = {record["what"]: record["actions"] for record in lines(body)}
        assert actions == {"Home address": ["entered", "ignored"], "First name": []}
    before = AuditEvent.objects.filter(event_type=EVENT).count()
    with_roles(monkeypatch, {"ministry_leader"}, frozenset({4}))
    code, body, document = export(run)
    assert code == 1 and document["error"]["code"] == "denied", document
    assert body == b"" and "result" not in document
    assert AuditEvent.objects.filter(event_type=EVENT).count() == before


def test_an_ended_session_exports_nothing(response_service, google, cli):  # noqa: F811
    """A revoked session is exit 5, before the choice check, with nothing
    written or recorded."""
    submitted(response_service)
    signed_in()
    run = cli()
    revoke(run.service, run.secret)
    for argv in ((), ("--timezone", "Mars/Olympus")):
        code, body, document = export(run, *argv)
        assert code == 5 and document["error"]["code"] == "session_ended", argv
        assert body == b""
    assert not AuditEvent.objects.filter(event_type=EVENT).exists()


def test_without_a_current_campaign(admin, google):  # noqa: F811
    """No current campaign: not_available, and nothing recorded."""
    code, body, document = stream(
        admin, "export", "census", secret=session(admin, scope="full")
    )
    assert code == 1 and document["error"]["code"] == "not_available", document
    assert body == b""
    assert not AuditEvent.objects.filter(event_type=EVENT).exists()
