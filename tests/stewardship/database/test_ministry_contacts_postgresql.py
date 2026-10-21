"""Member contact details on follow-up requests and the queue's export (#1157).

Real sign-in, the restricted web database login and the installed SQL: a
Ministry leader sees contact details only for the Ministries they lead, and
only published values; Staff see operational contact. Export these requests
downloads exactly the queue's filtered set with those details, and is
audited with Ministry scope but no values. The retired Members joining and
leaving pages, and their export form, answer 404 with no redirect.
"""

import csv
import io
import json
from uuid import uuid4

import pytest
from django.db import connection, transaction
from django.urls import reverse
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy import Principal
from parishkit.stewardship.audit.models import AuditContext, AuditEvent
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.reports import ministry_followup_views
from parishkit.stewardship.reports.export_models import MinistryExportSnapshot
from parishkit.stewardship.reports.ministries import MinistryQuery
from parishkit.stewardship.reports.ministry_contacts import (
    packet_contacts,
    request_contacts,
)
from parishkit.stewardship.reports.ministry_followup import FollowupQuery
from parishkit.stewardship.workflows.models import MinistryRequest

from ..census_factory import member
from .test_background_grants_postgresql import task_login
from .test_export_views_postgresql import post
from .test_information_followup_postgresql import search
from .test_live_ministries_postgresql import installed
from .test_ministry_exports_postgresql import create, leader
from .test_ministry_followup_postgresql import requests
from .test_ministry_packets_postgresql import packet, sections
from .test_ministry_reports_postgresql import actor, drop_leader, page, prepare, setup
from .test_ministry_responses_postgresql import ministry_source, respond, revisit
from .test_ministry_review_postgresql import decide, view, withdraw_leave
from .test_report_workspace_postgresql import read
from .test_response_http_postgresql import answers_for
from .test_source_families_postgresql import promote

pytestmark = pytest.mark.django_db(transaction=True)

EXPORT = "admin:ministry_followup_export"
# The harness Member's unpublished email, published home phone and address.
EMAIL = "valid@example.org"
PHONE = "+1 (202) 555-0123"


def item(browser, request):
    """One request page's body as text, after checking it was shown."""
    response, body = read(
        browser, reverse("admin:ministry_followup_item", args=[request.pk])
    )
    assert response.status_code == 200
    return body.decode()


def contact(body):
    """The page's Contact details section."""
    start = body.index('<section id="followup-contact"')
    return body[start : body.index("</section>", start)]


def export(browser, *, format="csv", **filters):
    """Download Export these requests for the queue view ``filters``."""
    fields = FollowupQuery(**filters).form_values() | {
        "format": format,
        "timezone": "UTC",
    }
    return search(browser, reverse(EXPORT), fields)


def table(body):
    """A CSV download's rows as dicts by heading."""
    return list(csv.DictReader(io.StringIO(body.decode())))


def exported():
    """Every Export these requests audit context, oldest first."""
    return [
        AuditContext.objects.get(event=row).context
        for row in AuditEvent.objects.filter(
            event_type="ministry_followup_exported"
        ).order_by("created_at")
    ]


def test_a_leader_sees_the_requesters_contacts_for_their_own_ministry_only(
    response_service, google, monkeypatch
):
    """The leader of Food pantry (9) sees its request's Member details,
    the unpublished email too since the Member asked (#1591) but not once
    the request is withdrawn, and is refused the Choir (4) request; the
    view is audited with the request's Ministry and no values."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google)
    join, leave = requests()
    real_principal = ministry_followup_views._principal

    def previously_staff(request, store, *, read_only=False):
        """An outer Staff decision the guarded leader check supersedes."""
        actor = real_principal(request, store, read_only=read_only)
        return actor if read_only else Principal(actor.identity, frozenset({"staff"}))

    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, join))
        assert "<dt>Email</dt><dd>Not published</dd>" not in section
        assert EMAIL in section and PHONE in section
        assert "1 Example Street" in section and "1960-01-01" not in section
        assert "<dt>Gender</dt><dd>Unspecified</dd>" in section
        # Contacts follow the reader's scope in SQL, never an outer
        # decision: an outer Staff decision shows no more than the leader's.
        with monkeypatch.context() as patch:
            patch.setattr(ministry_followup_views, "_principal", previously_staff)
            assert contact(item(browser, join)) == section
        path = reverse("admin:ministry_followup_item", args=[leave.pk])
        assert read(browser, path)[0].status_code == 403
    # On a standing request the leader's view and Staff's are the same, so
    # the scope check needs one where they differ: once the Family withdraws
    # the join, the leader sees the email Not published again, and the outer
    # Staff decision still shows no more than that.
    form = revisit(harness)
    answers = answers_for(form)
    answers["ministries"]["members"]["3"] = {"join": [], "leave": [4]}
    respond(harness, form, answers)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, join))
        assert "<dt>Email</dt><dd>Not published</dd>" in section
        assert EMAIL not in section and PHONE in section
        with monkeypatch.context() as patch:
            patch.setattr(ministry_followup_views, "_principal", previously_staff)
            assert contact(item(browser, join)) == section
    contexts = [
        AuditContext.objects.get(event=row).context
        for row in AuditEvent.objects.filter(event_type="ministry_followup_viewed")
    ]
    assert any(context.get("ministry_duid") == 9 for context in contexts)
    assert not any(context.get("ministry_duid") == 4 for context in contexts)
    assert "Example" not in json.dumps(contexts) and "valid@" not in json.dumps(
        contexts
    )


def test_staff_see_operational_contacts_and_a_leaves_role(response_service, google):
    """Staff see the unpublished email; a leave request shows the Member's
    current role in the Ministry and the packet's email and phones."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google, roles=("staff",))
    join, leave = requests()
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, join))
        assert EMAIL in section and PHONE in section
        section = contact(item(browser, leave))
        assert "<dt>Current Ministry role</dt><dd>Chairperson</dd>" in section
        assert EMAIL in section
        assert "Gender" not in section


def test_a_leaders_export_is_the_queue_view_in_their_scope(response_service, google):
    """A leader's file lists only their Ministry's requests, with the
    requesting Member's own values (#1591), in every format; the audit
    keeps the Ministry DUIDs, not values."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        response, body = export(browser, state="any")
        assert response.status_code == 200
        assert response["Content-Type"] == "text/csv"
        assert "stewardship-ministry-requests-" in response["Content-Disposition"]
        rows = table(body)
        assert [(row["Ministry"], row["Request"]) for row in rows] == [
            ("Food pantry", "Join")
        ]
        assert rows[0]["Email"] == EMAIL and rows[0]["Phones"]
        assert b"Choir" not in body
        # A Ministry filter outside the leader's scope reads nothing more.
        response, body = export(browser, state="any", ministry="4")
        assert response.status_code == 200 and table(body) == []
        response, body = export(browser, format="xlsx", state="any")
        sheet = load_workbook(io.BytesIO(body))["Ministry requests"]
        assert sheet["A1"].value == "Member" and sheet["F2"].value == "Food pantry"
        response, body = export(browser, format="pdf", state="any")
        assert response.status_code == 200 and body.startswith(b"%PDF-")
        # The form grammar is closed, and a query string is refused.
        fields = FollowupQuery().form_values() | {"format": "csv", "timezone": "UTC"}
        for invalid in ({"format": "doc"}, {"assignee": "any"}, {"state": "lost"}):
            response, _ = search(browser, reverse(EXPORT), fields | invalid)
            assert response.status_code == 400, invalid
        assert search(browser, reverse(EXPORT) + "?x=1", fields)[0].status_code == 400
        assert read(browser, reverse(EXPORT))[0].status_code == 405
    contexts = exported()
    assert contexts and all(
        context["ministry_duids"] in ([9], []) for context in contexts
    )
    assert all(context["ministry_operational"] is False for context in contexts)
    assert "Example" not in json.dumps(contexts) and "valid@" not in json.dumps(
        contexts
    )
    # ParishSoft no longer lists the leader's role: refused on the next file.
    drop_leader(harness)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert export(browser, state="any")[0].status_code == 403


def test_an_added_members_entered_phone_reads_as_a_number(response_service, google):
    """The Administrator's pk-local check (#1623): a Member the Family adds
    on the form with a mobile phone asks to join Food pantry (9). The form
    stores the phone as {display, normalized}; its leader's request page and
    Export these requests show it as a formatted number, never that JSON."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google)
    form = revisit(harness)
    answers = answers_for(form)
    added = str(uuid4())
    answers["proposed_members"][added] = member(
        first_name="New", last_name="Person", mobile_phone="5025551212"
    )
    answers["ministries"]["proposed_members"][added] = {"join": [9]}
    respond(harness, form, answers)
    request = MinistryRequest.objects.get(
        entity_kind="proposed_member", entity_key=added, action="join"
    )
    shown = "mobile: +1 (502) 555-1212"
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, request))
        assert shown in section and "normalized" not in section
        response, body = export(browser, state="any")
        assert response.status_code == 200
        (row,) = [row for row in table(body) if row["Member DUID"] == "New Member"]
        assert row["Phones"] == shown
        response, body = export(browser, format="xlsx", state="any")
        sheet = load_workbook(io.BytesIO(body))["Ministry requests"]
        cells = [str(cell.value) for cells in sheet.iter_rows() for cell in cells]
        assert shown in cells and not any("normalized" in cell for cell in cells)


def test_staff_export_every_ministry_with_operational_contacts(
    response_service, google
):
    """Staff's file holds both requests, the queue's filters applied."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google, roles=("staff",))
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        _, body = export(browser, state="any")
        rows = {row["Ministry"]: row for row in table(body)}
        assert set(rows) == {"Food pantry", "Choir"}
        assert rows["Food pantry"]["Email"] == EMAIL
        assert rows["Choir"]["Current Ministry role"] == "Chairperson"
        assert rows["Choir"]["Gender"] == ""
        _, body = export(browser, state="any", action="leave")
        assert [row["Ministry"] for row in table(body)] == ["Choir"]
    assert exported()[0]["ministry_duids"] == [4, 9]
    assert exported()[0]["ministry_operational"] is True


def test_the_retired_pages_are_gone(response_service, google):
    """No redirect, no alias: each old address is 404, by GET or POST."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google, roles=("staff",))
    report = reverse("admin:ministry_report")
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        # The follow-up queue's own address went with #1133.
        for path in (
            "joining/",
            "leaving/",
            "exports/",
            "joining",
            "leaving",
            "follow-up/",
            "follow-up",
        ):
            address = report + path
            assert browser.get(address).status_code == 404, address
            assert post(browser, address, {"ministry": "9"}).status_code == 404
        response, body = read(browser, report)
        # No retired list's form or Contact details control (#1157).
        assert response.status_code == 200 and b"/joining/" not in body
        assert b"Contact details of" not in body


def test_a_pdf_above_its_limit_is_refused_and_not_offered(
    response_service, google, monkeypatch
):
    """With the limit patched to one request, Staff's two requests are too
    many for a PDF: the queue offers no PDF and says why, a PDF asked for
    anyway is refused with the reason, and CSV and XLSX are not limited."""
    from parishkit.stewardship.reports import ministry_followup_exports

    monkeypatch.setattr(ministry_followup_exports, "EXPORT_PDF_LIMIT", 1)
    monkeypatch.setattr(ministry_followup_views, "EXPORT_PDF_LIMIT", 1)
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google, roles=("staff",))
    queue = reverse("admin:ministry_report")
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        response, body = search(browser, queue, {"stage": "all"})
        assert response.status_code == 200
        assert b'value="pdf" disabled data-server-disabled' in body
        assert b"The PDF is limited to 1 requests" in body
        response, body = export(browser, format="pdf", state="any")
        assert response.status_code == 400 and not body.startswith(b"%PDF")
        assert b"The PDF is limited to 1 requests." in body
        assert b"CSV or XLSX" in body
        for format in ("csv", "xlsx"):
            assert export(browser, format=format, state="any")[0].status_code == 200
        # One request fits: the PDF is offered and written.
        response, body = export(browser, format="pdf", state="any", action="join")
        assert response.status_code == 200 and body.startswith(b"%PDF-")


def test_a_cannot_participate_leave_hides_the_contact_on_an_earlier_join(
    response_service, google
):
    """Member 3 asks to join Food pantry (9), a ParishSoft refresh resolves it
    as joined, and the Family then checks the physical-limitations box, so
    Member 3 leaves 4 and 9. That leave is not consent (Administrator,
    #1591), and the old join row shares its chain: Food pantry's leader sees
    the email Not published on it, in the history selection, on its request
    page (request_contacts) and in a history capture of the report, and
    Export these requests lists only the leave; Staff still see it."""
    harness = setup(response_service)
    browser, actor, _, _ = leader(harness, google)
    join, _ = requests()
    data = ministry_source()
    data.ministry_type_memberships[9]["membership"] = [
        {
            "memberId": 3,
            "ministryRoleId": 1,
            "ministryRoleName": "Participant",
            "startDate": "2026-01-01",
            "endDate": None,
        }
    ]
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    # A refresh no longer resolves it (#1890), so the resolution a refresh
    # made before is forged as history, with the guards off.
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute("SET LOCAL session_replication_role = replica")
        cursor.execute(
            "UPDATE stewardship_ministry_request SET state='resolved',"
            " outcome='joined', resolved_at=now(), resolution_source_id=("
            "SELECT snapshot_id FROM stewardship_source_current WHERE singleton),"
            " version=version+1 WHERE id=%s",
            [join.pk],
        )
    join.refresh_from_db()
    assert (join.state, join.outcome) == ("resolved", "joined")
    # A resolved ordinary join still shares the email.
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert EMAIL in contact(item(browser, join))
    form = revisit(harness)
    answers = answers_for(form)
    answers["ministries"]["members"]["3"] = {"join": [], "leave": [4, 9]}
    answers["service"] = {
        "members": {"3": {"cannot_serve": True}},
        "proposed_members": {},
    }
    respond(harness, form, answers)
    assert MinistryRequest.objects.filter(
        ministry_duid=9, action="leave", state="new"
    ).exists()
    history = dict(history="all", state="any")
    rows = page(
        harness,
        roles=("ministry_leader",),
        scope=(9,),
        ministry=9,
        action="join",
        **history,
    )["rows"]
    (row,) = [row for row in rows if row["id"] == str(join.pk)]
    assert row["email_visibility"] == "not_published" and row["emails"] == []
    assert row["phone_visibility"] == "available"
    assert EMAIL not in json.dumps(rows, default=str)
    staff = page(harness, ministry=9, action="join", **history)["rows"]
    assert EMAIL in json.dumps(staff, default=str)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, join))
        assert "<dt>Email</dt><dd>Not published</dd>" in section
        assert EMAIL not in section and PHONE in section
        # The queue's file lists current requests only (#1131): the leave.
        response, body = export(browser, state="any")
        assert response.status_code == 200
        assert [row["Request"] for row in table(body)] == ["Leave (cannot participate)"]
        assert EMAIL not in body.decode()
    # A capture of the join report's history lists the old join row, Not
    # published too.
    query = MinistryQuery(history="all", state="any")
    request = create(harness, actor, query=query)
    snapshot = MinistryExportSnapshot.objects.get(pk=request.ministry_snapshot_id)
    assert snapshot.row_count == 1
    assert EMAIL not in json.dumps(snapshot.document)


# The rule before the reviewed Family form (#1705): forms drawn before
# its release, as one open across the deploy is.
@pytest.mark.usefixtures("earlier_rule")
def test_contacts_and_counts_follow_an_edit(response_service, google):
    """Staff edit Member 3's Food pantry (9) join into Ministry 4 (#1684):
    its contact details are read under Ministry 4 from then on (#1718), so
    Ministry 4's leader gets them, the Member's own unpublished email too
    (#1591), and Food pantry's leader, who still sees the request read
    only, gets none on its page, in its file or from its packet. By Ministry
    counts it under both, as the worklist's Ministry filter lists it."""
    harness = setup(response_service)
    browser, _, _, _ = leader(harness, google)
    staff = actor(harness, {"staff"}, ())
    fours = actor(harness, {"ministry_leader"}, (4,))
    pantry = actor(harness, {"ministry_leader"}, (9,))
    join = withdraw_leave(harness)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert EMAIL in contact(item(browser, join))
    decide(harness, staff, "edit", join, fields={"value": ["4"]})
    assert join.reviewed_ministry_duid == 4

    def details(principal):
        """The edited request's details as the worklist reads them."""
        row = view(harness, principal)[str(join.pk)]
        login = task_login(ServiceRole.WEB, exact=True, reconnect=True)
        with login, transaction.atomic():
            found = request_contacts(harness.campaign.pk, principal, [row])
        return found.get(row["id"])

    found = details(fours)
    assert found["ministry_duid"] == 4 and found["emails"]
    assert EMAIL in json.dumps(found)
    assert (found["email_visibility"], found["phone_visibility"]) == (
        "available",
        "available",
    )
    assert details(pantry) is None
    found = details(staff)
    assert EMAIL in json.dumps(found)
    assert (found["email_visibility"], found["phone_visibility"]) == (
        "available",
        "available",
    )
    # The join's earlier version (superseded when the Family answered
    # again) stays under Food pantry, but the latest's consent is Ministry
    # 4's now: Food pantry's leader reads it Not published, and Ministry 4's
    # leader lists only the edited version.
    earlier = MinistryRequest.objects.get(ministry_duid=9, state="superseded")
    history = dict(action="join", history="all", state="any")
    rows = page(harness, roles=("ministry_leader",), scope=(9,), ministry=9, **history)
    assert [row["id"] for row in rows["rows"]] == [str(earlier.pk)]
    assert rows["rows"][0]["email_visibility"] == "not_published"
    assert EMAIL not in json.dumps(rows, default=str)
    rows = page(harness, roles=("ministry_leader",), scope=(4,), ministry=4, **history)
    assert [row["id"] for row in rows["rows"]] == [str(join.pk)]
    # Once the Family says Member 3 cannot participate in any ministries,
    # the edited join is no longer consent to Ministry 4's leader either.
    form = revisit(harness)
    answers = answers_for(form)
    answers["ministries"]["members"]["3"] = {"join": [], "leave": [4]}
    answers["service"] = {
        "members": {"3": {"cannot_serve": True}},
        "proposed_members": {},
    }
    respond(harness, form, answers)
    # A decided request is not withdrawn by leaving it unticked (#1684), so
    # the change is the cannot-participate answer's alone.
    join.refresh_from_db()
    assert (join.state, join.reviewed_ministry_duid) == ("resolved", 4)
    found = details(fours)
    assert found["email_visibility"] == "not_published" and not found["emails"]
    assert EMAIL not in json.dumps(found)
    # The report lists the request under Ministry 4 alone, for every reader.
    for roles, scope, ministry, listed in (
        (("ministry_leader",), (4,), 4, True),
        (("ministry_leader",), (9,), 9, False),
        (("staff",), (), 9, False),
        (("staff",), (), 4, True),
    ):
        rows = page(
            harness,
            roles=roles,
            scope=scope,
            ministry=ministry,
            action="join",
            history="all",
            state="any",
        )["rows"]
        assert (str(join.pk) in {row["id"] for row in rows}) is listed, roles
    # So does the packet; its leave contacts are read from it.
    with task_login(ServiceRole.WEB, exact=True, reconnect=True), transaction.atomic():
        assert str(join.pk) in packet_contacts(harness.campaign.pk, staff, {4})[0]
        assert str(join.pk) not in packet_contacts(harness.campaign.pk, staff, {9})[0]
        assert str(join.pk) in packet_contacts(harness.campaign.pk, fours, {4})[0]
    # By Ministry counts it under both Ministries, as the worklist's Ministry
    # filter finds it under either.
    summaries = {
        row["duid"]: row["joining"]
        for row in page(harness, roles=("staff",), scope=())["summaries"]
    }
    assert summaries[9] == 1 and summaries[4] == 1
    assert str(join.pk) in view(harness, staff, ministry="9")
    assert str(join.pk) in view(harness, staff, ministry="4")
    # Food pantry's leader still opens the request, read only, with no
    # contact details, and its file reads them Unavailable.
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        section = contact(item(browser, join))
        assert "No contact details are available for this request." in section
        assert EMAIL not in section and PHONE not in section
        response, body = export(browser, state="any")
        assert response.status_code == 200
        rows = [row for row in table(body) if row["Ministry"] == "Food pantry"]
        assert [(row["Email"], row["Phones"]) for row in rows] == [
            ("Unavailable", "Unavailable")
        ]
        assert EMAIL not in body.decode()


def test_a_leaves_contacts_follow_its_edit(response_service):
    """The Family withdraws the Food pantry join, and Staff edit the leave of
    Ministry 4 into Food pantry (9): the leave's contacts, read from the
    packet, go to Food pantry's leader and no longer to Ministry 4's. Once
    Ministry 4 is removed from the campaign, the edited leave still keeps it
    listed, so a packet for it, which By Ministry offers, gets its (empty)
    section rather than failing."""
    harness = setup(response_service)
    staff = actor(harness, {"staff"}, ())
    fours = actor(harness, {"ministry_leader"}, (4,))
    pantry = actor(harness, {"ministry_leader"}, (9,))
    form = revisit(harness)
    answers = answers_for(form)
    answers["ministries"]["members"]["3"] = {"join": [], "leave": [4]}
    respond(harness, form, answers)
    leave = MinistryRequest.objects.get(ministry_duid=4, action="leave", state="new")
    decide(harness, staff, "edit", leave, fields={"value": ["9"]})
    assert leave.reviewed_ministry_duid == 9

    def details(principal):
        """The edited leave's details as the worklist reads them."""
        row = view(harness, principal)[str(leave.pk)]
        login = task_login(ServiceRole.WEB, exact=True, reconnect=True)
        with login, transaction.atomic():
            found = request_contacts(harness.campaign.pk, principal, [row])
        return found.get(row["id"])

    found = details(pantry)
    assert (found["kind"], found["ministry_duid"]) == ("leave", 9)
    assert EMAIL in json.dumps(found) and found["email_visibility"] == "available"
    assert details(fours) is None
    assert EMAIL in json.dumps(details(staff))
    assert installed(harness, [9]).state == "applied"
    by_ministry = page(harness)["summaries"]
    assert {(row["duid"], row["in_campaign"]) for row in by_ministry} == {
        (4, False),
        (9, True),
    }
    _, captured = sections(packet(harness, staff.identity, (4, 9), history=True))
    assert captured[4]["rows"] == []
    # Food pantry's section holds the edited leave (and, with the history
    # option, the withdrawn join).
    assert str(leave.pk) in {row["id"] for row in captured[9]["rows"]}
