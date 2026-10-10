"""Talents and limitations report under the real web role (#247)."""

import io
from uuid import uuid4

import pytest
from django.db import transaction
from django.urls import reverse
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy import Principal
from parishkit.stewardship.audit.models import AuditEvent
from parishkit.stewardship.campaigns.models import Campaign
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.reports.talents import TalentQuery, talents_report
from parishkit.stewardship.responses.service import DEFAULT_TALENTS
from parishkit.stewardship.source.models import SourceCurrent
from parishkit.stewardship.source.snapshot_names import snapshot_family_names

from .auth_builders import signed_in
from .response_builders import activate_response_service
from .test_background_grants_postgresql import task_login
from .test_export_views_postgresql import restricted_download_pool
from .test_information_followup_postgresql import search
from .test_ministry_followup_postgresql import read as followup
from .test_ministry_reports_postgresql import actor
from .test_ministry_responses_postgresql import respond, start
from .test_report_workspace_postgresql import read as get
from .test_response_http_postgresql import answers_for, load_form
from .test_runtime_auth_grants_postgresql import web_login

pytestmark = pytest.mark.django_db(transaction=True)

STAFF = Principal(uuid4(), frozenset({"staff"}), frozenset())
LEADER = Principal(uuid4(), frozenset({"ministry_leader"}), frozenset({4}))
PAINTER, OTHER = DEFAULT_TALENTS[0].id, DEFAULT_TALENTS[-1].id


def submit_live(harness, *, cannot_serve=False):
    """One live response with a talent, "Other" text and the welcome answer."""
    with web_login():
        form = load_form(harness)
        answers = answers_for(form)
        answers["cannot_attend"] = True
        answers["service"] = {
            "members": {
                "3": {
                    "cannot_serve": cannot_serve,
                    "talents": {PAINTER: "", OTHER: "Organ"},
                }
            },
            "proposed_members": {},
        }
        if cannot_serve:
            answers["ministries"]["members"]["3"] = {"join": [], "leave": [4]}
        return respond(harness, form, answers)


def report(harness, principal=STAFF, **values):
    """Read through the restricted web grants, not the migration owner."""
    campaign = Campaign.objects.select_related("active_configuration").get(
        pk=harness.campaign.pk
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True), transaction.atomic():
        return talents_report(
            campaign.pk,
            TalentQuery.parse(values),
            principal,
            configuration=campaign.active_configuration.values,
        )


def test_report_lists_talents_limitations_and_filters(response_service):
    """Rows come from the effective live response, worded and filterable."""
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness, cannot_serve=True)
    result = report(harness)
    assert len(result["members"]) == 1 and len(result["families"]) == 1
    row = result["members"][0]
    assert row["talents"] == ["Painter", "Other: Organ"]
    assert row["cannot_serve"] is True and row["member_name"]
    # The listed Member's ParishSoft DUID (#960).
    assert row["member_duid"] == 3
    # Search matches the Family as the page names it, "Surname, heads",
    # built in SQL as snapshot_names builds it (#960).
    duid = row["family_duid"]
    shown = snapshot_family_names(result["metadata"]["source_id"], {duid})[duid]
    assert ", " in shown, shown
    for text in (shown.upper(), shown.split(", ", 1)[1]):
        found = report(harness, search=text)
        assert [m["family_duid"] for m in found["members"]] == [duid], text
        assert [f["family_duid"] for f in found["families"]] == [duid], text
    assert result["summary"]["cannot_attend"] == 1
    assert dict(result["summary"]["talents"])["Painter"] == 1
    assert len(report(harness, talent=PAINTER)["members"]) == 1
    assert report(harness, talent=DEFAULT_TALENTS[1].id)["members"] == []
    only_families = report(harness, talent="cannot_attend")
    assert only_families["members"] == [] and len(only_families["families"]) == 1
    assert report(harness, search="no such name")["members"] == []
    with pytest.raises(PermissionError):
        report(harness, principal=LEADER)
    for invalid in ({"talent": "painter"}, {"extra": "x"}):
        with pytest.raises(ValueError):
            TalentQuery.parse(invalid)
    # The ministry follow-up queue explains why this Member is leaving.
    # The queue resolves scope from the actor in SQL (#389 L3), so it needs a
    # real portal user rather than the invented STAFF Principal.
    queue = followup(harness, actor(harness, ("staff",), ()))
    assert [(row["action"], row["cannot_serve"]) for row in queue["rows"]] == [
        ("leave", True)
    ]


def test_native_page_and_downloads(response_service, google, settings):
    """The page filters by CSRF POST; CSV and XLSX carry the same rows.

    Read with the real restricted web and download logins
    (``restricted_download_pool``): the download pool's login reads no
    campaign or source data, so a download must be read on the web
    connection, never there (#557).
    """
    settings.STEWARDSHIP_DOWNLOAD_POOL = None  # restricted_download_pool sets it
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness)
    route = reverse("admin:talents_report")
    browser, login = signed_in()
    assert login.status_code == 302
    with restricted_download_pool(settings):
        response, body = get(browser, route)
        assert response.status_code == 200 and response["Cache-Control"] == "no-store"
        assert b"Other: Organ" in body and b"Painter" in body
        assert get(browser, route + "?search=x")[0].status_code == 400
        response, body = search(browser, route, {"talent": "cannot_serve"})
        assert response.status_code == 200 and b"No matching Members." in body
        assert search(browser, route, {"talent": "bad"})[0].status_code == 400
        # Shared tables (#203): every column sorts on the server, both tables
        # page independently, and the controls are POST forms, never links.
        response, body = search(
            browser,
            route,
            {"talent": "any", "members_sort": "-latest", "families_size": "25"},
        )
        assert response.status_code == 200 and b"Page 1 of 1" in body
        assert b'aria-sort="descending"' in body and b'<a class="sort-link"' not in body
        assert b'name="families_size" value="25"' in body
        assert b'name="members_sort" value="latest"' in body
        # The Member DUID heading sorts too (#960).
        response, body = search(browser, route, {"members_sort": "-member_duid"})
        assert response.status_code == 200 and b'aria-sort="descending"' in body
        assert b'name="members_sort" value="member_duid"' in body
        for invalid in (
            {"members_sort": "submitted_at"},
            {"families_sort": "-talents"},
            {"members_size": "7"},
            {"families_page": "x"},
        ):
            assert search(browser, route, invalid)[0].status_code == 400
        export = reverse("admin:talents_export")
        response, body = search(
            browser, export, {"format": "csv", "timezone": "UTC", "talent": "any"}
        )
        assert response.status_code == 200
        assert response["Content-Type"] == "text/csv"
        assert response["Cache-Control"] == "no-store"
        assert "stewardship-talents-" in response["Content-Disposition"]
        assert b"Other: Organ" in body and b"Cannot attend Mass" in body
        response, body = search(browser, export, {"format": "xlsx", "timezone": "UTC"})
        assert response.status_code == 200
        assert response["Content-Disposition"].endswith(".xlsx")
        book = load_workbook(io.BytesIO(body))
        assert book.sheetnames == ["Members", "Families"]
        assert book["Members"]["D2"].value == "Painter; Other: Organ"
        assert search(browser, export, {"format": "pdf"})[0].status_code == 400
        # A talent filter no longer offered (here, never offered) reads as
        # "Everything": the tables, the re-posted forms and the download.
        removed = str(uuid4())
        response, body = search(
            browser, route, {"talent": removed, "members_sort": "-latest"}
        )
        assert response.status_code == 200 and b"Other: Organ" in body
        assert removed.encode() not in body
        assert b'name="talent" value="any"' in body
        response, body = search(
            browser, export, {"format": "csv", "timezone": "UTC", "talent": removed}
        )
        assert response.status_code == 200 and b"Other: Organ" in body
        _, body = get(browser, reverse("admin:participation"))
        assert route.encode() in body
        # The campaign reports page links the response dashboard too (#477).
        assert route.replace("talents", "responses").encode() in body
        # A search is audited as used, never as its text (#556).
        assert search(browser, route, {"search": "Organ"})[0].status_code == 200
    assert AuditEvent.objects.filter(event_type="talents_report_viewed").exists()
    assert AuditEvent.objects.filter(event_type="talents_report_exported").exists()
    # Each records its Show choice, whether the search box was used and the
    # ParishSoft snapshot read (#556); a removed talent reads as "any".
    snapshot = str(SourceCurrent.objects.get().snapshot_id)
    contexts = [
        (event.event_type, event.auditcontext.context)
        for event in AuditEvent.objects.filter(
            event_type__startswith="talents_report_"
        ).select_related("auditcontext")
    ]
    assert all(context["snapshot_id"] == snapshot for _, context in contexts)
    assert all("talent_option_id" not in context for _, context in contexts)
    assert ("talents_report_viewed", "cannot_serve", False) in {
        (kind, context["report_filter"], context["search_used"])
        for kind, context in contexts
    }
    assert {
        context["report_filter"]
        for kind, context in contexts
        if kind == "talents_report_exported"
    } == {"any"}
    assert ("talents_report_viewed", "any", True) in {
        (kind, context["report_filter"], context["search_used"])
        for kind, context in contexts
    }
    assert "Organ" not in str(contexts)


def test_both_tables_sort_every_column_in_memory():
    """The complete in-memory report sorts before paging, per table."""
    from datetime import UTC, datetime

    from parishkit.stewardship.reports.talent_views import tables

    def row(name, duid, day, talents=(), cannot=False, member=None):
        """One member/family row as the selection shapes it."""
        return {
            "member_name": name,
            "member_duid": member,
            "family_name": name,
            "family_duid": duid,
            "talents": list(talents),
            "cannot_serve": cannot,
            "submitted_at": datetime(2026, 9, day, tzinfo=UTC),
        }

    result = {
        "members": [row("b", 2, 1, ["Zither"]), row("A", 1, 3, [], True)],
        "families": [row("b", 2, 1), row("a", 1, 3)],
    }
    members, families = tables(
        result,
        TalentQuery(search="private"),
        {"members_sort": "-cannot_serve", "families_sort": "-duid"},
        "/r/",
    )
    assert [r["family_duid"] for r in members.rows] == [1, 2]
    assert [r["family_duid"] for r in families.rows] == [2, 1]
    assert members.method == families.method == "post" and members.action == "/r/"
    assert ("search", "private") in members.carried
    assert ("families_sort", "-duid") in members.carried
    members, _ = tables(result, TalentQuery(), {"members_sort": "talents"}, "/r/")
    assert [r["family_duid"] for r in members.rows] == [1, 2]
    # A Member added on the form (no DUID) sorts last either way (#960).
    result["members"] = [
        row("p", 3, 1),
        row("b", 2, 1, member=7),
        row("a", 1, 1, member=5),
    ]
    for sort, expected in (
        ("member_duid", [5, 7, None]),
        ("-member_duid", [7, 5, None]),
    ):
        members, _ = tables(result, TalentQuery(), {"members_sort": sort}, "/r/")
        assert [r["member_duid"] for r in members.rows] == expected, sort
