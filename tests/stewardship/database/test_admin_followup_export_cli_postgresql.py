"""``followup export``: the queue's Export these requests (#463, ADM-11 PR 11c).

Process admission is replaced by the test's own Django assembly; everything
after it is the real command line on the restricted web login, reading the
queue's whole filtered set through the page's own admission, choice check
and selection under the campaign read guard, with each Member's contact
details, and writing it through the page's own writer. Its CSV is the
page's download byte for byte, and it records the page's
``ministry_followup_exported`` with the same outcomes and context members.
"""

import csv
import hashlib
import io
import json

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from parishkit.stewardship.audit.models import AuditContext, AuditEvent
from parishkit.stewardship.reports import ministry_followup_exports
from parishkit.stewardship.reports.ministry_followup import FollowupQuery

from .auth_builders import signed_in
from .test_admin_export_cli_postgresql import read_only_session
from .test_admin_followup_cli_postgresql import member_names, with_roles
from .test_admin_logs_cli_postgresql import counted, stream
from .test_admin_report_cli_postgresql import cli, parity  # noqa: F401
from .test_admin_report_email_cli_postgresql import revoke
from .test_admin_status_cli_postgresql import admin, session  # noqa: F401
from .test_information_followup_postgresql import search
from .test_ministry_reports_postgresql import setup

pytestmark = pytest.mark.django_db(transaction=True)

EVENT = "ministry_followup_exported"
# The harness Member's unpublished email, which only operational readers see.
EMAIL = "valid@example.org"


def export(run, *argv, secret=None):
    """``followup export`` with ``argv``: (exit code, bytes, document)."""
    return stream(None, "followup", "export", *argv, secret=secret or run.secret)


def page_form(format="csv", zone="UTC", **filters):
    """The queue's export form for the view ``filters``, as the page posts it."""
    return FollowupQuery(**filters).form_values() | {
        "format": format,
        "timezone": zone,
    }


def table(body):
    """A CSV file's rows as dicts by heading."""
    return list(csv.DictReader(io.StringIO(body.decode())))


def contexts():
    """Every recorded download event's context, oldest first."""
    return [
        AuditContext.objects.get(event=row).context
        for row in AuditEvent.objects.filter(event_type=EVENT).order_by(
            "created_at", "id"
        )
    ]


def test_the_page_s_files(response_service, google, cli):  # noqa: F811
    """CSV is the page's download byte for byte; XLSX holds the page's
    sheet; PDF is a PDF; the filters narrow the file. Each records the
    page's event with the requests and Ministries in the file, and nothing
    else, and the document names no Member and no contact detail."""
    harness = setup(response_service)
    browser, _ = signed_in()
    run = cli()
    route = reverse("admin:ministry_followup_export")
    found = {}

    def page():
        """The page's download, as the browser posts it (in memory)."""
        response, found["page"] = search(
            browser, route, page_form(zone="America/New_York", state="any")
        )
        assert response.status_code == 200

    def command():
        """The command's CSV of the same view."""
        found["command"] = export(
            run, "--state", "any", "--timezone", "America/New_York", "--format", "csv"
        )

    parity(EVENT, page, command)
    code, body, document = found["command"]
    assert code == 0, document
    assert body == found["page"]
    rows = {row["Ministry"]: row for row in table(body)}
    assert set(rows) == {"Food pantry", "Choir"}
    assert rows["Food pantry"]["Email"] == EMAIL
    result = document["result"]
    assert result["campaign_id"] == str(harness.campaign.pk)
    assert result["format"] == "csv" and result["content_type"] == "text/csv"
    assert result["timezone"] == "America/New_York"
    assert result["count"] == 2
    assert result["size"] == len(body)
    assert result["sha256"] == hashlib.sha256(body).hexdigest()
    assert result["file_name"].startswith("stewardship-ministry-requests-")
    assert result["file_name"].endswith(".csv")
    text = json.dumps(document)
    assert EMAIL not in text and not any(name in text for name in member_names())
    assert counted(EVENT)[-2:] == [2, 2]
    assert contexts()[-1]["ministry_duids"] == [4, 9]
    assert contexts()[-1]["ministry_operational"] is True

    code, body, document = export(run, "--state", "any")
    assert code == 0, document
    assert document["result"]["format"] == "xlsx"
    assert document["result"]["content_type"].endswith("spreadsheetml.sheet")
    command_sheet = load_workbook(io.BytesIO(body))["Ministry requests"]
    response, page_body = search(browser, route, page_form(format="xlsx", state="any"))
    assert response.status_code == 200
    page_sheet = load_workbook(io.BytesIO(page_body))["Ministry requests"]
    assert [[cell.value for cell in line] for line in command_sheet.rows] == [
        [cell.value for cell in line] for line in page_sheet.rows
    ]

    code, body, document = export(run, "--state", "any", "--format", "pdf")
    assert code == 0 and body.startswith(b"%PDF-"), document
    assert document["result"]["content_type"] == "application/pdf"

    # The queue's filters, as on the page; only the page's event recorded.
    before = set(AuditEvent.objects.values_list("pk", flat=True))
    code, body, document = export(
        run, "--state", "any", "--action", "leave", "--format", "csv"
    )
    assert code == 0 and [row["Ministry"] for row in table(body)] == ["Choir"]
    code, body, document = export(
        run, "--state", "any", "--ministry", "9", "--format", "csv"
    )
    assert code == 0 and [row["Ministry"] for row in table(body)] == ["Food pantry"]
    code, body, document = export(run, "--state", "resolved", "--format", "csv")
    assert code == 0 and document["result"]["count"] == 0, document
    assert table(body) == []
    added = AuditEvent.objects.exclude(pk__in=before)
    assert set(added.values_list("event_type", flat=True)) == {EVENT}
    assert added.count() == 3


def test_a_pdf_above_the_limit_is_invalid(
    response_service,
    google,
    cli,  # noqa: F811
    monkeypatch,
):
    """With the page's limit at one request, two are too many for a PDF:
    ``invalid`` at ``format`` with the page's words, no byte written, and
    the page's failed event with the count reached; CSV is not limited."""
    monkeypatch.setattr(ministry_followup_exports, "EXPORT_PDF_LIMIT", 1)
    setup(response_service)
    signed_in()
    run = cli()
    code, body, document = export(run, "--state", "any", "--format", "pdf")
    assert code == 1 and body == b"", document
    assert document["error"]["code"] == "invalid"
    [field] = document["error"]["fields"]
    assert field["field"] == "format"
    assert field["message"] == (
        "The PDF is limited to 1 requests. Choose CSV or XLSX, or narrow the filters."
    )
    assert [context["outcome"] for context in contexts()] == ["failed"]
    assert counted(EVENT) == [2]
    code, body, document = export(run, "--state", "any", "--format", "csv")
    assert code == 0 and document["result"]["count"] == 2, document
    code, body, document = export(
        run, "--state", "any", "--action", "join", "--format", "pdf"
    )
    assert code == 0 and body.startswith(b"%PDF-"), document


def test_refusals_print_nothing_and_record_nothing(response_service, google, cli):  # noqa: F811
    """A choice the page refuses is invalid; a read-only session is refused;
    neither writes a byte or records the page's event."""
    setup(response_service)
    signed_in()
    run = cli()
    # Pairing the read-only session records its own events first.
    reader = read_only_session(run.service)
    before = AuditEvent.objects.count()
    for argv in (("--timezone", "Mars/Olympus"), ("--ministry", "0")):
        code, body, document = export(run, *argv)
        assert code == 1 and document["error"]["code"] == "invalid", document
        assert body == b""
    code, body, document = export(run, secret=reader)
    assert code == 1 and document["error"]["code"] == "denied", document
    assert body == b""
    assert AuditEvent.objects.count() == before


def test_staff_export_and_other_roles_are_refused(
    response_service,
    google,
    cli,  # noqa: F811
    monkeypatch,
):
    """Staff export every Ministry, as on the page. A session's holder
    turned Ministry leader whom ParishSoft does not list as one is refused
    by the selection's own scope check, inside the read, as on the page:
    nothing is written and the page's event records the failure. A
    principal with no role is refused before anything is read or recorded.

    A pk-admin session is approved only by an Administrator, so a leader's
    leader-scoped file is the page's own selection, proven on the page in
    ``test_ministry_contacts_postgresql.py``.
    """
    setup(response_service)
    signed_in()
    run = cli()
    with monkeypatch.context() as patch:
        with_roles(patch, {"staff"})
        code, body, document = export(run, "--state", "any", "--format", "csv")
        assert code == 0 and document["result"]["count"] == 2, document
        assert body.startswith(b"Member,")
    assert contexts()[-1]["ministry_operational"] is True
    with monkeypatch.context() as patch:
        with_roles(patch, {"ministry_leader"}, frozenset({9}))
        code, body, document = export(run, "--state", "any", "--format", "csv")
    assert code == 1 and document["error"]["code"] == "denied", document
    assert body == b"" and "result" not in document
    assert contexts()[-1]["outcome"] == "failed"
    before = AuditEvent.objects.filter(event_type=EVENT).count()
    with monkeypatch.context() as patch:
        with_roles(patch, set())
        code, body, document = export(run, "--state", "any")
    assert code == 1 and document["error"]["code"] == "denied", document
    assert body == b"" and "result" not in document
    assert AuditEvent.objects.filter(event_type=EVENT).count() == before


def test_an_ended_session_exports_nothing(response_service, google, cli):  # noqa: F811
    """A revoked session is exit 5, before the choice check, with nothing
    written or recorded."""
    setup(response_service)
    signed_in()
    run = cli()
    revoke(run.service, run.secret)
    for argv in ((), ("--timezone", "Mars/Olympus"), ("--ministry", "0")):
        code, body, document = export(run, *argv)
        assert code == 5 and document["error"]["code"] == "session_ended", argv
        assert body == b""
    assert not AuditEvent.objects.filter(event_type=EVENT).exists()


def test_without_a_current_campaign(admin, google):  # noqa: F811
    """No current campaign: not_available, and nothing recorded."""
    code, body, document = stream(
        admin, "followup", "export", secret=session(admin, scope="full")
    )
    assert code == 1 and document["error"]["code"] == "not_available", document
    assert body == b""
    assert not AuditEvent.objects.filter(event_type=EVENT).exists()
