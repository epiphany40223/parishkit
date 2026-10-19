"""Limitations report under the real web role (#247, #970).

The report was "Talents and limitations". ``submit_live`` sends talents as a
form opened before #970 would; they are dropped, not stored. Responses
submitted before then keep their stored talents (``store_old_talents``
recreates one), and the report must neither list nor count them.
"""

import io
from uuid import uuid4

import pytest
from django.db import connection, transaction
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy import Principal
from parishkit.stewardship.audit.models import AuditEvent
from parishkit.stewardship.campaigns.models import Campaign, FamilyCampaign
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.reports.limitations import (
    LimitationsQuery,
    limitations_report,
)
from parishkit.stewardship.responses.service import DEFAULT_TALENTS
from parishkit.stewardship.source.models import SourceCurrent

from .auth_builders import signed_in
from .family_link_checks import link_reads
from .response_builders import activate_response_service
from .test_background_grants_postgresql import task_login
from .test_export_views_postgresql import restricted_download_pool
from .test_information_followup_postgresql import search
from .test_ministry_followup_postgresql import read as followup
from .test_ministry_reports_postgresql import actor
from .test_ministry_responses_postgresql import respond, start
from .test_report_workspace_postgresql import read as get
from .test_response_http_postgresql import answers_for, load_form
from .test_runtime_auth_grants_postgresql import web_login

pytestmark = pytest.mark.django_db(transaction=True)

STAFF = Principal(uuid4(), frozenset({"staff"}), frozenset())
LEADER = Principal(uuid4(), frozenset({"ministry_leader"}), frozenset({4}))
PAINTER, OTHER = DEFAULT_TALENTS[0].id, DEFAULT_TALENTS[-1].id


def submit_live(harness, *, cannot_serve=False):
    """One live response with limitations, sent with talents as before #970."""
    with web_login():
        form = load_form(harness)
        answers = answers_for(form)
        answers["cannot_attend"] = True
        answers["service"] = {
            "members": {
                "3": {
                    "cannot_serve": cannot_serve,
                    "talents": {PAINTER: "", OTHER: "Organ"},
                }
            },
            "proposed_members": {},
        }
        if cannot_serve:
            answers["ministries"]["members"]["3"] = {"join": [], "leave": [4]}
        return respond(harness, form, answers)


def store_old_talents(submission, member, cannot_serve):
    """Rewrite one stored Member entry as a response from before #970 held it.

    The guard refuses talents on every new insert since migration 0056, and
    submissions are immutable, so the test turns triggers off for this one
    statement, as only a response recorded before then could look.
    """
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute("SET LOCAL session_replication_role = replica")
        cursor.execute(
            "UPDATE stewardship_submission SET answers=jsonb_set(answers,"
            "ARRAY['service','members',%s],%s::jsonb) WHERE id=%s",
            [
                member,
                f'{{"cannot_serve": {str(cannot_serve).lower()}, '
                f'"talents": {{"{PAINTER}": "", "{OTHER}": "Organ"}}}}',
                submission.pk,
            ],
        )
        cursor.execute("SET LOCAL session_replication_role = origin")


def report(harness, principal=STAFF, **values):
    """Read through the restricted web grants, not the migration owner."""
    campaign = Campaign.objects.get(pk=harness.campaign.pk)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True), transaction.atomic():
        return limitations_report(
            campaign.pk, LimitationsQuery.parse(values), principal
        )


def directory_name(duid):
    """The Family's name as the directory shows it, from the current snapshot."""
    from parishkit.stewardship.source.snapshot_models import SourceCurrent
    from parishkit.stewardship.source.snapshot_names import snapshot_family_names

    snapshot = SourceCurrent.objects.values_list("snapshot_id", flat=True).get()
    return snapshot_family_names(snapshot, [duid])[duid]


def test_report_lists_limitations_and_filters(response_service):
    """Rows come from the effective live response, filterable; no talents."""
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submission = submit_live(harness, cannot_serve=True)
    # Talents a form opened before #970 sends are dropped, not stored.
    assert submission.answers["service"]["members"]["3"]["talents"] == {}
    # A response from before then keeps its talents; the report ignores them,
    # and lists no Member only for a talent.
    store_old_talents(submission, "3", False)
    old = report(harness)
    assert old["members"] == [] and len(old["families"]) == 1
    assert old["summary"] == {"cannot_serve": 0, "cannot_attend": 1}
    store_old_talents(submission, "3", True)
    result = report(harness)
    assert len(result["members"]) == 1 and len(result["families"]) == 1
    row = result["members"][0]
    assert "talents" not in row
    assert row["cannot_serve"] is True and row["member_name"]
    # The Member's DUID comes from the SQL (#932); Member 3 is the fixture's.
    assert row["member_duid"] == 3 and row["proposed"] is False
    # Both tables name the Family as every Admin table does (#932).
    family = result["families"][0]
    assert ", " in row["family_name"] and family["family_name"] == row["family_name"]
    assert row["family_name"] == directory_name(row["family_duid"])
    assert result["summary"] == {"cannot_serve": 1, "cannot_attend": 1}
    only_families = report(harness, show="cannot_attend")
    assert only_families["members"] == [] and len(only_families["families"]) == 1
    assert report(harness, search="no such name")["members"] == []
    with pytest.raises(PermissionError):
        report(harness, principal=LEADER)
    for invalid in ({"show": PAINTER}, {"talent": "any"}, {"extra": "x"}):
        with pytest.raises(ValueError):
            LimitationsQuery.parse(invalid)
    # The ministry follow-up queue explains why this Member is leaving.
    # The queue resolves scope from the actor in SQL (#389 L3), so it needs a
    # real portal user rather than the invented STAFF Principal.
    queue = followup(harness, actor(harness, ("staff",), ()))
    assert [(row["action"], row["cannot_serve"]) for row in queue["rows"]] == [
        ("leave", True)
    ]


def test_native_page_and_downloads(response_service, google, settings):
    """The worklist filters and decides by CSRF POST; CSV, XLSX and PDF carry
    the rows at its Stage (#1685).

    Read with the real restricted web and download logins
    (``restricted_download_pool``): the download pool's login reads no
    campaign or source data, so a download must be read on the web
    connection, never there (#557).
    """
    settings.STEWARDSHIP_DOWNLOAD_POOL = None  # restricted_download_pool sets it
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submission = submit_live(harness, cannot_serve=True)
    duid = FamilyCampaign.objects.get(pk=submission.family_id).family_duid
    route = reverse("admin:limitations_report")
    assert route == "/admin/reports/limitations/"
    browser, login = signed_in()
    assert login.status_code == 302
    with restricted_download_pool(settings):
        with CaptureQueriesContext(connection) as queries:
            response, body = get(browser, route)
        assert response.status_code == 200 and response["Cache-Control"] == "no-store"
        page = body.decode()
        assert "<h1>Limitations</h1>" in page and "Organ" not in page
        assert "Painter" not in page
        # One worklist table, To review first: the Member's answer and the
        # Family's, with no Summary tiles, Baseline or Reviewer's value.
        assert ">To review (2)</option>" in page and ">Every stage (2)</option>" in page
        assert "stat-tile" not in page
        head = page.split("<thead>")[1].split("</thead>")[0]
        assert "Reviewer's value" not in head and "Baseline" not in head
        assert "Member DUID" in head and "Family's answer" in head
        assert "Cannot participate in any ministries" in page
        assert "Cannot attend Mass or prayer services" in page
        # Each Family name opens its Timeline (#1570): one record id read.
        assert '<th scope="row"><a href="/admin/reports/families/' in page
        assert len(link_reads(queries)) == 1
        assert get(browser, route + "?search=x")[0].status_code == 400
        response, body = search(browser, route, {"show": "cannot_attend"})
        assert response.status_code == 200 and b">To review (1)</option>" in body
        assert search(browser, route, {"show": "bad"})[0].status_code == 400
        # The filter's old name and a talent's id are refused (#970).
        assert search(browser, route, {"talent": "any"})[0].status_code == 400
        assert search(browser, route, {"show": str(uuid4())})[0].status_code == 400
        assert search(browser, route, {"stage": "history"})[0].status_code == 400
        # The shared table (#203): every column sorts on the server, and the
        # controls are POST forms, never links.
        response, body = search(browser, route, {"sort": "-duid", "size": "25"})
        assert response.status_code == 200 and b"Page 1 of 1" in body
        assert b'aria-sort="descending"' in body and b'<a class="sort-link"' not in body
        for invalid in ({"sort": "talents"}, {"size": "7"}, {"page": "x"}):
            assert search(browser, route, invalid)[0].status_code == 400
        # Accept the Family's answer: the notice names it by its DUID, and
        # it moves to To handle.
        decision = {
            "decide": "accept",
            "item": f"f{duid}",
            "version": "0",
            "request_key": str(uuid4()),
            "note": "",
        }
        response, body = search(browser, route, decision)
        assert response.status_code == 200
        assert f"The answer of Family DUID {duid} accepted.".encode() in body
        assert b">To handle (1)</option>" in body
        assert (
            search(browser, route, decision | {"decide": "ignore"})[0].status_code
            == 409
        )
        stale = decision | {"request_key": str(uuid4())}
        response, body = search(browser, route, stale)
        assert b": not recorded, changed by someone else." in body
        # Bulk: the Member's answer, ignored.
        bulk = {
            "bulk": "ignore",
            "bulk_key": str(uuid4()),
            "bulk_note": "Seen",
            "item": f"m{duid}-3:0",
        }
        response, body = search(browser, route, bulk)
        assert (
            f"The answer of Member DUID 3 in Family DUID {duid} ignored.".encode()
            in (body)
        )
        export = reverse("admin:limitations_export")
        response, body = search(
            browser,
            export,
            {"format": "csv", "timezone": "UTC", "show": "any", "stage": "all"},
        )
        assert response.status_code == 200
        assert response["Content-Type"] == "text/csv"
        assert response["Cache-Control"] == "no-store"
        assert "stewardship-limitations-" in response["Content-Disposition"]
        assert b"Organ" not in body
        # The Families table has the page's columns: no Cannot attend Mass,
        # and each row's Stage last. ParishSoft status follows each Family
        # name (#1216).
        assert (
            b"\r\nFamily,ParishSoft status,Family DUID,Latest response,Stage\r\n"
            in body
        )
        assert body.startswith(
            b"Member,Member DUID,Family,ParishSoft status,Family DUID,"
            b"Cannot participate in ministries,Latest response,Stage\r\n"
        )
        assert b",Ignored\r\n" in body and b",To handle\r\n" in body
        # The download applies the page's Stage.
        response, body = search(
            browser, export, {"format": "csv", "timezone": "UTC", "stage": "ignored"}
        )
        assert b",To handle\r\n" not in body and b",Ignored\r\n" in body
        assert (
            search(browser, export, {"format": "csv", "stage": "x"})[0].status_code
            == 400
        )
        response, body = search(browser, export, {"format": "xlsx", "timezone": "UTC"})
        assert response.status_code == 200
        assert response["Content-Disposition"].endswith(".xlsx")
        book = load_workbook(io.BytesIO(body))
        assert book.sheetnames == ["Members", "Families", "Report information"]
        # The workbook names its report, parish and campaign (#926).
        assert book.properties.title == "Limitations"
        details = dict(book["Report information"].iter_rows(values_only=True))
        assert details["Report"] == "Limitations"
        assert details["Parish"] and details["Campaign"]
        assert details["Members in this file"] == f"{book['Members'].max_row - 1:,}"
        assert [cell.value for cell in book["Members"][1]] == [
            "Member",
            "Member DUID",
            "Family",
            "ParishSoft status",
            "Family DUID",
            "Cannot participate in ministries",
            "Latest response",
            "Stage",
        ]
        assert book["Members"]["B2"].value == "3"
        assert book["Members"]["F2"].value == "Yes"
        # The PDF is offered too (#966), named and typed as one.
        response, body = search(browser, export, {"format": "pdf", "timezone": "UTC"})
        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert response["Content-Disposition"].endswith(".pdf")
        assert body.startswith(b"%PDF-")
        assert search(browser, export, {"format": "png"})[0].status_code == 400
        # The re-posted forms carry the Show and Stage choices.
        response, body = search(browser, route, {"sort": "-submitted"})
        assert response.status_code == 200 and b'name="show" value="any"' in body
        assert b'name="stage" value="to_review"' in body
        _, body = get(browser, reverse("admin:response_dashboard"))
        assert route.encode() in body
        # The dashboard's own choices name its address (#477, #1537).
        assert route.replace("limitations", "responses").encode() in body
        # The old address is gone, with no redirect (no legacy Admin URLs).
        assert browser.get("/admin/reports/talents/").status_code == 404
        assert browser.get("/admin/campaign/talents/").status_code == 404
        # A search is audited as used, never as its text (#556).
        assert search(browser, route, {"search": "Organ"})[0].status_code == 200
    assert AuditEvent.objects.filter(event_type="talents_report_viewed").exists()
    assert AuditEvent.objects.filter(event_type="talents_report_exported").exists()
    # Each records its Show choice, whether the search box was used and the
    # ParishSoft snapshot read (#556). The report keeps the events it
    # recorded as Talents and limitations (#970).
    snapshot = str(SourceCurrent.objects.get().snapshot_id)
    contexts = [
        (event.event_type, event.auditcontext.context)
        for event in AuditEvent.objects.filter(
            event_type__startswith="talents_report_"
        ).select_related("auditcontext")
    ]
    assert all(context["snapshot_id"] == snapshot for _, context in contexts)
    assert all("talent_option_id" not in context for _, context in contexts)
    assert ("talents_report_viewed", "cannot_attend", False) in {
        (kind, context["report_filter"], context["search_used"])
        for kind, context in contexts
    }
    assert {
        context["report_filter"]
        for kind, context in contexts
        if kind == "talents_report_exported"
    } == {"any"}
    assert ("talents_report_viewed", "any", True) in {
        (kind, context["report_filter"], context["search_used"])
        for kind, context in contexts
    }
    assert "Organ" not in str(contexts)


def test_the_worklist_sorts_every_column_in_memory():
    """The complete in-memory list sorts before paging, by every column."""
    from datetime import UTC, datetime

    from parishkit.stewardship.reports.limitation_worklist import SORTING
    from parishkit.stewardship.web.tables import paginate

    def row(name, duid, day, member=None):
        """One answer row as the selection shapes it."""
        return {
            "member_name": name,
            "family_name": name,
            "family_duid": duid,
            "member_duid": member,
            "submitted_at": datetime(2026, 9, day, tzinfo=UTC),
        }

    rows = [row("b", 2, 1, 7), row("A", 1, 3), row("c", 3, 2, 5)]
    for sort, order in (
        ("family", [1, 2, 3]),
        ("-duid", [3, 2, 1]),
        ("member_duid", [3, 2, 1]),
        ("-submitted", [1, 3, 2]),
    ):
        table = paginate(rows, {"sort": sort}, sorting=SORTING)
        assert [found["family_duid"] for found in table.rows] == order, sort
