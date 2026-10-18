"""Chosen-Family tests of ``pk-stewardship admin`` (ADM-11 PR 6c).

``test families-preview``, ``test families`` and ``test status`` through
the Send to chosen Families page's own functions, on the restricted web
login, over the ``family_test`` fixture of test_family_mail_test_postgresql.py
(source-backed Families and an invitation used by a schedule). The real
message goes only to the Testing recipient, rendered later by the worker;
these commands only write tickets. A live full-scope session admits them
and leaves the ``automation_fresh_gate`` trail (PR 5b).
"""

import csv
import io
import json
from uuid import uuid4

import pytest
from django.core import signing

from parishkit.stewardship.accounts import automation_sessions as automation
from parishkit.stewardship.accounts.automation_models import AutomationNotice
from parishkit.stewardship.accounts.campaign_family_test import SALT
from parishkit.stewardship.accounts.policy import current_principal
from parishkit.stewardship.audit.models import AuditEvent
from parishkit.stewardship.jobs.family_mail_models import FamilyMailTest
from parishkit.stewardship.jobs.models import TaskRun

from ..policy_factory import address
from .automation_builders import paired
from .test_admin_family_export_cli_postgresql import reports_root  # noqa: F401
from .test_admin_test_sample_cli_postgresql import runner
from .test_export_authorization_postgresql import add_policy
from .test_family_mail_test_postgresql import family_test, review  # noqa: F401
from .test_policy_postgresql import user
from .test_runtime_auth_grants_postgresql import web_login
from .test_setup_views_postgresql import post

pytestmark = pytest.mark.django_db(transaction=True)

EVENT = "admin_cmd_test_families"


@pytest.fixture
def admin(family_test, auth_service, monkeypatch):  # noqa: F811
    """Commands over the authentication service the Family test runs on."""
    return runner(auth_service, monkeypatch)


def revision(family_test):  # noqa: F811
    """The fixture's email revision, from its page path."""
    return family_test[3].rstrip("/").rsplit("/", 1)[-1]


def preview(admin, secret, family_test, *duids, key=None):  # noqa: F811
    """``test families-preview`` for ``duids``."""
    argv = ["test", "families-preview", revision(family_test)]
    for duid in duids:
        argv += ["--family", str(duid)]
    if key is not None:
        argv += ["--request-key", str(key)]
    return admin(*argv, secret=secret)


def send(admin, secret, token, answer=b""):
    """``test families`` with ``answer`` typed at the prompt."""
    return admin("test", "families", "--token", token, secret=secret, answer=answer)


def trail():
    """The command events, fresh-gate events and fresh-gated notices so far."""
    return (
        AuditEvent.objects.filter(event_type=EVENT).count(),
        AuditEvent.objects.filter(event_type="automation_fresh_gate").count(),
        AutomationNotice.objects.filter(kind="fresh_gated").count(),
    )


def test_chosen_families_are_previewed_then_sent_once(admin, family_test):  # noqa: F811
    """No names; the prompt; one ticket per Family; one trail; no resend."""
    _, secret, row = paired(admin.service)
    key = uuid4()
    code, document, _ = preview(admin, secret, family_test, 1, key=key)
    assert code == 0, document
    result = document["result"]
    assert result["families"] == [
        {"duid": 1, "eligible": True, "eligibility": "eligible"}
    ]
    assert result["credentials_ready"] is True and result["held"] is False
    assert result["request_key"] == str(key)
    # No Family name, no address; a preview records nothing.
    assert "name" not in json.dumps(result["families"])
    # The token is signed, not encrypted: its payload holds no address.
    payload = signing.loads(result["preview"]["token"], salt=SALT)
    assert "@" not in json.dumps(payload) and "@" not in json.dumps(result)
    assert trail() == (0, 0, 0) and not FamilyMailTest.objects.exists()
    token = result["preview"]["token"]
    tasks = TaskRun.objects.count()
    # Refused prompts change nothing: no ticket, task, event or notice.
    for answer in (b"", b"no\n"):
        code, refused, errors = send(admin, secret, token, answer)
        assert code == 4 and refused["error"]["code"] == "confirmation_required"
        assert "names and codes will be sent to the Testing recipient" in errors
        assert "1 Families" in errors
        assert not FamilyMailTest.objects.exists() and TaskRun.objects.count() == tasks
        assert trail() == (0, 0, 0)
    code, sent, _ = send(admin, secret, token, b"yes\n")
    assert code == 0, sent
    [ticket] = FamilyMailTest.objects.all()
    assert sent["result"] == {
        "created": True,
        "request_key": str(key),
        "tickets": [
            {
                "id": str(ticket.pk),
                "sequence": 1,
                "state": ticket.state,
                "task_id": str(ticket.task_id),
            }
        ],
    }
    assert ticket.requested_by_id == row.principal_id
    # The ticket records the session's sign-in instant.
    assert ticket.reauthenticated_at == row.authenticated_at
    assert trail() == (1, 1, 1)
    [event] = AuditEvent.objects.filter(event_type=EVENT)
    assert event.actor_id == row.principal_id and event.subject_id == row.pk
    # The same token again, from here or from the page, sends nothing more.
    code, again, _ = admin("test", "families", "--token", token, "--yes", secret=secret)
    assert code == 0 and again["result"]["created"] is False, again
    assert again["result"]["tickets"] == sent["result"]["tickets"]
    harness, browser, path, _ = family_test
    with web_login():
        response = post(
            browser, path, {"action": "confirm", "preview": token, "acknowledge": "on"}
        )
    assert response.status_code == 302, response.content
    assert FamilyMailTest.objects.count() == 1
    assert AuditEvent.objects.filter(event_type=EVENT).count() == 1
    # The status read lists the ticket without its DUID.
    _, reader, _ = paired(admin.service, scope="read-only")
    code, status, _ = admin("test", "status", secret=reader)
    assert code == 0, status
    [item] = status["result"]["tickets"]
    assert item["id"] == str(ticket.pk) and "duid" not in item


def test_a_page_review_is_sent_from_the_command_line(admin, family_test):  # noqa: F811
    """The page's signed review is the command line's token."""
    _, browser, path, _ = family_test
    _, secret, _ = paired(admin.service)
    token = review(browser, path, [1]).context["confirm"]["preview"].value()
    code, document, _ = send(admin, secret, token, b"yes\n")
    assert code == 0 and document["result"]["created"] is True, document
    assert FamilyMailTest.objects.count() == 1


def test_an_ineligible_family_is_refused(admin, family_test):  # noqa: F811
    """An unknown DUID is shown as such and cannot be sent."""
    _, secret, _ = paired(admin.service)
    code, document, _ = preview(admin, secret, family_test, 1, 999999)
    assert code == 0, document
    assert document["result"]["families"][1] == {
        "duid": 999999,
        "eligible": False,
        "eligibility": "unknown",
    }
    token = document["result"]["preview"]["token"]
    code, stale, _ = send(admin, secret, token, b"yes\n")
    assert code == 1 and stale["error"]["code"] == "stale_version", stale
    assert not FamilyMailTest.objects.exists() and trail() == (0, 0, 0)


def test_refusals_change_nothing(admin, family_test):  # noqa: F811
    """Read-only scope, bad input, an altered token and an ended session."""
    _, secret, row = paired(admin.service)
    _, reader, _ = paired(admin.service, scope="read-only")
    code, document, _ = preview(admin, reader, family_test, 1)
    assert code == 1 and document["error"]["code"] == "denied"
    code, document, _ = preview(admin, secret, family_test, *range(1, 12))
    assert code == 1 and document["error"]["code"] == "invalid", document
    code, document, _ = admin(
        "test", "families-preview", str(uuid4()), "--family", "1", secret=secret
    )
    assert code == 1 and document["error"]["code"] == "not_available", document
    token = preview(admin, secret, family_test, 1)[1]["result"]["preview"]["token"]
    code, document, _ = send(admin, reader, token, b"yes\n")
    assert code == 1 and document["error"]["code"] == "denied"
    code, document, _ = send(admin, secret, token[:-2] + "xx", b"yes\n")
    assert code == 1 and document["error"]["code"] == "invalid", document
    actor = current_principal(admin.service.store, row.principal_id)
    automation.revoke(row.pk, actor, reason="revoked_by_owner")
    code, document, _ = send(admin, secret, token, b"yes\n")
    assert code == 5 and document["error"]["code"] == "session_ended"
    assert not FamilyMailTest.objects.exists() and trail() == (0, 0, 0)


@pytest.mark.parametrize(
    ("change", "expected"),
    [("unscheduled", "denied"), ("gone", "stale_version")],
)
def test_an_email_changed_since_the_preview_is_refused(
    admin,
    family_test,  # noqa: F811
    monkeypatch,
    change,
    expected,
):
    """No schedule uses the email, or its revision left the configuration.

    The send's transaction rolled back and no ticket exists, so the command
    refuses (exit 1) rather than reporting an unknown outcome (exit 6).
    """
    from parishkit.stewardship.accounts import campaign_family_test
    from parishkit.stewardship.accounts.content_models import ContentVersion

    _, secret, _ = paired(admin.service)
    token = preview(admin, secret, family_test, 1)[1]["result"]["preview"]["token"]

    def changed(*args, **kwargs):
        """The page's checks, after the schedule or configuration changed."""
        if change == "unscheduled":
            raise LookupError("This email is not used by a current schedule.")
        raise ContentVersion.DoesNotExist

    monkeypatch.setattr(campaign_family_test, "prepare", changed)
    code, document, _ = send(admin, secret, token, b"yes\n")
    assert code == 1 and document["error"]["code"] == expected, document
    assert not FamilyMailTest.objects.exists() and trail() == (0, 0, 0)


# ------------------------------------------------- the names export (#817)

NAMES_EVENT = "admin_cmd_export_family_test_names"


def names_preview(admin, secret, family_test, *duids, key, fmt="csv", zone="UTC"):  # noqa: F811
    """``test families-preview --names`` for ``duids`` under ``key``.

    ``fmt`` is the ``--format``; the command's own default is XLSX (#1190), so
    these tests ask for CSV unless they test another format.
    ``zone`` is the ``--timezone``; None leaves it out (the parish's, #1127).
    """
    argv = ["test", "families-preview", revision(family_test)]
    for duid in duids:
        argv += ["--family", str(duid)]
    argv += ["--names", "--request-key", str(key)]
    if zone is not None:
        argv += ["--timezone", zone]
    argv += ["--format", fmt]
    return admin(*argv, secret=secret)


def page_names(*duids):
    """The names the page shows for ``duids`` (its ``_family_names``)."""
    from parishkit.stewardship.source.snapshot_models import SourceCurrent
    from parishkit.stewardship.source.snapshot_names import snapshot_family_names

    current = SourceCurrent.objects.get(singleton=True)
    return snapshot_family_names(current.snapshot_id, duids, "Family")


def names_trail():
    """The export requests, export_requested events and command events so far."""
    from parishkit.stewardship.reports.export_models import ExportRequest

    return (
        ExportRequest.objects.filter(report="family_test_names").count(),
        AuditEvent.objects.filter(event_type="export_requested").count(),
        AuditEvent.objects.filter(event_type=NAMES_EVENT).count(),
    )


def test_names_reach_only_the_exported_file(
    admin,
    family_test,  # noqa: F811
    reports_root,  # noqa: F811
    google,
):
    """--names: the preview's document, an export of the names, one trail.

    The names the page shows are in the file the worker renders and
    ``export download --stream`` writes (what ``export fetch`` runs), and in
    no document. A repeat returns the same export; the key for other DUIDs
    is refused.
    """
    from parishkit.stewardship.reports.export_models import (
        ExportRequest,
        FamilyTestNamesSnapshot,
    )

    from .test_admin_export_cli_postgresql import download
    from .test_admin_family_export_cli_postgresql import render

    _, secret, row = paired(admin.service)
    key = uuid4()
    names = page_names(1)
    assert names[1]
    code, document, errors = names_preview(
        admin, secret, family_test, 1, 999999, key=key
    )
    assert code == 0, document
    result = document["result"]
    # The review is the plain preview's, with no name anywhere printed.
    assert [item["duid"] for item in result["families"]] == [1, 999999]
    assert result["request_key"] == str(key) and result["preview"]["token"]
    assert names[1] not in json.dumps(document) and names[1] not in errors
    export = result["export"]
    made = ExportRequest.objects.get(pk=export["id"])
    assert export["report"] == "family_test_names" and export["format"] == "csv"
    assert export["state"] == "queued" and export["size"] is None
    assert made.request_key == key and made.requester_id == row.principal_id
    assert made.parameters == {"revision": revision(family_test), "duids": [1, 999999]}
    # The capture holds the page's own names, in the review's order; a DUID
    # the source does not know has none, as on the page.
    capture = FamilyTestNamesSnapshot.objects.get(pk=made.family_test_names_snapshot_id)
    assert capture.document == {
        "rows": [{"duid": 1, "name": names[1]}, {"duid": 999999, "name": ""}]
    }
    assert capture.row_count == 2 and capture.actor_id == row.principal_id
    assert names_trail() == (1, 1, 1)
    # The status page's trail runs through Send to chosen Families for this
    # email, so its crumb is the way back and no "Return to" repeats it
    # (#1345).
    from django.urls import reverse

    from parishkit.stewardship.deployment import ServiceRole

    from .auth_builders import signed_in
    from .test_background_grants_postgresql import task_login
    from .test_report_workspace_postgresql import read

    browser, _ = signed_in()
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        status, body = read(browser, reverse("admin:report_export", args=[made.pk]))
    assert status.status_code == 200
    families = reverse("admin:campaign_mail_families", args=[revision(family_test)])
    assert f'<li><a href="{families}">Send to chosen Families</a></li>'.encode() in body
    assert b"Return to " not in body
    [event] = AuditEvent.objects.filter(event_type=NAMES_EVENT)
    assert event.actor_id == row.principal_id and event.subject_id == row.pk
    # No ticket: the review sends nothing.
    assert not FamilyMailTest.objects.exists() and trail() == (0, 0, 0)

    # The same key and selection again returns the same export, recording
    # nothing new; the same key for other Families is the forms' 409.
    code, again, _ = names_preview(admin, secret, family_test, 1, 999999, key=key)
    assert code == 0 and again["result"]["export"]["id"] == str(made.pk), again
    assert names_trail() == (1, 1, 1)
    code, bound, _ = names_preview(admin, secret, family_test, 1, key=key)
    assert code == 1 and bound["error"]["code"] == "invalid", bound
    assert names_trail() == (1, 1, 1)

    # The real worker renders the CSV; the stream is that file, byte for byte.
    render(admin.service, reports_root, made.pk)
    code, body, streamed = download(made.pk, secret)
    assert code == 0, streamed
    assert streamed["result"]["file_name"] == "family_test_names.csv"
    assert streamed["result"]["count"] == 2
    assert streamed["result"]["size"] == len(body)
    text = body.decode()
    # Exactly the table, as the XLSX and PDF (#966): no metadata columns.
    rows = list(csv.reader(io.StringIO(text)))
    assert rows == [["Family", "Family DUID"], [names[1], "1"], ["", "999999"]]


def test_names_export_in_xlsx_and_pdf(admin, family_test, reports_root):  # noqa: F811
    """--format: the names as XLSX or PDF, with the CSV's columns (#966).

    The format binds the request key as the selection does; the worker
    renders each file.
    """
    from openpyxl import load_workbook

    from parishkit.stewardship.reports.export_models import ExportRequest
    from parishkit.stewardship.reports.family_test_names import HEADINGS

    from .test_admin_export_cli_postgresql import download
    from .test_admin_family_export_cli_postgresql import render

    _, secret, _ = paired(admin.service)
    names = page_names(1)
    made = {}
    for fmt in ("xlsx", "pdf"):
        key = uuid4()
        code, document, _ = names_preview(
            admin, secret, family_test, 1, key=key, fmt=fmt
        )
        assert code == 0, document
        assert document["result"]["export"]["format"] == fmt
        made[fmt] = ExportRequest.objects.get(pk=document["result"]["export"]["id"])
        # The same key in another format is another selection: refused.
        code, bound, _ = names_preview(admin, secret, family_test, 1, key=key)
        assert code == 1 and bound["error"]["code"] == "invalid", bound
        render(admin.service, reports_root, made[fmt].pk)
        code, body, streamed = download(made[fmt].pk, secret)
        assert code == 0, streamed
        assert streamed["result"]["file_name"] == f"family_test_names.{fmt}"
        assert streamed["result"]["count"] == 1
        if fmt == "xlsx":
            sheet = load_workbook(io.BytesIO(body)).worksheets[0]
            rows = list(sheet.iter_rows(values_only=True))
            assert rows == [HEADINGS, (names[1], "1")]
        else:
            assert body.startswith(b"%PDF-")
    assert names_trail() == (2, 2, 2)


def test_the_constraint_still_refuses_a_png_names_export(admin, family_test):  # noqa: F811
    """export_report_known admits the names only as CSV, XLSX or PDF (#966).

    Migration 0049 widened that one arm; PNG, which the format constraint
    allows for charts, stays refused for this kind. The probe copies a real
    names request with only its identities and format changed, with triggers
    off for that one savepoint (session_replication_role), so the request
    guard's binding check cannot answer first; CHECK constraints still run.
    """
    from django.db import IntegrityError, connection, transaction

    _, secret, _ = paired(admin.service)
    code, document, _ = names_preview(
        admin, secret, family_test, 1, key=uuid4(), fmt="pdf"
    )
    assert code == 0, document
    made = document["result"]["export"]["id"]
    copy = (
        "INSERT INTO stewardship_export_request SELECT (jsonb_populate_record("
        "r, jsonb_build_object('id', %s::uuid, 'request_key', %s::uuid,"
        " 'task_id', %s::uuid,"
        " 'format', %s))).* FROM stewardship_export_request r WHERE r.id=%s"
    )
    replica = "SET LOCAL session_replication_role = replica"
    # The savepoint is inside the raises block, so it rolls back first.
    with (
        pytest.raises(IntegrityError, match="export_report_known"),
        transaction.atomic(),
        connection.cursor() as cursor,
    ):
        cursor.execute(replica)
        cursor.execute(copy, [*(str(uuid4()) for _ in range(3)), "png", made])
    # The same copy as XLSX is admitted, so the refusal is the format's.
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute(replica)
        cursor.execute(copy, [*(str(uuid4()) for _ in range(3)), "xlsx", made])
        assert cursor.rowcount == 1
        transaction.set_rollback(True)


def test_names_without_a_zone_use_the_parish_zone(admin, family_test):  # noqa: F811
    """The names file has no times in its rows, so --timezone is optional.

    Without it the file's capture and request times use the parish's time
    zone (#1127): a host command has no browser zone to send.
    """
    from parishkit.stewardship.accounts.runtime_models import SystemConfiguration
    from parishkit.stewardship.reports.export_models import ExportRequest

    _, secret, _ = paired(admin.service)
    code, document, _ = names_preview(
        admin, secret, family_test, 1, key=uuid4(), zone=None
    )
    assert code == 0, document
    parish = SystemConfiguration.objects.values_list(
        "active_configuration__parish__timezone", flat=True
    ).get()
    request = ExportRequest.objects.get(report="family_test_names")
    assert request.browser_timezone == parish
    assert names_trail() == (1, 1, 1)


def test_names_need_a_testing_review(admin, family_test):  # noqa: F811
    """--timezone needs --names; refusals export nothing."""
    _, secret, _ = paired(admin.service)
    argv = ["test", "families-preview", revision(family_test), "--family", "1"]
    code, document, _ = admin(*argv, "--timezone", "UTC", secret=secret)
    assert code == 2 and document["error"]["code"] == "usage", document
    # --format chooses the names file's format, so it needs --names.
    code, document, _ = admin(*argv, "--format", "pdf", secret=secret)
    assert code == 2 and document["error"]["code"] == "usage", document
    code, document, _ = admin(
        *argv, "--names", "--timezone", "Not/AZone", secret=secret
    )
    assert code == 1 and document["error"]["code"] == "invalid", document
    code, document, _ = admin(
        "test",
        "families-preview",
        str(uuid4()),
        "--family",
        "1",
        "--names",
        "--timezone",
        "UTC",
        secret=secret,
    )
    assert code == 1 and document["error"]["code"] == "not_available", document
    assert names_trail() == (0, 0, 0)
    # Without --names the review is unchanged: no export, nothing recorded.
    code, document, _ = admin(*argv, secret=secret)
    assert code == 0 and document["result"]["export"] is None, document
    assert names_trail() == (0, 0, 0)


def test_the_names_capture_is_guarded_in_sql(admin, family_test):  # noqa: F811
    """SQL refuses a capture that is not the review's, or not an Administrator's.

    It also refuses one in Production mode, for a campaign that is not a
    draft, or by a Staff member; none leaves a capture or an export request.
    """
    from django.db import IntegrityError, InternalError, connection, transaction

    from parishkit.stewardship.accounts.policy_models import PortalUser
    from parishkit.stewardship.accounts.runtime_models import SystemConfiguration
    from parishkit.stewardship.reports.export_models import FamilyTestNamesSnapshot

    harness = family_test[0]
    runtime = SystemConfiguration.objects.get()
    _, _, row = paired(admin.service)
    good = {
        "campaign_id": harness.campaign.pk,
        "configuration_id": runtime.active_configuration_id,
        "actor_id": row.principal_id,
        "correlation_id": uuid4(),
        "parameters": {"revision": revision(family_test), "duids": [1]},
        "document": {"rows": [{"duid": 1, "name": "Example"}]},
        "row_count": 1,
    }
    bad = [
        {"actor_id": uuid4()},
        {"parameters": {"revision": revision(family_test), "duids": [1, 1]}},
        {"parameters": {"revision": "x", "duids": [1]}},
        {"document": {"rows": [{"duid": 2, "name": "Example"}]}},
        {"document": {"rows": [{"duid": 1, "name": 5}]}},
        {"document": {"rows": [{"duid": 1, "name": "A", "code": "x"}]}},
        {"document": {"rows": []}},
        {"row_count": 2},
    ]
    for change in bad:
        expected = "unavailable" if "actor_id" in change else "does not match"
        with (
            pytest.raises((IntegrityError, InternalError), match=expected),
            transaction.atomic(),
        ):
            FamilyTestNamesSnapshot.objects.create(**(good | change))
    # A well-formed capture passes the insert guard, but with no export
    # request it is refused at commit.
    with (
        pytest.raises((IntegrityError, InternalError), match="requires its request"),
        transaction.atomic(),
    ):
        FamilyTestNamesSnapshot.objects.create(**good)
    assert not FamilyTestNamesSnapshot.objects.exists()

    # The same well-formed capture is unavailable outside Testing mode or
    # for a campaign that is no longer a draft. The guard reads the mode
    # and the campaign state, so each is set, with the tables' own guards
    # off, inside the refused transaction, which rolls the change back.
    for table, change in (
        ("stewardship_system_configuration", "mode='production'"),
        ("stewardship_campaign", "state='active'"),
    ):
        with (
            pytest.raises((IntegrityError, InternalError), match="unavailable"),
            transaction.atomic(),
        ):
            with connection.cursor() as cursor:
                cursor.execute(f"ALTER TABLE {table} DISABLE TRIGGER USER")
                cursor.execute(f"UPDATE {table} SET {change}")
                cursor.execute(f"ALTER TABLE {table} ENABLE TRIGGER USER")
            FamilyTestNamesSnapshot.objects.create(**good)
    assert SystemConfiguration.objects.get().mode == "testing"
    harness.campaign.refresh_from_db()
    assert harness.campaign.state == "draft"

    # A Staff member, under a policy that grants the role, cannot capture
    # names: only an Administrator can. The Administrator, under the same
    # new configuration, still passes the insert guard (the control).
    admin_user = PortalUser.objects.get(pk=row.principal_id)
    add_policy(
        (admin.service.store, admin_user, None, None),
        address("staff@example.org", ("staff",)),
    )
    staff = user("staff@example.org")
    current = good | {
        "configuration_id": SystemConfiguration.objects.get().active_configuration_id
    }
    for actor, expected in (
        (staff.pk, "unavailable"),
        (row.principal_id, "requires its request"),
    ):
        with (
            pytest.raises((IntegrityError, InternalError), match=expected),
            transaction.atomic(),
        ):
            FamilyTestNamesSnapshot.objects.create(**(current | {"actor_id": actor}))
    # No refusal left a capture or an export request behind.
    assert not FamilyTestNamesSnapshot.objects.exists()
    assert names_trail()[0] == 0
