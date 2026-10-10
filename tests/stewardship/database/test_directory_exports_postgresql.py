"""Actual-role complete directory capture and the shared private export lifecycle."""

import csv
import io
import json
from datetime import timedelta
from uuid import uuid4

import pytest
from django.db import DatabaseError, connection, transaction
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from parishkit.stewardship.accounts.policy_models import PortalUser
from parishkit.stewardship.audit.models import AuditEvent
from parishkit.stewardship.campaigns.read_guards import DownloadPool, ReadLimits
from parishkit.stewardship.campaigns.runtime_models import CampaignWorkGate
from parishkit.stewardship.campaigns.work_locks import work_transaction
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.jobs.dispatch import execute_hint
from parishkit.stewardship.jobs.queues import WorkQueue
from parishkit.stewardship.reports.directories import DirectoryQuery
from parishkit.stewardship.reports.directory_exports import create_directory_export
from parishkit.stewardship.reports.export_models import (
    DirectoryExportSnapshot,
    ExportRequest,
)
from parishkit.stewardship.reports.export_services import (
    TASK_TYPE,
    ExportRequestBound,
    export_status,
)
from parishkit.stewardship.reports.export_tasks import export_handler

from ..export_urls import export_action
from ..policy_factory import address
from .auth_builders import signed_in, stale_sign_in
from .campaign_builders import change
from .response_builders import response_source
from .test_background_grants_postgresql import task_login
from .test_directories_postgresql import page
from .test_export_views_postgresql import post, restricted_download_pool
from .test_information_followup_postgresql import search
from .test_policy_postgresql import user
from .test_report_workspace_postgresql import read
from .test_source_families_postgresql import prepare, promote

pytestmark = pytest.mark.django_db(transaction=True)


def test_complete_directory_capture_is_private_immutable_and_source_pinned(
    live_response_service,
):
    """One promotion proves complete postal scope and stable code selection."""
    harness = live_response_service
    data = response_source()
    for duid in range(10, 61):
        data.families[duid] = data.families[1] | {
            "familyDUID": duid,
            "familyID": duid + 100,
            "firstName": "Repeated",
            "lastName": "Family",
        }
        data.members[1000 + duid] = data.members[3] | {
            "memberDUID": 1000 + duid,
            "familyDUID": duid,
            "memberType": "Other",
            "emailAddress": "",
        }
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    actor = user("admin@example.org").pk
    values = dict(
        campaign_id=harness.campaign.pk,
        query=DirectoryQuery(sort="duid"),
        postal=False,
        mac=harness.rings.mac,
        format="csv",
        browser_timezone="UTC",
        request_key=uuid4(),
    )
    interactive = page(harness, sort="duid")
    assert interactive["total"] == 52 and len(interactive["rows"]) == 50
    neither_page = page(harness, sort="duid", reach="neither")
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        request = create_directory_export(harness.service.store, actor, **values)
        retained = DirectoryExportSnapshot.objects.get(pk=request.directory_snapshot_id)
        assert retained.row_count == len(retained.document["rows"]) == 52
        assert [row["family_id"] for row in retained.document["rows"][:50]] == [
            row["family_id"] for row in interactive["rows"]
        ]
        assert harness.code not in json.dumps(retained.document)
        assert all("code" not in row for row in retained.document["rows"])
        # The code list keeps no address, phones or envelope (#388 L6); the
        # keys stay, emptied, so the document keeps its shape.
        assert interactive["rows"][0]["address"]
        for row in retained.document["rows"]:
            assert (row["address"], row["phones"], row["envelope"]) == ({}, [], None)
            assert row["heads"] is not None
        assert (
            create_directory_export(harness.service.store, actor, **values).pk
            == request.pk
        )
        with pytest.raises(ExportRequestBound):
            create_directory_export(
                harness.service.store, actor, **(values | {"postal": True})
            )
        postal = create_directory_export(
            harness.service.store,
            actor,
            **(values | {"postal": True, "request_key": uuid4()}),
        ).directory_snapshot
        # The mail merge leaves out the one Family with deliverable email
        # (#951): it always captures reach "mail".
        assert postal.row_count == 51
        assert postal.parameters["filters"]["reach"] == "mail"
        # The mail merge keeps the address only (#388 L6).
        rows = DirectoryExportSnapshot.objects.get(pk=postal.pk).document["rows"]
        assert any(row["address"] for row in rows)
        assert all(row["phones"] == [] and row["envelope"] is None for row in rows)
        # The code list for reach "neither" keeps the phones it renders.
        neither = create_directory_export(
            harness.service.store,
            actor,
            **(
                values
                | {
                    "query": DirectoryQuery(sort="duid", reach="neither"),
                    "request_key": uuid4(),
                }
            ),
        ).directory_snapshot
        neither = DirectoryExportSnapshot.objects.get(pk=neither.pk)
        assert neither.row_count == neither_page["total"] <= 50
        assert {
            row["family_id"]: row["phones"] for row in neither.document["rows"]
        } == {row["family_id"]: row["phones"] for row in neither_page["rows"]}
        assert all(
            row["address"] == {} and row["envelope"] is None
            for row in neither.document["rows"]
        )
        assert (
            create_directory_export(
                harness.service.store,
                actor,
                **(
                    values
                    | {
                        "postal": True,
                        "query": DirectoryQuery(sort="duid", reach="mail"),
                        "request_key": uuid4(),
                    }
                ),
            ).directory_snapshot.row_count
            == 51
        )
        selected = create_directory_export(
            harness.service.store,
            actor,
            **(
                values
                | {
                    "query": DirectoryQuery(exact_code=harness.code.lower()),
                    "request_key": uuid4(),
                }
            ),
        )
        assert selected.directory_snapshot.row_count == 1
        assert selected.parameters["exact"] and selected.parameters["family_id"]
        assert harness.code.lower() not in json.dumps(selected.parameters).lower()
        assert "exact_code" not in selected.parameters["filters"]
        empty = create_directory_export(
            harness.service.store,
            actor,
            **(
                values
                | {
                    "query": DirectoryQuery(exact_code="ZZZZZZZZ"),
                    "request_key": uuid4(),
                }
            ),
        )
        assert empty.directory_snapshot.row_count == 0 and empty.parameters["exact"]
        for statement in (
            "UPDATE stewardship_directory_export_snapshot SET document='{}'",
            "DELETE FROM stewardship_directory_export_snapshot",
        ):
            with (
                pytest.raises(DatabaseError),
                work_transaction(),
                connection.cursor() as cursor,
            ):
                cursor.execute(statement)
        with pytest.raises(DatabaseError), work_transaction():
            orphan = DirectoryExportSnapshot.objects.create(
                campaign_id=harness.campaign.pk,
                configuration_id=request.configuration_id,
                actor_id=actor,
                correlation_id=uuid4(),
                parameters=request.parameters,
                document={"rows": [{"code": "FORGED"}]},
            )
            orphan.refresh_from_db()
            assert "FORGED" not in json.dumps(orphan.document)
    before = retained.document
    data.families[1]["lastName"] = "Later source name"
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    retained.refresh_from_db()
    assert retained.document == before
    assert page(harness)["metadata"]["source_id"] != str(retained.source_id)


def test_native_directory_exports_render_download_and_regenerate_retained_inputs(
    live_response_service, google, tmp_path, settings, monkeypatch
):
    """One fixture covers all formats, guarded download and source-change replay."""
    harness = live_response_service
    browser, login = signed_in()
    assert login.status_code == 302
    actor = user("admin@example.org").pk
    root = tmp_path / "directory-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    route = reverse("admin:family_directory")
    first = None
    for format in ("csv", "xlsx", "pdf"):
        fields = DirectoryQuery(exact_code=harness.code).form_values() | dict(
            format=format, browser_timezone="UTC", request_key=str(uuid4())
        )
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            response, body = read(browser, route)
            assert response.status_code == 200 and b"Queue complete export" in body
            assert (
                browser.post(
                    reverse("admin:family_directory_export"), fields
                ).status_code
                == 403
            )
            if format == "csv":
                assert (
                    browser.get(reverse("admin:family_directory_export")).status_code
                    == 405
                )
                for invalid in (
                    {"page": "2"},
                    {"family_id": str(uuid4())},
                    {"format": ["csv", "pdf"]},
                ):
                    rejected = post(
                        browser,
                        reverse("admin:family_directory_export"),
                        fields | invalid,
                    )
                    assert (
                        rejected.status_code == 400
                        and harness.code.encode() not in rejected.content
                    )
                    assert rejected["Cache-Control"] == "no-store"
            response = post(browser, reverse("admin:family_directory_export"), fields)
            assert response.status_code == 302
            request = ExportRequest.objects.get(request_key=fields["request_key"])
            first = first or request
            assert (
                post(browser, reverse("admin:family_directory_export"), fields)[
                    "Location"
                ]
                == response["Location"]
            )
            # The same one-time key reused for another format (a page back
            # from the browser cache) is refused clearly, and queues nothing.
            other = "pdf" if format != "pdf" else "csv"
            before = ExportRequest.objects.count()
            reused = post(
                browser,
                reverse("admin:family_directory_export"),
                fields | {"format": other},
            )
            assert reused.status_code == 409
            assert b"already used for a different export" in reused.content
            assert ExportRequest.objects.count() == before
            job_route = response["Location"]
            with CaptureQueriesContext(connection) as queries:
                response, body = read(browser, job_route)
            assert all(
                '"stewardship_directory_export_snapshot"."document"' not in item["sql"]
                for item in queries.captured_queries
            )
            assert response.status_code == 200 and b"Family-directory export" in body
            # A code-list export with no reach filter returns to the plain page.
            assert f'href="{route}">'.encode() in body
            assert b"mail-merge export" not in body
        with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
            assert execute_hint(
                request.task_id,
                queue=WorkQueue.GENERAL,
                worker_id=uuid4(),
                handlers={
                    TASK_TYPE: export_handler(
                        store=harness.service.store,
                        root=root,
                        general=harness.rings.general,
                    )
                },
            )
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            assert (
                export_status(harness.service.store, actor, request.pk)["state"]
                == "ready"
            )
        with restricted_download_pool(settings):
            response, body = search(browser, export_action(job_route, "download"), {})
            assert response.status_code == 200
            assert body.startswith(
                {
                    "csv": b"Family,ParishSoft DUID,Family code,Family head emails\r\n",
                    "xlsx": b"PK",
                    "pdf": b"%PDF",
                }[format]
            )
            assert "family_directory." + format in response["Content-Disposition"]
            if format == "csv":
                # The code list names Families, codes and head emails (read
                # from the captured source, #604), not addresses.
                assert harness.code.encode() in body
                assert b"1 Example Street" not in body
                assert body.endswith(b",Member Example: valid@example.org\r\n")
    data = response_source()
    data.families[1]["lastName"] = "Later changed name"
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    from parishkit.stewardship.reports import export_services

    with work_transaction():
        now = export_services.database_now()
    monkeypatch.setattr(
        export_services, "database_now", lambda: now + timedelta(days=8)
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        fields = {"request_key": str(uuid4())}
        response = post(
            browser, reverse("admin:report_export_regenerate", args=[first.pk]), fields
        )
        assert response.status_code == 302
        regenerated = ExportRequest.objects.get(request_key=fields["request_key"])
        assert regenerated.directory_snapshot_id == first.directory_snapshot_id
        assert (
            regenerated.directory_snapshot.document["rows"][0]["family_name"]
            == "Example"
        )


def test_directory_export_staff_gates_and_service_boundaries(
    live_response_service, google
):
    """Check Staff capture, preparation gates and metadata-only service roles."""
    harness = live_response_service
    store = harness.service.store
    rule = address("reader@example.org", roles=("staff", "ministry_leader"))
    change(
        store,
        store.active(),
        uuid4(),
        [{"operation": "add", "section": "login_rules", **rule}],
    )
    google[0]["email"] = "reader@example.org"
    browser, login = signed_in()
    assert login.status_code == 302
    actor = PortalUser.objects.get(google_subject=google[0]["sub"]).pk
    values = dict(
        campaign_id=harness.campaign.pk,
        query=DirectoryQuery(),
        postal=True,
        format="csv",
        browser_timezone="UTC",
        request_key=uuid4(),
        mac=harness.rings.mac,
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        request = create_directory_export(store, actor, **values)
        # The mail merge lists only Families without deliverable head email
        # (#951); the one Family here has it, so the file has no rows.
        assert (
            request.report == "postal_outreach"
            and request.directory_snapshot.row_count == 0
        )
    with (
        task_login(ServiceRole.SCHEDULER, exact=True, reconnect=True),
        connection.cursor() as cursor,
    ):
        cursor.execute("SELECT id,row_count FROM stewardship_directory_export_snapshot")
        assert cursor.fetchone() == (request.directory_snapshot_id, 0)
        with pytest.raises(DatabaseError):
            cursor.execute("SELECT document FROM stewardship_directory_export_snapshot")
    for role in (ServiceRole.WORKER, ServiceRole.SCHEDULER):
        with (
            task_login(role, exact=True, reconnect=True),
            pytest.raises(DatabaseError),
            work_transaction(),
        ):
            DirectoryExportSnapshot.objects.create(
                campaign_id=harness.campaign.pk,
                configuration_id=request.configuration_id,
                actor_id=actor,
                correlation_id=uuid4(),
                parameters=request.parameters,
            )
    # Simulate only the eventual purge owner's sentinel in this disposable DB.
    # No application purge executor or retained user data is involved.
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute(
            "ALTER TABLE stewardship_campaign_work_gate DISABLE TRIGGER USER"
        )
        CampaignWorkGate.objects.create(
            campaign=harness.campaign,
            request_id=uuid4(),
            initiated_by_id=uuid4(),
            state="preparing",
        )
        cursor.execute("SET CONSTRAINTS ALL IMMEDIATE")
        cursor.execute("ALTER TABLE stewardship_campaign_work_gate ENABLE TRIGGER USER")
    route = reverse("admin:family_directory")
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        with pytest.raises(PermissionError):
            create_directory_export(store, actor, **(values | {"request_key": uuid4()}))
        response, body = read(browser, route + "?mailing=yes")
        assert response.status_code == 200 and b"<fieldset disabled>" in body
        assert b"Other campaign work" in body
        assert (
            read(browser, reverse("admin:report_export", args=[request.pk]))[
                0
            ].status_code
            == 200
        )
    change(
        store,
        store.active(),
        uuid4(),
        [
            {
                "operation": "update",
                "section": "login_rules",
                "id": rule["id"],
                "values": {
                    "roles": ["ministry_leader"],
                    "grants": {
                        "ministry_leader": rule["values"]["grants"]["ministry_leader"]
                    },
                },
            }
        ],
    )
    fields = DirectoryQuery().form_values() | {
        "mailing": "yes",
        "format": "csv",
        "browser_timezone": "UTC",
        "request_key": str(uuid4()),
    }
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert (
            post(browser, reverse("admin:family_directory_export"), fields).status_code
            == 403
        )
        # Mailing columns (addresses) stay with the roles that could open the
        # old postal page.
        assert read(browser, route + "?mailing=yes")[0].status_code == 403
        denied = post(browser, route, {"mailing": "yes", "reach": "email"})
        assert denied.status_code == 403 and b"Example Street" not in denied.content
        # The old postal routes are retired with no redirect (#758, #864).
        legacy = f"/admin/reports/{harness.campaign.pk}/postal/"
        assert browser.get(legacy).status_code == 404
        assert post(browser, legacy, {}).status_code == 404
        assert post(browser, legacy + "export", fields).status_code == 404
        assert (
            read(browser, reverse("admin:report_export", args=[request.pk]))[
                0
            ].status_code
            == 403
        )


def test_postal_mail_merge_leaves_out_families_without_a_mailing_address(
    live_response_service, google, tmp_path, settings
):
    """The postal file holds the Families postal invitations are for (#951).

    None of the three Families has a head email. Two have a street line but
    no city, so no usable mailing address: postal mail cannot reach them, so
    they are on the "neither" list, not in the mail merge, and a partial
    address never prints. The page and the file list the same one Family.
    """
    harness = live_response_service
    data = response_source()
    data.members[3]["emailAddress"] = ""
    for duid in (10, 11):
        data.families[duid] = data.families[1] | {
            "familyDUID": duid,
            "familyID": duid + 100,
            "lastName": "Unmailable",
            "primaryCity": "",
        }
        data.members[1000 + duid] = data.members[3] | {
            "memberDUID": 1000 + duid,
            "familyDUID": duid,
            "emailAddress": "",
        }
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    browser, _ = signed_in()
    root = tmp_path / "postal-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    route = reverse("admin:family_directory")
    fields = DirectoryQuery().form_values() | dict(
        mailing="yes", format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        response, body = read(browser, route + "?mailing=yes")
        # A campaign-wide fact, not a claim about this file's contents.
        assert (
            b"Across all active Families, 2 Families have neither a deliverable" in body
        )
        # The "neither" list holds them; mailing columns always list "By
        # postal mail only", so asking them for "neither" lists only Family 1.
        response, neither = search(browser, route, {"reach": "neither"})
        assert response.status_code == 200 and b"Unmailable" in neither
        response, mailing = search(
            browser, route, {"mailing": "yes", "reach": "neither"}
        )
        assert response.status_code == 200 and b"Unmailable" not in mailing
        assert b"Matching Families: 1." in mailing
        # The mailing columns show what the mail-merge file will hold.
        assert b"Matching Families: 1." in body and b"Unmailable" not in body
        for text in (
            b"Addressee",
            b"Mailing address",
            b"Every listed Family is in the mail-merge file: the active "
            b"Families with no deliverable head email and a usable mailing "
            b"address.",
            b"<td>Member Example</td>",
            b'name="mailing" value="yes" checked',
            b'<input type="hidden" name="mailing" value="yes">',
            b"ParishSoft DUID, Family, Addressee, Family heads,",
        ):
            assert text in body
        response = post(browser, reverse("admin:family_directory_export"), fields)
        assert response.status_code == 302
    request = ExportRequest.objects.get(request_key=fields["request_key"])
    with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
        assert execute_hint(
            request.task_id,
            queue=WorkQueue.GENERAL,
            worker_id=uuid4(),
            handlers={
                TASK_TYPE: export_handler(
                    store=harness.service.store,
                    root=root,
                    general=harness.rings.general,
                )
            },
        )
    with restricted_download_pool(settings):
        response, body = search(
            browser, export_action(response["Location"], "download"), {}
        )
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(body.decode())))
    assert rows[0][:4] == ["ParishSoft DUID", "Family", "Addressee", "Family heads"]
    assert [row[:2] for row in rows[1:]] == [["1", "Example"]]
    assert "40000" in ",".join(rows[1]) and harness.code in rows[1]
    # The mail merge ends with the head emails column (#604); this head has
    # none, and the row still lists them.
    assert rows[0][-1] == "Family head emails" and len(rows[0]) == 12
    assert rows[1][11] == "Member Example: (no email)"
    assert request.directory_snapshot.row_count == 1
    assert request.parameters["filters"]["reach"] == "mail"
    assert request.report == "postal_outreach"
    # The old postal export address is retired with no redirect (#758,
    # #864) and queues nothing.
    legacy = DirectoryQuery().form_values() | dict(
        format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    old_route = f"/admin/reports/{harness.campaign.pk}/postal/export"
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert post(browser, old_route, legacy).status_code == 404
    assert not ExportRequest.objects.filter(request_key=legacy["request_key"]).exists()


def test_mail_merge_always_lists_by_postal_mail_only(
    live_response_service, google, tmp_path, settings
):
    """Mailing columns always use "By postal mail only" (#951).

    The one Family has no head email and a mailing address. Asked for "By
    email" with mailing columns on, the page, the capture and the mail
    merge all use reach ``mail`` and hold it; the page keeps the reach asked
    for, and the status page returns with the mailing columns alone.
    """
    harness = live_response_service
    data = response_source()
    data.members[3]["emailAddress"] = ""
    snapshot, claim = prepare(data)
    promote(snapshot, claim, harness.campaign, harness.rings)
    browser, _ = signed_in()
    root = tmp_path / "mailing-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    route = reverse("admin:family_directory")
    filters = DirectoryQuery(reach="email").form_values() | {"mailing": "yes"}
    fields = filters | dict(
        format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        response, body = search(browser, route, filters)
        assert response.status_code == 200 and b"Matching Families: 1." in body
        assert b"<td>1 Example Street<br>" in body
        assert b'<option value="mail" selected>' in body
        # The export form carries the page's own reach; the service applies
        # "mail" when it binds the selection.
        assert b'<input type="hidden" name="reach" value="email">' in body
        response = post(browser, reverse("admin:family_directory_export"), fields)
        assert response.status_code == 302
    request = ExportRequest.objects.get(request_key=fields["request_key"])
    assert request.report == "postal_outreach"
    assert request.parameters["postal"] is True
    assert request.parameters["filters"]["reach"] == "mail"
    assert request.directory_snapshot.row_count == 1
    # The status page names the mail merge and returns with the mailing
    # columns on (closed presets only), which apply reach "mail" themselves.
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        status, body = read(browser, response["Location"])
    assert status.status_code == 200
    assert b"Family-directory mail-merge export" in body
    assert f'href="{route}?mailing=yes"'.encode() in body
    with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
        assert execute_hint(
            request.task_id,
            queue=WorkQueue.GENERAL,
            worker_id=uuid4(),
            handlers={
                TASK_TYPE: export_handler(
                    store=harness.service.store,
                    root=root,
                    general=harness.rings.general,
                )
            },
        )
    with restricted_download_pool(settings):
        response, body = search(
            browser, export_action(response["Location"], "download"), {}
        )
    assert response.status_code == 200
    lines = body.decode().splitlines()
    assert lines[0].startswith("ParishSoft DUID,Family,Addressee,Family heads,")
    assert len(lines) == 2 and lines[1].startswith("1,Example,Member Example,")
    assert "1 Example Street" in lines[1] and harness.code in lines[1]
    assert lines[1].endswith(",Member Example: (no email)")
    events = AuditEvent.objects.filter(event_type="postal_outreach_viewed")
    assert events.exists()


def _compact(monkeypatch, snapshot_id):
    """Compact one superseded source snapshot now, as the refresh would later.

    Real compaction keeps recent, anchored and referenced snapshots for a
    while; the 15-minute refresh compacts an unchanged previous snapshot at
    once. This narrows the real owner's candidate rules to ``snapshot_id``
    only, so the test needs no waiting and no other snapshot is touched.
    """
    from django.db.models import Q

    from parishkit.stewardship.source import compaction
    from parishkit.stewardship.source.snapshot_models import SourceSnapshot

    from .test_source_compaction_postgresql import cleanup

    def eligible(*args, **kwargs):
        """The "older than the recent window" rule selects just this one."""
        if "promoted_at__lt" in kwargs:
            return Q(pk=snapshot_id)
        return Q(*args, **kwargs)

    with monkeypatch.context() as patched:
        patched.setattr(compaction, "Q", eligible)
        patched.setattr(compaction, "_retained_anchors", lambda *args: set())
        patched.setattr(
            compaction,
            "LIVE_REFERENCES",
            "SELECT id FROM stewardship_source_snapshot "
            f"WHERE id<>'{snapshot_id}'::uuid",
        )
        while cleanup().snapshot_count:
            pass
    assert SourceSnapshot.objects.get(pk=snapshot_id).compacted_at is not None


def test_regenerated_and_retried_exports_read_current_head_emails_after_compaction(
    live_response_service, google, tmp_path, settings, monkeypatch
):
    """Regeneration and retry still render once the capture's source is gone.

    The 15-minute refresh compacts a capture's source long before a file
    expires (7 days). The same Families, names and codes are rendered from
    the capture; the same heads' emails come from the current ParishSoft data,
    and the file's report details say "Head emails as of" its refresh (#604).
    """
    from openpyxl import load_workbook

    from parishkit.stewardship.jobs.models import TaskRun
    from parishkit.stewardship.jobs.storage import _status
    from parishkit.stewardship.reports import export_services, export_tasks
    from parishkit.stewardship.reports.directory_documents import HEAD_EMAILS_DETAIL
    from parishkit.stewardship.reports.export_services import retry_export
    from parishkit.stewardship.source.snapshot_models import SourceCurrent

    from .test_taskrun_postgresql import act

    harness = live_response_service
    browser, _ = signed_in()
    actor = user("admin@example.org").pk
    root = tmp_path / "directory-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    captured = SourceCurrent.objects.get().snapshot_id

    def queue():
        """Queue one XLSX code list (its report details are a sheet)."""
        fields = DirectoryQuery().form_values() | dict(
            format="xlsx", browser_timezone="UTC", request_key=str(uuid4())
        )
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            response = post(browser, reverse("admin:family_directory_export"), fields)
        assert response.status_code == 302
        request = ExportRequest.objects.get(request_key=fields["request_key"])
        assert request.directory_snapshot.source_id == captured
        return request, response["Location"]

    def render(run_id):
        """The real worker renders and publishes one export task run."""
        with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
            return execute_hint(
                run_id,
                queue=WorkQueue.GENERAL,
                worker_id=uuid4(),
                handlers={
                    TASK_TYPE: export_handler(
                        store=harness.service.store,
                        root=root,
                        general=harness.rings.general,
                    )
                },
            )

    def download(job_route):
        """The file's Families sheet rows and its report details."""
        with restricted_download_pool(settings):
            response, body = search(browser, export_action(job_route, "download"), {})
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(body))
        rows = [list(row) for row in book["Families"].iter_rows(values_only=True)]
        details = {
            row[0].value: row[1].value for row in book["Report information"].iter_rows()
        }
        book.close()
        return rows, details

    # A file rendered while its capture's source exists uses that source and
    # has no "as of" detail.
    first, first_route = queue()
    assert render(first.task_id)
    rows, details = download(first_route)
    assert rows[1][-1] == "Member Example: valid@example.org"
    assert HEAD_EMAILS_DETAIL not in details
    # A second export's first attempt fails after rendering from the capture.
    retried, retried_route = queue()

    def crash(*args):
        """The first attempt dies before publishing its file."""
        raise RuntimeError("synthetic render failure")

    with monkeypatch.context() as patched:
        patched.setattr(export_tasks, "write_artifact", crash)
        with pytest.raises(RuntimeError, match="synthetic render failure"):
            render(retried.task_id)
    with work_transaction():
        act(_status(TaskRun.objects.get(pk=retried.task_id)), "permanent_failure")
    # The next refresh changes the head's email, and the capture's source is
    # compacted.
    data = response_source()
    data.members[3]["emailAddress"] = "changed@example.org"
    snapshot, claim = prepare(data)
    current = promote(snapshot, claim, harness.campaign, harness.rings)
    _compact(monkeypatch, captured)
    # The retry renders: same Family and code, current email, dated.
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        retry = retry_export(
            harness.service.store, actor, retried.pk, request_key=uuid4()
        )
    assert render(retry.run_id)
    rows, details = download(retried_route)
    assert rows[1][:3] == ["Example, Member", "1", harness.code]
    assert rows[1][-1] == "Member Example: changed@example.org"
    # The refresh that promoted the current data (XLSX keeps milliseconds).
    as_of = current.promoted_at.replace(tzinfo=None)
    assert abs(details[HEAD_EMAILS_DETAIL] - as_of) < timedelta(milliseconds=1)
    # Regenerating the expired first file renders the same way.
    with work_transaction():
        now = export_services.database_now()
    # Only the regeneration request sees the clock past the first file's
    # expiry; its new file is then fresh at the real time.
    with (
        monkeypatch.context() as patched,
        task_login(ServiceRole.WEB, exact=True, reconnect=True),
    ):
        patched.setattr(
            export_services, "database_now", lambda: now + timedelta(days=8)
        )
        fields = {"request_key": str(uuid4())}
        response = post(
            browser, reverse("admin:report_export_regenerate", args=[first.pk]), fields
        )
    assert response.status_code == 302
    regenerated = ExportRequest.objects.get(request_key=fields["request_key"])
    assert regenerated.directory_snapshot_id == first.directory_snapshot_id
    assert render(regenerated.task_id)
    rows, details = download(response["Location"])
    assert rows[1][:3] == ["Example, Member", "1", harness.code]
    assert rows[1][-1] == "Member Example: changed@example.org"
    assert HEAD_EMAILS_DETAIL in details


def test_directory_export_and_regenerate_need_a_fresh_sign_in(
    live_response_service, google, tmp_path, settings, monkeypatch
):
    """A stale sign-in queues nothing and returns to the page the form came from.

    After the step-up the same form queues exactly one export, and posting it
    again returns that export rather than a second one (#547). Regenerate of a
    directory export asks for the same fresh sign-in, returning to the
    export's status page.
    """
    harness = live_response_service
    browser, _ = signed_in()
    route = reverse("admin:family_directory")
    export = reverse("admin:family_directory_export")
    fields = DirectoryQuery(exact_code=harness.code).form_values() | dict(
        format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    stale_sign_in()
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        for path, mailing, back in (
            (export, "no", route),
            (export, "yes", route + "?mailing=yes"),
        ):
            values = dict(fields)
            values["mailing"] = mailing
            refused = post(browser, path, values, HTTP_ACCEPT="text/html")
            assert refused.status_code == 403
            page = refused.content.decode()
            assert "Confirm with Google" in page and "Nothing was done" in page
            assert f'name="next" value="{back}"' in page
            assert (
                "You will then return to Active parishioner family directory." in page
            )
            assert harness.code not in page
        assert not ExportRequest.objects.exists()
    signed_in(browser)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        created = post(browser, export, fields)
        assert created.status_code == 302
        assert post(browser, export, fields)["Location"] == created["Location"]
        assert ExportRequest.objects.count() == 1
        request = ExportRequest.objects.get()
    stale_sign_in()
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        refused = post(
            browser,
            reverse("admin:report_export_regenerate", args=[request.pk]),
            {"request_key": str(uuid4())},
            HTTP_ACCEPT="text/html",
        )
        assert refused.status_code == 403
        page = refused.content.decode()
        assert f'name="next" value="/admin/reports/exports/{request.pk}/"' in page
        assert "You will then return to Report export." in page
        assert ExportRequest.objects.count() == 1
    # Render the file and let it expire; after the step-up, regenerating works.
    root = tmp_path / "directory-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
        assert execute_hint(
            request.task_id,
            queue=WorkQueue.GENERAL,
            worker_id=uuid4(),
            handlers={
                TASK_TYPE: export_handler(
                    store=harness.service.store,
                    root=root,
                    general=harness.rings.general,
                )
            },
        )
    from parishkit.stewardship.reports import export_services

    with work_transaction():
        now = export_services.database_now()
    monkeypatch.setattr(
        export_services, "database_now", lambda: now + timedelta(days=8)
    )
    signed_in(browser)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        key = str(uuid4())
        regenerated = post(
            browser,
            reverse("admin:report_export_regenerate", args=[request.pk]),
            {"request_key": key},
        )
        assert regenerated.status_code == 302
        assert ExportRequest.objects.get(request_key=key).pk != request.pk
