"""The Campaign comparison downloads (#1018, slice 6), on PostgreSQL.

The comparison scenario (``test_campaign_comparison_postgresql``: a real
financial campaign with live pledges, a ready fact set and synthetic
imported archives; every value invented) is downloaded through the page's
form in each format, rendered by the real export worker and downloaded,
under the restricted web, worker and download logins. The capture's SQL
guard, the export kinds and the frozen file's closing check are exercised
directly.
"""

import csv
import io
from datetime import timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from django.db import DatabaseError, connection, transaction
from django.db.models import F
from django.urls import reverse
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy_models import PortalUser
from parishkit.stewardship.accounts.runtime_models import SystemConfiguration
from parishkit.stewardship.audit.models import AuditContext
from parishkit.stewardship.campaigns.read_guards import DownloadPool, ReadLimits
from parishkit.stewardship.campaigns.work_locks import work_transaction
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.jobs.dispatch import execute_hint
from parishkit.stewardship.jobs.queues import WorkQueue
from parishkit.stewardship.reports import export_services
from parishkit.stewardship.reports.comparison import PAGE_ONLY, WITHHELD
from parishkit.stewardship.reports.comparison_exports import create_comparison_export
from parishkit.stewardship.reports.export_models import (
    ComparisonExportSnapshot,
    ExportRequest,
)
from parishkit.stewardship.reports.export_services import TASK_TYPE, export_status
from parishkit.stewardship.reports.export_tasks import export_handler
from parishkit.stewardship.web.dates import parish_zone

from ..export_urls import export_action
from .auth_builders import signed_in
from .campaign_builders import campaign_clock
from .test_admin_report_cli_postgresql import cli  # noqa: F401
from .test_background_grants_postgresql import task_login
from .test_campaign_comparison_postgresql import (  # noqa: F401
    CURRENT_YEAR,
    archive_bytes,
    import_archive,
    scenario,
    stewardship_year,
)
from .test_export_views_postgresql import post, restricted_download_pool
from .test_information_followup_postgresql import search
from .test_report_workspace_postgresql import read

pytestmark = pytest.mark.django_db(transaction=True)

ROOT = Path(__file__).resolve().parents[3]
FROZEN = (
    ROOT / "src/parishkit/stewardship/schema/migrations/0071_comparison_exports.sql"
)
# Note texts the synthetic archives carry, which the page shows and no
# download may.
NOTES = (b"Invented test cohort", b"Synthetic data for tests")
ROUTE = "/admin/reports/comparison/exports/"


def imported(scenario):  # noqa: F811
    """Last year (3 pledging Families, so withheld) and the year before."""
    harness, _, _, actor = scenario
    store = harness.service.store
    import_archive(store, actor, archive_bytes(CURRENT_YEAR - 1, [1, 2, 99]))
    import_archive(
        store, actor, archive_bytes(CURRENT_YEAR - 2, [1, 99, 500, 501, 502, 503])
    )


def day_two(harness):
    """The clock at noon of the campaign's day 3, so day 2 is compared."""
    opened = harness.campaign.active_configuration.starts_at
    return campaign_clock(opened + timedelta(days=2, hours=12))


def work(harness, request, root):
    """Render one export with the real worker, as the worker login."""
    with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
        assert execute_hint(
            request.task_id,
            queue=WorkQueue.GENERAL,
            worker_id=uuid4(),
            handlers={
                TASK_TYPE: export_handler(store=harness.service.store, root=root)
            },
        )


def test_download_form_queues_each_format_from_a_protected_capture(
    scenario,  # noqa: F811
    settings,
    tmp_path,
    monkeypatch,
):
    """The page's form, the worker's three files, the audit and regeneration."""
    harness, _, browser, actor = scenario
    imported(scenario)
    root = tmp_path / "comparison-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    assert reverse("admin:campaign_comparison_export") == ROUTE
    first = None
    for format in ("csv", "xlsx", "pdf"):
        fields = dict(
            day="2", format=format, browser_timezone="UTC", request_key=str(uuid4())
        )
        with day_two(harness), task_login(ServiceRole.WEB, exact=True, reconnect=True):
            response, body = read(browser, reverse("admin:campaign_comparison"))
            assert response.status_code == 200
            assert b"Queue download" in body and b'name="day" value="2"' in body
            assert b"<fieldset disabled" not in body
            # No CSRF header, a GET, extra or repeated fields, a query, or a
            # day the campaign has no figures for are refused.
            assert browser.post(ROUTE, fields).status_code == 403
            if format == "csv":
                assert browser.get(ROUTE).status_code == 405
                for invalid in (
                    {"extra": "x"},
                    {"format": ["csv", "pdf"]},
                    {"day": "9"},
                    {"day": "02"},
                    {"format": "png"},
                ):
                    rejected = post(browser, ROUTE, fields | invalid)
                    assert rejected.status_code == 400, invalid
                    assert rejected["Cache-Control"] == "no-store"
                assert post(browser, ROUTE + "?day=2", fields).status_code == 400
            response = post(browser, ROUTE, fields)
            assert response.status_code == 302
            request = ExportRequest.objects.get(request_key=fields["request_key"])
            first = first or request
            # The same form again is the same export.
            assert post(browser, ROUTE, fields)["Location"] == response["Location"]
            job = response["Location"]
            response, body = read(browser, job)
            assert response.status_code == 200
            assert b"Campaign comparison download" in body
            assert b"Comparison day" in body
        snapshot = request.comparison_snapshot
        assert (request.report, request.parameters) == (
            "campaign_comparison",
            {"day": 2},
        )
        assert snapshot.parameters == {"day": 2}
        # This campaign and the two imported years.
        assert snapshot.row_count == 3
        stored = str(snapshot.document).encode()
        assert not any(note in stored for note in NOTES)
        work(harness, request, root)
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            assert export_status(harness.service.store, actor, request.pk)["state"] == (
                "ready"
            )
        with restricted_download_pool(settings):
            response, body = search(browser, export_action(job, "download"), {})
        assert response.status_code == 200
        assert not any(note in body for note in NOTES)
        if format == "csv":
            rows = list(csv.reader(io.StringIO(body.decode())))
            assert rows[0][:4] == ["Year", "Day 1", "Days", "Families on the day"]
            table = {row[0]: dict(zip(rows[0], row, strict=True)) for row in rows[1:]}
            current = table[next(name for name in table if "(so far)" in name)]
            assert current["Families on the day"] == "4"
            assert current["Median"] == PAGE_ONLY
            # Last year had 3 pledging Families: its averages and quartiles
            # are withheld, its totals stay.
            last = table[str(CURRENT_YEAR - 1)]
            assert last["Average pledge on the day"] == WITHHELD
            assert last["Median"] == WITHHELD
            assert last["Final pledges"] == "600.00"
        elif format == "xlsx":
            book = load_workbook(io.BytesIO(body))
            sheet = book["Every number"]
            assert sheet.cell(1, 1).value == "Year"
            assert sheet.cell(2, 4).value == 4
            assert "Report information" in book.sheetnames
        else:
            assert body.startswith(b"%PDF")
    contexts = list(
        AuditContext.objects.filter(
            event__event_type="campaign_comparison_exported"
        ).values_list("context", flat=True)
    )
    assert len(contexts) == 3
    assert all(c == {"outcome": "started", "count": 2} for c in contexts)
    # Regeneration after expiry renders the same capture, never a fresh read.
    with work_transaction():
        now = export_services.database_now()
    monkeypatch.setattr(
        export_services, "database_now", lambda: now + timedelta(days=8)
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        fields = {"request_key": str(uuid4())}
        response = post(
            browser, reverse("admin:report_export_regenerate", args=[first.pk]), fields
        )
        assert response.status_code == 302
        again = ExportRequest.objects.get(request_key=fields["request_key"])
        assert again.comparison_snapshot_id == first.comparison_snapshot_id
        assert again.parameters == {"day": 2}


def capture_row(harness, actor, document=None, **values):
    """A capture as the web would store it, with ``document`` changed."""
    configuration = SystemConfiguration.objects.get().active_configuration_id
    payload = {
        "campaign": "Stewardship 2055",
        "day": 2,
        "report_date": "2054-10-05",
        "financial": True,
        "updating": False,
        "compared_with": "",
        "tiles": [{"label": "Families responded", "value": "4", "lines": []}],
        "columns": ["Year", "Day 1", "Days"],
        "rows": [["2055", {"date": "2054-10-04"}, 2]],
        "charts": [],
    } | (document or {})
    return ComparisonExportSnapshot(
        campaign_id=harness.campaign.pk,
        configuration_id=configuration,
        actor_id=actor,
        correlation_id=uuid4(),
        parameters=values.get("parameters", {"day": 2}),
        document=payload,
    )


@pytest.mark.parametrize(
    "document, parameters",
    [
        # A key a note could ride in.
        ({"notes": ["Invented test cohort"]}, None),
        # The document's day is not the request's.
        ({"day": 3}, None),
        ({"day": None, "report_date": None}, {"day": 0}),
        # A row with a cell too many, and a malformed amount.
        ({"rows": [["2055", {"date": "2054-10-04"}, 2, 4]]}, None),
        ({"rows": [["2055", {"money": "12.5"}, 2]]}, None),
        ({"rows": [["2055", {"money": "1.00", "note": "x"}, 2]]}, None),
        # A tile or chart outside its shape.
        ({"tiles": [{"label": "x", "value": "1", "lines": [], "note": "y"}]}, None),
        (
            {
                "charts": [
                    {
                        "title": "t",
                        "axis": "a",
                        "kind": "count",
                        "last_day": 2,
                        "lines": [{"role": "family", "name": "x", "points": []}],
                    }
                ]
            },
            None,
        ),
    ],
)
def test_the_capture_guard_refuses_a_document_outside_its_shape(
    scenario,  # noqa: F811
    document,
    parameters,
):
    """The web login cannot store a capture the downloads would not make."""
    harness, _, _, actor = scenario
    row = capture_row(
        harness, actor, document, **({"parameters": parameters} if parameters else {})
    )
    with (
        task_login(ServiceRole.WEB, exact=True, reconnect=True),
        pytest.raises(DatabaseError, match="Invalid comparison export"),
        transaction.atomic(),
    ):
        row.save(force_insert=True)


def test_a_capture_needs_its_request_and_never_changes(scenario):  # noqa: F811
    """A lone capture is refused at commit; the worker stores none; rows are fixed."""
    harness, _, _, actor = scenario
    # The binding is checked when the transaction commits, so the refusal
    # comes as the atomic block (inside the raises) ends.
    with (
        task_login(ServiceRole.WEB, exact=True, reconnect=True),
        pytest.raises(DatabaseError, match="requires its request"),
        transaction.atomic(),
    ):
        capture_row(harness, actor).save(force_insert=True)
    with (
        task_login(ServiceRole.WORKER, exact=True, reconnect=True),
        pytest.raises(DatabaseError),
        transaction.atomic(),
    ):
        capture_row(harness, actor).save(force_insert=True)
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        request = create_comparison_export(
            harness.service.store,
            actor,
            campaign_id=harness.campaign.pk,
            day=2,
            format="csv",
            browser_timezone="UTC",
            request_key=uuid4(),
            document=capture_row(harness, actor).document,
        )
        snapshot = request.comparison_snapshot
        assert snapshot.row_count == 1
    # The web login has no UPDATE or DELETE grant; even the owner's are
    # refused by the capture's own trigger.
    for statement in (
        "UPDATE stewardship_comparison_export_snapshot SET row_count=9",
        "DELETE FROM stewardship_comparison_export_snapshot",
    ):
        with (
            pytest.raises(DatabaseError, match="immutable"),
            transaction.atomic(),
            connection.cursor() as cursor,
        ):
            cursor.execute(statement + " WHERE id=%s", [snapshot.pk])


def test_download_form_takes_the_browsers_zone_unasked(
    scenario,  # noqa: F811
    settings,
    tmp_path,
):
    """The form asks for no time zone: the browser's, else the parish's (#1161).

    The files hold campaign days and dates, no action times, so the page
    sends the browser's zone in a hidden field; a missing, empty or unknown
    one falls back to the parish's time zone instead of being refused.
    """
    harness, _, browser, _ = scenario
    imported(scenario)
    root = tmp_path / "comparison-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    with day_two(harness), task_login(ServiceRole.WEB, exact=True, reconnect=True):
        response, body = read(browser, reverse("admin:campaign_comparison"))
        assert response.status_code == 200
        hidden = b'<input type="hidden" name="browser_timezone" data-browser-zone>'
        assert hidden in body
        assert b'id="comparison-timezone"' not in body
        parish = parish_zone()
        for zone, expected in (
            (None, parish),
            ("", parish),
            ("Etc/Unknown", parish),
            ("Asia/Tokyo", "Asia/Tokyo"),
        ):
            fields = {"day": "2", "format": "csv", "request_key": str(uuid4())}
            if zone is not None:
                fields["browser_timezone"] = zone
            assert post(browser, ROUTE, fields).status_code == 302, zone
            request = ExportRequest.objects.get(request_key=fields["request_key"])
            assert request.browser_timezone == expected, zone
        # A repeated zone is still a malformed form.
        fields = {
            "day": "2",
            "format": "csv",
            "request_key": str(uuid4()),
            "browser_timezone": ["UTC", "Asia/Tokyo"],
        }
        assert post(browser, ROUTE, fields).status_code == 400


def test_cli_export_comparison_is_the_page_form(scenario, cli):  # noqa: F811
    """``export comparison`` makes the same capture and records its events."""
    harness, _, _, actor = scenario
    imported(scenario)
    run = cli()
    with day_two(harness):
        code, document = run(
            "export", "comparison", "--format", "pdf", "--timezone", "UTC"
        )
        assert code == 0, document
        result = document["result"]
        assert result["created"] is True
        request = ExportRequest.objects.get(pk=result["export"]["id"])
        assert request.report == "campaign_comparison"
        assert request.parameters == {"day": 2}
        assert request.browser_timezone == "UTC"
        # The file holds days, not times, so the zone is optional (#1161):
        # without it the parish's time zone is used.
        code, document = run("export", "comparison", "--format", "csv")
        assert code == 0 and document["result"]["created"], document
        request = ExportRequest.objects.get(pk=document["result"]["export"]["id"])
        assert request.browser_timezone == parish_zone()
        code, document = run(
            "export",
            "comparison",
            "--day",
            "1",
            "--format",
            "csv",
            "--timezone",
            "UTC",
        )
        assert code == 0 and document["result"]["created"], document
        code, document = run(
            "export",
            "comparison",
            "--day",
            "9",
            "--format",
            "csv",
            "--timezone",
            "UTC",
        )
        assert code == 1 and document["error"]["code"] == "not_available", document
    assert (
        AuditContext.objects.filter(
            event__event_type="admin_cmd_export_comparison"
        ).count()
        == 3
    )
    assert {
        context["count"]
        for context in AuditContext.objects.filter(
            event__event_type="campaign_comparison_exported"
        ).values_list("context", flat=True)
    } == {1, 2}


def test_migration_check_refuses_a_missing_part():
    """The frozen file's DO block fails without its trigger, kind or guard."""
    text = FROZEN.read_text(encoding="utf-8")
    check = text[text.index("DO $check$") : text.index("$check$;") + len("$check$;")]
    with transaction.atomic(), connection.cursor() as cursor:
        # 0084 (#1018 part (a)) widens export_report_known with the campaign
        # archive, which 0071's block, written before it, does not count:
        # put 0071's own check back first (rolled back below).
        for line in text.splitlines():
            if line.startswith(
                tuple(
                    f"ALTER TABLE public.stewardship_export_request {verb} "
                    "CONSTRAINT export_report_known"
                    for verb in ("DROP", "ADD")
                )
            ):
                cursor.execute(line)
        cursor.execute(check)  # The installed schema passes.
        for change in (
            "DROP TRIGGER comparison_export_binding"
            " ON stewardship_comparison_export_snapshot",
            "ALTER TABLE stewardship_export_request"
            " DROP CONSTRAINT export_report_known",
            "ALTER FUNCTION public.stewardship_export_request_guard_v1()"
            " SECURITY DEFINER",
            "DROP INDEX export_comparison_snapshot",
        ):
            with (
                pytest.raises(DatabaseError, match="Migration 0071"),
                transaction.atomic(),
            ):
                cursor.execute(change)
                cursor.execute(check)
        cursor.execute(check)
        transaction.set_rollback(True)


def test_a_ministry_leader_and_a_disabled_user_cannot_export(
    scenario,  # noqa: F811
    cli,  # noqa: F811
    google,
):
    """Only a current Administrator or Staff member exports, by every path.

    The fixture's Choir Chairperson (a Ministry leader) cannot sign in to
    this campaign, which does not ask about Ministries, so the form never
    reaches them, and neither the service nor the capture trigger accepts
    them. Once the Administrator is disabled, the form and the command are
    refused and the trigger refuses their capture: nothing is stored.
    """
    harness, _, browser, actor = scenario
    imported(scenario)
    run = cli()
    google[0].update(email="valid@example.org", sub="chairperson-subject")
    leader_browser, login = signed_in()
    assert login.status_code == 403
    leader = PortalUser.objects.get(email="valid@example.org").pk
    fields = dict(
        day="2", format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    with day_two(harness), task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert post(leader_browser, ROUTE, fields).status_code in {302, 403}
        with pytest.raises(PermissionError):
            create_comparison_export(
                harness.service.store,
                leader,
                campaign_id=harness.campaign.pk,
                day=2,
                format="csv",
                browser_timezone="UTC",
                request_key=uuid4(),
                document=capture_row(harness, leader).document,
            )
        # The capture trigger's own policy refuses, not a grant or a binding.
        with (
            pytest.raises(DatabaseError, match="capture is unavailable"),
            transaction.atomic(),
        ):
            capture_row(harness, leader).save(force_insert=True)
    PortalUser.objects.filter(pk=actor).update(disabled=True, version=F("version") + 1)
    with day_two(harness):
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            refused = post(browser, ROUTE, fields)
            assert refused.status_code in {302, 403}
            assert "/admin/reports/exports/" not in refused.get("Location", "")
            with (
                pytest.raises(DatabaseError, match="capture is unavailable"),
                transaction.atomic(),
            ):
                capture_row(harness, actor).save(force_insert=True)
        code, document = run(
            "export", "comparison", "--format", "csv", "--timezone", "UTC"
        )
        assert code != 0 and not document["ok"], document
    assert not ExportRequest.objects.exists()
    assert not ComparisonExportSnapshot.objects.exists()
