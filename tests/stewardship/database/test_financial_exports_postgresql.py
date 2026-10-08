"""Complete financial captures, immutable money and ownership under real roles."""

import json
from dataclasses import asdict
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from django.db import DatabaseError, connection
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from parishkit.stewardship.audit.models import AuditContext
from parishkit.stewardship.campaigns.read_guards import DownloadPool, ReadLimits
from parishkit.stewardship.campaigns.work_locks import work_transaction
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.jobs.dispatch import execute_hint
from parishkit.stewardship.jobs.queues import WorkQueue
from parishkit.stewardship.reports import export_services, financial_exports
from parishkit.stewardship.reports.export_models import (
    ExportRequest,
    FinancialExportSnapshot,
)
from parishkit.stewardship.reports.export_services import TASK_TYPE, export_status
from parishkit.stewardship.reports.export_tasks import export_handler
from parishkit.stewardship.reports.financial import FinancialQuery
from parishkit.stewardship.reports.financial_exports import create_financial_export

from ..test_financial_answers import CHECK, OPTIONS
from .auth_builders import signed_in
from .campaign_builders import campaign_clock
from .response_builders import activate_response_service
from .test_background_grants_postgresql import task_login
from .test_export_views_postgresql import post, restricted_download_pool
from .test_financial_report_postgresql import family_session, head, pledge, report
from .test_financial_source_postgresql import financial_source
from .test_information_followup_postgresql import search
from .test_ministry_exports_postgresql import leader
from .test_policy_postgresql import user
from .test_report_workspace_postgresql import read
from .test_response_http_postgresql import load_form
from .test_runtime_auth_grants_postgresql import web_login

pytestmark = pytest.mark.django_db(transaction=True)


def three_families(harness):
    """Three real households with distinct pledges, built as the page's cases are."""
    financial_source(
        harness,
        options=map(asdict, OPTIONS),
        extra_families={
            6: dict(familyDUID=6, registeredOrganizationID=5, lastName="Zeta")
        },
        extra_members={21: head(21, 2), 22: head(22, 6)},
    )
    harness = activate_response_service(harness)
    opened = harness.campaign.active_configuration.starts_at
    with campaign_clock(opened + timedelta(hours=1)), web_login():
        pledge(harness, load_form(harness), shares={CHECK: ""})
    with campaign_clock(opened + timedelta(hours=2)), web_login():
        second = family_session(harness, 2)
        pledge(second, load_form(second), annual_pledge="500", frequency="weekly")
    with campaign_clock(opened + timedelta(hours=3)), web_login():
        third = family_session(harness, 6)
        pledge(third, load_form(third), annual_pledge="0", frequency="")
    return harness


def capture(harness, actor, **values):
    """Queue one export through the real service under the web role."""
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        return create_financial_export(
            harness.service.store,
            actor,
            campaign_id=harness.campaign.pk,
            **(
                dict(
                    query=FinancialQuery(),
                    format="csv",
                    browser_timezone="UTC",
                    request_key=uuid4(),
                )
                | values
            ),
        )


def test_complete_capture_is_the_whole_result_with_proven_money(
    response_service, monkeypatch
):
    """The capture is unpaged, proven, canonical and immutable; replay is exact."""
    harness = three_families(response_service)
    page = report(harness, page_size=2)
    assert page["total"] == 3 and len(page["rows"]) == 2
    actor = user("admin@example.org").pk
    key = uuid4()
    request = capture(
        harness, actor, query=FinancialQuery(sort="pledge_desc"), request_key=key
    )
    # The service brings back only the capture's header; the document is read
    # here as the render owner reads it, never into the web request.
    assert request.financial_snapshot.document == {}
    snapshot = FinancialExportSnapshot.objects.get(pk=request.financial_snapshot_id)
    document = snapshot.document
    assert snapshot.row_count == document["total"] == len(document["rows"]) == 3
    assert [row["family_duid"] for row in document["rows"]] == [1, 2, 6]
    # Money crosses the capture as canonical text, with the proof honored: a
    # proven read with no matching source row is a real zero.
    rows = {row["family_duid"]: row for row in document["rows"]}
    assert rows[1]["annual_pledge"] == "1234.50"
    assert rows[1]["pledge_total"] == "1200.00"
    assert rows[2]["contribution_total"] == "8888.00"
    assert rows[6]["pledge_total"] == "0.00" and rows[6]["frequency"] is None
    assert request.parameters["proof"] is not None
    assert request.parameters["filters"]["sort"] == "pledge_desc"
    assert str(snapshot.source_id) == document["metadata"]["source_id"]
    assert document["summary"]["families"] == 3
    # The same key replays the same request; a different selection is refused.
    replayed = capture(
        harness, actor, query=FinancialQuery(sort="pledge_desc"), request_key=key
    )
    assert replayed.pk == request.pk
    # Still the same export when the proof would now come out differently, as
    # after a promotion between two identical submissions: nothing is recomputed.
    with monkeypatch.context() as patch:
        patch.setattr(financial_exports, "giving_proof", lambda campaign: None)
        again = capture(
            harness, actor, query=FinancialQuery(sort="pledge_desc"), request_key=key
        )
    assert again.pk == request.pk and again.parameters["proof"] is not None
    with pytest.raises(ValueError):
        capture(harness, actor, query=FinancialQuery(sort="name"), request_key=key)
    # Captures are immutable in SQL itself: even the schema owner, who holds
    # every privilege, is refused by the trigger rather than by a grant.
    with pytest.raises(DatabaseError, match="immutable"), connection.cursor() as cursor:
        cursor.execute(
            "UPDATE stewardship_financial_export_snapshot SET row_count=0 WHERE id=%s",
            [snapshot.pk],
        )
    # The worker holds no insert grant at all; the trigger's own guard against
    # it stands behind that grant, never instead of it.
    with (
        task_login(ServiceRole.WORKER, exact=True, reconnect=True),
        pytest.raises(DatabaseError, match="permission denied"),
    ):
        FinancialExportSnapshot.objects.create(
            campaign_id=harness.campaign.pk,
            configuration_id=snapshot.configuration_id,
            actor_id=actor,
            correlation_id=uuid4(),
            parameters=snapshot.parameters,
        )
    # A filtered capture is exactly the filtered result, with its own summary.
    filtered = capture(harness, actor, query=FinancialQuery(amount="nonzero"))
    assert filtered.financial_snapshot.row_count == 2
    stored = FinancialExportSnapshot.objects.get(pk=filtered.financial_snapshot_id)
    assert stored.document["summary"]["families"] == 2
    contexts = list(
        AuditContext.objects.filter(event__event_type="export_requested").values_list(
            "context", flat=True
        )
    )
    assert [context["count"] for context in contexts] == [3, 2]
    assert "1234" not in str(contexts) and "Zeta" not in str(contexts)


def test_native_financial_exports_use_real_worker_and_guarded_downloads(
    response_service, google, tmp_path, settings, monkeypatch
):
    """One fixture proves three formats, native ownership, regeneration and denial."""
    harness = response_service
    financial_source(harness, modules=["financial"], options=map(asdict, OPTIONS))
    harness = activate_response_service(harness)
    with web_login():
        pledge(harness, load_form(harness), shares={CHECK: ""})
    name = report(harness)["rows"][0]["family_name"].encode()
    # Pages must not show the Family's row; the bare surname ("Example")
    # also appears in unrelated page text.
    row = b'<th scope="row">' + name + b"<br>"
    browser, login = signed_in()
    assert login.status_code == 302
    actor = user("admin@example.org").pk
    root = tmp_path / "financial-reports"
    root.mkdir(mode=0o700)
    settings.STEWARDSHIP_REPORTS_ROOT = root
    settings.STEWARDSHIP_DOWNLOAD_POOL = DownloadPool(ReadLimits(process_pool_size=1))
    route = f"/admin/reports/{harness.campaign.pk}/financial/"
    query = FinancialQuery()
    first = None
    for format in ("csv", "xlsx", "pdf"):
        fields = query.form_values() | dict(
            format=format, browser_timezone="UTC", request_key=str(uuid4())
        )
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            response, body = read(browser, route)
            assert response.status_code == 200 and b"Queue complete export" in body
            assert b'name="request_key"' in body and b"<fieldset disabled" not in body
            assert browser.post(route + "export", fields).status_code == 403  # No CSRF.
            if format == "csv":
                assert browser.get(route + "export").status_code == 405
                for invalid in (
                    {"page": "2"},
                    {"extra": "x"},
                    {"sort": "random"},
                    {"format": ["csv", "pdf"]},
                ):
                    rejected = post(browser, route + "export", fields | invalid)
                    assert rejected.status_code == 400
                    assert row not in rejected.content
                    assert rejected["Cache-Control"] == "no-store"
                # Filters are refused where a query string sent them, and another
                # campaign is indistinguishable from none.
                queried = post(browser, route + "export?search=x", fields)
                assert queried.status_code == 400
                wrong = f"/admin/reports/{uuid4()}/financial/export"
                # Until #145 any campaign but the current one is gone (410).
                assert post(browser, wrong, fields).status_code == 410
            response = post(browser, route + "export", fields)
            assert response.status_code == 302
            request = ExportRequest.objects.get(request_key=fields["request_key"])
            first = first or request
            assert (
                post(browser, route + "export", fields)["Location"]
                == response["Location"]
            )
            job_route = response["Location"]
            with CaptureQueriesContext(connection) as queries:
                response, body = read(browser, job_route)
            assert queries.captured_queries
            assert all(
                '"stewardship_financial_export_snapshot"."document"' not in q["sql"]
                for q in queries.captured_queries
            )
            assert response.status_code == 200
            assert b"Financial stewardship export" in body
            assert b"Matching Families" in body and row not in body
        with task_login(ServiceRole.WORKER, exact=True, reconnect=True):
            assert execute_hint(
                request.task_id,
                queue=WorkQueue.GENERAL,
                worker_id=uuid4(),
                handlers={
                    TASK_TYPE: export_handler(store=harness.service.store, root=root)
                },
            )
        with task_login(ServiceRole.WEB, exact=True, reconnect=True):
            assert (
                export_status(harness.service.store, actor, request.pk)["state"]
                == "ready"
            )
        with restricted_download_pool(settings):
            response, body = search(browser, job_route + "download", {})
            assert response.status_code == 200
            assert body.startswith(
                {"csv": b"Family,", "xlsx": b"PK", "pdf": b"%PDF"}[format]
            )
            assert "financial." + format in response["Content-Disposition"]
            if format == "csv":
                assert name in body and b"$1,234.50" in body and b"$1,200.00" in body
                assert b"Financial stewardship detail" in body
                # The comparison columns carry the page's labels (#404).
                assert b"ParishSoft pledged,ParishSoft contributed" in body
            elif format == "xlsx":
                # The real worker's file holds summable dollar-formatted numbers.
                sheet = load_workbook(BytesIO(body))["Financial detail"]
                assert sheet.cell(1, 8).value == "ParishSoft pledged"
                for column, expected in ((4, 1234.5), (8, 1200)):
                    cell = sheet.cell(2, column)
                    assert (cell.data_type, cell.value) == ("n", expected)
                    assert cell.number_format == '"$"#,##0.00'
    # Regeneration after expiry keeps the retained capture, never a fresh read.
    with work_transaction():
        now = export_services.database_now()
    monkeypatch.setattr(
        export_services, "database_now", lambda: now + timedelta(days=8)
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        fields = {"request_key": str(uuid4())}
        response = post(
            browser, f"/admin/reports/exports/{first.pk}/regenerate", fields
        )
        assert response.status_code == 302
        regenerated = ExportRequest.objects.get(request_key=fields["request_key"])
        assert regenerated.financial_snapshot_id == first.financial_snapshot_id
        assert regenerated.parameters == first.parameters
    # A Ministry leader never captures money: over HTTP, in the service or in SQL.
    other, leader_id, *_ = leader(harness, google)
    fields = query.form_values() | dict(
        format="csv", browser_timezone="UTC", request_key=str(uuid4())
    )
    with task_login(ServiceRole.WEB, exact=True, reconnect=True):
        assert post(other, route + "export", fields).status_code == 403
        with pytest.raises(PermissionError):
            create_financial_export(
                harness.service.store,
                leader_id,
                campaign_id=harness.campaign.pk,
                query=query,
                format="csv",
                browser_timezone="UTC",
                request_key=uuid4(),
            )
        # The capture trigger's own policy refuses, not a grant or a binding.
        with pytest.raises(DatabaseError, match="capture is unavailable"):
            FinancialExportSnapshot.objects.create(
                campaign_id=harness.campaign.pk,
                configuration_id=first.configuration_id,
                actor_id=leader_id,
                correlation_id=uuid4(),
                parameters=first.parameters,
            )
    assert FinancialExportSnapshot.objects.count() == 3


def staff_entered(submission_id, admin_id):
    """Mark one live submission as entered by Staff, as Open form records it.

    Open form's own path is tested in test_open_form_postgresql; here only
    the projection reading the marker matters, so the immutable submission
    gets its #529 column directly with its guard bypassed for this one
    statement (the test login is a superuser).
    """
    with work_transaction(), connection.cursor() as cursor:
        cursor.execute("SET LOCAL session_replication_role = replica")
        cursor.execute(
            "UPDATE stewardship_submission SET entered_by_id = %s WHERE id = %s",
            [admin_id, submission_id],
        )


def test_page_and_capture_say_whether_staff_entered_the_response(
    response_service,
):
    """Entered by Staff for the effective response, on the page and in a capture.

    Family 2's response is marked as Staff-entered; the others are their own.
    The flag is a yes or no and never names the Staff member (#794).
    """
    from parishkit.stewardship.campaigns.credential_models import FamilyCampaign

    harness = three_families(response_service)
    actor = user("admin@example.org").pk
    second = FamilyCampaign.objects.get(campaign=harness.campaign, family_duid=2)
    staff_entered(second.effective_submission_id, actor)
    page = report(harness)
    flags = {row["family_duid"]: row["staff_entered"] for row in page["rows"]}
    assert flags == {1: False, 2: True, 6: False}
    request = capture(harness, actor)
    document = FinancialExportSnapshot.objects.get(
        pk=request.financial_snapshot_id
    ).document
    assert {row["family_duid"]: row["staff_entered"] for row in document["rows"]} == {
        1: False,
        2: True,
        6: False,
    }
    assert str(actor) not in str(document)


def test_only_the_current_response_counts_as_entered_by_staff(response_service):
    """The column follows each Family's current response, not its first (#794).

    - Family 2: Staff entered its first response, then the Family submitted
      again as itself, so its current response reads No.
    - Family 1: it submitted twice; only the later, current response is
      marked as Staff-entered, so it reads Yes.
    - Family 6: a Staff-marked Testing response stays out of the report; its
      own live response reads No.
    """
    from parishkit.stewardship.campaigns.credential_models import FamilyCampaign
    from parishkit.stewardship.responses.models import Submission

    harness = three_families(response_service)
    actor = user("admin@example.org").pk
    opened = harness.campaign.active_configuration.starts_at

    def current(duid):
        return FamilyCampaign.objects.get(
            campaign=harness.campaign, family_duid=duid
        ).effective_submission_id

    first_of_two = current(2)
    staff_entered(first_of_two, actor)
    with campaign_clock(opened + timedelta(hours=4)), web_login():
        again = family_session(harness, 2)
        pledge(again, load_form(again), annual_pledge="600", frequency="weekly")
    assert current(2) != first_of_two
    with campaign_clock(opened + timedelta(hours=5)), web_login():
        repeat = family_session(harness, 1)
        pledge(repeat, load_form(repeat), shares={CHECK: ""})
    staff_entered(current(1), actor)
    # A Testing response for Family 6, marked as Staff-entered: a copy of its
    # live row in Testing mode (guards bypassed for this fixture row only).
    live = Submission.objects.get(pk=current(6))
    with work_transaction(), connection.cursor() as cursor:
        cursor.execute("SELECT * FROM stewardship_submission WHERE id = %s", [live.pk])
        columns = [column.name for column in cursor.description]
        values = dict(zip(columns, cursor.fetchone(), strict=True))
        cursor.execute(
            "SELECT max(campaign_sequence) + 1 FROM stewardship_submission "
            "WHERE campaign_id = %s",
            [live.campaign_id],
        )
        values.update(
            id=uuid4(),
            mode="test",
            rehearsal_epoch_id=uuid4(),
            entered_by_id=actor,
            campaign_sequence=cursor.fetchone()[0],
            baseline_id=uuid4(),
            prior_submission_id=None,
        )
        values = {
            key: json.dumps(value) if isinstance(value, dict | list) else value
            for key, value in values.items()
        }
        cursor.execute("SET LOCAL session_replication_role = replica")
        cursor.execute(
            f"INSERT INTO stewardship_submission ({', '.join(values)}) "
            f"VALUES ({', '.join(['%s'] * len(values))})",
            list(values.values()),
        )
    assert Submission.objects.filter(mode="test", entered_by_id=actor).count() == 1
    flags = {
        row["family_duid"]: row["staff_entered"] for row in report(harness)["rows"]
    }
    assert flags == {1: True, 2: False, 6: False}
    request = capture(harness, actor)
    document = FinancialExportSnapshot.objects.get(
        pk=request.financial_snapshot_id
    ).document
    assert {row["family_duid"]: row["staff_entered"] for row in document["rows"]} == {
        1: True,
        2: False,
        6: False,
    }
