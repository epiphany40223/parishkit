"""Census changes worklist under the real web role (#528)."""

import csv
import io
from datetime import datetime
from uuid import uuid4

import pytest
from django.db import connection, transaction
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from parishkit.stewardship.accounts.policy import Principal
from parishkit.stewardship.accounts.runtime_models import SystemConfiguration
from parishkit.stewardship.audit.models import AuditContext, AuditEvent
from parishkit.stewardship.deployment import ServiceRole
from parishkit.stewardship.reports.census_changes import CensusQuery, census_changes
from parishkit.stewardship.responses.models import ProposedChange
from parishkit.stewardship.source.source_status import NAME_NOT_RECORDED

from ..census_factory import address
from ..policy_factory import address as rule_address
from .auth_builders import signed_in, unguarded
from .campaign_builders import change
from .family_link_checks import link_reads, timeline_link
from .family_name_cases import DEFAULT, NAME_CASES, seed_name_cases
from .leader_builders import promote_leaders
from .response_builders import activate_response_service
from .test_background_grants_postgresql import task_login
from .test_export_views_postgresql import restricted_download_pool
from .test_information_followup_postgresql import search
from .test_ministry_responses_postgresql import ministry_source, respond, start
from .test_report_workspace_postgresql import read as get
from .test_response_http_postgresql import answers_for, load_form
from .test_runtime_auth_grants_postgresql import web_login

pytestmark = pytest.mark.django_db(transaction=True)

STAFF = Principal(uuid4(), frozenset({"staff"}), frozenset())
ADMIN = Principal(uuid4(), frozenset({"administrator"}), frozenset())
LEADER = Principal(uuid4(), frozenset({"ministry_leader"}), frozenset({4}))


def submit_live(harness):
    """One response that renames Member 3, an automatic change, and gives the
    Family a new home address, a By-hand change (no verified address read)."""
    with web_login():
        form = load_form(harness)
        answers = answers_for(form)
        answers["members"]["3"]["first_name"] = "Requested"
        answers["family"]["home_address"] = address(line1="9 Census Way")
        return respond(harness, form, answers)


def worklist(harness, principal=STAFF, **values):
    """Read through the restricted web grants, not the migration owner."""
    with task_login(ServiceRole.WEB, exact=True, reconnect=True), transaction.atomic():
        return census_changes(harness.campaign.pk, CensusQuery.parse(values), principal)


def directory_name(duid):
    """The Family's name as the directory shows it, from the current snapshot."""
    from parishkit.stewardship.source.snapshot_models import SourceCurrent
    from parishkit.stewardship.source.snapshot_names import snapshot_family_names

    snapshot = SourceCurrent.objects.values_list("snapshot_id", flat=True).get()
    return snapshot_family_names(snapshot, [duid])[duid]


def test_worklist_reads_proposals_with_their_status(response_service):
    """A live response's change is listed, worded, and To review only for an
    Administrator; Staff find it with the status filter; leaders never."""
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness)
    # The Family address is By hand and To do, with no ParishSoft value.
    (home,) = worklist(harness)["rows"]
    assert (home["who"], home["label"]) == ("Family", "Home address")
    assert home["route_label"] == "By hand" and home["status"] == "to_do"
    assert home["parishsoft_now"] == "" and "9 Census Way" in home["family_answer"]
    rows = [row for row in worklist(harness, ADMIN)["rows"] if row["who"] != "Family"]
    assert len(rows) == 1
    row = rows[0]
    assert row["label"] == "First name" and row["family_answer"] == "Requested"
    assert row["status"] == "to_review" and row["route_label"] == "Automatic"
    assert row["who"] and row["family_duid"]
    # Named as in every Admin table: surname, then heads (#932).
    assert ", " in row["family_name"]
    assert row["family_name"] == directory_name(row["family_duid"])
    # Member 3's change carries its Member DUID; the Family's has none (#932).
    assert row["member_duid"] == 3 and home["member_duid"] is None
    assert len(worklist(harness, status="to_review")["rows"]) == 1
    by_hand = worklist(harness, status="all", route="by_hand")["rows"]
    assert [item["label"] for item in by_hand] == ["Home address"]
    assert worklist(harness, status="all", search="no such name")["rows"] == []
    with pytest.raises(PermissionError):
        worklist(harness, LEADER)


def test_family_names_follow_the_shared_name_cases(response_service):
    """The Family name comes from the shared SQL lookup, so every shared edge
    case (Unicode padding, inactive or missing heads, surname fallbacks) reads
    as snapshot_family_names gives it, and a Family missing from the current
    snapshot has no name (None), which the page marks Not in ParishSoft
    (#932, #1216)."""
    from parishkit.stewardship.source.snapshot_models import SourceCurrent

    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness)
    (home,) = worklist(harness)["rows"]
    duid = home["family_duid"]
    snapshot = SourceCurrent.objects.values_list("snapshot_id", flat=True).get()

    def drop_family():
        """Remove the Family's row from the current snapshot."""
        # Snapshot rows are immutable through the ORM; only raw SQL in
        # replica mode, as a superuser could, removes one.
        with unguarded(), connection.cursor() as cursor:
            cursor.execute(
                "DELETE FROM stewardship_snapshot_family "
                "WHERE snapshot_id=%s AND source_key=%s",
                [snapshot, str(duid)],
            )

    drop_family()
    assert worklist(harness)["rows"][0]["family_name"] is None
    for case in NAME_CASES:
        # Each case's head Member DUIDs are distinct, so earlier cases' heads
        # can stay in the snapshot.
        drop_family()
        (expected,) = seed_name_cases(snapshot, (case,), first=duid).values()
        if expected == DEFAULT:
            expected = NAME_NOT_RECORDED
        assert worklist(harness)["rows"][0]["family_name"] == expected, case[0]


def parish_name():
    """The active configuration's parish name, as the files state it."""
    system = SystemConfiguration.objects.select_related(
        "active_configuration__parish"
    ).get()
    return system.active_configuration.parish.name


def test_native_page_and_downloads(response_service, google, settings):
    """The page filters by CSRF POST; CSV, XLSX and PDF carry the same rows
    under the same headings, and the view and download are audited."""
    settings.STEWARDSHIP_DOWNLOAD_POOL = None  # restricted_download_pool sets it
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness)
    route = "/admin/reports/census/"
    # The rows the page shows by default, read before the download pool
    # takes over the web login.
    rows = worklist(harness, ADMIN)["rows"]
    browser, login = signed_in()
    assert login.status_code == 302
    with restricted_download_pool(settings):
        with CaptureQueriesContext(connection) as queries:
            response, body = get(browser, route)
        assert response.status_code == 200 and response["Cache-Control"] == "no-store"
        assert b"Requested" in body and b"To review" in body
        # Each row's Family name opens its Timeline (#1570), the page's
        # record ids read once, however many rows it shows.
        duid = rows[0]["family_duid"]
        link = timeline_link(harness.campaign, duid, directory_name(duid))
        # The worklist draws the Family as each row's header cell.
        cell = f'<th scope="row">{link}</th>'.encode()
        assert len(rows) == body.count(cell) == 2
        assert len(link_reads(queries)) == 1
        assert get(browser, route + "?search=x")[0].status_code == 400
        response, body = search(browser, route, {"stage": "conflict"})
        assert response.status_code == 200 and b"Nothing is at this stage." in body
        for invalid in (
            {"stage": "bad"},
            {"status": "bad"},
            {"sort": "id"},
            {"size": "7"},
        ):
            assert search(browser, route, invalid)[0].status_code == 400
        response, body = search(browser, route, {"sort": "-who"})
        assert response.status_code == 200 and b'aria-sort="descending"' in body
        # Member DUID has its own column beside Who, and sorts in place (#932).
        response, body = search(browser, route, {"stage": "all", "sort": "member_duid"})
        assert response.status_code == 200 and b'aria-sort="ascending"' in body
        assert b'<td class="numeric">3</td>' in body
        export = route + "exports/"
        response, body = search(
            browser, export, {"format": "csv", "timezone": "UTC", "stage": "all"}
        )
        assert response.status_code == 200 and response["Content-Type"] == "text/csv"
        assert "stewardship-census-changes-" in response["Content-Disposition"]
        assert b"Requested" in body and b"How it reaches ParishSoft" in body
        table = list(csv.reader(io.StringIO(body.decode())))
        # The page's order (#932): ParishSoft status after Family (#1216),
        # Member DUID right after Who, then Submitted and the latest
        # resolution's three columns (#528).
        assert table[0][:2] == ["Family", "ParishSoft status"]
        assert table[0][3:5] == ["Who", "Member DUID"]
        assert table[0][11:] == ["Submitted", "Recorded by", "Recorded at", "Note"]
        assert any(values[4] == "3" for values in table[1:])
        # Every Family is current: no status to mark.
        assert {values[1] for values in table[1:]} == {""}
        response, body = search(
            browser, export, {"format": "xlsx", "timezone": "UTC", "stage": "all"}
        )
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(body))
        assert book.sheetnames == ["Census changes", "Report information"]
        # The shared sheet names the download (#926): only names, a time, a
        # zone, a count and the privacy line, never an internal id (#956).
        information = list(book["Report information"].iter_rows(values_only=True))
        assert [label for label, _ in information] == [
            "Report",
            "Parish",
            "Campaign",
            "Downloaded at",
            "Display time zone",
            "Changes in this file",
            "Privacy",
        ]
        values = dict(information)
        assert values["Report"] == "Census changes"
        assert values["Parish"] == parish_name()
        assert values["Campaign"] == harness.campaign.active_configuration.name
        assert isinstance(values["Downloaded at"], datetime)
        assert values["Display time zone"] == "UTC"
        assert values["Changes in this file"] == str(len(table) - 1)
        sheet = [
            ["" if value is None else value for value in values]
            for values in book["Census changes"].iter_rows(values_only=True)
        ]
        # One column set and the same rows in both files (#966).
        assert sheet == table
        assert "Requested" in [values[7] for values in sheet[1:]]
        response, body = search(
            browser, export, {"format": "pdf", "timezone": "UTC", "stage": "all"}
        )
        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert ".pdf" in response["Content-Disposition"]
        assert body.startswith(b"%PDF-")
        assert search(browser, export, {"format": "png"})[0].status_code == 400
    assert AuditEvent.objects.filter(event_type="census_changes_viewed").exists()
    assert AuditEvent.objects.filter(event_type="census_changes_exported").exists()
    contexts = [
        AuditContext.objects.get(event=event).context
        for event in AuditEvent.objects.filter(
            event_type__in=("census_changes_viewed", "census_changes_exported")
        )
    ]
    # Access evidence is a count only: never a name, value or filter.
    assert contexts and all(
        set(context) == {"outcome", "count"} for context in contexts
    )


def test_testing_responses_are_excluded(response_service):
    """A Testing-mode response's proposals never reach the worklist."""
    harness = response_service
    start(harness)
    submission = submit_live(harness)
    assert submission.mode != "live"
    assert ProposedChange.objects.filter(submission=submission).exists()
    assert worklist(harness, ADMIN, status="all", history="yes")["rows"] == []
    # Positive control: the same rows appear once the response is live, so
    # the exclusion above is the mode filter, not a query that reads nothing.
    # Submissions are immutable, so the flip bypasses their triggers, as the
    # migration owner only.
    with transaction.atomic(), connection.cursor() as cursor:
        cursor.execute("SET LOCAL session_replication_role = replica")
        cursor.execute(
            "UPDATE stewardship_submission SET mode='live', rehearsal_epoch_id=NULL"
            " WHERE id=%s",
            [submission.pk],
        )
        cursor.execute("SET LOCAL session_replication_role = origin")
    rows = worklist(harness, ADMIN, status="all", history="yes")["rows"]
    assert {row["label"] for row in rows} == {"First name", "Home address"}


@pytest.mark.parametrize(
    ("roles", "allowed"), [(("staff",), True), (("ministry_leader",), False)]
)
def test_staff_may_open_the_page_and_ministry_leaders_may_not(
    response_service, google, roles, allowed
):
    """Over HTTP, Staff read the worklist; a Ministry leader is refused."""
    harness = response_service
    start(harness)
    harness = activate_response_service(harness)
    submit_live(harness)
    store = harness.service.store
    if roles == ("ministry_leader",):
        # A Ministry leader leads through a ParishSoft role (#922).
        promote_leaders(harness, ministry_source(), {"reader@example.org": [9]})
    rule = rule_address("reader@example.org", roles=roles)
    change(
        store,
        store.active(),
        uuid4(),
        [{"operation": "add", "section": "login_rules", **rule}],
    )
    google[0]["email"] = "reader@example.org"
    browser, login = signed_in()
    assert login.status_code == 302
    route = "/admin/reports/census/"
    response, body = get(browser, route)
    if allowed:
        assert response.status_code == 200 and b"9 Census Way" in body
        # Staff may open the Timeline (its summary), so the name links.
        assert b'<th scope="row"><a href="/admin/reports/families/' in body
    else:
        assert response.status_code in {403, 404} and b"9 Census Way" not in body
        # A real CSRF POST to the live route: refused by access, not CSRF.
        response, body = search(browser, route + "exports/", {"format": "csv"})
        assert response.status_code == 403 and b"9 Census Way" not in body
