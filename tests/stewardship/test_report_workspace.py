"""Report query safety, shared display parity and typed spreadsheet outputs."""

from contextlib import ExitStack
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock
from uuid import uuid4

import pytest
from django.http import QueryDict
from openpyxl import load_workbook

from parishkit.stewardship.reports.digest_presentation import (
    participation_context,
    snapshot_context,
    tooltip_dollars,
)
from parishkit.stewardship.reports.money import MoneyAmount
from parishkit.stewardship.reports.spreadsheets import participation_xlsx
from parishkit.stewardship.reports.workspace import RenderedReport, ReportQuery
from parishkit.stewardship.reports.workspace_views import statistics_panel

from .test_daily_digest_content import document


@pytest.mark.parametrize(
    "system_open,known",
    [(False, False), (False, True), (True, False), (True, True)],
)
def test_read_admission_classifies_only_known_campaign_closure(
    monkeypatch, system_open, known
):
    """Global/unknown denials never turn into a lifecycle retry or acceptance."""
    from parishkit.stewardship.campaigns.read_guards import ReadUnavailable
    from parishkit.stewardship.reports import read_admission

    system, campaign = Mock(), Mock()
    system.objects.filter.return_value.exists.return_value = system_open
    campaign.objects.filter.return_value.exists.return_value = known
    monkeypatch.setattr(read_admission, "SystemConfiguration", system)
    monkeypatch.setattr(read_admission, "Campaign", campaign)
    admission = Mock(side_effect=PermissionError("private refusal"))
    monkeypatch.setattr(read_admission, "admit_campaign", admission)
    identifier = uuid4()
    # The campaign is the current one; only its lifecycle admission fails.
    system.objects.values_list.return_value.first.return_value = identifier
    error = ReadUnavailable if system_open and known else PermissionError
    with pytest.raises(error, match="unavailable") as caught:
        read_admission.admit_report_read(identifier)
    assert "private refusal" not in str(caught.value)
    admission.assert_called_once_with(identifier, mutating=False)
    system.objects.filter.assert_called_once_with(
        active_configuration_id__isnull=False, restore_review_required=False
    )
    if system_open:
        campaign.objects.filter.assert_called_once_with(pk=identifier)
    else:
        campaign.objects.filter.assert_not_called()


@pytest.mark.parametrize("current", [None, "other"])
def test_read_admission_refuses_a_campaign_that_is_not_current(monkeypatch, current):
    """Until #145 only the current campaign is reported (rule 10): 410, no read.

    With no current campaign at all the refusal says so plainly instead of
    calling the requested campaign "no longer" current.
    """
    from parishkit.stewardship.reports import read_admission
    from parishkit.stewardship.web.refusals import UserFacingGone

    system = Mock()
    system.objects.values_list.return_value.first.return_value = (
        uuid4() if current else None
    )
    monkeypatch.setattr(read_admission, "SystemConfiguration", system)
    admission = Mock()
    monkeypatch.setattr(read_admission, "admit_campaign", admission)
    with pytest.raises(UserFacingGone) as caught:
        read_admission.admit_report_read(uuid4())
    assert str(caught.value) == (
        "This campaign is no longer the current campaign."
        if current
        else "There is no current campaign."
    )
    system.objects.values_list.assert_called_once_with("current_campaign_id", flat=True)
    admission.assert_not_called()


@pytest.mark.parametrize(
    "query",
    [
        "scope=all",
        "page=0",
        "page=01",
        "page=10001",
        "page=-1",
        "page=1&page=2",
        "page=",
        "page=１２",
        "sort=name",
        "email=private@example.org",
        # An unknown zone falls back to UTC instead (#1127,
        # test_export_time_zones.py); a repeated one is still refused.
        "timezone=UTC&timezone=Asia/Tokyo",
        "size=7",
        "size=1000",
        "size=",
        "mode=live",
        "mode=testing&mode=testing",
    ],
)
def test_queries_are_bounded_nonidentifying_and_closed(query):
    with pytest.raises(ValueError):
        ReportQuery.parse(QueryDict(query))


def test_campaign_links_preserve_only_validated_state():
    value = ReportQuery.parse(
        QueryDict(
            "scope=current&timezone=America%2FDetroit&inactive=yes&grain=day"
            "&page=2&sort=date_asc&mode=testing"
        )
    )
    url = value.url(page=3)
    assert url.startswith("/admin/reports/responses/?")
    assert "scope=current" in url and "page=3" in url and "mode=testing" in url
    assert "timezone=America%2FDetroit" in url and "sort=date_asc" in url
    # The removed inactive subtotal's and activity grain's parameters still
    # load, and are dropped (#728, #1537).
    assert "inactive" not in url and "grain" not in url
    assert "mode=" not in value.url(mode="production")
    assert "timezone=" not in ReportQuery.parse(QueryDict()).url()
    assert "timezone=UTC" in ReportQuery.parse(QueryDict("timezone=UTC")).url()


def test_daily_table_opens_newest_first_with_five_rows():
    """The defaults (#1534, #1537): newest first, 5 rows, Production; the
    address leaves each default out, and a saved oldest-first link still
    opens oldest first."""
    default = ReportQuery.parse(QueryDict())
    assert (default.sort, default.size, default.mode) == (
        "date_desc",
        "5",
        "production",
    )
    assert default.url() == "/admin/reports/responses/?scope=historical&page=1"
    assert ReportQuery.parse(QueryDict("sort=date_asc")).sort == "date_asc"
    # 5 is offered beside the shared sizes; the old default 50 still loads.
    for size in ("5", "25", "50", "100", "250", "all"):
        assert ReportQuery.parse(QueryDict(f"size={size}")).size == size
    assert "size=50" in ReportQuery.parse(QueryDict("size=50")).url()


@pytest.mark.parametrize("consume", [False, True])
def test_rendered_report_releases_generation_even_before_first_byte(consume):
    closed = []
    stack = ExitStack()
    stack.callback(closed.append, True)
    response = RenderedReport(b"report", stack)
    if consume:
        assert next(response) == b"report"
    response.close()
    response.close()
    assert closed == [True]


def test_tooltip_points_are_three_short_lines_per_date():
    """The visible tooltip is a date and short label/value lines (#575).

    The fuller sentence remains in ``labels`` for screen readers; ``points``
    carries only what the tooltip shows, formatted exactly as the table is.
    """
    chart = document().participation
    interaction = participation_context(chart)["chart_interaction"]
    assert len(interaction["points"]) == len(interaction["labels"]) == len(chart.days)
    first, day = interaction["points"][0], chart.days[0]
    assert first["date"] == participation_context(chart)["rows"][0][0]
    assert [label for label, _ in first["rows"]] == [
        "Families",
        "New today",
        "Pledges",
    ]
    assert first["rows"][0][1] == f"{day.cumulative_responses:,}"
    assert first["rows"][1][1] == f"{day.first_responses:,}"
    # Whole dollars in the tooltip; the table row keeps the cents.
    assert day.pledge_total == Decimal("1234.56")
    assert first["rows"][2][1] == "$1,235"
    assert participation_context(chart)["rows"][0][3] == "$1,234.56"
    assert all(len(value) < 20 for _, value in first["rows"])
    plain = participation_context(replace(chart, financial_enabled=False))
    assert [
        [label for label, _ in point["rows"]]
        for point in plain["chart_interaction"]["points"]
    ] == [["Families", "New today"]] * len(chart.days)


def test_tooltip_on_an_unavailable_day_says_unavailable():
    """A missing day is labeled Unavailable, never shown as zero."""
    chart = document().participation
    missing = replace(
        chart.days[0],
        first_responses=0,
        cumulative_responses=0,
        cohort_denominator=0,
        source_generation=None,
        source_as_of=None,
        population_available=False,
        pledge_available=False,
        pledge_total=None,
    )
    chart = replace(chart, days=(missing, *chart.days[1:]))
    point = participation_context(chart)["chart_interaction"]["points"][0]
    assert point["rows"] == [
        ["Families", "Unavailable"],
        ["New today", "Unavailable"],
        ["Pledges", "Unavailable"],
    ]


@pytest.mark.parametrize(
    ("amount", "shown"),
    [
        ("0", "$0"),
        ("0.49", "$0"),
        ("0.50", "$1"),
        ("2.50", "$3"),
        ("760409.50", "$760,410"),
        ("1234567890123456.78", "$1,234,567,890,123,457"),
    ],
)
def test_tooltip_pledges_are_whole_dollars_rounding_halves_up(amount, shown):
    """Halves round up (not to even), and large totals keep every digit."""
    assert tooltip_dollars(Decimal(amount)) == shown


def test_live_presentation_is_the_digest_presentation_without_population_mixing():
    value = document()
    live = participation_context(value.participation)
    retained = snapshot_context(
        value, mode="testing", chart_url="/chart", download_url="/download"
    )
    for field in ("rows", "headings", "chart_interaction", "chart_last_index"):
        assert live[field] == retained[field]
    assert all(
        "Historical as of day" in label for label in live["chart_interaction"]["labels"]
    )
    missing = replace(
        value.statistics, active=None, comparison_pledge_all=MoneyAmount(None)
    )
    assert all(
        tile.text == "—" and tile.unknown_reason
        for tile in statistics_panel(missing)["tiles"]
    )
    # #728: one comparison figure, the parish-wide total, ends the tiles.
    whole = replace(value.statistics, comparison_pledge_all=MoneyAmount(98265700))
    tiles = statistics_panel(whole)["tiles"]
    assert (tiles[0].label, tiles[0].text) == ("Campaign Families", "1,000")
    assert (tiles[-1].label, tiles[-1].text) == (
        "Last year's pledges (all Families)",
        "$982,657.00",
    )
    assert [tile.label for tile in tiles if "Last year" in tile.label] == [
        tiles[-1].label
    ]


def test_statistics_panel_is_the_active_population_row_only():
    """The Summary's Active population row: four tiles in order, no progress
    chart and no responded or email-reach tiles (#1537). Without an
    observation every tile reads "—" with the reason, never zero (#1540)."""
    statistics = document().statistics
    panel = statistics_panel(statistics)
    assert set(panel) == {"tiles"}
    assert [tile.label for tile in panel["tiles"]] == [
        "Campaign Families",
        "Active Members",
        "Current annual pledges",
        "Last year's pledges (all Families)",
    ]
    missing = statistics_panel(
        replace(statistics, active=None, comparison_pledge_all=MoneyAmount(None))
    )
    assert [(tile.key, tile.text) for tile in missing["tiles"]] == [
        (key, "—")
        for key in (
            "active_families",
            "active_members",
            "annual_pledges",
            "comparison_pledges_all",
        )
    ]
    assert {tile.unknown_reason for tile in missing["tiles"]} == {
        "No ParishSoft data has been loaded for this campaign yet."
    }
    without_giving = statistics_panel(replace(statistics, financial_enabled=False))
    assert [tile.key for tile in without_giving["tiles"]] == [
        "active_families",
        "active_members",
    ]


def test_each_unknown_pledge_total_gives_the_reason_for_its_cause():
    """No giving definition, no full set of last year's pledges and an
    unknown answer amount each say what is missing, never a promise that a
    later refresh will fix it (#1540)."""
    statistics = document().statistics
    unknown = MoneyAmount(None)
    statistics = replace(
        statistics,
        active=replace(statistics.active, annual_pledge=unknown),
        comparison_pledge_all=unknown,
    )
    reasons = lambda value: {  # noqa: E731 - a tiny local reader
        tile.key: tile.unknown_reason for tile in statistics_panel(value)["tiles"]
    }
    nothing = reasons(replace(statistics, financial=None))
    assert nothing["annual_pledges"] == nothing["comparison_pledges_all"]
    assert "no giving definition" in nothing["annual_pledges"]
    # Any definition will do: only its presence is read here.
    partial = reasons(replace(statistics, financial=object()))
    assert "every responding Family" in partial["annual_pledges"]
    assert "does not hold every pledge" in partial["comparison_pledges_all"]
    assert "refresh" not in partial["comparison_pledges_all"]


def test_xlsx_is_complete_typed_formula_safe_and_keeps_original_request():
    value = replace(
        document().participation, parish_name="=HYPERLINK(1)", campaign_name="+SUM(1)"
    )
    stream = BytesIO()
    participation_xlsx(value, stream)
    stream.seek(0)
    book = load_workbook(stream)
    sheet = book["Daily participation"]
    assert sheet.max_row == len(value.days) + 1
    assert sheet["C2"].data_type == "n" and sheet["C2"].value == 1
    assert sheet["H2"].data_type == "n" and sheet["H2"].value == 1234.56
    assert sheet["A2"].is_date
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:H4"
    assert sheet.print_title_rows == "$1:$1"
    metadata = book["Report information"]
    assert metadata["B1"].value == "=HYPERLINK(1)" and metadata["B1"].data_type == "s"
    assert value.as_of_label == metadata["B6"].value
    # No internal UUIDs or cutoff numbers (#966).
    labels = {row[0].value for row in metadata.iter_rows()}
    assert not {"Campaign UUID", "Fact generation UUID", "Submission cutoff"} & labels
    assert not any(
        cell.data_type == "f" for page in book for row in page for cell in row
    )
    book.close()


def test_large_money_stays_exact_instead_of_excel_rounding():
    value = document().participation
    day = replace(value.days[0], pledge_total=Decimal("1234567890123456.78"))
    stream = BytesIO()
    participation_xlsx(replace(value, days=(day, *value.days[1:])), stream)
    stream.seek(0)
    book = load_workbook(stream)
    assert book["Daily participation"]["H2"].value == "1234567890123456.78"
    assert book["Daily participation"]["H2"].data_type == "s"
    book.close()


def test_zero_pledge_total_is_a_number_not_text():
    """A zero Decimal is falsy but still an exact, summable amount."""
    value = document().participation
    day = replace(value.days[0], pledge_total=Decimal("0"))
    stream = BytesIO()
    participation_xlsx(replace(value, days=(day, *value.days[1:])), stream)
    stream.seek(0)
    book = load_workbook(stream)
    cell = book["Daily participation"]["H2"]
    assert cell.data_type == "n" and cell.value == 0
    assert cell.number_format == '"$"#,##0.00'
    book.close()


def test_daily_table_size_is_carried_only_when_not_default():
    """The default size and sort stay out of links; chosen ones survive."""
    assert "size=" not in ReportQuery.parse(QueryDict()).url()
    assert "sort=" not in ReportQuery.parse(QueryDict()).url()
    chosen = ReportQuery.parse(QueryDict("size=all&sort=date_asc"))
    assert chosen.size == "all" and "size=all" in chosen.url(page=1)
    assert chosen.sort == "date_asc" and "sort=date_asc" in chosen.url()
    # The shared table carries the sort itself, so the filters leave it out.
    assert all(name != "sort" for name, _ in chosen.carried())


def test_daily_table_sorts_every_column_on_the_server():
    """Each heading sorts the whole in-memory document by its exact value;
    the established date tokens still parse and anything else is refused."""
    from django.template.loader import render_to_string

    from parishkit.stewardship.reports.workspace_views import daily_table

    chart = document().participation
    presented = participation_context(chart)

    def dates(sort):
        """The campaign dates in the order one sort token shows them."""
        query = ReportQuery.parse(QueryDict(f"sort={sort}&size=all"))
        table = daily_table(chart, presented, query)["table"]
        return [day.local_date for day, _cells in table.rows]

    by_date = sorted(day.local_date for day in chart.days)
    assert dates("date_asc") == by_date and dates("date_desc") == by_date[::-1]
    most = dates("-first")
    counts = {day.local_date: day.first_responses for day in chart.days}
    assert [counts[day] for day in most] == sorted(counts.values(), reverse=True)
    for token in ("local_date", "first_responses", "-date"):
        with pytest.raises(ValueError):
            ReportQuery.parse(QueryDict(f"sort={token}"))
    context = daily_table(chart, presented, ReportQuery.parse(QueryDict("sort=-first")))
    assert len(context["columns"]) == len(presented["headings"])
    html = render_to_string(
        "stewardship/table-navigator.html",
        {"table": context["table"], "label": "Daily table pages"},
    )
    assert '<input type="hidden" name="sort" value="-first">' in html
    # 5 rows by default, and 5 is the first rows-per-page choice (#1537).
    assert context["table"].size == 5
    assert [value for value, _label, _chosen in context["table"].size_choices] == [
        "5",
        "25",
        "50",
        "100",
        "250",
        "all",
    ]
    # Without a sort the newest date comes first (#1534).
    newest = daily_table(chart, presented, ReportQuery.parse(QueryDict()))
    assert newest["table"].rows[0][0].local_date == max(by_date)
