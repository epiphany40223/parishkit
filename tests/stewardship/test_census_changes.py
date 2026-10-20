"""Census changes worklist wording, status and filters (#528), without a database."""

import csv
import io
from datetime import UTC, datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.reports import pdf_design
from parishkit.stewardship.reports.census_change_views import SORTING
from parishkit.stewardship.reports.census_changes import (
    HEADINGS,
    PDF_WEIGHTS,
    PRIVACY,
    CensusQuery,
    census_csv,
    census_pdf,
    census_xlsx,
    display,
    family_answer,
    select,
    shape,
    status,
)
from parishkit.stewardship.web.dates import UnknownZone

SUBMITTED = datetime(2054, 10, 5, 14, 30, tzinfo=UTC)


def workbook(result, zone):
    """The Census changes XLSX of ``result``, downloaded at ``SUBMITTED``."""
    body = census_xlsx(
        result,
        zone,
        parish="Example Parish",
        campaign="Stewardship 2055",
        as_of=SUBMITTED,
    )
    return load_workbook(io.BytesIO(body))


def row(**values):
    """One proposal row as the query returns it, with overrides."""
    return {
        "id": "1",
        "entity_kind": "member",
        "entity_key": "41",
        "field": "mobile_phone",
        "baseline_available": True,
        "baseline_value": "555-0100",
        "submitted_value": "555-0199",
        "current_available": True,
        "current_value": "555-0100",
        "admin_value_set": False,
        "admin_value": None,
        "handling": "api",
        "decision": "unreviewed",
        "execution": "pending",
        "submitted_at": SUBMITTED,
        "family_duid": 7001,
        "family_name": "Sample",
        "member_name": "Pat Sample",
        "version": 1,
        "resolution_action": None,
        "resolution_note": None,
        "resolution_at": None,
        "resolution_by": None,
    } | values


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        # A terminal execution is the outcome, whatever the decision says.
        ({"execution": "published", "decision": "ignored"}, "published"),
        ({"execution": "resolved_upstream"}, "already"),
        ({"execution": "resolved_external", "handling": "manual"}, "entered"),
        ({"execution": "superseded"}, "superseded"),
        ({"execution": "cancelled"}, "cancelled"),
        # Then an ignored decision, then the execution.
        ({"decision": "ignored", "execution": "conflict"}, "ignored"),
        ({"decision": "approved", "execution": "queued"}, "being_published"),
        ({"execution": "conflict"}, "conflict"),
        ({"decision": "approved", "execution": "failed"}, "to_do"),
        # A pending automatic change waits for an Administrator's decision.
        ({}, "to_review"),
        ({"decision": "approved"}, "to_do"),
        ({"handling": "manual", "field": "suffix"}, "to_do"),
        # A Family address is worked by hand: no verified address read exists.
        ({"entity_kind": "family", "field": "home_address"}, "to_do"),
    ],
)
def test_status_follows_the_spec_mapping(values, expected):
    """The status comes from decision and execution in the spec's order."""
    assert status(row(**values)) == expected


def test_display_words_each_value_shape():
    """Addresses, Members, phones, choices and nothing read as plain text."""
    address = {
        "line1": "1 Main St",
        "line2": "",
        "city": "Springfield",
        "region": "IL",
        "postal_code": "62701",
        "country": "US",
    }
    assert display(address) == "1 Main St, Springfield, IL, 62701, US"
    assert display({"display": "(555) 010-0199", "normalized": "+15550100199"}) == (
        "(555) 010-0199"
    )
    assert display(True) == "Yes" and display(False) == "No"
    assert display(None) == ""
    member = {"first_name": "Lee", "last_name": "Sample", "email": "lee@example.org"}
    assert display(member) == "Lee Sample; Email address: lee@example.org"


def test_family_answer_words_an_unknown_birth_date():
    """A birth date of Unknown reads Unknown; any other removal stays blank."""
    assert family_answer({"field": "birth_date", "submitted_value": None}) == (
        "Unknown"
    )
    assert family_answer({"field": "email", "submitted_value": ""}) == ""
    assert family_answer({"field": "nickname", "submitted_value": None}) == ""
    assert family_answer({"field": "birth_date", "submitted_value": "2001-02-03"}) == (
        "2001-02-03"
    )


def test_shape_names_who_and_hides_unread_family_values():
    """A Family field has no ParishSoft value; a new Member is named."""
    family = shape(
        row(
            entity_kind="family",
            field="email_opt_out",
            handling="manual",
            current_available=True,
            current_value=False,
            submitted_value=True,
        )
    )
    assert family["who"] == "Family" and family["parishsoft_now"] == ""
    assert family["label"] == "Opt out of all parish emails"
    assert family["route_label"] == "By hand"
    new = shape(
        row(
            entity_kind="proposed_member",
            entity_key="p1",
            field="new_member",
            handling="manual",
            current_available=False,
            submitted_value={"first_name": "Ari", "last_name": "Sample"},
        )
    )
    assert new["who"] == "New Member: Ari Sample" and new["kind"] == "new_member"
    edited = shape(row(admin_value_set=True, admin_value="555-0123"))
    assert edited["edited"] == "555-0123" and edited["route_label"] == "Automatic"


def rows():
    """One row of each status the default view weighs."""
    return [
        shape(row(id="review")),
        shape(row(id="todo", handling="manual", field="suffix")),
        shape(row(id="conflict", execution="conflict")),
        shape(row(id="done", execution="published", decision="approved")),
        shape(row(id="old", execution="superseded")),
        shape(row(id="moved", field="moved_household", handling="manual")),
    ]


def ids(result):
    """The kept rows' ids."""
    return [item["id"] for item in result["rows"]]


def test_default_view_shows_open_work_and_to_review_only_to_administrators():
    """Staff see To do and Conflict; an Administrator also sees To review."""
    query = CensusQuery()
    assert ids(select(rows(), query, administrator=False)) == [
        "todo",
        "conflict",
        "moved",
    ]
    assert ids(select(rows(), query, administrator=True)) == [
        "review",
        "todo",
        "conflict",
        "moved",
    ]


def test_filters_narrow_the_rows():
    """History, route, kind and search each narrow the list."""
    every = CensusQuery(status="all")
    assert "old" not in ids(select(rows(), every, administrator=True))
    history = CensusQuery(status="all", history="yes")
    assert "old" in ids(select(rows(), history, administrator=True))
    by_hand = CensusQuery(status="all", route="by_hand")
    assert ids(select(rows(), by_hand, administrator=True)) == ["todo", "moved"]
    moved = CensusQuery(status="all", kind="moved")
    assert ids(select(rows(), moved, administrator=True)) == ["moved"]
    found = CensusQuery(status="all", search="7001")
    assert len(ids(select(rows(), found, administrator=True))) == 5
    assert ids(select(rows(), CensusQuery(search="nobody"), administrator=True)) == []
    # Every change is counted by status, history only when it is shown.
    counts = select(rows(), every, administrator=True)["status_counts"]
    assert counts["to_do"] == 2 and counts["superseded"] == 0
    counts = select(rows(), history, administrator=True)["status_counts"]
    assert counts["superseded"] == 1


def test_dates_are_days_in_the_browser_zone():
    """From and Through are whole local days; a date needs a known zone."""
    query = CensusQuery.parse(
        {"start": "2054-10-05", "end": "2054-10-05", "zone": "America/Chicago"}
    )
    assert ids(select(rows(), query, administrator=True))
    later = CensusQuery.parse({"start": "2054-10-06", "zone": "America/Chicago"})
    assert ids(select(rows(), later, administrator=True)) == []
    with pytest.raises(UnknownZone):
        CensusQuery.parse({"start": "2054-10-05"})
    assert CensusQuery.parse({"zone": "Nowhere/Zone"}).zone == ""


@pytest.mark.parametrize(
    "values",
    [
        {"status": "superseded"},
        {"route": "api"},
        {"kind": "pledge"},
        {"history": "1"},
        {"start": "2054-13-01", "zone": "UTC"},
        {"start": "2054-10-06", "end": "2054-10-05", "zone": "UTC"},
        {"unknown": "x"},
    ],
)
def test_parse_refuses_values_outside_the_closed_choices(values):
    """Only the offered choices are accepted, as single values."""
    with pytest.raises(ValueError):
        CensusQuery.parse(values)


def test_csv_follows_the_page_column_order_in_the_chosen_zone():
    """The download has the page's columns in the page's order (#932):
    the Member's ParishSoft status (#1752) and Member DUID right after Who,
    then Submitted and the latest resolution's three columns (#528), times
    in the chosen zone."""
    result = select(rows(), CensusQuery(), administrator=False)
    text = census_csv(result, ZoneInfo("America/Chicago")).decode()
    table = list(csv.reader(io.StringIO(text)))
    assert tuple(table[0]) == HEADINGS
    assert HEADINGS[:2] == ("Family", "ParishSoft status")
    assert HEADINGS[3:6] == ("Who", "Member's ParishSoft status", "Member DUID")
    assert HEADINGS[12:] == ("Submitted", "Recorded by", "Recorded at", "Note")
    assert table[1][3] == "Pat Sample" and table[1][12] == "2054-10-05T09:30:00-05:00"
    # A current Family and Member have no status; no resolution yet: its
    # three columns are blank.
    assert table[1][1] == table[1][4] == ""
    assert table[1][13:] == ["", "", ""]
    assert len(table) == 4


def test_downloads_put_a_missing_family_s_status_in_its_own_column():
    """A Family the data no longer has: a blank Family cell and the shared
    words in ParishSoft status, never the words in place of a name (#1216)."""
    result = {"rows": [shape(row(family_name=None)), shape(row(id="2"))]}
    zone = ZoneInfo("UTC")
    table = list(csv.reader(io.StringIO(census_csv(result, zone).decode())))
    assert [line[:3] for line in table] == [
        ["Family", "ParishSoft status", "Family DUID"],
        ["", "Not in ParishSoft", "7001"],
        ["Sample", "", "7001"],
    ]
    sheet = workbook(result, zone)["Census changes"]
    assert [sheet.cell(2, column).value for column in (1, 2)] == [
        None,
        "Not in ParishSoft",
    ]


@pytest.mark.parametrize(
    ("status", "words"),
    [
        ("moved", "Not in ParishSoft"),
        ("missing", "Not in ParishSoft"),
        ("inactive", "Inactive in ParishSoft"),
        ("deceased", "Deceased in ParishSoft"),
    ],
)
def test_downloads_put_a_member_s_status_beside_who(status, words):
    """The Who cell's badge (rule 4, #1738) is the shared words in Member's
    ParishSoft status, right after Who, in every format (#1752); Who keeps
    the name, and a current Member or a household change leaves it blank."""
    left = shape(row(member_status=status))
    family = shape(
        row(id="2", entity_kind="family", entity_key="7001", field="home_address")
    )
    result = {"rows": [left, shape(row(id="3", member_status="current")), family]}
    zone = ZoneInfo("UTC")
    table = list(csv.reader(io.StringIO(census_csv(result, zone).decode())))
    assert [line[3:6] for line in table] == [
        ["Who", "Member's ParishSoft status", "Member DUID"],
        ["Pat Sample", words, "41"],
        ["Pat Sample", "", "41"],
        ["Family", "", ""],
    ]
    sheet = workbook(result, zone)["Census changes"]
    assert [sheet.cell(line, 5).value for line in (1, 2)] == [
        "Member's ParishSoft status",
        words,
    ]


def test_downloads_word_member_duid_as_the_page_does():
    """CSV and XLSX give Member DUID right after Who: the DUID, "New Member"
    for a Member added on the form, and blank for a household change (#932)."""

    proposed = row(
        id="2",
        entity_kind="proposed_member",
        entity_key="p-1",
        field="new_member",
        submitted_value={"first_name": "Lee", "last_name": "Sample"},
        member_name=None,
    )
    family = row(id="3", entity_kind="family", entity_key="7001", field="home_address")
    result = {"rows": [shape(item) for item in (row(), proposed, family)]}
    zone = ZoneInfo("UTC")
    table = list(csv.reader(io.StringIO(census_csv(result, zone).decode())))
    assert [line[5] for line in table] == ["Member DUID", "41", "New Member", ""]
    sheet = workbook(result, zone)["Census changes"]
    assert sheet.max_column == len(HEADINGS)
    column = HEADINGS.index("Member DUID") + 1
    assert [sheet.cell(line, column).value for line in range(1, 4)] == [
        "Member DUID",
        "41",
        "New Member",
    ]
    assert sheet.cell(4, column).value in (None, "")


def test_family_duid_column_sorts_both_ways():
    """Family DUID has its own column, and it sorts both ways (#932)."""
    shaped = [shape(row(id=str(duid), family_duid=duid)) for duid in (7002, 13, 7001)]
    assert SORTING.tokens["duid"] == ("duid", False)
    up = SORTING.sort_rows(shaped, "duid")
    assert [item["family_duid"] for item in up] == [13, 7001, 7002]
    down = SORTING.sort_rows(shaped, "-duid")
    assert [item["family_duid"] for item in down] == [7002, 7001, 13]


def test_member_duid_comes_from_a_member_rows_key():
    """A Member change shows its DUID; New Members and the Family have none."""
    assert shape(row())["member_duid"] == 41
    proposed = row(
        entity_kind="proposed_member",
        entity_key="p-1",
        field="new_member",
        submitted_value={"first_name": "Lee", "last_name": "Sample"},
        member_name=None,
    )
    assert shape(proposed)["member_duid"] is None
    family = row(entity_kind="family", entity_key="7001", field="home_address")
    assert shape(family)["member_duid"] is None
    # A key that is not a plain integer is left blank, never a broken page.
    assert shape(row(entity_key="x41"))["member_duid"] is None


def test_member_duid_column_sorts_with_blank_rows_after_members():
    """Member DUID sorts in place both ways (#932); rows without one follow
    the Members on the first click."""
    shaped = [
        shape(row(id="a", entity_key="900")),
        shape(row(id="b", entity_kind="family", entity_key="7001")),
        shape(row(id="c", entity_key="41")),
    ]
    assert SORTING.tokens["member_duid"] == ("member_duid", False)
    up = SORTING.sort_rows(shaped, "member_duid")
    assert [item["id"] for item in up] == ["c", "a", "b"]
    down = SORTING.sort_rows(shaped, "-member_duid")
    assert [item["id"] for item in down] == ["b", "a", "c"]


def test_xlsx_report_information_states_what_the_pdf_frame_states():
    """The workbook's "Report information" sheet (#926) names the download.

    As the Limitations workbook's does, it states the parish, campaign,
    download time in the chosen zone, the row count and the privacy line
    the PDF's page frame shows, and no internal id (#956); the CSV stays
    the table alone.
    """
    zone = ZoneInfo("America/Chicago")
    result = select(rows(), CensusQuery(), administrator=False)
    book = workbook(result, zone)
    assert book.sheetnames == ["Census changes", "Report information"]
    sheet = book["Report information"]
    assert list(sheet.iter_rows(values_only=True)) == [
        ("Report", "Census changes"),
        ("Parish", "Example Parish"),
        ("Campaign", "Stewardship 2055"),
        # A native date-time cell, in the download's zone (9:30 AM Chicago).
        ("Downloaded at", datetime(2054, 10, 5, 9, 30)),
        ("Display time zone", "America/Chicago"),
        ("Changes in this file", f"{len(result['rows']):,}"),
        ("Privacy", PRIVACY),
    ]
    assert sheet["B4"].is_date
    assert sheet["A1"].font.b and sheet.page_setup.orientation == "portrait"
    assert book.properties.title == "Census changes"
    assert book.properties.created == datetime(2054, 10, 5, 14, 30)
    assert b"Report information" not in census_csv(result, zone)


def test_csv_xlsx_and_pdf_carry_the_same_columns_and_rows():
    """One column set in every format (#966): CSV, XLSX and the PDF table.

    The XLSX cells are the CSV's; the PDF table has the same headings and
    values, only its Submitted time worded in the parish's date format.
    """
    zone = ZoneInfo("America/Chicago")
    # A formula-like name: the CSV guards it with an apostrophe, the XLSX
    # stores it as literal text and the PDF shows it as typed.
    given = rows()
    given[1] = shape(
        row(id="todo", handling="manual", field="suffix", family_name="=1+2")
    )
    result = select(given, CensusQuery(), administrator=False)
    table = list(csv.reader(io.StringIO(census_csv(result, zone).decode())))
    book = workbook(result, zone)
    sheet = [
        ["" if value is None else value for value in values]
        for values in book["Census changes"].iter_rows(values_only=True)
    ]
    assert "'=1+2" in [values[0] for values in table]
    table = [[cell.removeprefix("'") for cell in values] for values in table]
    assert sheet == table
    formula_like = [
        cell
        for values in book["Census changes"].iter_rows()
        for cell in values
        if cell.value == "=1+2"
    ]
    assert formula_like and all(cell.data_type == "s" for cell in formula_like)
    drawn = {}
    original = pdf_design.table_pages

    def capture(pdf_table, pdf_rows, frame):
        """Keep the table and rows the PDF lays out, then lay them out."""
        pdf_rows = list(pdf_rows)
        drawn.update(table=pdf_table, rows=pdf_rows, frame=frame)
        return original(pdf_table, pdf_rows, frame)

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(pdf_design, "table_pages", capture)
        body = census_pdf(
            result,
            zone,
            parish="Example Parish",
            campaign="Stewardship 2055",
            as_of=SUBMITTED,
        )
    assert body.startswith(b"%PDF-")
    assert drawn["table"].headings == HEADINGS == tuple(table[0])
    # Every cell matches except the times, which the PDF words as dates.
    submitted = HEADINGS.index("Submitted")
    times = {submitted, HEADINGS.index("Recorded at")}

    def without_submitted(values):
        """The row's cells other than its Submitted and Recorded at times."""
        return [cell for index, cell in enumerate(values) if index not in times]

    assert [without_submitted(values) for values in drawn["rows"]] == [
        without_submitted(values) for values in table[1:]
    ]
    assert drawn["rows"][0][submitted] == "Oct 5, 2054 9:30 AM"
    assert drawn["frame"].eyebrow == "Example Parish · Stewardship 2055"
    assert drawn["frame"].details == ("3 changes in this file.",)
    # Every heading has a PDF column width, in percent, and the columns fill
    # the page.
    assert set(PDF_WEIGHTS) == set(HEADINGS)
    assert sum(PDF_WEIGHTS.values()) == 100
    assert sum(drawn["table"].widths) == pytest.approx(pdf_design.BODY_WIDTH)
    # A compact time keeps its date on one line and its time on the next,
    # the longest stage breaks only between words, and a whole DUID and
    # "Automatic" each stay on one line.
    for heading, text, lines in (
        ("Submitted", "Sep 30, 2026 12:04 PM", ["Sep 30, 2026", "12:04 PM"]),
        ("Recorded at", "Sep 30, 2026 12:04 PM", ["Sep 30, 2026", "12:04 PM"]),
        ("Stage", "Updated in ParishSoft", ["Updated", "in", "ParishSoft"]),
        ("Family DUID", "1234567", ["1234567"]),
        ("Member DUID", "1234567", ["1234567"]),
        ("How it reaches ParishSoft", "Automatic", ["Automatic"]),
    ):
        values = [text if name == heading else "" for name in HEADINGS]
        cell = drawn["table"].cells(values)[HEADINGS.index(heading)]
        assert [line.strip() for line in cell] == lines, heading


@pytest.mark.parametrize(
    ("values", "resolve", "reopen"),
    [
        # A By-hand change still to do, or in conflict, takes the tick and Ignore.
        ({"handling": "manual", "field": "suffix"}, True, False),
        ({"handling": "manual", "execution": "conflict"}, True, False),
        ({"entity_kind": "family", "field": "home_address"}, True, False),
        # An automatic change is the Administrator's to publish.
        ({}, False, False),
        # Ignored: only reopen (offered to Administrators by the page).
        ({"handling": "manual", "decision": "ignored"}, False, True),
        # Final outcomes take nothing.
        ({"handling": "manual", "execution": "resolved_external"}, False, False),
        ({"handling": "manual", "execution": "superseded"}, False, False),
    ],
)
def test_shape_offers_only_the_actions_the_row_allows(values, resolve, reopen):
    """The page draws exactly the actions the service would accept."""
    shown = shape(row(**values))
    assert (shown["can_resolve"], shown["can_reopen"]) == (resolve, reopen)


@pytest.mark.parametrize(
    ("values", "reviews"),
    [
        # To review: approve or ignore it.
        ({}, ["approve", "ignore_automatic"]),
        ({"decision": "approved"}, ["ignore_automatic", "to_review"]),
        ({"decision": "ignored"}, ["approve", "to_review"]),
        # A failed publication may be approved again; a conflict is resolved
        # before it is approved.
        ({"execution": "failed", "decision": "ignored"}, ["approve", "to_review"]),
        ({"execution": "conflict"}, ["ignore_automatic"]),
        # Queued for publication, or final: nothing.
        ({"execution": "queued", "decision": "approved"}, []),
        ({"execution": "published", "decision": "approved"}, []),
        ({"execution": "resolved_upstream"}, []),
        # By-hand rows are reviewed nowhere; they take Record outcome.
        ({"handling": "manual", "field": "suffix"}, []),
        ({"entity_kind": "family", "field": "home_address"}, []),
    ],
)
def test_shape_offers_reviews_only_on_open_automatic_rows(values, reviews):
    """An Automatic row's Review offers exactly what the service accepts
    (#528, ADM-09): a different decision, while the change is open."""
    assert shape(row(**values))["review_actions"] == reviews


def test_shape_reads_the_latest_resolution():
    """The latest resolution's who, when and note reach the page and files."""
    when = datetime(2054, 10, 6, 15, tzinfo=UTC)
    shown = shape(
        row(
            handling="manual",
            execution="resolved_external",
            resolution_action="entered",
            resolution_note="Typed it in",
            resolution_at=when.isoformat(),
            resolution_by="staff@example.org",
        )
    )
    assert shown["status_label"] == "Entered by hand"
    assert shown["resolution_at"] == when
    text = census_csv({"rows": [shown]}, ZoneInfo("UTC")).decode()
    line = list(csv.reader(io.StringIO(text)))[1]
    assert line[13:] == [
        "staff@example.org",
        "2054-10-06T15:00:00+00:00",
        "Typed it in",
    ]
    # A resolution whose portal user is gone still says someone made it.
    gone = shape(row(resolution_action="ignored", resolution_at=when))
    assert gone["resolution_by"] == "Former portal user"


def test_pdf_keeps_a_seven_digit_duid_on_one_line():
    """The narrow DUID columns still hold a whole DUID on one line."""
    from parishkit.stewardship.reports.pdf_design import Table

    table = Table.weighted(HEADINGS, [PDF_WEIGHTS[name] for name in HEADINGS])
    for heading in ("Family DUID", "Member DUID"):
        values = ["1234567" if name == heading else "" for name in HEADINGS]
        assert table.cells(values)[HEADINGS.index(heading)] == ["1234567"], heading


def test_a_bulk_review_covers_one_campaign():
    """A batch whose changes span two campaigns is refused before anything
    is admitted or written (only a crafted post can mix them)."""
    from types import SimpleNamespace

    from parishkit.stewardship.responses.census_review import one_campaign

    def proposal(campaign):
        """A stand-in proposal of ``campaign``."""
        return SimpleNamespace(submission=SimpleNamespace(campaign_id=campaign))

    assert one_campaign([proposal("a"), proposal("a")]) == "a"
    with pytest.raises(ValueError):
        one_campaign([proposal("a"), proposal("b")])
