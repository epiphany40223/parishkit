"""The Campaign comparison downloads (#1018, slice 6), without a database.

The capture (``comparison_exports.capture``) is what leaves the Admin
portal, so these tests hold the campaign-history spec's privacy rules on it:
aggregates only, no note text, the small-count protection, and no current
campaign (or partial year) quartile. The files
(``comparison_documents``) carry the page's Every number columns in every
format, typed in CSV and XLSX, and the PDF is the one-page summary (tiles
and day-by-day charts) followed by the table, inside the shared page frame.
Every figure here is invented.
"""

import csv
import io
import json
import re
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.reports import pdf_design
from parishkit.stewardship.reports.comparison import (
    PAGE_ONLY,
    WITHHELD,
    Band,
    Comparison,
    Day,
    Distribution,
    Renewal,
    Year,
    tiles,
)
from parishkit.stewardship.reports.comparison_documents import (
    NOMINAL,
    PAGE_ONLY_NOTE,
    WITHHELD_NOTE,
    comparison_document,
    render_comparison,
)
from parishkit.stewardship.reports.comparison_exports import capture, headings
from parishkit.stewardship.reports.money import MoneyAmount

from .test_file_design import assert_framed, frames

START = date(2054, 10, 4)
REQUESTED = datetime(2054, 10, 20, 14, 5, tzinfo=UTC)
# A note an archive might carry; it must never reach a download.
NOTE = "Private conversion note <about one Family>"


def series(families, *, active=200, per=100, pledging=True):
    """Days 1 to len(families): each Family pledging ``per`` dollars."""
    return tuple(
        Day(
            number,
            START + timedelta(days=number - 1),
            count,
            active,
            count if pledging else None,
            Decimal(per * count) if pledging else None,
        )
        for number, count in enumerate(families, start=1)
    )


def bands(*counts):
    """A distribution with these band counts, $100 per Family per band."""
    rows = tuple(
        Band(index + 1, count, Decimal(100 * (index + 1) * count))
        for index, count in enumerate(counts)
    )
    total = sum((row.total for row in rows), Decimal(0))
    quartiles = (Decimal("150.00"), Decimal("250.00"), Decimal("400.00"))
    return Distribution(sum(counts), total, rows, quartiles)


def year(value, days, *, day=3, final=True, distribution=None, current=False):
    """A year with its series; ``at_day`` is ``day`` (carried when past it)."""
    if day <= len(days):
        at = days[day - 1]
    elif final and days:
        at = replace(days[-1], carried=True)
    else:
        at = None
    return Year(
        value,
        f"Stewardship {value}",
        current=current,
        final=final,
        start_date=START,
        days=len(days),
        at_day=at,
        end=days[-1] if days else None,
        renewal=Renewal(4, 10) if current else Renewal(reason="No past year."),
        distribution=distribution,
        notes=(("Source", NOTE),),
        series=days,
    )


def comparison(*, financial=True, day=3, withheld_last=True):
    """Day 3 of 2055 against four imported years.

    2054 (last year) has a band of 4, so it is withheld unless
    ``withheld_last`` is false; 2053 is plain; 2052 is partial (its
    quartiles stay inside); 2051 ended after two days, so its day-3
    figures are carried "(final)", and has 3 pledging Families.
    """
    last = bands(10, 4, 6) if withheld_last else bands(10, 12, 6)
    imported = (
        year(2054, series([10, 20, 30, 40]), day=day, distribution=last),
        year(2053, series([5, 10, 15]), day=day, distribution=bands(5, 6, 7)),
        year(2052, series([8, 16, 24]), day=day, final=False, distribution=bands(9)),
        year(2051, series([1, 3], per=50), day=day, distribution=bands(1, 1, 1)),
    )
    current = year(
        2055,
        series([12, 24, 36], per=200),
        day=day,
        final=False,
        current=True,
        distribution=bands(3, 4, 5),
    )
    if not financial:
        strip = lambda found: replace(  # noqa: E731
            found,
            distribution=None,
            series=tuple(replace(d, pledging=None, pledged=None) for d in found.series),
            at_day=found.at_day and replace(found.at_day, pledging=None, pledged=None),
            end=found.end and replace(found.end, pledging=None, pledged=None),
        )
        imported = tuple(strip(found) for found in imported)
        current = strip(current)
    return Comparison(
        "Stewardship 2055",
        day,
        START + timedelta(days=day - 1),
        financial,
        "current",
        current,
        imported,
        imported[0],
        latest_day=3,
    )


def document(found=None, **options):
    """The worker's document for a capture, after a JSON round trip.

    The capture is stored as JSON, so it must survive one unchanged.
    """
    payload = json.loads(json.dumps(capture(found or comparison(**options))))
    return comparison_document(
        payload,
        parish_name="Sample Parish",
        requested_at=REQUESTED,
        timezone="America/New_York",
    )


def render(found, format):
    """The bytes of one format."""
    output = io.BytesIO()
    render_comparison(found, output, format=format)
    return output.getvalue()


# ----------------------------------------------------------------- capture


def test_headings_are_the_page_columns_in_its_order():
    """The Every number columns, as the page's template shows them."""
    assert headings(comparison()) == [
        "Year",
        "Day 1",
        "Days",
        "Families on the day",
        "Percent on the day",
        "Pledges on the day",
        "Average pledge on the day",
        "Renewal on the day",
        "Final Families",
        "Final percent",
        "Final pledges",
        "Final average pledge",
        "First quartile",
        "Median",
        "Third quartile",
    ]
    assert headings(comparison(financial=False)) == [
        "Year",
        "Day 1",
        "Days",
        "Families on the day",
        "Percent on the day",
        "Renewal on the day",
        "Final Families",
        "Final percent",
    ]
    # Before day 1 has ended there is no "on the day" column.
    early = replace(comparison(), day=None)
    assert "Families on the day" not in headings(early)


def test_the_capture_holds_aggregates_only_with_the_protection_applied():
    """Typed cells; withheld averages; quartiles only for final years of 5+.

    Rows are the current campaign first, then the imported years newest
    first, as on the page.
    """
    found = capture(comparison())
    rows = {
        row[0]: dict(zip(found["columns"], row, strict=True)) for row in found["rows"]
    }
    current = rows["Stewardship 2055 (so far)"]
    assert current["Day 1"] == {"date": "2054-10-04"}
    assert current["Families on the day"] == 36
    assert current["Percent on the day"] == "18%"
    assert current["Pledges on the day"] == {"money": "7200.00"}
    # The current campaign's averages stay; its quartiles stay inside.
    assert current["Average pledge on the day"] == {"money": "200.00"}
    assert current["Renewal on the day"] == "40%"
    assert [current[q] for q in ("First quartile", "Median", "Third quartile")] == [
        PAGE_ONLY
    ] * 3
    # Last year has a band of 4: every average is withheld, at any day, but
    # its totals and final quartiles stay.
    last = rows["2054"]
    assert last["Average pledge on the day"] == WITHHELD
    assert last["Final average pledge"] == WITHHELD
    assert last["Pledges on the day"] == {"money": "3000.00"}
    assert last["Median"] == {"money": "250.00"}
    plain = rows["2053"]
    assert plain["Average pledge on the day"] == {"money": "100.00"}
    assert plain["Final average pledge"] == {"money": "100.00"}
    # A partial year's quartiles stay inside.
    assert rows["2052"]["Median"] == PAGE_ONLY
    # 2051 ended after day 2: day 3 keeps the page's "(final)" mark, and
    # with 3 pledging Families its averages and quartiles are withheld.
    old = rows["2051"]
    assert old["Families on the day"] == "3 (final)"
    assert old["Pledges on the day"] == "$150.00 (final)"
    assert old["Average pledge on the day"] == WITHHELD
    assert old["Median"] == WITHHELD
    assert old["Renewal on the day"] is None


def test_the_capture_carries_no_note_and_no_distribution():
    """No archive note, no band or quartile of the current campaign, no median."""
    found = capture(comparison())
    text = json.dumps(found)
    assert NOTE not in text and "Private" not in text
    assert set(found) == {
        "campaign",
        "charts",
        "columns",
        "compared_with",
        "day",
        "financial",
        "report_date",
        "rows",
        "tiles",
        "updating",
    }
    assert "Median so far" not in text and "band" not in text.lower()
    # Only the three day-by-day charts leave the portal.
    assert [chart["title"] for chart in found["charts"]] == [
        "Families responded by campaign day",
        "Percent of active Families by campaign day",
        "Annual pledges by campaign day",
    ]
    assert found["compared_with"] == "2054 and the 2051–2054 average"


def test_outside_tiles_withhold_the_average_pledge_comparisons():
    """Last year's and the average's average pledge are withheld outside.

    The page's own tiles still show every figure and the median so far.
    """
    found = comparison()
    inside = {tile["key"]: tile for tile in tiles(found)}
    outside = {tile["key"]: tile for tile in tiles(found, outside=True)}
    assert inside["average"]["median"] == "$250.00"
    assert outside["average"]["median"] is None
    assert outside["average"]["value"] == "$200.00"
    assert outside["average"]["lines"] == [
        f"Last year on day 3: {WITHHELD}",
        f"2051–2054 average: {WITHHELD}",
    ]
    assert not any(WITHHELD in line for line in inside["average"]["lines"])
    # The other tiles are the page's.
    assert outside["families"] == inside["families"]
    # Without a withheld year nothing is withheld.
    clear = comparison(withheld_last=False)
    clear = replace(clear, imported=clear.imported[:3])
    lines = {tile["key"]: tile for tile in tiles(clear, outside=True)}
    assert not any(WITHHELD in line for line in lines["average"]["lines"])


# ------------------------------------------------------------------ files


def test_csv_is_the_table_with_plain_amounts_and_blank_unknowns():
    """One header row and one row per year; money canonical; no "—"."""
    found = document()
    rows = list(csv.reader(io.StringIO(render(found, "csv").decode())))
    assert rows[0] == list(found.headings)
    assert len(rows) == 1 + 5
    current = dict(zip(rows[0], rows[1], strict=True))
    assert current["Pledges on the day"] == "7200.00"
    assert current["Day 1"] == "2054-10-04"
    assert current["Median"] == PAGE_ONLY
    old = dict(zip(rows[0], rows[5], strict=True))
    assert old["Renewal on the day"] == ""
    body = render(found, "csv").decode()
    assert "—" not in body and NOTE not in body


def test_xlsx_cells_are_typed_and_its_information_names_the_notes():
    """Counts are numbers, money dollar numbers, day 1 a date; notes explained."""
    book = load_workbook(io.BytesIO(render(document(), "xlsx")))
    try:
        sheet = book["Every number"]
        header = [cell.value for cell in sheet[1]]
        current = {name: cell for name, cell in zip(header, sheet[2], strict=True)}
        assert current["Families on the day"].value == 36
        assert current["Pledges on the day"].value == 7200
        assert current["Pledges on the day"].number_format == '"$"#,##0.00'
        assert current["Day 1"].is_date
        info = {
            row[0]: row[1]
            for row in book["Report information"].iter_rows(values_only=True)
        }
        notes = [
            row[1]
            for row in book["Report information"].iter_rows(values_only=True)
            if row[0] == "Note"
        ]
        assert NOMINAL in notes and WITHHELD_NOTE in notes and PAGE_ONLY_NOTE in notes
        assert info["Compared"].startswith("Day 3 of Stewardship 2055")
        assert NOTE not in str(list(book["Report information"].values))
    finally:
        book.close()


def test_pdf_is_the_summary_then_the_charts_then_the_table(monkeypatch):
    """Page 1: tiles and a chart; page 2: two charts; then the table."""
    found = document()
    pages = frames(
        monkeypatch, lambda output: render_comparison(found, output, format="pdf")
    )
    assert len(pages) == 3
    assert_framed(pages, "Campaign comparison")
    summary, charts, table = pages
    titles = [chart["title"] for chart in found.charts]
    assert titles[0] in summary
    assert titles[1] in charts and titles[2] in charts
    assert "Families responded" in summary and "$200.00" in summary
    assert any(WITHHELD in value for value in summary)
    assert any(value.startswith("Day 3 of Stewardship 2055") for value in summary)
    # Table cells wrap to their columns, so read the page's text as a whole.
    text = " ".join(table)
    assert "Stewardship 2055" in text and "(so far)" in text
    assert "3 (final)" in text and "—" in table and PAGE_ONLY.split()[0] in text
    drawn = " ".join(value for page in pages for value in page)
    assert NOTE not in drawn and "Median so far" not in drawn


def test_pdf_summary_stays_inside_the_page_body(monkeypatch):
    """Every chart, with its tick labels and axis titles, fits the body.

    Measured after drawing, in points, against the frame's margins and
    body, with the widest amounts a campaign might have.
    """
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.backends.backend_pdf import PdfPages

    found = comparison()
    rich = replace(
        found,
        current=replace(
            found.current,
            series=tuple(
                replace(day, pledged=day.pledged * 10_000)
                for day in found.current.series
            ),
        ),
    )
    report = document(rich)
    frame = pdf_design.PdfFrame.for_document(report, details=(report.data_line,))
    seen = []

    def save(pdf, figure, **kwargs):
        """Measure each page's axes before saving it."""
        renderer = FigureCanvasAgg(figure).get_renderer()
        scale = 72 / figure.dpi
        for axes in figure.axes:
            box = axes.get_tightbbox(renderer)
            seen.append(
                (
                    box.x0 * scale,
                    pdf_design.PAGE_HEIGHT - box.y1 * scale,
                    box.x1 * scale,
                    pdf_design.PAGE_HEIGHT - box.y0 * scale,
                )
            )

    monkeypatch.setattr(PdfPages, "savefig", save)
    render_comparison(report, io.BytesIO(), format="pdf")
    assert len(seen) == 3
    # Each chart's plot spans nearly the whole body width, as on the page.
    for left, _, right, _ in seen:
        assert right - left > 0.9 * pdf_design.BODY_WIDTH
    for left, top, right, bottom in seen:
        assert left >= pdf_design.MARGIN - 0.5
        assert right <= pdf_design.PAGE_WIDTH - pdf_design.MARGIN + 0.5
        assert top >= frame.body_top
        assert bottom <= frame.body_bottom


@pytest.mark.parametrize("financial", [True, False])
def test_charts_are_full_width_and_never_overlap(financial):
    """Every chart box is the body's width, inside the body, apart.

    On the summary page the charts sit below the tiles and above the notes;
    on a later page they start below the header and end above the footer.
    """
    from parishkit.stewardship.reports.comparison_documents import (
        LEGEND_ROW,
        NOTE_PITCH,
        TILE_PAD,
        chart_layout,
        notes_lines,
        tile_width,
        tiles_height,
    )

    report = document(financial=financial)
    frame = pdf_design.PdfFrame.for_document(report, details=(report.data_line,))
    tiles_top, notes_top, pages = chart_layout(report, frame)
    tiles_bottom = tiles_top + tiles_height(
        report.tiles, tile_width(report.tiles) - 2 * TILE_PAD
    )
    assert notes_top + NOTE_PITCH * len(notes_lines(report)) <= frame.body_bottom
    placed = [entry for _, entries in pages for entry in entries]
    assert [chart for chart, _ in placed] == list(report.charts)
    assert pages[0][1], "the summary page holds at least one chart"
    for index, (legend_top, entries) in enumerate(pages):
        floor = tiles_bottom if index == 0 else frame.body_top
        ceiling = notes_top if index == 0 else frame.body_bottom
        assert legend_top >= floor
        previous = legend_top + LEGEND_ROW
        for _, (x, y, width, height) in entries:
            assert x == pdf_design.MARGIN and width == pdf_design.BODY_WIDTH
            # The chart's title is drawn just above its box.
            assert y >= previous and y + height <= ceiling
            previous = y + height
    assert len(pages) == 2


def test_two_charts_fill_a_later_page_despite_float_rounding():
    """A body whose bounds round badly still holds two charts per page.

    The two charts end exactly at the body's bottom; without a tolerance
    float rounding (as with these bounds) put each on a page of its own.
    """
    from types import SimpleNamespace

    from parishkit.stewardship.reports.comparison_documents import chart_layout

    report = document()
    top, bottom = 42.093, 560.166
    frame = SimpleNamespace(body_top=top, body_bottom=bottom, body_height=bottom - top)
    _, _, pages = chart_layout(report, frame)
    assert len(pages[1][1]) == 2


def test_without_financial_stewardship_no_pledge_figure_leaves():
    """No pledge column, tile or chart, and no dollars note."""
    found = document(financial=False)
    assert "Final pledges" not in found.headings
    assert [tile["label"] for tile in found.tiles] == [
        "Families responded",
        "Percent of active Families",
        "Renewal of last year's responders",
    ]
    assert len(found.charts) == 2
    assert NOMINAL not in found.notes
    assert not any(
        isinstance(value, MoneyAmount) for row in found.rows for value in row
    )


def test_before_day_one_the_files_hold_final_results_only():
    """No day is compared: no tiles or day columns; the data line says why."""
    early = comparison()
    early = replace(
        early,
        day=None,
        report_date=None,
        current=replace(early.current, at_day=None),
        imported=tuple(replace(found, at_day=None) for found in early.imported),
    )
    found = document(early)
    assert found.day is None and found.tiles == ()
    assert "Day 1 has not ended yet" in found.data_line
    # The summary page, a page of charts, then the table.
    assert len(re.findall(rb"/Type /Page\b", render(found, "pdf"))) == 3


@pytest.mark.parametrize(
    "value", [{"money": "12.50"}, {"date": "2054-10-04"}, 7, "x", None]
)
def test_cells_round_trip(value):
    """Each stored cell becomes the renderer's value."""
    from parishkit.stewardship.reports.comparison_documents import cell

    result = cell(value)
    if isinstance(value, dict) and "money" in value:
        assert result == MoneyAmount(1250)
    elif isinstance(value, dict):
        assert result == date(2054, 10, 4)
    else:
        assert result == value


def test_updating_figures_say_so_in_every_file(monkeypatch):
    """A comparison read from the last ready figures says it is updating."""
    found = document(replace(comparison(), status="updating"))
    labels = dict(found.metadata)
    assert labels["Figures"].startswith("Updating:")
    pages = frames(
        monkeypatch, lambda output: render_comparison(found, output, format="pdf")
    )
    assert all(any(value.startswith("Updating:") for value in drawn) for drawn in pages)


def test_a_small_final_year_is_withheld_whatever_its_reason():
    """A final year under 5 pledging Families withholds its averages outside.

    Even when its bands cannot be shown (an older band set), its pledging
    count alone is small enough to single Families out, so every average
    pledge of it is withheld, and so is the prior-years average pledge.
    A year of 5 or more with the same reason keeps its averages.
    """
    from parishkit.stewardship.reports.comparison import (
        OLDER_BANDS,
        outside_document,
        withheld_years,
    )

    def older(pledging):
        """An older-band-set distribution of ``pledging`` Families."""
        return Distribution(pledging, Decimal(100 * pledging), (), None, OLDER_BANDS)

    found = comparison(withheld_last=False)
    small = year(2050, series([1, 2, 3]), distribution=older(3))
    large = year(2049, series([5, 6, 7]), distribution=older(7))
    found = replace(found, imported=(*found.imported[:2], small, large))
    assert withheld_years(found) == {2050}
    rows = {row[0]: row for row in capture(found)["rows"]}
    columns = capture(found)["columns"]
    average = columns.index("Average pledge on the day")
    final = columns.index("Final average pledge")
    assert rows["2050"][average] == WITHHELD and rows["2050"][final] == WITHHELD
    assert rows["2049"][average] == {"money": "100.00"}
    printed = {entry["year"]: entry for entry in outside_document(found)["years"]}
    assert printed[2050]["at_day"]["average_pledge"] is None
    assert printed[2049]["at_day"]["average_pledge"] == "100.00"
    assert outside_document(found)["averages"]["average"] is None
    lines = {tile["key"]: tile for tile in tiles(found, outside=True)}["average"]
    assert lines["lines"][-1].endswith(WITHHELD)
