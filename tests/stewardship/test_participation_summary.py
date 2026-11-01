"""The Response dashboard's Summary in its export files (#1557), no database.

The export captures the page's Summary (``participation_summary``) when it
is requested, and every format carries it beside the daily rows: the CSV's
Summary rows, the XLSX's Summary sheet and the PDF's Summary table.
"""

import csv
import io
import json
from dataclasses import replace
from datetime import datetime
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.reports import charts, pdf_design
from parishkit.stewardship.reports.money import MoneyAmount
from parishkit.stewardship.reports.participation import participation_csv
from parishkit.stewardship.reports.participation_summary import (
    COLUMNS,
    ParticipationSummary,
    summary_document,
)
from parishkit.stewardship.reports.response_metrics import Stage
from parishkit.stewardship.reports.spreadsheets import participation_xlsx

from .test_daily_digest_content import document as digest_document
from .test_participation_rendering import csv_table, document, summary

STAGES = (
    Stage("invited", 10),
    Stage("link_followed", 6, "Includes mail-scanner prefetches."),
    Stage("form_opened", 5),
    Stage("progressed", 4),
    Stage("submitted", 3),
)


def captured(statistics, stages=STAGES, scope="current"):
    """The capture document for ``statistics`` and a funnel of ``stages``."""
    return summary_document(statistics, SimpleNamespace(stages=list(stages)), scope)


def test_the_capture_is_the_pages_summary_as_plain_json():
    """Labels, counts, exact amounts and the Data as of instant, no identity.

    The document is plain JSON (the SQL shape check reads it) and holds
    counts and totals only: no campaign, source or Family identifier (#966).
    """
    statistics = digest_document().statistics
    payload = captured(statistics)
    assert set(payload) == {"observed_at", "population", "stages"}
    assert payload["observed_at"] == statistics.observed_at.isoformat()
    assert [row["label"] for row in payload["population"]] == [
        "Families",
        "Active Members",
        "Annual pledges",
        "Last year's pledges (all Families)",
    ]
    assert payload["population"][0] == {
        "label": "Families",
        "value": statistics.active.families,
        "note": "",
    }
    pledges = payload["population"][2]["value"]
    assert pledges == {"money": statistics.active.annual_pledge.canonical}
    assert [row["label"] for row in payload["stages"]] == [
        "Invited",
        "Link followed",
        "Form opened",
        "Progressed past the first step",
        "Submitted",
    ]
    assert payload["stages"][1] == {
        "label": "Link followed",
        "count": 6,
        "note": "Includes mail-scanner prefetches.",
    }
    text = json.dumps(payload)
    assert str(statistics.campaign_id) not in text
    assert str(statistics.configuration_id) not in text
    assert ParticipationSummary.from_document(json.loads(text)).stages[0] == (
        "Invited",
        10,
        "",
    )


def test_the_capture_counts_the_exported_population():
    """The historical capture (#2324) holds the cohort's figures, and the
    Members tile's note says which Members it counts, as the page does."""
    statistics = digest_document().statistics
    population = captured(statistics, scope="historical")["population"]
    assert [row["value"] for row in population[:2]] == [
        statistics.cohort.families,
        statistics.cohort.active_members,
    ]
    assert population[1]["note"] == "Members of the Families still in ParishSoft"
    assert population[2]["value"] == {
        "money": statistics.cohort.annual_pledge.canonical
    }
    assert {row["note"] for row in population[2:]} == {""}


def test_an_unknown_figure_carries_the_pages_reason_never_zero():
    """With no ParishSoft data each figure is null with the page's reason."""
    statistics = replace(
        digest_document().statistics,
        active=None,
        comparison_pledge_all=MoneyAmount(None),
    )
    population = captured(statistics)["population"]
    assert {row["value"] for row in population} == {None}
    assert {row["note"] for row in population} == {
        "No ParishSoft data has been loaded for this campaign yet."
    }
    without_giving = captured(replace(statistics, financial_enabled=False))
    assert len(without_giving["population"]) == 2


def test_rows_add_data_as_of_and_give_shares_only_once_invited():
    """Data as of follows the population; Invited itself has no share, and no
    stage has one before an invitation was delivered, as on the page."""
    rows = summary().rows()
    assert [row[1] for row in rows[:5]] == [
        "Campaign Families",
        "Active Members",
        "Annual pledges",
        "Last year's pledges (all Families)",
        "Data as of",
    ]
    assert {row[0] for row in rows[:5]} == {"Population"}
    assert {row[0] for row in rows[5:]} == {"Families by stage"}
    assert [row[3] for row in rows[5:]] == [
        None,
        (600, 1000),
        (500, 1000),
        (450, 1000),
        (409, 1000),
    ]
    nobody = replace(summary(), stages=(("Invited", 0, ""), ("Submitted", 2, "")))
    assert {row[3] for row in nobody.rows()} == {None}


def test_csv_leads_with_the_summary_rows_under_the_section_column():
    """One table: Section, the Summary's columns, then the daily columns.

    Counts are plain numbers, an amount its exact decimal text, a share the
    page's "41%", Data as of ISO 8601 in the requester's zone, an unknown
    value blank (never zero); a daily row leaves the Summary's cells blank.
    """
    stream = io.BytesIO()
    participation_csv(replace(document(), summary=summary()), stream)
    reader = csv.DictReader(io.StringIO(csv_table(stream.getvalue())))
    rows = list(reader)
    assert reader.fieldnames[: 1 + len(COLUMNS)] == ["Section", *COLUMNS]
    sections = [row["Section"] for row in rows]
    assert sections == ["Summary"] * 10 + ["Daily participation"] * 3
    by_figure = {row["Figure"]: row for row in rows[:10]}
    assert by_figure["Campaign Families"]["Value"] == "1234"
    assert by_figure["Annual pledges"]["Value"] == "123456.78"
    assert by_figure["Last year's pledges (all Families)"]["Value"] == ""
    assert (
        by_figure["Last year's pledges (all Families)"]["Note"]
        == "The total is not added up."
    )
    assert by_figure["Data as of"]["Value"] == "2026-11-01 21:30:00-08:00"
    assert by_figure["Invited"]["Share compared with Invited"] == ""
    assert by_figure["Submitted"]["Share compared with Invited"] == "41%"
    assert by_figure["Link followed"]["Note"] == "Includes mail-scanner prefetches."
    assert all(row["Campaign date"] == "" for row in rows[:10])
    assert all(row["Group"] == "" and row["Campaign date"] for row in rows[10:])


def test_xlsx_summary_sheet_keeps_native_types():
    """Counts, dollars, a date-time and a percentage; unknown reads Unavailable."""
    stream = io.BytesIO()
    participation_xlsx(replace(document(), summary=summary()), stream)
    book = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        assert book.sheetnames == [
            "Summary",
            "Daily participation",
            "Report information",
        ]
        sheet = book["Summary"]
        # Under the header block (R59 on #2324), the table's headings.
        assert [cell.value for cell in sheet[5]] == list(COLUMNS)
        rows = {row[1].value: row for row in sheet.iter_rows(min_row=6)}
        families = rows["Campaign Families"][2]
        assert families.value == 1234 and families.number_format == "#,##0"
        pledges = rows["Annual pledges"][2]
        assert pledges.value == 123456.78 and pledges.number_format == '"$"#,##0.00'
        assert rows["Last year's pledges (all Families)"][2].value == "Unavailable"
        observed = rows["Data as of"][2]
        assert observed.is_date
        assert observed.value == datetime(2026, 11, 1, 21, 30)
        share = rows["Submitted"][3]
        assert share.value == 0.409 and share.number_format == "0%"
        assert rows["Invited"][3].value is None
        assert not any(
            cell.data_type == "f" for page in book for row in page for cell in row
        )
    finally:
        book.close()


def test_pdf_summary_table_reads_as_the_page(monkeypatch):
    """The PDF's Summary table holds every column in the page's words."""
    seen = []
    pages = pdf_design.table_pages

    def keep(table, rows, frame):
        """Keep each table's headings and rows, then paginate as usual."""
        seen.append((table.headings, rows))
        return pages(table, rows, frame)

    monkeypatch.setattr(pdf_design, "table_pages", keep)
    output = io.BytesIO()
    value = replace(document(), summary=summary())
    charts.render_participation(value, output, format="pdf")
    (headings, rows), (daily, _) = seen
    assert headings == COLUMNS
    shown = {row[1]: row for row in rows}
    assert shown["Campaign Families"][2] == "1,234"
    assert shown["Annual pledges"][2] == "$123,456.78"
    assert shown["Last year's pledges (all Families)"][2] == "Unavailable"
    assert shown["Submitted"][3] == "41%"
    assert shown["Data as of"][2] == datetime(
        2026, 11, 1, 21, 30, tzinfo=ZoneInfo("America/Los_Angeles")
    )
    assert daily[0] == "Campaign date"


@pytest.mark.parametrize(
    ("scope", "counted"),
    [
        ("historical", "Including Families that left"),
        ("current", "In this campaign today"),
    ],
)
def test_every_format_heads_its_file_with_the_captured_population(
    scope, counted, monkeypatch
):
    """R59 on #2324: the CSV, each XLSX table sheet and every PDF page say
    "Families counted: …" under the report title, beside Source as of, from
    the request's captured scope; the Population row and column stay."""
    value = replace(document(), summary=summary(), population_scope=scope)
    source = value.source_as_of.astimezone(ZoneInfo(value.browser_timezone))
    stream = io.BytesIO()
    participation_csv(value, stream)
    records = list(csv.reader(io.StringIO(stream.getvalue().decode())))
    assert records[:4] == [
        ["Report", "Family participation"],
        ["Families counted", counted],
        ["Source as of", source.isoformat(sep=" ")],
        [],
    ]
    assert records[4][0] == "Section"
    # The daily rows' Population column still names it.
    population = records[4].index("Population")
    assert {row[population] for row in records[5:] if row[0] != "Summary"} == {counted}
    stream = io.BytesIO()
    participation_xlsx(value, stream)
    book = load_workbook(io.BytesIO(stream.getvalue()))
    try:
        for name, heading in (
            ("Summary", "Group"),
            ("Daily participation", "Campaign date"),
        ):
            sheet = book[name]
            assert [
                (sheet.cell(r, 1).value, sheet.cell(r, 2).value) for r in (1, 2)
            ] == [
                ("Report", "Family participation"),
                ("Families counted", counted),
            ], name
            assert sheet["A3"].value == "Source as of"
            assert sheet["B3"].value == source.replace(tzinfo=None)
            assert sheet["A4"].value is None and sheet["A5"].value == heading
            # The heading row stays the frozen and printed one.
            assert sheet.freeze_panes == "A6" and sheet.print_title_rows == "$5:$5"
        information = {
            cells[0].value: cells[1].value
            for cells in book["Report information"].iter_rows()
        }
        assert information["Population"] == counted
    finally:
        book.close()
    lines, draw = [], pdf_design.Canvas.text

    def spy(self, x, y, text, *args, **kwargs):
        """Record one drawn line, then draw it."""
        if text:
            lines.append(text)
        return draw(self, x, y, text, *args, **kwargs)

    monkeypatch.setattr(pdf_design.Canvas, "text", spy)
    output = io.BytesIO()
    charts.render_participation(value, output, format="pdf")
    assert output.getvalue().startswith(b"%PDF")
    title = f"Family participation — {counted}"
    starts = [index for index, line in enumerate(lines) if line == title]
    # Every page (the Summary, the chart and the daily table) has the line
    # first under its title, then its section.
    assert len(starts) == 3
    assert [lines[index + 1] for index in starts] == [
        f"Families counted: {counted}"
    ] * 3
    assert [lines[index + 2] for index in starts] == [
        "Summary",
        "Daily participation",
        "Daily participation",
    ]
