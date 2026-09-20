"""Multi-Ministry packet mapping, privacy and CSV/XLSX/PDF section boundaries."""

import csv
import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.reports.information_rendering import PAGE_LINES
from parishkit.stewardship.reports.ministry_exports import (
    MAX_PACKET_MINISTRIES,
    packet_parameters,
)
from parishkit.stewardship.reports.ministry_packets import (
    HEADINGS,
    MAX_CELL_CHARACTERS,
    packet_document,
    packet_pages,
    render_packet,
    sheet_names,
)

MOMENT = datetime(2026, 9, 20, 2, 30, tzinfo=UTC)
OUTCOME = HEADINGS.index("Outcome")


def item(**values):
    """One captured request row, as SQL projects it, with hostile spreadsheet text."""
    return (
        dict(
            id="r",
            entity_kind="member",
            action="join",
            state="new",
            outcome=None,
            submitted_at=MOMENT.isoformat(),
            member_name="=Example Member",
            member_duid=12345,
            proposed_id=None,
            email_visibility="available",
            emails=[{"value": "member@example.org"}, {"value": None}],
            phone_visibility="available",
            phones={"home": "202-555-0123", "mobile": None},
            email_contact_at=None,
            phone_contact_at=None,
            notes=None,
        )
        | values
    )


def packet(
    sections=None,
    *,
    ministries=None,
    history=False,
    timezone="UTC",
    year_label=None,
):
    """Build a document from a closed capture payload."""
    if sections is None:
        sections = [dict(duid=9, name="Choir", chairs=["Pat Lee"], rows=[item()])]
    payload = dict(
        total=sum(len(section["rows"]) for section in sections),
        sections=sections,
        metadata=dict(
            id="campaign",
            name="Stewardship 2027",
            source_id="source",
            source_generation=1234,
            source_as_of=MOMENT.isoformat(),
            observed_at=MOMENT.isoformat(),
            start_date="2026-09-01",
            end_date="2026-10-31",
            year_label=year_label,
        ),
    )
    return packet_document(
        payload,
        packet_parameters(ministries, history=history),
        parish_name="Sample Parish",
        requested_at=MOMENT,
        timezone=timezone,
    )


def rendered(document, format):
    """Render to bytes through the same closed dispatcher the worker uses."""
    output = io.BytesIO()
    render_packet(document, output, format=format)
    return output.getvalue()


def test_selection_is_canonical_and_every_other_filter_is_neutral():
    """One packet capture has one meaning; SQL enforces the same closed shape."""
    everything = packet_parameters(None, history=False)
    assert everything["action"] == "packet" and everything["ministry"] is None
    assert everything["ministries"] is None
    assert everything["filters"] == dict(
        search="",
        activity="any",
        history="current",
        state="any",
        start="",
        end="",
        sort="name",
    )
    chosen = packet_parameters((4, 9), history=True)
    assert chosen["ministries"] == [4, 9] and chosen["filters"]["history"] == "all"
    for invalid in (
        (),
        (9, 4),
        (9, 9),
        (0,),
        (2**31,),
        ("9",),
        [4, 9],
        tuple(range(1, MAX_PACKET_MINISTRIES + 2)),
    ):
        with pytest.raises(ValueError):
            packet_parameters(invalid, history=False)
    with pytest.raises(ValueError):
        packet_parameters(None, history="all")


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        ({}, ""),
        ({"state": "in_progress"}, ""),
        ({"state": "cancelled"}, ""),
        ({"state": "resolved", "outcome": "joined"}, "Joined ministry"),
        ({"state": "resolved", "outcome": "leave_confirmed"}, "Left ministry"),
        (
            {"state": "resolved", "outcome": "declined"},
            "Declined / no longer interested",
        ),
        ({"state": "closed_no_response", "outcome": "no_response"}, "No response"),
        ({"state": "resolved", "outcome": "duplicate"}, "Duplicate request"),
        (
            {"state": "resolved", "outcome": "other", "notes": "Moved to St. Mark"},
            "Other: Moved to St. Mark",
        ),
        ({"state": "resolved", "outcome": "other"}, "Other"),
        # Notes are private unless the outcome is Other, even if a capture had them.
        (
            {"state": "resolved", "outcome": "joined", "notes": "PRIVATE"},
            "Joined ministry",
        ),
    ],
)
def test_exact_outcome_mapping_and_unresolved_is_blank(values, expected):
    """The specification's labels, with blank left for completion by hand."""
    section = dict(duid=9, name="Choir", chairs=[], rows=[item(**values)])
    assert packet([section]).sections[0].rows[0][OUTCOME] == expected


def test_contact_dates_prefill_in_the_display_zone_and_privacy_holds():
    """Recorded dates are local calendar dates; unpublished contacts never render."""
    rows = [
        item(email_contact_at=MOMENT.isoformat()),
        item(
            member_name="Proposed Person",
            member_duid=None,
            proposed_id="00000000-0000-0000-0000-000000000001",
            entity_kind="proposed_member",
            action="leave",
            email_visibility="not_published",
            emails=None,
            phone_visibility="not_published",
            phones=None,
            phone_contact_at=MOMENT.isoformat(),
        ),
    ]
    document = packet(
        [dict(duid=9, name="Choir", chairs=["Pat Lee", "Sam Roe"], rows=rows)],
        timezone="America/Detroit",
    )
    first, second = document.sections[0].rows
    # 02:30 UTC is still the previous evening in Detroit.
    assert first == (
        "=Example Member",
        "12345",
        "Join",
        "New",
        "member@example.org",
        "2026-09-19",
        "home: 202-555-0123",
        "",
        "",
    )
    assert second[1] == "" and second[2] == "Leave"
    assert second[4] == second[6] == "Not published" and second[7] == "2026-09-19"
    assert dict(document.sections[0].details)["Chairs"] == "Pat Lee, Sam Roe"
    assert dict(document.sections[0].details)["Stewardship period"] == (
        "2026-09-01 to 2026-10-31"
    )
    # No configured label: the campaign-year rule falls back to the start year.
    assert dict(document.sections[0].details)["Stewardship year"] == "2026"
    assert dict(document.metadata)["Stewardship year"] == "2026"
    assert "hidden" not in str(document) and document.item_count == 2


def test_incomplete_or_naive_captures_are_refused():
    """A packet never silently renders fewer rows than SQL captured."""
    section = dict(duid=9, name="Choir", chairs=[], rows=[item()])
    with pytest.raises(ValueError):
        packet_document(
            dict(packet_payload(), total=2, sections=[section]),
            packet_parameters(None, history=False),
            parish_name="P",
            requested_at=MOMENT,
            timezone="UTC",
        )
    with pytest.raises(ValueError):
        packet([dict(section, rows=[item(email_contact_at="2026-09-20T02:30:00")])])


def packet_payload():
    """The capture metadata alone, for malformed-payload cases."""
    return dict(
        metadata=dict(
            id="c",
            name="n",
            source_id="s",
            source_generation=1,
            source_as_of=MOMENT.isoformat(),
            observed_at=MOMENT.isoformat(),
            start_date="2026-09-01",
            end_date="2026-10-31",
        )
    )


def three_ministries():
    """Captured order is authoritative; one Ministry is selected but empty."""
    return [
        dict(duid=4, name="Altar Servers", chairs=[], rows=[item(action="leave")]),
        dict(duid=7, name="Bereavement", chairs=["Pat Lee"], rows=[]),
        dict(duid=9, name="Choir", chairs=[], rows=[item(), item(member_name="Zed")]),
    ]


def test_csv_has_a_blank_row_and_repeated_headings_between_ministries():
    """One file, formula-safe, with each Ministry self-describing."""
    rows = list(
        csv.reader(io.StringIO(rendered(packet(three_ministries()), "csv").decode()))
    )
    assert rows[0] == ["Report", "Ministry follow-up packet"]
    assert rows.count(list(HEADINGS)) == 3 and rows.count([]) == 3
    starts = [index for index, row in enumerate(rows) if row[:1] == ["Ministry"]]
    assert [rows[index][1] for index in starts] == [
        "Altar Servers",
        "Bereavement",
        "Choir",
    ]
    # Every Ministry block is preceded by its blank separator row.
    assert all(rows[index - 1] == [] for index in starts)
    assert ["Chairs", "None recorded"] in rows and ["Chairs", "Pat Lee"] in rows
    assert [row[0] for row in rows if row and row[0].endswith("Example Member")] == [
        "'=Example Member",
        "'=Example Member",
    ]


def test_xlsx_has_one_literal_sheet_per_ministry_with_safe_distinct_names():
    """Sheet titles survive duplicates, truncation and forbidden characters."""
    long = "Saint Example Ministry of Extraordinary Length"
    names = ["Choir", "choir", "A/B: [Youth]?*", long, long + " Two", "'", ""]
    titles = sheet_names(names)
    assert len({title.lower() for title in titles}) == len(names)
    assert all(
        0 < len(title) <= 31 and not set(title) & set("[]:*?/\\") for title in titles
    )
    assert titles[:3] == ["Choir", "choir (2)", "A B   Youth"]
    assert titles[-2:] == ["Ministry", "Ministry (2)"]
    assert "Report information" not in sheet_names(["Report Information"])
    # The application reserves History, and truncation can expose an apostrophe.
    assert sheet_names(["History", "history"]) == ["History (2)", "history (3)"]
    cut = sheet_names(["A" * 30 + "'tail", "'" * 40, "B" * 29 + "''"])
    assert cut == ["A" * 30, "Ministry", "B" * 29]
    # Leading forbidden characters become spaces, which must not use the budget.
    assert sheet_names(["[*] " + "C" * 40]) == ["C" * 31]
    assert all(not title.startswith("'") and not title.endswith("'") for title in cut)

    book = load_workbook(io.BytesIO(rendered(packet(three_ministries()), "xlsx")))
    assert book.sheetnames == [
        "Altar Servers",
        "Bereavement",
        "Choir",
        "Report information",
    ]
    choir = book["Choir"]
    # Pinned independently of the renderer: six header lines, one blank row,
    # then headings on row 8. A header line added or lost must fail here.
    assert [choir.cell(row, 1).value for row in range(1, 8)] == [
        "Ministry",
        "Ministry DUID",
        "Chairs",
        "Stewardship campaign",
        "Stewardship year",
        "Stewardship period",
        None,
    ]
    assert [cell.value for cell in choir[8]] == list(HEADINGS)
    assert choir.cell(9, 1).value == "=Example Member"
    assert choir.cell(9, 1).data_type == "s"  # A literal string, never a formula.
    # A selected but empty Ministry has its header and headings, no request rows.
    assert book["Bereavement"].max_row == 8
    # An empty packet is still a valid workbook carrying its provenance.
    assert load_workbook(io.BytesIO(rendered(packet([]), "xlsx"))).sheetnames == [
        "Report information"
    ]


def test_pdf_starts_each_ministry_on_a_new_page_and_keeps_all_text():
    """Overflow adds pages; it never truncates or merges two Ministries."""
    many = [item(member_name=f"Member {number:03d}") for number in range(40)]
    long = item(state="resolved", outcome="other", notes="word " * 400)
    sections = [
        dict(duid=4, name="Altar Servers", chairs=[], rows=many),
        dict(duid=9, name="Choir", chairs=[], rows=[long]),
    ]
    document = packet(sections)
    pages = list(packet_pages(document))
    assert all(0 < len(page) <= PAGE_LINES for page in pages)
    firsts = [page[0] for page in pages]
    assert firsts[0].startswith("Report: ")
    assert sum(line.startswith("Ministry: ") for line in firsts) == 2
    assert [line for line in firsts if line.startswith("Ministry: ")] == [
        "Ministry: Altar Servers",
        "Ministry: Choir",
    ]
    text = "\n".join(line for page in pages for line in page)
    assert all(f"Member {number:03d}" in text for number in range(40))
    assert text.count("word") == 400
    assert rendered(document, "pdf").startswith(b"%PDF")
    with pytest.raises(ValueError):
        render_packet(document, io.BytesIO(), format="zip")


def test_xlsx_refuses_a_value_beyond_the_cell_limit_instead_of_truncating():
    """The library would cut it silently; an incomplete packet must not publish."""
    limit = MAX_CELL_CHARACTERS
    fits = [dict(duid=9, name="Choir", chairs=["Q" * limit], rows=[])]
    book = load_workbook(io.BytesIO(rendered(packet(fits), "xlsx")))
    assert len(book["Choir"].cell(3, 2).value) == limit
    over = [dict(duid=9, name="Choir", chairs=["Q" * (limit + 1)], rows=[])]
    with pytest.raises(ValueError):
        rendered(packet(over), "xlsx")
    # CSV and PDF have no such limit and keep the complete value, even as one
    # unbroken token that the PDF paginator must wrap across many pages.
    assert ("Q" * (limit + 1)).encode() in rendered(packet(over), "csv")
    wrapped = "".join(line for page in packet_pages(packet(over)) for line in page)
    assert wrapped.count("Q") == limit + 1


def test_stewardship_year_is_the_configured_label_not_a_derived_date():
    """An autumn campaign funding next year is labelled by its configuration."""
    configured = packet(year_label="2027")
    assert dict(configured.metadata)["Stewardship year"] == "2027"
    assert dict(configured.sections[0].details)["Stewardship year"] == "2027"
    assert "Stewardship year,2027" in rendered(configured, "csv").decode()
    # Without a label the shared campaign-year rule falls back to the start year,
    # as it also must for a blank label or a capture that has no such key.
    assert dict(packet().metadata)["Stewardship year"] == "2026"
    assert dict(packet(year_label="").metadata)["Stewardship year"] == "2026"
    section = dict(duid=9, name="Choir", chairs=[], rows=[])
    payload = dict(packet_payload(), total=0, sections=[section])
    assert "year_label" not in payload["metadata"]
    keyless = packet_document(
        payload,
        packet_parameters(None, history=False),
        parish_name="P",
        requested_at=MOMENT,
        timezone="UTC",
    )
    assert dict(keyless.metadata)["Stewardship year"] == "2026"
