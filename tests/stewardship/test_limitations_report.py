"""Limitations report shaping and exports (#247 follow-up, #970).

The report was "Talents and limitations"; since #970 it lists only the
limitations, and its SQL (``stewardship_limitations_report_v2`` since #1685)
neither lists nor counts older answers' stored talents. Since #1685 the page
is the cannot-participate stream's review worklist, and each download row
carries its Stage.
"""

import copy
import csv
import io
from datetime import UTC, datetime
from uuid import UUID
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from parishkit.stewardship.reports import pdf_design
from parishkit.stewardship.reports.limitations import (
    FAMILY_CAPTION,
    MEMBER_CAPTION,
    PDF_WEIGHTS,
    PRIVACY,
    LimitationsQuery,
    limitations_csv,
    limitations_pdf,
    limitations_xlsx,
    shape_result,
)
from parishkit.stewardship.responses.service import DEFAULT_TALENTS

PAINTER = DEFAULT_TALENTS[0].id
WHEN = "2026-10-05T14:00:00+00:00"
# ParishSoft status follows each Family name (#1216).
MEMBER_HEADINGS = [
    "Member",
    "Member DUID",
    "Family",
    "ParishSoft status",
    "Family DUID",
    "Cannot participate in ministries",
    "Latest response",
    "Stage",
]
FAMILY_HEADINGS = [
    "Family",
    "ParishSoft status",
    "Family DUID",
    "Limitation",
    "Latest response",
    "Stage",
]
# The writers read ``zone.key``, so tests pass a ZoneInfo, as the view does.
UTC_ZONE = ZoneInfo("UTC")
# The download context the view lends the XLSX and PDF writers.
CONTEXT = dict(
    parish="Example Parish",
    campaign="Stewardship 2055",
    as_of=datetime.fromisoformat(WHEN),
)


def workbook(shaped, zone=UTC_ZONE):
    """The XLSX download of ``shaped``, opened."""
    return load_workbook(io.BytesIO(limitations_xlsx(shaped, zone, **CONTEXT)))


def result():
    """A raw SQL projection with one Member and one Family."""
    return {
        "members": [
            {
                "family_name": "Example",
                "family_duid": 10,
                "member_name": "Alex Example",
                "member_duid": 3,
                "proposed": False,
                "cannot_serve": True,
                "submitted_at": WHEN,
                "kind": "member",
                "member_key": "3",
                "version": 0,
                "decision": None,
                "in_force": False,
                "decided_by": None,
                "decided_at": None,
                "note": None,
            }
        ],
        "families": [
            {
                "family_name": "Example",
                "family_duid": 10,
                "submitted_at": WHEN,
                "kind": "family",
                "member_key": "",
                "version": 1,
                "decision": "accepted",
                "in_force": True,
                "decided_by": "staff@example.org",
                "decided_at": WHEN,
                "note": "",
            }
        ],
        "summary": {"cannot_serve": 1, "cannot_attend": 1},
    }


def test_shaping_parses_response_times_and_keeps_the_summary():
    """The SQL rows pass through with parsed times; the summary is as read."""
    shaped = shape_result(result())
    assert shaped["members"][0]["submitted_at"].isoformat() == WHEN
    assert shaped["families"][0]["submitted_at"].isoformat() == WHEN
    assert shaped["summary"] == {"cannot_serve": 1, "cannot_attend": 1}


def test_a_family_no_longer_in_parishsoft_has_a_blank_name_and_its_status():
    """The download splits the page's badge out of the name cell (#1216)."""
    raw = result()
    raw["members"][0]["family_name"] = None
    raw["families"][0]["family_name"] = None
    rows = list(
        csv.reader(io.StringIO(limitations_csv(shape_result(raw), UTC).decode()))
    )
    assert rows[1][2:4] == ["", "Not in ParishSoft"]
    assert rows[4][:2] == ["", "Not in ParishSoft"]


def test_downloads_match_the_page_columns():
    """CSV and XLSX hold the page's columns in its order (#932, #970).

    A proposed Member has no DUID yet, so its Member DUID cell is empty.
    """
    raw = result()
    raw["members"].append(
        raw["members"][0]
        | {"member_name": "New Example", "member_duid": None, "proposed": True}
    )
    shaped = shape_result(raw)
    rows = list(csv.reader(io.StringIO(limitations_csv(shaped, UTC).decode())))
    assert rows[0] == MEMBER_HEADINGS
    member = ["Alex Example", "3", "Example", "", "10", "Yes", WHEN, "To review"]
    assert rows[1] == member
    assert rows[2] == ["New Example (proposed)", "", *member[2:]]
    rows.pop(2)
    assert rows[2] == [] and rows[3] == FAMILY_HEADINGS
    # The Family's answer was accepted, and its decision still holds.
    assert rows[4] == [
        "Example",
        "",
        "10",
        "Cannot attend Mass or prayer services",
        WHEN,
        "To handle",
    ]
    book = workbook(shaped)
    assert [cell.value for cell in book["Members"][1]] == MEMBER_HEADINGS
    assert [cell.value for cell in book["Families"][1]] == FAMILY_HEADINGS


def test_pdf_prints_both_tables_with_the_csv_columns_and_rows(monkeypatch):
    """The PDF (#966) is the CSV's two tables, Members then Families.

    Each table has the CSV's headings and values, only Latest response
    worded in the parish's date format; the Families table starts on its
    own page, and one frame names the parish, campaign and both counts.
    """
    zone = ZoneInfo("America/Chicago")
    raw = result()
    raw["members"][0]["member_name"] = "=cmd"
    shaped = shape_result(raw)
    drawn, original = [], pdf_design.table_pages

    def capture(table, rows, frame):
        """Keep each table and its rows, then lay them out."""
        rows = list(rows)
        drawn.append((table, rows, frame))
        return original(table, rows, frame)

    monkeypatch.setattr(pdf_design, "table_pages", capture)
    body = limitations_pdf(
        shaped,
        zone,
        parish="Example Parish",
        campaign="Stewardship 2055",
        as_of=datetime.fromisoformat(WHEN),
    )
    # Two pages: the Families table does not share the Members' page.
    assert body.startswith(b"%PDF-") and b"/Count 2" in body
    csv_rows = list(csv.reader(io.StringIO(limitations_csv(shaped, zone).decode())))
    (members, member_rows, frame), (families, family_rows, _) = drawn
    assert list(members.headings) == csv_rows[0] == MEMBER_HEADINGS
    assert list(families.headings) == csv_rows[3] == FAMILY_HEADINGS
    # The formula-like name prints as typed; the CSV guards it.
    # The Member DUID follows the Member, as on the page (#932).
    assert member_rows[0][:6] == ("=cmd", "3", "Example", "", "10", "Yes")
    assert member_rows[0][6] == family_rows[0][4] == "Oct 5, 2026 9:00 AM"
    assert family_rows[0][:4] == (
        "Example",
        "",
        "10",
        "Cannot attend Mass or prayer services",
    )
    assert frame.eyebrow == "Example Parish · Stewardship 2055"
    # The tables' frame adds two blank lines of room for their captions.
    assert frame.details == (
        "1 Member cannot participate in ministries; 1 Family cannot attend Mass.",
        "",
        "",
    )
    assert set(PDF_WEIGHTS) == {*MEMBER_HEADINGS, *FAMILY_HEADINGS}
    for table in (members, families):
        assert sum(table.widths) == pytest.approx(pdf_design.BODY_WIDTH)


def test_xlsx_report_information_states_what_the_pdf_frame_states():
    """The workbook's "Report information" sheet (#926) names the download.

    Like the other reports' workbooks, it states the parish, campaign,
    download time in the chosen zone, both counts and the privacy line the
    PDF's page frame shows; the CSV stays the two tables alone.
    """
    zone = ZoneInfo("America/Chicago")
    book = workbook(shape_result(result()), zone)
    assert book.sheetnames == ["Members", "Families", "Report information"]
    sheet = book["Report information"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows == [
        ("Report", "Limitations"),
        ("Parish", "Example Parish"),
        ("Campaign", "Stewardship 2055"),
        # A native date-time cell, in the download's zone (9 AM Chicago).
        ("Downloaded at", datetime(2026, 10, 5, 9, 0)),
        ("Display time zone", "America/Chicago"),
        ("Members in this file", "1"),
        ("Families in this file", "1"),
        ("Privacy", PRIVACY),
    ]
    assert sheet["B4"].is_date
    assert sheet["A1"].font.b and sheet.page_setup.orientation == "portrait"
    assert book.properties.title == "Limitations"
    assert book.properties.created == datetime(2026, 10, 5, 14, 0)


def test_exports_keep_formula_text_literal():
    """A parish-typed name such as =cmd is never a spreadsheet formula."""
    shaped = shape_result(result())
    shaped["members"][0]["member_name"] = "=cmd"
    rows = list(csv.reader(io.StringIO(limitations_csv(shaped, UTC).decode())))
    assert rows[1][0] == "'=cmd"
    book = workbook(shaped)
    value = book["Members"]["A2"]
    assert value.data_type == "s" and value.value == "=cmd"


def test_xlsx_escapes_control_characters_in_source_names():
    """A vertical tab pasted into a ParishSoft name cannot break the workbook.

    openpyxl rejects characters XML 1.0 forbids, so a raw write raised and the
    download failed. The shared writer escapes them reversibly instead.
    """
    raw = result()
    raw["members"][0]["member_name"] = "Alex\x0bExample"
    raw["families"][0]["family_name"] = "Example\x0b"
    shaped = shape_result(raw)
    book = workbook(shaped)
    member = book["Members"]["A2"]
    assert member.data_type == "s" and member.value == "Alex\\u000bExample"
    assert book["Families"]["A2"].value == "Example\\u000b"


def page(raw, paging=None, stage="all"):
    """Render the worklist page from a shaped result, as the view does."""
    from django.template.loader import render_to_string

    from parishkit.stewardship.reports.limitation_worklist import (
        SORTING,
        WorklistQuery,
        answer_rows,
        at_stage,
        worklist_page,
    )
    from parishkit.stewardship.web.tables import paginate

    shaped = shape_result(raw)
    query = WorklistQuery(stage=stage)
    shown, counts = at_stage(answer_rows(shaped), stage)
    table = paginate(shown, paging or {}, sorting=SORTING)
    worklist, table = worklist_page(counts, query, table, "/report", mutable=True)
    return render_to_string(
        "stewardship/limitations-report.html",
        {
            "worklist": worklist,
            "table": table,
            "campaign_id": "00000000-0000-4000-8000-000000000000",
            "query": query,
            "query_fields": query.form_values(),
            "mutable": True,
        },
    )


def test_report_page_is_named_limitations_and_never_shows_talents():
    """The page is "Limitations", with no talent column, count or filter."""
    html = page(result())
    assert "<h1>Limitations</h1>" in html
    assert "talent" not in html.casefold()
    assert '<select id="limitations-filter" name="show">' in html
    assert '<option value="cannot_serve"' in html
    assert '<option value="cannot_attend"' in html


def test_the_show_filter_takes_only_the_limitation_words():
    """The Show filter is closed, and the SQL takes the same parameters."""
    for word in ("any", "cannot_serve", "cannot_attend"):
        query = LimitationsQuery.parse({"search": "Alex", "show": word})
        assert query.form_values() == {"search": "Alex", "show": word}
    # A talent's id (the filter's old values) and the old field name are
    # refused, as any other unknown filter is.
    for invalid in ({"show": PAINTER}, {"talent": "any"}, {"show": "talents"}):
        with pytest.raises(ValueError):
            LimitationsQuery.parse(invalid)


def table_body(raw, paging):
    """The rendered worklist table of the page (shaping a copy of ``raw``)."""
    html = page(copy.deepcopy(raw), paging)
    return html[html.index("<table") : html.index("</table>")]


def test_the_table_leads_with_the_family_then_the_member_and_their_duids():
    """Family, Family DUID, Member, Member DUID, then the shared columns
    (#932, #1685); each DUID has its own right-aligned cell and sorts."""
    raw = result()
    raw["members"].append(
        raw["members"][0]
        | {"family_name": "Other", "family_duid": 7, "member_name": "B"}
        | {"member_duid": 41, "member_key": "41"}
    )
    table = table_body(raw, {"sort": "duid"})
    columns = ["family", "duid", "member", "member_duid"]
    positions = [table.index(f'data-sort-column="{key}"') for key in columns]
    assert positions == sorted(positions)
    assert 'data-sort-column="talents"' not in table
    assert '<td class="numeric">10</td><td>Alex Example</td>' in table
    assert '<td>Alex Example</td><td class="numeric">3</td>' in table
    assert "Cannot participate in any ministries" in table
    assert "Cannot attend Mass or prayer services" in table
    # Sorted by Family DUID: 7 comes before 10.
    assert table.index(">7</td>") < table.index(">10</td>")
    # Sorted by Member DUID, descending: 41 comes before 3.
    table = table_body(raw, {"sort": "-member_duid"})
    assert table.index(">41</td>") < table.index(">3</td>")


def test_a_proposed_member_reads_new_member_and_sorts_last():
    """A proposed Member has no DUID; its cell reads New Member and sorts
    after the Members' DUIDs."""
    raw = result()
    raw["members"].append(
        raw["members"][0]
        | {"member_name": "New", "member_duid": None, "proposed": True}
        | {"kind": "proposed_member", "member_key": str(UUID(int=5))}
    )
    table = table_body(raw, {"sort": "member_duid"})
    assert '<td>New</td><td class="numeric">New Member</td>' in table
    assert table.index(">3</td>") < table.index("New Member")


def test_the_stage_filter_lists_and_counts_each_stage():
    """The Family's accepted answer is To handle; the Member's is To review."""
    html = page(result(), stage="to_review")
    assert ">To review (1)</option>" in html and ">To handle (1)</option>" in html
    assert ">Every stage (2)</option>" in html
    table = html[html.index("<table") : html.index("</table>")]
    assert "Alex Example" in table and "To handle" not in table


def captured_pdf(monkeypatch, shaped):
    """Render the PDF; return its body, table layouts and drawn captions."""
    drawn, captions, original = [], [], pdf_design.table_pages

    def capture(table, rows, frame):
        """Keep each table and its rows, then lay them out."""
        rows = list(rows)
        drawn.append((table, rows, frame))
        return original(table, rows, frame)

    text = pdf_design.Canvas.text

    def keep_caption(canvas, x, y, value, size, face="regular", *args, **kwargs):
        """Note each caption drawn, then draw the text."""
        if value in (MEMBER_CAPTION, FAMILY_CAPTION):
            captions.append((value, y))
        return text(canvas, x, y, value, size, face, *args, **kwargs)

    monkeypatch.setattr(pdf_design, "table_pages", capture)
    monkeypatch.setattr(pdf_design.Canvas, "text", keep_caption)
    body = limitations_pdf(
        shaped,
        ZoneInfo("America/Chicago"),
        parish="Example Parish",
        campaign="Stewardship 2055",
        as_of=datetime.fromisoformat(WHEN),
    )
    return body, drawn, captions


def test_pdf_names_each_table_with_its_page_heading(monkeypatch):
    """Every PDF page names its table as the page's table heading does."""
    _, drawn, captions = captured_pdf(monkeypatch, shape_result(result()))
    assert [caption for caption, _ in captions] == [MEMBER_CAPTION, FAMILY_CAPTION]
    # Each caption sits in the room the tables leave below the header rule
    # (two detail lines), above the table's heading row.
    (_, _, body), _ = drawn
    room = body.body_top - 2 * pdf_design.SMALL_PITCH
    assert all(room < y < body.body_top for _, y in captions)


def test_pdf_with_no_families_still_prints_both_tables(monkeypatch):
    """No Family who cannot attend: the Families page says no rows match.

    The Families table still gets its own captioned page, so the file says
    plainly that the list is empty rather than leaving it out.
    """
    raw = result()
    raw["families"] = []
    raw["summary"]["cannot_attend"] = 0
    body, drawn, captions = captured_pdf(monkeypatch, shape_result(raw))
    assert body.startswith(b"%PDF-") and b"/Count 2" in body
    (_, members, frame), (families, rows, _) = drawn
    assert len(members) == 1 and rows == []
    assert list(families.headings) == FAMILY_HEADINGS
    assert [caption for caption, _ in captions] == [MEMBER_CAPTION, FAMILY_CAPTION]
    assert frame.details[0].endswith("0 Families cannot attend Mass.")
