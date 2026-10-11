"""The chosen-Family test names export's file and access (#817).

Pure tests: the CSV the worker writes from a retained capture (every
reviewed DUID in order and a name that looks like a spreadsheet formula
kept as text), the report information the XLSX and PDF state, and that
only the page's configure capability may read the export. The command
and the SQL guards against a real database are in
database/test_admin_test_families_cli_postgresql.py.
"""

import csv
import io
from datetime import UTC, datetime
from types import SimpleNamespace
from uuid import UUID

import pytest

from parishkit.stewardship.reports import export_services, family_test_names
from parishkit.stewardship.reports.family_test_names import names_document
from parishkit.stewardship.reports.information_rendering import render_information
from parishkit.stewardship.reports.pdf_design import PdfFrame

CAMPAIGN = UUID("00000000-0000-4000-8000-000000000041")
REVISION = UUID("00000000-0000-4000-8000-000000000042")
REQUESTER = UUID("00000000-0000-4000-8000-000000000043")
CAPTURED = datetime(2054, 10, 6, 15, 30, tzinfo=UTC)
CAMPAIGN_NAME = "Stewardship 2055"


@pytest.fixture(autouse=True)
def named_campaign(monkeypatch):
    """The capture's campaign name, without a database."""
    monkeypatch.setattr(
        family_test_names, "campaign_name", lambda request: CAMPAIGN_NAME
    )


def request(rows, *, row_count=None):
    """A ``family_test_names`` export of ``rows`` (``(duid, name)`` pairs)."""
    snapshot = SimpleNamespace(
        document={"rows": [{"duid": duid, "name": name} for duid, name in rows]},
        row_count=len(rows) if row_count is None else row_count,
        parameters={"revision": str(REVISION), "duids": [duid for duid, _ in rows]},
        created_at=CAPTURED,
    )
    return SimpleNamespace(
        family_test_names_snapshot=snapshot,
        browser_timezone="America/New_York",
        configuration=SimpleNamespace(parish=SimpleNamespace(name="Example Parish")),
        campaign_id=CAMPAIGN,
        created_at=CAPTURED,
    )


def test_the_file_lists_every_reviewed_family_in_order():
    """DUID and name per row, and nothing else (#966); the report
    information (for the XLSX and PDF) says what was captured and when,
    with no internal references."""
    rows = [(1234, "Squyres, Tracy and Jeff"), (99, ""), (5, "=HYPERLINK(1)")]
    document = names_document(request(rows))
    assert document.item_count == 3
    metadata = dict(document.metadata)
    assert metadata["Report"] == "Chosen-Family test names"
    assert metadata["Families"] == "3"
    assert not {"Email revision", "Campaign reference"} & set(metadata)
    # Instants are in the export's time zone.
    assert metadata["Captured at"].isoformat().startswith("2054-10-06T11:30")
    output = io.BytesIO()
    render_information(document, output, format="csv")
    lines = list(csv.reader(io.StringIO(output.getvalue().decode())))
    assert lines[0] == ["Family DUID", "Family name"]
    assert lines[1:] == [
        ["1234", "Squyres, Tracy and Jeff"],
        ["99", ""],
        # A name that looks like a formula stays text, as every export's.
        ["5", "'=HYPERLINK(1)"],
    ]


def test_the_pdf_frame_names_the_campaign():
    """Like every report, the file names its campaign; the PDF eyebrow shows it."""
    document = names_document(request([(1, "A")]))
    assert ("Campaign", CAMPAIGN_NAME) in document.metadata
    assert PdfFrame.for_document(document).eyebrow == (
        f"Example Parish · {CAMPAIGN_NAME}"
    )


def test_a_capture_missing_a_row_is_refused():
    """The file holds every reviewed Family or is not written."""
    with pytest.raises(ValueError):
        names_document(request([(1, "A")], row_count=2))


@pytest.mark.parametrize(
    ("roles", "allowed"),
    [(frozenset({"administrator"}), True), (frozenset({"staff"}), False)],
)
def test_only_the_configure_capability_reads_the_names(monkeypatch, roles, allowed):
    """Campaign reporting alone (Staff) is not the page's capability."""
    from parishkit.stewardship.accounts.policy import Principal
    from parishkit.stewardship.reports.export_models import ExportRequest

    monkeypatch.setattr(
        export_services,
        "current_principal",
        lambda store, user_id: Principal(identity=user_id, roles=roles),
    )
    export = ExportRequest(report="family_test_names", requester_id=REQUESTER)
    if allowed:
        assert export_services.authorize(None, REQUESTER, request=export)
    else:
        with pytest.raises(PermissionError):
            export_services.authorize(None, REQUESTER, request=export)


def test_xlsx_and_pdf_carry_the_csvs_columns():
    """One column set in every format (#966): the XLSX sheet is the CSV's
    whole table, and each PDF card holds exactly a row's DUID and name."""
    from openpyxl import load_workbook

    from parishkit.stewardship.reports.family_test_names import FORMATS, HEADINGS
    from parishkit.stewardship.reports.information_rendering import (
        information_records,
    )

    assert FORMATS == ("csv", "xlsx", "pdf")
    rows = [(1234, "Squyres, Tracy and Jeff"), (99, "")]
    document = names_document(request(rows))
    output = io.BytesIO()
    render_information(document, output, format="csv")
    table = list(csv.reader(io.StringIO(output.getvalue().decode())))
    output = io.BytesIO()
    render_information(document, output, format="xlsx")
    book = load_workbook(io.BytesIO(output.getvalue()))
    sheet = [
        ["" if value is None else value for value in values]
        for values in book["Families"].iter_rows(values_only=True)
    ]
    assert sheet == table and tuple(sheet[0]) == HEADINGS
    _, *cards = information_records(document)
    assert [(card.title, card.fields) for card in cards] == [
        ("Squyres, Tracy and Jeff (DUID 1234)", ()),
        ("DUID 99", ()),
    ]
    output = io.BytesIO()
    render_information(document, output, format="pdf")
    assert output.getvalue().startswith(b"%PDF-")
